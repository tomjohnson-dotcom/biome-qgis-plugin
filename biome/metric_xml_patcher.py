# metric_xml_patcher.py
# Writes QGIS habitat data into the statutory BNG metric .xlsm by patching
# the worksheet XML parts directly, instead of round-tripping the workbook
# through openpyxl.
#
# Why: openpyxl's keep_vba save silently strips the metric's macro-assigned
# navigation buttons, chart styling, drawings, logos and conditional-
# formatting extensions. Patching only the target sheet XML leaves every
# other part of the workbook byte-identical, so the exported file looks and
# behaves exactly like the official template.
#
# Drop-in usage from xlsm_exporter._export_metric:
#
#     from .metric_xml_patcher import patch_metric, validate_rows
#
#     warnings = validate_rows(_BUNDLED_TEMPLATE, [
#         ("area",  rows_area_bl + rows_area_cr),
#         ("hedge", rows_hedge_bl + rows_hedge_cr),
#         ("water", rows_water_bl + rows_water_cr),
#     ])
#     patch_metric(_BUNDLED_TEMPLATE, out_path, [
#         ("A-1 On-Site Habitat Baseline", rows_area_bl,  _A1_WRITE_COLS, _A1_DATA_START_ROW),
#         ("A-2 On-Site Habitat Creation", rows_area_cr,  _A2_WRITE_COLS, _A2_DATA_START_ROW),
#         ("B-1 On-Site Hedge Baseline",   rows_hedge_bl, _B1_WRITE_COLS, _B1_DATA_START_ROW),
#         ("B-2 On-Site Hedge Creation",   rows_hedge_cr, _B2_WRITE_COLS, _B2_DATA_START_ROW),
#         ("C-1 On-Site WaterC' Baseline", rows_water_bl, _C1_WRITE_COLS, _C1_DATA_START_ROW),
#         ("C-2 On-Site WaterC' Creation", rows_water_cr, _C2_WRITE_COLS, _C2_DATA_START_ROW),
#     ])
#
# `rows` and the (col_num, field, fmt) write-column tuples are exactly the
# structures _export_metric already builds — no changes needed upstream.
# fmt is treated as "numeric" when not None, matching existing behaviour.
#
# Pure stdlib for the patching itself; openpyxl (read-only) is used only by
# validate_rows to read the template's own lookup lists.

import re
import zipfile
from xml.sax.saxutils import escape


class MetricPatchError(RuntimeError):
    """Raised when the template does not match expectations."""


# ── Internal helpers ─────────────────────────────────────────────────────────

def _col_letter(n):
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def _sheet_parts(zf):
    """Map sheet name → zip part path, via workbook.xml + its rels."""
    wb_xml = zf.read("xl/workbook.xml").decode("utf-8")
    rels = zf.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    rel_map = dict(re.findall(
        r'Id="(rId\d+)"[^>]*Target="(worksheets/[^"]+)"', rels))
    out = {}
    for m in re.finditer(r'<sheet name="([^"]+)"[^>]*r:id="(rId\d+)"', wb_xml):
        name, rid = m.group(1), m.group(2)
        if rid in rel_map:
            out[name] = "xl/" + rel_map[rid]
    return out


def _cell_xml(ref, style, val):
    """Serialise a value into a styled cell. Numbers as <v>, text inline.

    ``xml:space="preserve"`` is REQUIRED on the <t> element: several metric
    lookup keys carry significant leading/trailing spaces (e.g. the off-site
    "Spatial risk category" strings in G-3, and some habitat/condition names).
    Without it Excel silently trims that whitespace on load, so the value no
    longer matches its VLOOKUP key and the row scores zero / errors. Harmless
    for values that have no edge whitespace.
    """
    if isinstance(val, (int, float)):
        return f'<c r="{ref}" s="{style}"><v>{val!r}</v></c>'
    return (f'<c r="{ref}" s="{style}" t="inlineStr">'
            f'<is><t xml:space="preserve">{escape(str(val))}</t></is></c>')


def _capacity(sheet_xml, first_col_letter, start_row):
    """Count contiguous pre-styled empty input rows from start_row."""
    refs = re.findall(
        rf'<c r="{first_col_letter}(\d+)" s="\d+"\s*/>', sheet_xml)
    rows = sorted(int(r) for r in refs if int(r) >= start_row)
    n = 0
    for i, r in enumerate(rows):
        if r == start_row + i:
            n += 1
        else:
            break
    return n


def _patch_sheet(sheet_xml, sheet_name, rows, write_cols, start_row):
    """Return sheet XML with row values written into empty input cells.

    Mirrors the openpyxl writer's behaviour: empty/None values leave the
    cell untouched; cells already holding a formula are skipped; numeric
    fields (fmt is not None) are coerced to float or skipped.
    """
    cap = _capacity(sheet_xml, _col_letter(write_cols[0][0]), start_row)
    if len(rows) > cap:
        raise MetricPatchError(
            f"{sheet_name}: {len(rows)} rows to write but the template only "
            f"has {cap} input rows (rows {start_row}–{start_row + cap - 1}). "
            f"Aggregate further or split the export.")

    for row_idx, row_data in enumerate(rows, start_row):
        for col_num, field, fmt in write_cols:
            val = row_data.get(field)
            if val is None or val == "":
                continue
            if fmt is not None:
                try:
                    val = float(val)
                except (ValueError, TypeError):
                    continue
            ref = f"{_col_letter(col_num)}{row_idx}"

            empty_pat = re.compile(rf'<c r="{ref}" s="(\d+)"\s*/>')
            m = empty_pat.search(sheet_xml)
            if m:
                sheet_xml = empty_pat.sub(
                    _cell_xml(ref, m.group(1), val), sheet_xml, count=1)
                continue

            # Cell exists with content — skip formulas (shaded cells),
            # error on anything else so we never silently misplace data.
            existing = re.search(rf'<c r="{ref}"[^>]*>(.*?)</c>',
                                 sheet_xml, re.S)
            if existing and "<f>" in existing.group(1):
                continue
            raise MetricPatchError(
                f"{sheet_name}: input cell {ref} not found as an empty "
                f"styled cell — template layout may have changed.")
    return sheet_xml


# ── Public API ───────────────────────────────────────────────────────────────

def patch_metric(template_path, out_path, sheets):
    """Copy template_path to out_path with data written into input cells.

    sheets: list of (sheet_name, rows, write_cols, data_start_row), using the
    exact rows / _*_WRITE_COLS / _*_DATA_START_ROW structures from
    xlsm_exporter. Everything outside the patched sheet parts is copied
    byte-for-byte; calcChain.xml is dropped and fullCalcOnLoad is set so
    Excel rebuilds all formula values on first open.
    """
    with zipfile.ZipFile(template_path) as zin:
        parts = _sheet_parts(zin)
        patched = {}
        for sheet_name, rows, write_cols, start_row in sheets:
            if sheet_name not in parts:
                raise MetricPatchError(
                    f"Sheet not found in template: {sheet_name!r}")
            if not rows:
                continue
            part = parts[sheet_name]
            xml = patched.get(part) or zin.read(part).decode("utf-8")
            patched[part] = _patch_sheet(
                xml, sheet_name, rows, write_cols, start_row)

        wb_part = "xl/workbook.xml"
        wb_xml = zin.read(wb_part).decode("utf-8")
        if "fullCalcOnLoad" not in wb_xml:
            wb_xml = re.sub(r"<calcPr ", '<calcPr fullCalcOnLoad="1" ',
                            wb_xml, count=1)
        patched[wb_part] = wb_xml

        with zipfile.ZipFile(out_path, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                if item.filename == "xl/calcChain.xml":
                    continue
                data = patched.get(item.filename)
                if data is not None:
                    zout.writestr(item, data)
                else:
                    zout.writestr(item, zin.read(item.filename))


def validate_rows(template_path, groups):
    """Check habitat-type strings against the template's own lookup lists.

    groups: list of (kind, rows) where kind is "area", "hedge" or "water"
    and rows are the dicts headed for the metric sheets. Returns a list of
    human-readable warnings; an unmatched name means that row will score
    zero units in the metric, so surface these to the user before export.
    """
    from openpyxl import load_workbook

    wb = load_workbook(template_path, read_only=True, data_only=True)

    def _col_values(sheet, col, min_row, max_row):
        ws = wb[sheet]
        return {str(r[0]).strip()
                for r in ws.iter_rows(min_row=min_row, max_row=max_row,
                                      min_col=col, max_col=col,
                                      values_only=True)
                if r[0]}

    valid = {
        "area":  _col_values("G-1 All Habitats", 1, 3, 140),
        "hedge": _col_values("G-6 Hedgerow Data", 2, 3, 20),
        "water": _col_values("G-7 WaterC' Data", 2, 4, 12),
    }
    wb.close()

    warnings = []
    for kind, rows in groups:
        names = valid.get(kind, set())
        seen = set()
        for row in rows:
            sub = str(row.get("sub_habitat_type") or "").strip()
            if sub and sub not in names and sub not in seen:
                seen.add(sub)
                warnings.append(
                    f"{kind}: \u201c{sub}\u201d is not in the metric's "
                    f"habitat list — its rows will score zero units.")
    return warnings
