# desk_study_tab.py — DeskStudyMixin, extracted verbatim from main_panel.py (2026-07-22 split).
# Methods run on BNGMainPanel via mixin inheritance; `self` is the panel.


from qgis.PyQt.QtWidgets import (QWidget, QVBoxLayout, QHBoxLayout, QPushButton, QLabel,
    QGridLayout, QMessageBox, QSlider, QMenu)
from qgis.PyQt.QtCore import (Qt, QSettings, QTimer, QSize, QObject, QEvent,
    pyqtSignal)
from qgis.PyQt.QtGui import QCursor, QActionGroup
from qgis.core import Qgis
from . import theme
from .base_layers import (load_gcn_for_point, remove_gcn_layers, _GCN_GROUP_NAME)
from . import region as _region
from .designated_sites import (DESIGNATED_LAYERS, load_designated_layer,
    ds_label, all_ds_labels,
    remove_designated_layer, query_designated_sites_2km, query_priority_habitats_2km,
    add_2km_buffer, remove_2km_buffer, query_epsl_2km, epsl_impact_summary, load_epsl_layer,
    set_epsl_label_mode, load_epsl_distance_lines, apply_epsl_star_renderer,
    apply_epsl_emoji_renderer, apply_epsl_recent_historic_renderer, query_tpo_2km,
    load_tpo_layer, load_tpo_distance_lines, query_tpo_zones_2km, load_tpo_zones_layer,
    query_conservation_areas_2km, load_conservation_areas_layer, get_lpa_for_bng,
    get_tpo_council_portal, PHI_HABITAT_GROUPS, load_phi_group_layer, hide_phi_group_layer,
    load_ancient_woodland_layer, hide_ancient_woodland_layer, load_irz_layer)
from .council_maps import (build_tpo_deeplink, tpo_data_download, lpa_homepage, lpa_full_name,
    tposearch_link)
from .pond_mapper import (get_or_create_pond_layer, remove_pond_layer, add_pond,
    remove_pond_distance_lines, set_pond_lines_visible,
    load_pond_distance_lines, POND_LAYER_NAME, POND_LINES_LAYER_NAME,
    apply_pond_label_mode, pond_label_mode, ensure_pond_fields)
from .config import (get_os_api_key, get_os_ngd_api_key,
                     get_google_maps_api_key)
from .config import get_phi_group_nearby, set_phi_group_nearby


from .panel_widgets import (_dbg, _NumericSortItem, _svg_pixmap, _svg_icon, _dot_pixmap,
    _banner)
from .panel_map_tools import (_PondPlaceTool, _PondMoveTool, _FeatureInfoTool)


class _PondAbortFilter(QObject):
    """App-level filter alive ONLY while pond placement is armed: pressing any
    OTHER button in the Biome panel means the user has moved on to another
    activity, so finish placing before that button's action runs. Tool
    switches, tab switches, Escape and right-click already disarm via other
    paths — this catches panel actions that don't change the map tool
    (searches, reports, layer toggles). Never consumes the event; parented to
    the panel so Qt auto-removes it if the panel dies while installed."""

    def __init__(self, panel):
        super().__init__(panel)
        self._panel = panel

    def eventFilter(self, obj, event):
        try:
            if event.type() != QEvent.Type.MouseButtonPress:
                return False
            from qgis.PyQt.QtWidgets import QAbstractButton
            if not isinstance(obj, QAbstractButton):
                return False
            panel = self._panel
            # Belt-and-braces: never fire (→ pan trigger) unless actually armed.
            tool = getattr(panel, "_pond_tool", None)
            if tool is None or panel.iface.mapCanvas().mapTool() is not tool:
                return False
            # The place button itself toggles off normally and the lines toggle
            # is part of the same activity. Other actions, including opening the
            # marker manager, finish placement before they run.
            if obj is getattr(panel, "_pond_place_btn", None) \
                    or obj is getattr(panel, "_pond_lines_btn", None):
                return False
            w = obj
            while w is not None and w is not panel:
                w = w.parentWidget()
            if w is panel:
                panel._finish_pond_placing()
        except Exception:
            pass
        return False


def _flash_copied(btn, original_label):
    """Brief 'Copied' feedback on report buttons that copy invisible text."""
    from qgis.PyQt.QtCore import QTimer
    from qgis.PyQt import sip
    btn.setText("Copied")
    def _restore():
        if not sip.isdeleted(btn):
            btn.setText(original_label)
    QTimer.singleShot(1200, _restore)


def _phi_type_keys(row):
    """Return the distinct underlying habitat types represented by a PHI row.

    Combined PHI records such as ``DWOOD, LMEAD`` describe two existing
    habitat types, not a third combined classification.  Codes are preferred
    because habitat names can contain commas themselves.
    """
    codes = []
    for raw_code in str(row.get("hab_code", "") or "").split(","):
        code = raw_code.strip().upper()
        if code and code != "NMHAB" and code not in codes:
            codes.append(code)
    if codes:
        return tuple(codes)

    # Defensive fallback for sources without PHI codes, including the
    # separately sourced Ancient Woodland row.
    labels = []
    for raw_label in str(row.get("habitat", "") or "").split("/"):
        label = " ".join(raw_label.split()).casefold()
        if label and label not in labels:
            labels.append(label)
    return tuple(labels)


def build_landing_page(site_name, phone_url, port, token, tunnel, lan_ip):
    """The page the desktop browser opens: a QR code pointing at the phone URL.

    Module-level and pure so it can be rendered and checked without starting a
    server. The QR is drawn by the plugin's own generator rather than a CDN
    script — this page has to work on a machine with no internet, which is
    exactly the case where the tunnel also failed.
    """
    # qr_generator (the plugin's own, inline SVG, no internet needed) is tried
    # first, but QGIS does NOT bundle the `qrcode` package on either 3.44 or
    # 4.2 despite what that module's docstring claimed, so in practice this
    # falls through to the CDN renderer \u2014 hence the explicit onerror handling,
    # which turns a silently blank card into a readable instruction.
    qr_svg = ''
    try:
        from .qr_generator import url_to_qr_svg
        qr_svg = url_to_qr_svg(phone_url, cell=8)
    except Exception:
        qr_svg = ''

    if tunnel:
        note = ('<div class="note ok">Secure HTTPS tunnel is up, so GPS, the '
                'camera and offline map caching all work on the phone.</div>')
    elif lan_ip == '127.0.0.1':
        note = ('<div class="note bad">No network address could be found, so '
                'this QR code will not reach your phone. Connect this computer '
                'to Wi-Fi and press the button again, or let it download '
                'cloudflared for a secure tunnel.</div>')
    else:
        note = ('<div class="note">Serving over local Wi-Fi without HTTPS. Your '
                'phone must be on the <b>same network</b>, and iOS will block '
                'GPS, the camera and offline map caching. Press the button '
                'again and accept the cloudflared download to enable them.</div>')

    css = """
:root{--muted:#54635A;--accent:#2E6B4F;--dark:#1F4A37;
      --surface:#F7F8F6;--sunken:#F1F4F0;--border:#E3E7E1}
*{box-sizing:border-box}
body{font-family:-apple-system,BlinkMacSystemFont,'Segoe UI',sans-serif;
     background:var(--surface);display:flex;align-items:center;
     justify-content:center;min-height:100vh;margin:0;padding:24px}
.card{background:#fff;border-radius:18px;padding:30px 34px 26px;max-width:400px;
      width:100%;box-shadow:0 6px 28px rgba(20,40,30,0.10);text-align:center;
      border:1px solid var(--border)}
.mark{width:42px;height:42px;border-radius:12px;background:var(--accent);
      margin:0 auto 12px;display:flex;align-items:center;justify-content:center}
h1{font-size:19px;color:var(--dark);margin:0 0 3px;font-weight:650}
.site{font-size:13px;color:var(--muted);margin:0 0 18px}
.qr{display:flex;justify-content:center;margin:6px 0 14px}
.qr svg{width:236px;height:236px;border:12px solid #fff;border-radius:10px;
        box-shadow:0 2px 14px rgba(20,40,30,0.14)}
.url-box{font-size:11.5px;color:var(--accent);word-break:break-all;
  background:var(--sunken);border-radius:9px;padding:10px 13px;margin:2px 0 14px;
  font-family:ui-monospace,SFMono-Regular,Menlo,monospace}
.open-btn{display:inline-block;background:var(--dark);color:#fff;
  border-radius:11px;padding:12px 26px;font-size:14px;font-weight:600;
  text-decoration:none}
.open-btn:hover{background:var(--accent)}
.copy-btn{display:block;width:100%;margin:0 0 12px;background:#fff;
  color:var(--accent);border:1px solid var(--border);border-radius:10px;
  padding:9px;font-size:13px;font-weight:600;cursor:pointer;
  font-family:inherit}
.copy-btn:hover{background:var(--sunken)}
.qr-fail{font-size:12.5px;color:#8F2A23;background:#FBEDEB;border-radius:10px;
  padding:16px 14px;line-height:1.55;margin:0}
.note{font-size:11.5px;color:#7a5800;background:#FBF4E4;border-radius:9px;
  padding:10px 13px;margin-top:10px;line-height:1.55;text-align:left}
.note.ok{color:#175E3B;background:#E9F4EC}
.note.bad{color:#8F2A23;background:#FBEDEB}
"""

    # Placeholder substitution, not %-formatting or .format(): the CSS contains
    # both "width:100%" and braces, either of which breaks those.
    template = (
        '<!DOCTYPE html>\n<html><head><meta charset="utf-8"/>\n'
        '<meta name="viewport" content="width=device-width,initial-scale=1"/>\n'
        '<title>Biome Field Mapper</title>\n<style>@CSS@</style></head><body>\n'
        '<div class="card">\n'
        '  <div class="mark"><svg width="24" height="24" viewBox="0 0 24 24" '
        'fill="none" stroke="#fff" stroke-width="1.8" stroke-linecap="round" '
        'stroke-linejoin="round">'
        '<path d="M5 21c.5-4.5 2.5-8 7-10"/>'
        '<path d="M9 18c6.2 0 10.5-3.3 11-12v-2h-4c-9 0-12 4-12 9c0 1 0 3 2 5h3"/>'
        '</svg></div>\n'
        '  <h1>Biome Field Mapper</h1>\n'
        '  <p class="site">@SITE@ &middot; scan with your phone camera</p>\n'
        '  <div class="qr" id="qr">@QR@</div>\n'
        '  <div class="url-box" id="url">@URL@</div>\n'
        '  <button class="copy-btn" id="copy">Copy link</button>\n'
        '  <a class="open-btn" href="http://127.0.0.1:@PORT@/?v=@TOKEN@" '
        'target="_blank">Open on this computer instead</a>\n'
        '  <div class="note">Leave QGIS running: the phone loads the app from '
        'it, and &ldquo;Send straight to QGIS&rdquo; needs it open at the end of '
        'the survey too. Records are saved on the phone either way.</div>\n'
        '  @NOTE@\n</div>\n@SCRIPT@\n</body></html>'
    )
    script = _LANDING_SCRIPT if not qr_svg else _COPY_SCRIPT
    for key, value in (('@CSS@', css), ('@QR@', qr_svg), ('@NOTE@', note),
                       ('@SITE@', _esc(site_name)), ('@URL@', _esc(phone_url)),
                       ('@PORT@', str(int(port))), ('@TOKEN@', _esc(token)),
                       ('@SCRIPT@', script)):
        template = template.replace(key, value)
    return template


# Copy-to-clipboard only — used when the QR was drawn server-side.
_COPY_SCRIPT = """<script>
document.getElementById('copy').onclick = function () {
  var t = document.getElementById('url').textContent.trim(), b = this;
  navigator.clipboard.writeText(t).then(function () {
    b.textContent = 'Copied'; setTimeout(function () { b.textContent = 'Copy link'; }, 1600);
  }, function () { b.textContent = 'Copy failed'; });
};
</script>"""

# CDN renderer plus a real failure path: if the script cannot load (no
# internet, blocked CDN) the card says so and points at the address, instead
# of showing an empty white square.
_LANDING_SCRIPT = """<script>
function qrFailed(msg) {
  var el = document.getElementById('qr');
  if (el) el.innerHTML = '<p class="qr-fail">' + msg +
    '<br>Type or copy the address below into your phone instead.</p>';
}
var s = document.createElement('script');
s.src = 'https://cdn.jsdelivr.net/npm/qrcodejs@1.0.0/qrcode.min.js';
s.onerror = function () { qrFailed('No internet on this computer, so the QR code could not be drawn.'); };
s.onload = function () {
  try {
    new QRCode(document.getElementById('qr'), {
      text: document.getElementById('url').textContent.trim(),
      width: 236, height: 236, colorDark: '#000000', colorLight: '#ffffff',
      correctLevel: QRCode.CorrectLevel.L});
  } catch (e) { qrFailed('The QR code could not be drawn.'); }
};
document.head.appendChild(s);
setTimeout(function () {
  var el = document.getElementById('qr');
  if (el && !el.childElementCount) qrFailed('The QR code took too long to load.');
}, 6000);
document.getElementById('copy').onclick = function () {
  var t = document.getElementById('url').textContent.trim(), b = this;
  navigator.clipboard.writeText(t).then(function () {
    b.textContent = 'Copied'; setTimeout(function () { b.textContent = 'Copy link'; }, 1600);
  }, function () { b.textContent = 'Copy failed'; });
};
</script>"""


def _esc(s):
    return (str(s or '').replace('&', '&amp;').replace('<', '&lt;')
            .replace('>', '&gt;').replace('"', '&quot;'))

class _FieldSyncBridge(QObject):
    """Carries a survey pushed by the phone from the HTTP worker thread to the
    GUI thread. Created in the GUI thread, so ``emit`` from a worker becomes a
    queued call — nothing touches QGIS off-thread."""
    received = pyqtSignal(object)


class _LpaLookupBridge(QObject):
    """Deliver a completed background LPA lookup back to the GUI thread."""
    resolved = pyqtSignal(object)


class DeskStudyMixin:
    def _build_desk_study_tab(self):
        """Build the Desk Study tab with redline stats and ecology overlay toggles."""
        from qgis.PyQt.QtWidgets import QScrollArea, QFrame

        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setFrameShape(QFrame.Shape.NoFrame)
        scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        scroll.setStyleSheet("QScrollArea { background:#F7F8F6; }")

        # ── Basemap selector (same segmented bar as the Mapping tab) ──────
        # Pinned ABOVE the scroll area so it stays put while the rest of the
        # tab scrolls: outer = [pinned bar strip] + [scroll]. The strip's
        # margins mirror the inner layout's so the bar lines up with the
        # cards below it; a hairline marks where the content slides under.
        outer = QWidget()
        outer.setStyleSheet("background:#F7F8F6;")
        ol = QVBoxLayout(outer)
        ol.setContentsMargins(0, 0, 0, 0)
        ol.setSpacing(0)
        pin = QWidget()
        pin.setStyleSheet("background:#F7F8F6;")
        pl = QVBoxLayout(pin)
        pl.setContentsMargins(10, 12, 10, 8)
        pl.setSpacing(0)
        pl.addWidget(self._make_basemap_bar())
        ol.addWidget(pin)
        pin_rule = QFrame()
        pin_rule.setFixedHeight(1)
        pin_rule.setStyleSheet("background:#E3E7E1; border:none;")
        ol.addWidget(pin_rule)
        ol.addWidget(scroll, 1)

        inner = QWidget()
        inner.setStyleSheet(
            "background:#F7F8F6;")
        vl = QVBoxLayout()
        vl.setContentsMargins(10, 8, 10, 12)
        vl.setSpacing(8)

        # ── Shared helpers ─────────────────────────────────────────────────

        BTN_H = 44   # standard button height throughout

        def _section_hdr(text, color="#717E75"):
            lbl = QLabel(text.upper())
            lbl.setStyleSheet(
                f"font-size:11px; font-weight:600; color:#3F5148;"
                " letter-spacing:0.6px; padding:17px 3px 7px 3px;"
                " background:transparent;"
                " font-family:'Inter SemiBold','Inter','Segoe UI',sans-serif;"
            )
            return lbl

        def _section_group():
            group = QWidget()
            group.setObjectName("deskSectionGroup")
            group.setStyleSheet(
                "QWidget#deskSectionGroup { background:#EEF2EF; border:none;"
                " border-radius:11px; }")
            gl = QVBoxLayout(group)
            gl.setContentsMargins(6, 6, 6, 6)
            gl.setSpacing(4)
            return group, gl

        def _grid_widget(btns, cols=2, height=BTN_H):
            """Pack a flat list of QPushButtons into a cols-wide grid."""
            w = QWidget()
            g = QGridLayout(w)
            g.setContentsMargins(0, 0, 0, 0)
            g.setSpacing(4)
            for i, b in enumerate(btns):
                b.setMinimumHeight(height)
                b.setMaximumHeight(height + 8)
                g.addWidget(b, i // cols, i % cols)
            for c in range(cols):
                g.setColumnStretch(c, 1)
            return w

        def _std_btn(label, fg, border, hover, checkable=False, tooltip="",
                     icon_name=None):
            btn = QPushButton(label)
            btn.setCheckable(checkable)
            btn.setCursor(QCursor(Qt.CursorShape.PointingHandCursor))
            btn.setAttribute(Qt.WidgetAttribute.WA_AlwaysShowToolTips)
            if tooltip:
                self._wire_hint(btn, tooltip)
            if icon_name:
                from qgis.PyQt.QtCore import QSize
                # Off = tinted to fg; On (checked) = white to read on the fill.
                btn.setIcon(_svg_icon(icon_name, fg, 18, on_color="white"))
                btn.setIconSize(QSize(18, 18))
            _align = "left" if icon_name else "center"
            _pad = "8px 8px 8px 11px" if icon_name else "8px 6px"
            btn.setStyleSheet(f"""
                QPushButton {{
                    background:white; color:{fg};
                    border:1px solid #E3E7E1; border-radius:10px;
                    font-size:12px; font-weight:500;
                    padding:{_pad}; text-align:{_align};
                    letter-spacing:0px;
                    font-family: 'Inter Medium','Inter','Segoe UI',sans-serif;
                }}
                QPushButton:checked {{ background:{fg}; color:white; border-color:{fg}; }}
                QPushButton:hover:!checked {{ background:{hover}; border-color:{border}; }}
                QPushButton:pressed {{ background:{border}; color:white; }}
            """)
            return btn

        def _report_btn(label, tooltip, fg, border, hover):
            btn = _std_btn(label, fg, border, hover, tooltip=tooltip)
            return btn

        TOOL_FG = "#48616B"   # neutral blue-grey for ordinary tools and links

        def _tool_btn(label, icon_name, color=TOOL_FG, checkable=False,
                      tooltip="", hover="#F4F6F2", trailing=None,
                      trailing_click=False):
            """Icon-left card button: tinted line icon + wrapping label, both
            vertically centred (built as child QLabels so multi-line text aligns
            properly). `color` tints the icon/text and is the checked-fill colour
            (toggling flips icon+text white). `btn._label_lbl` is exposed so
            stateful handlers can change the text WITHOUT drawing over the child
            label (never call btn.setText()). With `trailing_click` the trailing
            glyph becomes a small QToolButton (exposed as `btn._trail_btn` for
            the caller to wire) instead of a decorative label."""
            btn = QPushButton()
            btn.setCheckable(checkable)
            btn.setCursor(QCursor(Qt.CursorShape.PointingHandCursor))
            btn.setMinimumHeight(BTN_H)
            if tooltip:
                self._wire_hint(btn, tooltip)
            btn.setStyleSheet(
                "QPushButton { background:white; border:1px solid transparent;"
                " border-radius:10px; }"
                f" QPushButton:checked {{ background:{color}; border-color:{color}; }}"
                f" QPushButton:hover:!checked {{ background:{hover}; border-color:{color}; }}")
            row = QHBoxLayout(btn)
            row.setContentsMargins(12, 6, 10, 6)
            row.setSpacing(9)
            ico = QLabel()
            ico.setFixedSize(18, 18)
            ico.setAttribute(Qt.WidgetAttribute.WA_TransparentForMouseEvents)
            ico.setStyleSheet("background:transparent; border:none;")
            ico.setPixmap(_svg_pixmap(icon_name, color, 18))
            _css = ("color:%s; font-size:12px; font-weight:500;"
                    " background:transparent; border:none;"
                    " font-family:'Inter Medium','Inter','Segoe UI',sans-serif;")
            txt = QLabel(label)
            txt.setWordWrap(True)
            txt.setAttribute(Qt.WidgetAttribute.WA_TransparentForMouseEvents)
            txt.setStyleSheet(_css % color)
            trail = None
            if trailing and trailing_click:
                from qgis.PyQt.QtWidgets import QToolButton
                trail = QToolButton()
                trail.setFixedSize(24, 24)
                trail.setCursor(QCursor(Qt.CursorShape.PointingHandCursor))
                trail.setIcon(_svg_icon(trailing, "#717E75", 15))
                trail.setIconSize(QSize(15, 15))
                trail.setStyleSheet(
                    "QToolButton { border:none; background:transparent;"
                    " border-radius:6px; }"
                    " QToolButton:hover { background:#E3EAE4; }")
                btn._trail_btn = trail
            elif trailing:
                trail = QLabel()
                trail.setFixedSize(15, 15)
                trail.setAttribute(Qt.WidgetAttribute.WA_TransparentForMouseEvents)
                trail.setStyleSheet("background:transparent; border:none;")
                trail.setPixmap(_svg_pixmap(trailing, "#717E75", 15))
            row.addWidget(ico)
            row.addWidget(txt, 1)
            if trail is not None:
                row.addWidget(trail)
            if checkable:
                def _sync(checked, _i=icon_name, _ico=ico, _txt=txt, _c=color,
                          _tr=trail, _trn=trailing):
                    col = "white" if checked else _c
                    _ico.setPixmap(_svg_pixmap(_i, col, 18))
                    _txt.setStyleSheet(_css % col)
                    if _tr is not None and hasattr(_tr, "setPixmap"):
                        _tr.setPixmap(_svg_pixmap(_trn, "white" if checked else "#717E75", 15))
                btn.toggled.connect(_sync)
                # Programmatic state restoration sometimes has to block Qt
                # signals to avoid running the button action. Expose the purely
                # visual half so those callers can still recolour icon + text.
                btn._sync_checked_style = _sync
            btn._label_lbl = txt
            btn._icon_lbl = ico
            # Redirect setText() to the child label (stripping any leading emoji
            # and newlines) so the existing stateful-button setText() calls
            # update the label instead of drawing the button's own text over it.
            def _set_label(s, _t=txt):
                t = str(s).replace("\n", " ")
                i = 0
                while i < len(t) and not (t[i].isalnum() or t[i] == "("):
                    i += 1
                _t.setText(t[i:].strip())
            btn.setText = _set_label
            return btn

        def _collapsible_section(header_text, header_color, content_widget,
                                 collapsed=True, icon=None):
            """Collapsible constraint row: a clean white header with a
            module-colour cue, the title, and a line-icon chevron that flips
            on expand. `icon` names an icons/*.svg tinted to `header_color`;
            without one the original small dot is drawn (kept for NESTED
            sections, whose lighter cue marks them as a sub-level)."""
            container = QWidget()
            container.setStyleSheet("background:transparent;")
            cvl = QVBoxLayout(container)
            cvl.setContentsMargins(0, 0, 0, 0)
            cvl.setSpacing(4)

            hdr_btn = QPushButton()
            hdr_btn.setCheckable(True)
            hdr_btn.setChecked(not collapsed)
            hdr_btn.setCursor(QCursor(Qt.CursorShape.PointingHandCursor))
            hdr_btn.setStyleSheet(
                "QPushButton { background:white; border:1px solid transparent;"
                " border-radius:9px; }"
                " QPushButton:hover { background:#F7F9F7; border-color:#C9D1C9; }")
            hdr_btn.setMinimumHeight(38)
            hrow = QHBoxLayout(hdr_btn)
            hrow.setContentsMargins(11, 7, 10, 7)
            hrow.setSpacing(10)
            if icon:
                dot = QLabel(); dot.setFixedSize(20, 20)
                dot.setPixmap(_svg_pixmap(icon, header_color, 20))
            else:
                dot = QLabel(); dot.setFixedSize(10, 10)
                dot.setPixmap(_dot_pixmap(header_color, 10))
            dot.setAttribute(Qt.WidgetAttribute.WA_TransparentForMouseEvents)
            dot.setStyleSheet("background:transparent; border:none;")
            title = QLabel(header_text)
            title.setAttribute(Qt.WidgetAttribute.WA_TransparentForMouseEvents)
            title.setStyleSheet(
                "color:#1A2420; font-size:12px; font-weight:500; background:transparent;"
                " border:none; font-family:'Inter Medium','Inter','Segoe UI',sans-serif;")
            chev = QLabel(); chev.setFixedSize(16, 16)
            chev.setAttribute(Qt.WidgetAttribute.WA_TransparentForMouseEvents)
            chev.setStyleSheet("background:transparent; border:none;")
            chev.setPixmap(_svg_pixmap(
                "chevron-down" if not collapsed else "chevron-right", "#717E75", 16))
            hrow.addWidget(dot)
            hrow.addWidget(title, 1)
            hrow.addWidget(chev)

            # Group the section's buttons in a tinted, bordered well so that when
            # the section is open they read as belonging to it, instead of the
            # white buttons blending into the column of white section rows.
            well = QWidget()
            well.setObjectName("dsWell")
            well.setStyleSheet(
                "QWidget#dsWell { background:#ECF1EC; border:1px solid #DBE2DA;"
                " border-radius:9px; }")
            _wl = QVBoxLayout(well)
            _wl.setContentsMargins(7, 7, 7, 7)
            _wl.setSpacing(0)
            content_widget.setStyleSheet("background:transparent;")
            _wl.addWidget(content_widget)
            well.setVisible(not collapsed)

            def _toggle(checked, _c=chev, _w=well):
                _w.setVisible(checked)
                _c.setPixmap(_svg_pixmap(
                    "chevron-down" if checked else "chevron-right", "#717E75", 16))

            hdr_btn.clicked.connect(_toggle)
            cvl.addWidget(hdr_btn)
            cvl.addWidget(well)
            return container

        # ── Contextual hint — shown until redline is drawn ────────────────────
        self._ds_hint_card = QWidget()
        hint_vl = QVBoxLayout(self._ds_hint_card)
        hint_vl.setContentsMargins(0, 0, 0, 0)
        hint_vl.setSpacing(0)
        hint_vl.addWidget(_banner(
            "info-circle",
            "<b>Getting started with the Desk Study</b><br/><br/>"
            "1. Go to <b>Map &amp; Tools</b> and draw your <b>Red Line Boundary</b> first.<br/>"
            "2. Come back here and click <b>Search buffer</b> to show your search area "
            "(right-click it to change the distance).<br/>"
            "3. Expand the <b>Constraints</b> sections below to load designated sites, "
            "priority habitats, GCN risk zones, EPSL locations, and standing water bodies.<br/>"
            "4. Use the <b>report buttons</b> inside each section to generate distance tables.",
            fg="#1F4A37", bg="#F1F4F0", border="1px solid #2E6B4F", font_px=11))
        # This is onboarding, not an empty-project warning. Existing users who
        # have completed the plugin's first-run flow should never see it, and a
        # new user sees it in their first session only.
        _ds_settings = QSettings()
        _onboarding_complete = _ds_settings.value(
            "biome/onboarding_complete", False, type=bool)
        _desk_hint_seen = _ds_settings.value(
            "biome/desk_study/getting_started_seen", False, type=bool)
        self._ds_show_first_run_hint = (
            not _onboarding_complete and not _desk_hint_seen)
        self._ds_hint_card.setVisible(self._ds_show_first_run_hint)
        if self._ds_show_first_run_hint:
            _ds_settings.setValue(
                "biome/desk_study/getting_started_seen", True)
        vl.addWidget(self._ds_hint_card)

        # ── Redline stats card ─────────────────────────────────────────────
        stats_card = QFrame()
        stats_card.setObjectName("dsStatsCard")
        stats_card.setStyleSheet(
            "QFrame#dsStatsCard { background:#F6F5F2; border:1px solid #E1DED7;"
            " border-radius:10px; }")
        stats_vl = QVBoxLayout()
        stats_vl.setContentsMargins(12, 10, 12, 12)
        stats_vl.setSpacing(8)

        # Header — line icon + sentence-case title + a quiet copy hint
        _sh = QWidget(); _sh.setStyleSheet("background:transparent;")
        _shl = QHBoxLayout(_sh); _shl.setContentsMargins(0, 0, 0, 0); _shl.setSpacing(7)
        _shi = QLabel(); _shi.setFixedSize(18, 18)
        _shi.setStyleSheet("background:transparent; border:none;")
        _shi.setPixmap(_svg_pixmap("map-pin", "#1F7A4D", 18))
        _shl.addWidget(_shi)
        stats_header = QLabel("Red Line boundary")
        stats_header.setStyleSheet(
            "font-size:13px; font-weight:500; color:#1A2420; background:transparent;"
            " border:none; font-family:'Inter Medium','Inter','Segoe UI',sans-serif;")
        _shl.addWidget(stats_header)
        _shl.addStretch()
        _copy_hint = QLabel("tap to copy")
        _copy_hint.setStyleSheet(
            "font-size:11px; color:#717E75; background:transparent; border:none;")
        _shl.addWidget(_copy_hint)
        stats_vl.addWidget(_sh)

        # The three area units now share one tile: left-click copies the value
        # shown, right-click cycles ha → m² → km² (hectares by default).
        self._DS_AREA_UNITS = ("Area (ha)", "Area (m²)", "Area (km²)")
        self._ds_area_vals  = ["—", "—", "—"]          # ha / m² / km², cycle order
        try:
            self._ds_area_unit_idx = int(
                QSettings().value("biome/desk_study/area_unit_idx", 0)) % 3
        except Exception:
            self._ds_area_unit_idx = 0

        self._ds_area_value = QLabel("—")
        self._ds_centroid   = QLabel("—")
        self._ds_lpa        = QLabel("—")   # local planning authority (from centroid)

        def _stat_card(label, value_lbl, val_size=15, on_right=None, tooltip=None):
            card = QFrame()
            card.setObjectName("dsTile")
            card.setStyleSheet(
                "QFrame#dsTile { background:#F1F4F0; border:none; border-radius:8px; }")
            card.setCursor(QCursor(Qt.CursorShape.PointingHandCursor))
            if tooltip:
                card.setToolTip(tooltip)
            cl = QVBoxLayout(card)
            cl.setContentsMargins(10, 8, 10, 9)
            cl.setSpacing(2)
            lbl = QLabel(label)
            lbl.setStyleSheet(
                "font-size:11px; color:#717E75; font-weight:400; background:transparent;"
                " border:none; font-family:'Inter','Segoe UI',sans-serif;")
            value_lbl.setStyleSheet(
                "font-size:%dpx; font-weight:500; color:#1A2420; background:transparent;"
                " border:none; font-family:'Inter Medium','Inter','Segoe UI',sans-serif;"
                % val_size)
            value_lbl.setWordWrap(True)
            cl.addWidget(lbl)
            cl.addWidget(value_lbl)

            def _flash():
                card.setStyleSheet(
                    "QFrame#dsTile { background:#D9E8DE; border:none; border-radius:8px; }")
                _base_bg = getattr(card, "_ds_base_bg", "#F1F4F0")
                QTimer.singleShot(260, lambda: card.setStyleSheet(
                    "QFrame#dsTile { background:%s; border:none; border-radius:8px; }"
                    % _base_bg))

            def _on_click(e):
                if on_right is not None and e.button() == Qt.MouseButton.RightButton:
                    on_right()
                    _flash()
                    return
                txt = value_lbl.text()
                if txt and txt not in ("—", "No redline boundary", "No valid geometry"):
                    from qgis.PyQt.QtWidgets import QApplication
                    QApplication.clipboard().setText(txt)
                _flash()

            card.mousePressEvent = _on_click
            return card, lbl

        area_card, self._ds_area_caption = _stat_card(
            "", self._ds_area_value, on_right=self._cycle_ds_area_unit,
            tooltip="Left-click to copy · Right-click to switch unit (ha / m² / km²)")
        centroid_card, _ = _stat_card("Centroid", self._ds_centroid)
        lpa_card, _ = _stat_card(
            "Local planning authority", self._ds_lpa, val_size=14,
            tooltip="The LPA whose boundary contains the redline centroid "
                    "(planning.data.gov.uk) — left-click to copy")

        grid = QGridLayout()
        grid.setSpacing(8)
        grid.setContentsMargins(0, 0, 0, 0)
        grid.addWidget(area_card,     0, 0)
        grid.addWidget(centroid_card, 0, 1)
        grid.addWidget(lpa_card,      1, 0, 1, 2)   # full-width row (names are long)
        grid.setColumnStretch(0, 1)
        grid.setColumnStretch(1, 1)
        stats_vl.addLayout(grid)
        for _card in (area_card, centroid_card, lpa_card):
            _card._ds_base_bg = "#EEEDE8"
            _card.setStyleSheet(
                "QFrame#dsTile { background:#EEEDE8; border:none; border-radius:8px; }")

        self._render_ds_area()

        stats_card.setLayout(stats_vl)
        vl.addWidget(stats_card)

        # ── Blue-line stats card ────────────────────────────────────────
        # Only shown when a Blue Line Boundary exists.  It mirrors the
        # redline's useful geometry summary but deliberately has no LPA row.
        self._ds_blue_stats_card = QFrame()
        self._ds_blue_stats_card.setObjectName("dsBlueStatsCard")
        self._ds_blue_stats_card.setStyleSheet(
            "QFrame#dsBlueStatsCard { background:#F1F7FE; border:1px solid #C8DCF4;"
            " border-radius:10px; }")
        blue_vl = QVBoxLayout(self._ds_blue_stats_card)
        blue_vl.setContentsMargins(12, 10, 12, 12)
        blue_vl.setSpacing(8)

        _bh = QWidget(); _bh.setStyleSheet("background:transparent;")
        _bhl = QHBoxLayout(_bh); _bhl.setContentsMargins(0, 0, 0, 0); _bhl.setSpacing(7)
        _bhi = QLabel(); _bhi.setFixedSize(18, 18)
        _bhi.setStyleSheet("background:transparent; border:none;")
        _bhi.setPixmap(_svg_pixmap("map-pin", "#0066CC", 18))
        _bhl.addWidget(_bhi)
        blue_header = QLabel("Blue Line boundary")
        blue_header.setStyleSheet(
            "font-size:13px; font-weight:500; color:#1A2420; background:transparent;"
            " border:none; font-family:'Inter Medium','Inter','Segoe UI',sans-serif;")
        _bhl.addWidget(blue_header)
        _bhl.addStretch()
        _blue_copy_hint = QLabel("tap to copy")
        _blue_copy_hint.setStyleSheet(
            "font-size:11px; color:#717E75; background:transparent; border:none;")
        _bhl.addWidget(_blue_copy_hint)
        blue_vl.addWidget(_bh)

        self._ds_blue_area_vals = ["—", "—", "—"]
        self._ds_blue_area_value = QLabel("—")
        self._ds_blue_centroid = QLabel("—")
        blue_area_card, self._ds_blue_area_caption = _stat_card(
            "", self._ds_blue_area_value, on_right=self._cycle_ds_area_unit,
            tooltip="Left-click to copy · Right-click to switch unit (ha / m² / km²)")
        blue_centroid_card, _ = _stat_card("Centroid", self._ds_blue_centroid)
        blue_grid = QGridLayout()
        blue_grid.setSpacing(8)
        blue_grid.setContentsMargins(0, 0, 0, 0)
        blue_grid.addWidget(blue_area_card, 0, 0)
        blue_grid.addWidget(blue_centroid_card, 0, 1)
        blue_grid.setColumnStretch(0, 1)
        blue_grid.setColumnStretch(1, 1)
        blue_vl.addLayout(blue_grid)
        for _card in (blue_area_card, blue_centroid_card):
            _card._ds_base_bg = "#E6F0FC"
            _card.setStyleSheet(
                "QFrame#dsTile { background:#E6F0FC; border:none; border-radius:8px; }")
        self._ds_blue_stats_card.setVisible(False)
        vl.addWidget(self._ds_blue_stats_card)

        # ── Mobile View Link ───────────────────────────────────────────────
        self._mobile_link_btn = _tool_btn(
            "Open field mapper", "device-mobile",
            tooltip="Generate a Biome Field Mapper file with the redline pre-loaded — place features in the field and import back into QGIS"
        )

        def _open_field_mapper():
            """Serve the Field Mapper on a free port, then expose it via a
            Cloudflare Quick Tunnel so the phone gets a real HTTPS URL
            (required for GPS on iOS).  Falls back to plain LAN HTTP if
            cloudflared is not available."""
            from qgis.core import (QgsCoordinateReferenceSystem,
                                   QgsCoordinateTransform, QgsProject)
            from qgis.PyQt.QtWidgets import QMessageBox, QApplication
            from qgis.PyQt.QtGui import QDesktopServices
            from qgis.PyQt.QtCore import QUrl
            from .field_mapper_export import (build_field_mapper_html,
                                              build_service_worker)
            import json, socket, threading, os, re, sys, shutil, subprocess
            import http.server, socketserver

            redline_layer = self.layer_manager.get_layer("redline")
            if redline_layer is None or redline_layer.featureCount() == 0:
                QMessageBox.warning(self, "No Red Line",
                    "Please draw a Red Line Boundary first.")
                return

            wgs84  = QgsCoordinateReferenceSystem("EPSG:4326")
            to_wgs = QgsCoordinateTransform(redline_layer.crs(), wgs84, QgsProject.instance())

            features_geojson, all_lons, all_lats = [], [], []
            for feat in redline_layer.getFeatures():
                geom = feat.geometry()
                if geom.isEmpty():
                    continue
                geom.transform(to_wgs)
                gj = json.loads(geom.asJson())
                features_geojson.append({
                    "type": "Feature", "geometry": gj,
                    "properties": {"name": "Red Line Boundary"}
                })
                def _collect(coords):
                    if isinstance(coords[0], list):
                        for c in coords: _collect(c)
                    else:
                        all_lons.append(coords[0]); all_lats.append(coords[1])
                _collect(gj["coordinates"])

            if not features_geojson:
                QMessageBox.warning(self, "No Geometry", "No valid redline geometry found.")
                return

            # Name the survey after the project so the phone's title bar and
            # the exported filename both identify the site.
            site_name = site_ref = surveyor = ""
            try:
                from .xlsm_exporter import load_remembered_details
                _d = load_remembered_details() or {}
                site_name = (_d.get("site_address") or _d.get("applicant") or "").strip()
                site_ref = (_d.get("project_name") or "").strip()
                surveyor = (_d.get("completed_by") or "").strip()
            except Exception:
                pass
            if not site_name:
                site_name = os.path.splitext(os.path.basename(
                    QgsProject.instance().fileName() or ""))[0] or "Field Mapper"

            app_html = build_field_mapper_html(
                json.dumps({"type": "FeatureCollection", "features": features_geojson}),
                (min(all_lats) + max(all_lats)) / 2,
                (min(all_lons) + max(all_lons)) / 2,
                site_name=site_name, site_ref=site_ref, surveyor=surveyor,
            )
            sw_js = build_service_worker(os.urandom(4).hex())

            # ── Shut down previous server / tunnel ────────────────────────────
            prev = getattr(self, '_fm_server', None)
            if prev:
                t = threading.Thread(target=prev.shutdown, daemon=True)
                t.start(); t.join(timeout=3)
                self._fm_server = None
            prev_t = getattr(self, '_fm_tunnel_proc', None)
            if prev_t:
                try: prev_t.terminate(); prev_t.wait(timeout=5)
                except Exception:
                    try: prev_t.kill()
                    except Exception: pass
                self._fm_tunnel_proc = None

            # ── LAN IP ────────────────────────────────────────────────────────
            try:
                with socket.socket(socket.AF_INET, socket.SOCK_DGRAM) as _s:
                    _s.connect(('8.8.8.8', 80))
                    lan_ip = _s.getsockname()[0]
            except Exception:
                lan_ip = '127.0.0.1'

            # ── HTTP server on a free port ────────────────────────────────────
            class _Server(socketserver.ThreadingTCPServer):
                allow_reuse_address = True
                daemon_threads      = True

            httpd_probe = _Server(('', 0), http.server.BaseHTTPRequestHandler)
            actual_port = httpd_probe.server_address[1]
            httpd_probe.server_close()

            token       = os.urandom(4).hex()
            app_bytes   = app_html.encode('utf-8')

            # ── Find or download cloudflared ──────────────────────────────────
            def _find_cf():
                found = shutil.which('cloudflared')
                if found: return found
                # Store in biome_data/ next to the plugin folder — survives reinstalls
                plugin_dir  = os.path.dirname(__file__)
                plugins_dir = os.path.dirname(plugin_dir)
                data_dir    = os.path.join(plugins_dir, 'biome_data')
                ext = '.exe' if sys.platform.startswith('win') else ''
                cached = os.path.join(data_dir, f'cloudflared{ext}')
                if os.path.isfile(cached) and os.access(cached, os.X_OK):
                    return cached
                return None

            def _download_cf():
                plugin_dir  = os.path.dirname(__file__)
                plugins_dir = os.path.dirname(plugin_dir)
                data_dir    = os.path.join(plugins_dir, 'biome_data')
                os.makedirs(data_dir, exist_ok=True)
                import urllib.request, tempfile
                if sys.platform.startswith('win'):
                    import struct
                    arch = 'amd64' if struct.calcsize('P') * 8 == 64 else '386'
                    url  = (f'https://github.com/cloudflare/cloudflared/releases/'
                            f'latest/download/cloudflared-windows-{arch}.exe')
                    dest = os.path.join(data_dir, 'cloudflared.exe')
                elif sys.platform == 'darwin':
                    import platform as _pf
                    arch = 'arm64' if _pf.machine() == 'arm64' else 'amd64'
                    url  = (f'https://github.com/cloudflare/cloudflared/releases/'
                            f'latest/download/cloudflared-darwin-{arch}.tgz')
                    dest = os.path.join(data_dir, 'cloudflared')
                else:
                    import platform as _pf
                    m = _pf.machine().lower()
                    arch = 'arm64' if ('aarch64' in m or 'arm64' in m) else 'amd64'
                    url  = (f'https://github.com/cloudflare/cloudflared/releases/'
                            f'latest/download/cloudflared-linux-{arch}')
                    dest = os.path.join(data_dir, 'cloudflared')
                try:
                    tmp = tempfile.mktemp(dir=data_dir)
                    urllib.request.urlretrieve(url, tmp)
                    if url.endswith('.tgz'):
                        import tarfile
                        with tarfile.open(tmp, 'r:gz') as tar:
                            for mem in tar.getmembers():
                                if mem.name.endswith('cloudflared'):
                                    mem.name = os.path.basename(mem.name)
                                    tar.extract(mem, data_dir)
                        os.remove(tmp)
                    else:
                        if os.path.exists(dest): os.remove(dest)
                        os.rename(tmp, dest)
                    if not sys.platform.startswith('win'):
                        os.chmod(dest, 0o755)
                    return dest
                except Exception:
                    try: os.remove(tmp)
                    except Exception: pass
                    return None

            cf_bin = _find_cf()
            tunnel_url = None

            if cf_bin is None:
                reply = QMessageBox.question(
                    self, "GPS requires HTTPS",
                    "To enable GPS on your phone, a small one-time download "
                    "of cloudflared (~30 MB) from Cloudflare is needed.\n\n"
                    "Download now?",
                    QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No)
                if reply == QMessageBox.StandardButton.Yes:
                    self.iface.messageBar().pushMessage(
                        "Biome", "Downloading cloudflared…",
                        level=Qgis.MessageLevel.Info, duration=0)
                    QApplication.processEvents()
                    cf_bin = _download_cf()
                    self.iface.messageBar().clearWidgets()
                    if cf_bin is None:
                        QMessageBox.warning(self, "Download Failed",
                            "Could not download cloudflared.\n"
                            "GPS will not work on iPhone.\n\n"
                            "Install manually from:\n"
                            "https://developers.cloudflare.com/"
                            "cloudflare-one/connections/"
                            "connect-networks/downloads/")

            if cf_bin:
                try:
                    # ── HTTP server must be up BEFORE cloudflared connects ────
                    _app_bytes_ref = app_bytes
                    _sw_bytes_ref  = sw_js.encode('utf-8')
                    class _EarlyHandler(http.server.BaseHTTPRequestHandler):
                        def log_message(self, *a): pass
                        def log_error(self, *a):   pass
                        def do_GET(self):
                            path = self.path.split('?')[0].rstrip('/') or '/'
                            if path == '/':
                                data  = _app_bytes_ref
                                ctype = 'text/html; charset=utf-8'
                            elif path == '/sw.js':
                                data  = _sw_bytes_ref
                                ctype = 'application/javascript; charset=utf-8'
                            else:
                                self.send_response(404); self.end_headers(); return
                            self.send_response(200)
                            self.send_header('Content-Type', ctype)
                            self.send_header('Content-Length', str(len(data)))
                            self.send_header('Cache-Control', 'no-store')
                            self.end_headers()
                            self.wfile.write(data)

                    httpd_early = _Server(('', actual_port), _EarlyHandler)
                    threading.Thread(target=httpd_early.serve_forever,
                                     daemon=True).start()
                    self._fm_server = httpd_early

                    proc = subprocess.Popen(
                        [cf_bin, 'tunnel', '--url',
                         f'http://127.0.0.1:{actual_port}'],
                        stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                        text=True)
                    self._fm_tunnel_proc = proc
                    _url_box = [None]
                    _ready   = threading.Event()
                    def _drain():
                        try:
                            for line in proc.stderr:
                                m = re.search(
                                    r'https://[a-z0-9-]+\.trycloudflare\.com',
                                    line)
                                if m:
                                    _url_box[0] = m.group(0)
                                    _ready.set()
                            for _ in proc.stderr: pass
                        except Exception: pass
                    threading.Thread(target=_drain, daemon=True).start()
                    _ready.wait(timeout=15)
                    tunnel_url = _url_box[0]
                    if not tunnel_url:
                        try: proc.terminate()
                        except Exception: pass
                        self._fm_tunnel_proc = None
                except Exception:
                    tunnel_url = None

            # phone URL carries the token → serves the app
            # desktop URL is bare → serves the landing page with QR
            # The desktop ALWAYS gets the landing page: without it there is no
            # QR to scan, and getting the app onto a phone is the entire point
            # of the button. It used to hand the desktop the tokened URL
            # whenever no tunnel came up, which just opened the field app in the
            # wrong browser.
            phone_url   = (f'{tunnel_url}/?v={token}' if tunnel_url
                           else f'http://{lan_ip}:{actual_port}/?v={token}')
            desktop_url = f'http://127.0.0.1:{actual_port}/'

            landing_bytes = build_landing_page(
                site_name=site_name, phone_url=phone_url, port=actual_port,
                token=token, tunnel=bool(tunnel_url), lan_ip=lan_ip,
            ).encode('utf-8')


            # Update handler now that landing_bytes is defined — restart server
            # so it serves the correct landing page with tunnel URL in QR code.
            # NOTE: when using cloudflared, ALL requests arrive from 127.0.0.1
            # (the tunnel daemon forwards locally), so we can't use client IP
            # to distinguish desktop vs phone. Instead we use the token: the
            # landing page omits ?v=TOKEN, so requests WITH the token are from
            # the phone (via QR) and requests WITHOUT are the desktop.
            prev_early = getattr(self, '_fm_server', None)
            if prev_early:
                t2 = threading.Thread(target=prev_early.shutdown, daemon=True)
                t2.start(); t2.join(timeout=3)
                self._fm_server = None

            sw_bytes = sw_js.encode('utf-8')

            # Cross-thread hand-off: the HTTP handler runs on a worker thread
            # and must never touch QGIS objects. Emitting a signal on an object
            # that lives in the GUI thread gives us a queued call for free.
            bridge = getattr(self, '_fm_bridge', None)
            if bridge is None:
                bridge = _FieldSyncBridge()
                bridge.received.connect(self._on_field_survey_pushed)
                self._fm_bridge = bridge

            class _Handler(http.server.BaseHTTPRequestHandler):
                def log_message(self, *a): pass
                def log_error(self, *a):   pass

                def _send(self, data, ctype, status=200):
                    self.send_response(status)
                    self.send_header('Content-Type', ctype)
                    self.send_header('Content-Length', str(len(data)))
                    self.send_header('Cache-Control',
                                     'no-store, no-cache, must-revalidate')
                    self.send_header('Pragma', 'no-cache')
                    self.end_headers()
                    self.wfile.write(data)

                def do_GET(self):
                    raw_path = self.path
                    path = raw_path.split('?')[0].rstrip('/') or '/'
                    has_token = f'v={token}' in raw_path
                    if path == '/':
                        # Phone QR URL contains the token; desktop URL is bare /
                        self._send(app_bytes if has_token else landing_bytes,
                                   'text/html; charset=utf-8')
                    elif path == '/sw.js':
                        self._send(sw_bytes, 'application/javascript; charset=utf-8')
                    else:
                        self.send_response(404)
                        self.send_header('Content-Length', '0')
                        self.end_headers()

                def do_POST(self):
                    """Accept a survey pushed straight from the phone."""
                    path = self.path.split('?')[0].rstrip('/') or '/'
                    if path != '/sync':
                        self.send_response(404)
                        self.send_header('Content-Length', '0')
                        self.end_headers()
                        return
                    try:
                        length = int(self.headers.get('Content-Length') or 0)
                        # Generous but finite: a long survey with photos is a
                        # few tens of MB, an unbounded read is a DoS.
                        if length <= 0 or length > 256 * 1024 * 1024:
                            raise ValueError('bad length')
                        raw = self.rfile.read(length)
                        payload = json.loads(raw.decode('utf-8'))
                        n = len(payload.get('features') or [])
                    except Exception as exc:
                        body = json.dumps({'error': str(exc)}).encode('utf-8')
                        self._send(body, 'application/json', status=400)
                        return
                    # Answer before importing — the import shows its own dialog
                    # and the phone should not wait on a modal message box.
                    body = json.dumps({'received': n}).encode('utf-8')
                    self._send(body, 'application/json')
                    bridge.received.emit(payload)

            httpd = _Server(('', actual_port), _Handler)
            threading.Thread(target=httpd.serve_forever, daemon=True).start()
            self._fm_server = httpd

            QDesktopServices.openUrl(QUrl(desktop_url))

        self._mobile_link_btn.clicked.connect(_open_field_mapper)


        vl.addWidget(_section_hdr("Mapping Tools"))

        from .config import (get_buffer_radius, get_search_radius,
                             get_pond_inner_radius)
        buffers_btn = _tool_btn(
            "Local buffers\n%s–%s m" % (
                get_pond_inner_radius(),
                get_search_radius("standing_water")),
            "circles", checkable=True,
            tooltip="Show shorter search areas and other buffers around the Red Line. "
                    "TPO uses a 200 m outer-only favourite. Standing water uses "
                    "250 m inner and 500 m outer. These switch automatically when "
                    "the relevant layers are shown. Right-click to customise.")
        buffers_btn.clicked.connect(self._on_redline_buffers_toggle)
        buffers_btn.setContextMenuPolicy(Qt.ContextMenuPolicy.CustomContextMenu)
        buffers_btn.customContextMenuRequested.connect(self._redline_buffers_menu)
        self._buffers_btn = buffers_btn

        self._buffer_2km_btn = _tool_btn(
            "Desk study area\n%s" % self._buffer_preset_label(get_buffer_radius()),
            "circle-dashed", checkable=True,
            tooltip="Show the shared search area used by designated sites, priority "
                    "habitats, ancient woodland, EPSL, conservation areas, GCN eDNA "
                    "records, and NBN Atlas. Right-click to change the distance.")
        self._buffer_2km_btn.clicked.connect(self._on_buffer_2km_toggle)
        self._buffer_2km_btn.setContextMenuPolicy(Qt.ContextMenuPolicy.CustomContextMenu)
        self._buffer_2km_btn.customContextMenuRequested.connect(self._buffer_radius_menu)

        # Feature info + Measure-from-redline are now on the map right-click menu
        # (plan mode) — see _canvas_context_menu — so the standalone tool buttons
        # were removed.

        self._prop_btn = _tool_btn("Property boundaries", "map-2", checkable=True,
                                   tooltip="Toggle HM Land Registry INSPIRE property boundary polygons")
        self._prop_btn.clicked.connect(self._on_property_boundaries_toggle)

        mapping_group, mapping_vl = _section_group()
        mapping_vl.addWidget(_grid_widget([buffers_btn, self._buffer_2km_btn]))
        mapping_vl.addWidget(_grid_widget([self._prop_btn, self._mobile_link_btn]))
        vl.addWidget(mapping_group)

        # ── Constraints ────────────────────────────────────────────────────
        vl.addWidget(_section_hdr("Constraints"))
        constraints_group, constraints_vl = _section_group()
        constraints_vl.setSpacing(3)
        vl.addWidget(constraints_group)

        # ── Designated Sites (collapsible) ─────────────────────────────
        _DS_COLOURS = {
            "sssi":         ("#C24038", "#FBEDEB"),
            "sac":          ("#2E6B4F", "#F1F4F0"),
            "spa":          ("#1F4A37", "#EBF2F8"),
            "ramsar":       ("#00838f", "#e0f7fa"),
            "nnr":          ("#1F7A4D", "#E9F4EC"),
            "lnr":          ("#558b2f", "#F0F7F1"),
            "aonb":         ("#B07D00", "#FBF6E6"),
            "nat_parks":    ("#175E3B", "#E9F4EC"),
        }
        self._ds_btns = {}
        ds_individual_btns = []
        for key, label, *_rest in DESIGNATED_LAYERS:
            fg, bg = _DS_COLOURS.get(key, ("#1A2420", "#fff"))
            btn = _tool_btn(label, "map-pin", color=fg, checkable=True,
                            tooltip=f"Toggle {label} overlay")
            btn.clicked.connect(
                lambda checked, k=key: self._on_ds_layer_toggle(k, checked)
            )
            btn.setContextMenuPolicy(Qt.ContextMenuPolicy.CustomContextMenu)
            btn.customContextMenuRequested.connect(
                lambda pos, k=key, b=btn: self._constraint_buffer_menu("ds", k, b, pos)
            )
            self._ds_btns[key] = btn
            ds_individual_btns.append(btn)

        self._ds_toggle_all_btn = _tool_btn("Show all designated sites", "eye", color="#1F4A37",
                                            checkable=True,
                                            tooltip="Load all designated site layers at once")
        self._ds_toggle_all_btn.clicked.connect(self._on_ds_toggle_all)

        self._ds_report_btn = _tool_btn(
            "Designated sites report", "file-text", color="#1F4A37",
            tooltip="Query all statutory designated sites within the search radius",
        )
        self._ds_report_btn.clicked.connect(self._run_ds_report)

        ds_content = QWidget()
        ds_vl = QVBoxLayout(ds_content)
        ds_vl.setContentsMargins(4, 4, 4, 4)
        ds_vl.setSpacing(4)
        ds_vl.addWidget(_grid_widget([self._ds_toggle_all_btn, self._ds_report_btn]))
        ds_vl.addWidget(_collapsible_section("Designation Types", "#1F4A37",
                                             _grid_widget(ds_individual_btns), collapsed=True))
        self._ds_map_btn = _tool_btn("Designated sites map", "map", color="#3A463F",
                                     trailing="pencil", trailing_click=True,
                                     tooltip="Create a print layout showing all toggled designated site overlays")
        self._ds_map_btn.clicked.connect(lambda: self._create_desk_study_map("Designated Sites Map"))
        ds_vl.addWidget(self._ds_map_btn)

        constraints_vl.addWidget(_collapsible_section("Designated Sites", "#1F4A37", ds_content, collapsed=True,
                                          icon="shield"))

        # ── SSSI Impact Risk Zones (collapsible) — standalone constraint ──
        self._irz_btn = _tool_btn("Show SSSI impact risk zones", "eye", color="#8F2A23",
                                  checkable=True,
                                  tooltip="Load SSSI Impact Risk Zones intersecting the redline — shows which SSSIs and development types require Natural England consultation")
        self._irz_btn.clicked.connect(self._on_irz_toggle)

        irz_content = QWidget()
        irz_vl = QVBoxLayout(irz_content)
        irz_vl.setContentsMargins(4, 4, 4, 4)
        irz_vl.setSpacing(4)
        irz_vl.addWidget(self._irz_btn)

        self._irz_url_btn = _tool_btn(
            "Open NE consultation page", "external-link", color="#1F4A37",
            tooltip="Open the NE IRZ tool for this zone — uses the direct URL from the ArcGIS dataset"
        )
        self._irz_url_btn.setEnabled(False)
        self._irz_url_btn.clicked.connect(self._open_irz_url)
        irz_vl.addWidget(self._irz_url_btn)

        from qgis.PyQt.QtWidgets import QLabel as _QL
        _irz_info = _QL(
            "<span style='font-size:10px;color:#54635A'>"
            "Load zones first. The IRZ Tool button opens the correct NE consultation "
            "page for the most relevant zone at this location."
            "</span>"
        )
        _irz_info.setWordWrap(True)
        _irz_info.setTextFormat(Qt.TextFormat.RichText)
        irz_vl.addWidget(_irz_info)

        self._irz_map_btn = _tool_btn("SSSI IRZ map", "map", color="#3A463F",
                                      trailing="pencil", trailing_click=True,
                                      tooltip="Create a print layout showing SSSI Impact Risk Zones intersecting the redline")
        self._irz_map_btn.clicked.connect(lambda: self._create_desk_study_map("SSSI IRZ Map"))
        irz_vl.addWidget(self._irz_map_btn)

        constraints_vl.addWidget(_collapsible_section("SSSI Impact Risk Zones", "#8F2A23", irz_content, collapsed=True,
                                          icon="status-alert"))

        # ── Local Wildlife Sites (collapsible) ─────────────────────────
        self._lws_add_btn = _tool_btn("Add wildlife site", "map-pin-plus", color="#00897b",
                                      tooltip="Enter a Local Wildlife Site name, grid reference and area to add it to the map")
        self._lws_add_btn.clicked.connect(self._on_lws_add)

        self._lws_report_btn = _tool_btn(
            "LWS report", "file-text", color="#00897b",
            tooltip="Generate a table of all entered Local Wildlife Sites with distance and direction from the redline",
        )
        self._lws_report_btn.clicked.connect(self._run_lws_report)

        self._lws_import_btn = _tool_btn(
            "Add LWS from file", "file-text", color="#00897b",
            tooltip="Import Local Wildlife Site boundaries from a GeoPackage or other spatial file",
        )
        self._lws_import_btn.clicked.connect(self._on_lws_import)

        self._lws_hide_btn = _tool_btn("Hide sites", "eye", color="#00897b",
                                       checkable=True,
                                       tooltip="Show or hide all Local Wildlife Site markers on the map")
        self._lws_hide_btn.setChecked(False)
        self._lws_hide_btn.setEnabled(False)
        self._lws_hide_btn.clicked.connect(self._on_lws_visibility_toggle)

        self._lws_clear_btn = _tool_btn("Clear all sites", "trash", color="#922b21",
                                        tooltip="Remove all manually entered Local Wildlife Sites from the map")
        self._lws_clear_btn.clicked.connect(self._on_lws_clear)

        # Small list widget to show entered sites
        from qgis.PyQt.QtWidgets import QListWidget
        self._lws_list = QListWidget()
        self._lws_list.setMaximumHeight(100)
        self._lws_list.setStyleSheet(
            "QListWidget { font-size:10px; color:#004d40; background:#f0faf8;"
            " border:1px solid #b2dfdb; border-radius:4px; }"
            "QListWidget::item { padding:2px 4px; }"
            "QListWidget::item:selected { background:#00897b; color:white; }"
        )
        self._lws_list.setToolTip("Double-click a site to remove it")
        self._lws_list.itemDoubleClicked.connect(self._on_lws_list_double_click)
        self._lws_feature_ids = []   # parallel list of feature ids

        self._lws_map_btn = _tool_btn("Local wildlife sites map", "map", color="#3A463F",
                                      trailing="pencil", trailing_click=True,
                                      tooltip="Create a print layout showing entered Local Wildlife Sites")
        self._lws_map_btn.clicked.connect(lambda: self._create_desk_study_map("LWS Map"))

        lws_content = QWidget()
        lws_vl = QVBoxLayout(lws_content)
        lws_vl.setContentsMargins(4, 4, 4, 4)
        lws_vl.setSpacing(4)
        lws_vl.addWidget(_grid_widget([self._lws_add_btn, self._lws_report_btn]))
        lws_vl.addWidget(_grid_widget([self._lws_import_btn, self._lws_hide_btn]))
        lws_vl.addWidget(self._lws_clear_btn)
        lws_vl.addWidget(self._lws_map_btn)

        constraints_vl.addWidget(_collapsible_section("Local Wildlife Sites", "#00897b", lws_content, collapsed=True,
                                          icon="leaf"))

        # ── Priority Habitats (collapsible) ────────────────────────────
        self._phi_btns = {}
        phi_group_btns = []
        for gkey, glabel, gemoji, fg, bg, _habs in PHI_HABITAT_GROUPS:
            members = ", ".join(_habs)
            btn = _tool_btn(glabel, "plant-2", color=fg, checkable=True,
                            tooltip=f"Toggle {glabel} priority habitat polygons within the search radius\nIncludes: {members}")
            btn.clicked.connect(
                lambda checked, k=gkey: self._on_phi_group_toggle(k, checked)
            )
            btn.setContextMenuPolicy(Qt.ContextMenuPolicy.CustomContextMenu)
            btn.customContextMenuRequested.connect(
                lambda pos, k=gkey, b=btn: self._constraint_buffer_menu("phi", k, b, pos)
            )
            self._phi_btns[gkey] = btn
            phi_group_btns.append(btn)

        # Ancient Woodland — uses dedicated WFS (separate from PHI ArcGIS dataset)
        self._aw_btn = _tool_btn("Ancient woodland", "tree", color="#4e342e",
                                 checkable=True,
                                 tooltip="Toggle Ancient Woodland polygons within the search radius\n(Source: Natural England WFS)")
        self._aw_btn.clicked.connect(self._on_ancient_woodland_toggle)

        self._phi_toggle_all_btn = _tool_btn("Show all habitats", "eye", color="#175E3B",
                                             checkable=True,
                                             tooltip="Load all priority habitat groups at once")
        self._phi_toggle_all_btn.clicked.connect(self._on_phi_toggle_all)

        self._phi_report_btn = _tool_btn(
            "Priority habitats report", "file-text", color="#175E3B",
            tooltip="Query the Natural England Priority Habitats Inventory within the search radius"
                    "\n(Scottish sites: NatureScot Habitat Map of Scotland)",
        )
        self._phi_report_btn.clicked.connect(self._run_phi_report)

        phi_content = QWidget()
        phi_vl = QVBoxLayout(phi_content)
        phi_vl.setContentsMargins(4, 4, 4, 4)
        phi_vl.setSpacing(4)
        phi_vl.addWidget(_grid_widget([self._phi_toggle_all_btn, self._phi_report_btn]))
        phi_vl.addWidget(_collapsible_section("Habitat Groups", "#175E3B",
                                              _grid_widget([self._aw_btn] + phi_group_btns), collapsed=True))
        self._phi_map_btn = _tool_btn("Priority habitats map", "map", color="#3A463F",
                                      trailing="pencil", trailing_click=True,
                                      tooltip="Create a print layout showing all loaded priority habitat layers")
        self._phi_map_btn.clicked.connect(lambda: self._create_desk_study_map("Priority Habitats Map"))
        phi_vl.addWidget(self._phi_map_btn)

        constraints_vl.addWidget(_collapsible_section("Priority Habitats", "#175E3B", phi_content, collapsed=True,
                                          icon="plant-2"))

        # ── Standing Water Bodies / Pond Mapper (collapsible) ──────────
        self._pond_place_btn = _tool_btn("Place manually", "droplet", color="#1F4A37",
                                         checkable=True,
                                         tooltip="Place the required numbered marker for each water body to include")
        self._pond_place_btn.clicked.connect(self._on_pond_place_toggle)

        self._pond_auto_place_btn = _tool_btn(
            "Place automatically", "droplet", color="#1F4A37",
            tooltip="Place one marker on each highlighted water body at its "
                    "closest point to the redline")
        self._pond_auto_place_btn.clicked.connect(self._on_pond_auto_place)

        self._pond_manage_btn = _tool_btn(
            "Manage markers", "layers", color="#1F4A37",
            tooltip="Move or delete individual water-body markers")
        self._pond_manage_btn.clicked.connect(self._open_pond_marker_manager)

        self._pond_lines_btn = _tool_btn("Distance lines: off", "line", color="#1F4A37",
                                         checkable=True,
                                         tooltip="Show or hide distance lines from redline to each pond")
        self._pond_lines_btn.setChecked(False)
        self._pond_lines_btn.setEnabled(False)
        self._pond_lines_btn.clicked.connect(self._on_pond_lines_toggle)

        self._pond_labels_btn = _tool_btn(
            "Labels: Name + distance", "tag", color="#1F4A37",
            checkable=True, trailing="chevron-down", trailing_click=True,
            tooltip="Show or hide marker labels; use the arrow to choose their content")
        self._pond_labels_btn.clicked.connect(self._on_pond_labels_toggle)
        self._pond_labels_btn._trail_btn.clicked.connect(
            lambda: self._open_pond_labels_menu(self._pond_labels_btn))

        self._pond_map_btn = _tool_btn("Standing water body map", "map", color="#3A463F",
                                       trailing="pencil", trailing_click=True,
                                       tooltip="Create a print layout showing placed ponds")
        self._pond_map_btn.clicked.connect(lambda: self._create_desk_study_map("Pond Map"))

        self._os_topo_btn = _tool_btn("Highlight known ponds & lakes", "droplet", color="#1F4A37",
                                      checkable=True,
                                      tooltip="Highlight reference mapping within the search radius; this does not "
                                              "create the required water-body markers. Sources: "
                                              "OS Zoomstack surface water + OpenStreetMap (+ the Habitat Map of "
                                              "Scotland's lochs and ponds for a Scottish site)")
        self._os_topo_btn.clicked.connect(self._toggle_os_topo_layer)

        pond_content = QWidget()
        pond_vl = QVBoxLayout(pond_content)
        pond_vl.setContentsMargins(4, 4, 4, 4)
        pond_vl.setSpacing(6)

        self._pond_search_range_lbl = QLabel()
        self._pond_search_range_lbl.setStyleSheet(
            "color:#3F5148;font-size:10px;font-weight:500;"
            "background:transparent;border:none;padding:3px 3px 4px 3px;")
        pond_vl.addWidget(self._pond_search_range_lbl)

        def _pond_section_header(text, status_attr):
            row = QWidget()
            hl = QHBoxLayout(row)
            hl.setContentsMargins(3, 3, 3, 0)
            hl.setSpacing(6)
            heading = QLabel(text)
            heading.setStyleSheet(
                "color:#3F5148;font-size:10px;font-weight:600;"
                "background:transparent;border:none;")
            status = QLabel()
            status.setAlignment(Qt.AlignmentFlag.AlignCenter)
            status.setStyleSheet(
                "color:#3F5148;font-size:10px;font-weight:500;"
                "background:#F7F8F6;border:1px solid #E3E7E1;"
                "border-radius:7px;padding:3px 8px;")
            setattr(self, status_attr, status)
            hl.addWidget(heading, 1)
            hl.addWidget(status, 0)
            return row

        pond_vl.addWidget(_pond_section_header(
            "Find and correct pond outlines", "_pond_known_status_lbl"))
        pond_vl.addWidget(self._os_topo_btn)
        mapped_hint = QLabel(
            "Reference only. Small or seasonal water bodies may be absent. "
            "Check the aerial imagery and place a marker for each water body "
            "you want to include.")
        mapped_hint.setWordWrap(True)
        mapped_hint.setStyleSheet(
            "color:#54635A;font-size:10px;background:transparent;"
            "padding:2px 3px 5px 3px;")
        self._pond_reference_hint = mapped_hint
        pond_vl.addWidget(mapped_hint)

        self._pond_edit_known_btn = _tool_btn(
            "Correct pond outlines", "pencil", color="#315B83",
            checkable=True,
            tooltip="Move or delete vertices where the mapped OS NGD water "
                    "outline does not match the current aerial imagery")
        self._pond_edit_known_btn.toggled.connect(
            self._toggle_known_water_editing)
        pond_vl.addWidget(self._pond_edit_known_btn)

        pond_vl.addWidget(_pond_section_header(
            "Place water-body markers", "_pond_placed_status_lbl"))
        pond_vl.addWidget(_grid_widget([
            self._pond_place_btn, self._pond_auto_place_btn]))
        pond_vl.addWidget(_grid_widget([
            self._pond_labels_btn, self._pond_lines_btn]))
        pond_vl.addWidget(self._pond_manage_btn)
        pond_vl.addWidget(self._pond_map_btn)

        self._refresh_pond_controls()

        constraints_vl.addWidget(_collapsible_section("Standing Water Bodies", "#1F4A37", pond_content, collapsed=True,
                                          icon="droplet"))

        # ── Tree Preservation Orders (collapsible) ──────────────────────
        self._tpo_btn = _tool_btn("Show TPO locations", "tree", color="#175E3B",
                                  checkable=True,
                                  tooltip="Load Tree Preservation Orders within the search radius — individual trees "
                                          "(points) and woodland/area/group zones (polygons)")
        self._tpo_btn.clicked.connect(self._on_tpo_toggle)

        self._tpo_report_btn = _tool_btn(
            "TPO report", "file-text", color="#175E3B",
            tooltip="Show full table of Tree Preservation Orders within the search radius",
        )
        self._tpo_report_btn.clicked.connect(self._run_tpo_report)

        self._tpo_lines_btn = _tool_btn("Distance lines: off", "line", color="#175E3B",
                                         checkable=True,
                                         tooltip="Show or hide distance lines from redline to each TPO point")
        self._tpo_lines_btn.setChecked(False)
        self._tpo_lines_btn.setEnabled(False)
        self._tpo_lines_btn.clicked.connect(self._on_tpo_lines_toggle)

        self._tpo_label_mode = "reference"
        self._tpo_label_btn  = _tool_btn("Label: reference", "tag", color="#175E3B",
                                         tooltip="Cycle TPO point label: Reference → Name → Address → Off")
        self._tpo_label_btn.setEnabled(False)
        self._tpo_label_btn.clicked.connect(self._on_tpo_label_cycle)

        # Conservation areas — related tree protection (trees within need 6
        # weeks' notice before works). Separate designation, own polygon layer.
        self._ca_btn = _tool_btn("Conservation areas", "building-bank", color="#7A5200",
                                 checkable=True,
                                 tooltip="Load conservation areas within the search radius — trees within are protected "
                                         "(6 weeks' written notice to the LPA required before works)")
        self._ca_btn.clicked.connect(self._on_conservation_toggle)

        self._ca_report_btn = _tool_btn("CA report", "file-text", color="#7A5200",
                                        tooltip="Show full table of conservation areas within the search radius")
        self._ca_report_btn.clicked.connect(self._run_ca_report)

        self._tpo_map_btn = _tool_btn("TPO map", "map", color="#3A463F",
                                      trailing="pencil", trailing_click=True,
                                      tooltip="Create a print layout showing TPO trees and zones within the search radius")
        self._tpo_map_btn.clicked.connect(lambda: self._create_desk_study_map("TPO Map"))

        self._ca_map_btn = _tool_btn("CA map", "map", color="#3A463F",
                                     trailing="pencil", trailing_click=True,
                                     tooltip="Create a print layout showing conservation areas within the search radius")
        self._ca_map_btn.clicked.connect(lambda: self._create_desk_study_map("Conservation Areas Map"))

        tpo_content = QWidget()
        tpo_vl = QVBoxLayout(tpo_content)
        tpo_vl.setContentsMargins(4, 4, 4, 4)
        tpo_vl.setSpacing(4)
        tpo_vl.addWidget(_grid_widget([self._tpo_btn, self._tpo_report_btn]))
        tpo_vl.addWidget(_grid_widget([self._tpo_lines_btn, self._tpo_label_btn]))
        tpo_vl.addWidget(_grid_widget([self._ca_btn, self._ca_report_btn]))
        # Map buttons full width, like every other section's.
        tpo_vl.addWidget(self._tpo_map_btn)
        tpo_vl.addWidget(self._ca_map_btn)

        constraints_vl.addWidget(_collapsible_section("TPOs & Conservation Areas", "#175E3B", tpo_content, collapsed=True,
                                          icon="tree"))

        # ── EPSL (collapsible) ─────────────────────────────────────────
        self._epsl_btn = _tool_btn("Show EPSL locations", "license", color="#8A4B12",
                                   checkable=True,
                                   tooltip="Load granted EPSL locations within the search radius as map points")
        self._epsl_btn.clicked.connect(self._on_epsl_toggle)

        self._epsl_report_btn = _tool_btn(
            "EPSL report", "file-text", color="#8A4B12",
            tooltip="Show full table of granted EPSL licences within the search radius",
        )
        self._epsl_report_btn.clicked.connect(self._run_epsl_report)

        self._epsl_lines_btn = _tool_btn("Distance lines: off", "line", color="#8A4B12",
                                         checkable=True,
                                         tooltip="Show or hide distance lines from redline to each EPSL point")
        self._epsl_lines_btn.setChecked(False)
        self._epsl_lines_btn.setEnabled(False)
        self._epsl_lines_btn.clicked.connect(self._on_epsl_lines_toggle)

        self._epsl_label_mode = "species"
        self._epsl_label_btn  = _tool_btn("Labels: Species", "tag", color="#8A4B12",
                                          tooltip="Choose which labels appear beside EPSL locations",
                                          trailing="chevron-down")
        self._epsl_label_btn.setEnabled(False)
        label_menu = QMenu(self._epsl_label_btn)
        label_menu.setMinimumWidth(220)
        label_menu.setStyleSheet(
            "QMenu { background:#FFFFFF; color:#1A2420;"
            " border:1px solid #B7BFB9; border-radius:8px; padding:5px;"
            " font-size:11px; }"
            "QMenu::item { color:#1A2420; background:transparent;"
            " padding:7px 24px 7px 10px; border-radius:5px; }"
            "QMenu::item:selected { background:#FBF4E4; color:#8A4B12; }"
            "QMenu::item:checked { background:#F4E7C6; color:#6F3B0D;"
            " font-weight:600; }")
        label_group = QActionGroup(label_menu)
        label_group.setExclusive(True)
        self._epsl_label_actions = {}
        for text, mode in (("Species", "species"),
                           ("Licence reference", "licence"),
                           ("Species and reference", "both"),
                           ("None", "off")):
            action = label_menu.addAction(text)
            action.setCheckable(True)
            action.setData(mode)
            action.setChecked(mode == "species")
            action.triggered.connect(self._on_epsl_label_selected)
            label_group.addAction(action)
            self._epsl_label_actions[mode] = action
        self._epsl_label_btn.setMenu(label_menu)

        self._epsl_filter_mode = "all"
        self._epsl_filter_btn  = _tool_btn("Species: All species", "filter", color="#8A4B12",
                                           checkable=True,
                                           tooltip="Toggle between all species and bats only")
        self._epsl_filter_btn.setEnabled(False)
        self._epsl_filter_btn.clicked.connect(self._on_epsl_filter_toggle)

        # Single Style cycle replacing the old 'Stylised icons' / 'Recent-
        # Historic' checkboxes — the three are mutually exclusive marker styles,
        # so one cycle button reads cleaner and can't conflict with each other.
        # It composes with the Filter rather than resetting it.
        self._epsl_style_mode = "icons"
        self._epsl_style_btn  = _tool_btn("Display: Species icons", "palette", color="#8A4B12",
                                          tooltip="Choose how EPSL locations are displayed",
                                          trailing="chevron-down")
        self._epsl_style_btn.setEnabled(False)
        display_menu = QMenu(self._epsl_style_btn)
        display_menu.setMinimumWidth(200)
        display_menu.setStyleSheet(
            "QMenu { background:#FFFFFF; color:#1A2420;"
            " border:1px solid #B7BFB9; border-radius:8px; padding:5px;"
            " font-size:11px; }"
            "QMenu::item { color:#1A2420; background:transparent;"
            " padding:7px 24px 7px 10px; border-radius:5px; }"
            "QMenu::item:selected { background:#FBF4E4; color:#8A4B12; }"
            "QMenu::item:checked { background:#F4E7C6; color:#6F3B0D;"
            " font-weight:600; }")
        display_group = QActionGroup(display_menu)
        display_group.setExclusive(True)
        self._epsl_display_actions = {}
        for text, mode in (("Standard", "star"),
                           ("Species icons", "icons"),
                           ("Recent/historic", "recent")):
            action = display_menu.addAction(text)
            action.setCheckable(True)
            action.setData(mode)
            action.setChecked(mode == "icons")
            action.triggered.connect(self._on_epsl_display_selected)
            display_group.addAction(action)
            self._epsl_display_actions[mode] = action
        self._epsl_style_btn.setMenu(display_menu)

        epsl_content = QWidget()
        epsl_vl = QVBoxLayout(epsl_content)
        epsl_vl.setContentsMargins(4, 4, 4, 4)
        epsl_vl.setSpacing(4)
        epsl_vl.addWidget(_grid_widget([self._epsl_btn, self._epsl_report_btn]))
        epsl_vl.addWidget(_grid_widget([self._epsl_lines_btn, self._epsl_label_btn]))
        epsl_vl.addWidget(_grid_widget([self._epsl_filter_btn, self._epsl_style_btn]))
        self._epsl_map_btn = _tool_btn("EPSL map", "map", color="#3A463F",
                                       trailing="pencil", trailing_click=True,
                                       tooltip="Create a print layout showing EPSL licence locations within the search radius")
        self._epsl_map_btn.clicked.connect(lambda: self._create_desk_study_map("EPSL Map"))
        epsl_vl.addWidget(self._epsl_map_btn)

        constraints_vl.addWidget(_collapsible_section("EPSL Locations", "#8A4B12", epsl_content, collapsed=True,
                                          icon="license"))

        # ── GCN District Licensing (collapsible) ───────────────────────
        self._gcn_btn = _tool_btn("Show GCN risk zones", "eye", color="#175E3B",
                                  checkable=True,
                                  tooltip="Toggle Great Crested Newt DLL Risk Zones")
        self._gcn_btn.clicked.connect(self._on_gcn_toggle)

        self._gcn_coverage_btn = _tool_btn("DLL scheme coverage", "map", color="#175E3B",
                                           checkable=True,
                                           tooltip="Show/hide the true outlines of the areas covered by a GCN DLL scheme")
        self._gcn_coverage_btn.clicked.connect(self._on_gcn_coverage_toggle)

        gcn_opacity_row = QWidget()
        gcn_opacity_hl  = QHBoxLayout(gcn_opacity_row)
        gcn_opacity_hl.setContentsMargins(2, 2, 2, 2)
        gcn_opacity_hl.setSpacing(6)
        gcn_opacity_lbl = QLabel("Opacity")
        gcn_opacity_lbl.setStyleSheet("font-size:10px; color:#54635A;")
        gcn_opacity_hl.addWidget(gcn_opacity_lbl)
        self._gcn_opacity_slider = QSlider(Qt.Orientation.Horizontal)
        self._gcn_opacity_slider.setMinimum(10)
        self._gcn_opacity_slider.setMaximum(100)
        self._gcn_opacity_slider.setValue(65)
        self._gcn_opacity_slider.setTickInterval(10)
        self._wire_hint(self._gcn_opacity_slider, "Adjust GCN Risk Zone layer opacity")
        self._gcn_opacity_slider.setStyleSheet("""
            QSlider::groove:horizontal { height:4px; background:#D8ECDE; border-radius:2px; }
            QSlider::handle:horizontal { background:#1F7A4D; border:none; width:14px; height:14px; margin:-5px 0; border-radius:7px; }
            QSlider::sub-page:horizontal { background:#1F7A4D; border-radius:2px; }
        """)
        self._gcn_opacity_slider.valueChanged.connect(self._on_gcn_opacity_changed)
        gcn_opacity_hl.addWidget(self._gcn_opacity_slider)
        self._gcn_opacity_val_lbl = QLabel("65%")
        self._gcn_opacity_val_lbl.setStyleSheet("font-size:10px; color:#54635A; min-width:28px;")
        gcn_opacity_hl.addWidget(self._gcn_opacity_val_lbl)

        gcn_content = QWidget()
        gcn_vl = QVBoxLayout(gcn_content)
        gcn_vl.setContentsMargins(4, 4, 4, 4)
        gcn_vl.setSpacing(4)
        gcn_vl.addWidget(_grid_widget([self._gcn_btn, self._gcn_coverage_btn]))
        gcn_vl.addWidget(gcn_opacity_row)

        # National eDNA pond surveys (within the search radius) + NatureSpace deep-link.
        self._edna_btn = _tool_btn(
            "eDNA ponds", "droplet", color="#175E3B", checkable=True,
            tooltip="Toggle Natural England GCN eDNA pond survey results within "
                    "the search radius of the redline (present / absent / inconclusive). "
                    "Covers all of England — including outside the DLL risk-zone "
                    "areas — and rides along on the GCN risk zones map.")
        self._edna_btn.clicked.connect(self._on_edna_toggle)
        self._naturespace_btn = _tool_btn(
            "NatureSpace map", "external-link", color="#1F4A37",
            tooltip="Open NatureSpace's public Impact Risk Zone map — the DLL "
                    "areas run by NatureSpace rather than Natural England. "
                    "Their terms apply.")
        self._naturespace_btn.clicked.connect(self._on_naturespace_link)
        gcn_vl.addWidget(_grid_widget([self._edna_btn, self._naturespace_btn]))

        self._gcn_hsi_btn = _tool_btn(
            "GCN HSI calculator", "calculator", color="#175E3B",
            tooltip="Open the Great Crested Newt Habitat Suitability Index "
                    "calculator (ARG UK Advice Note 5)")
        self._gcn_hsi_btn.clicked.connect(self._on_gcn_hsi_calculator)
        gcn_vl.addWidget(self._gcn_hsi_btn)
        self._gcn_hsi_map_btn = _tool_btn(
            "Map HSI scores", "map", color="#3A463F",
            tooltip="Create a pond assessment map with grouped HSI scores and survey statuses")
        self._gcn_hsi_map_btn.clicked.connect(self._create_gcn_hsi_map)
        gcn_vl.addWidget(self._gcn_hsi_map_btn)
        self._gcn_map_btn = _tool_btn("GCN risk zones map", "map", color="#3A463F",
                                      trailing="pencil", trailing_click=True,
                                      tooltip="Create a print layout showing the loaded GCN Risk Zone layer")
        self._gcn_map_btn.clicked.connect(lambda: self._create_desk_study_map("GCN Risk Zones Map"))
        gcn_vl.addWidget(self._gcn_map_btn)

        constraints_vl.addWidget(_collapsible_section("Great Crested Newts", "#175E3B", gcn_content, collapsed=True,
                                          icon="newt"))

        # Keep the visible search areas in step with the layers that actually
        # use them. A zero-delay sync runs after each layer's own toggle handler,
        # so failed loads that untick themselves cannot leave a misleading ring.
        shared_radius_buttons = (
            list(self._ds_btns.values()) + list(self._phi_btns.values()) +
            [self._aw_btn, self._epsl_btn, self._ca_btn, self._edna_btn]
        )
        for _radius_btn in shared_radius_buttons:
            _radius_btn.toggled.connect(
                lambda _checked=False: (
                    None if getattr(self, "_syncing_ds_buttons", False)
                    else QTimer.singleShot(
                        0, lambda: self._sync_auto_search_areas(
                            sync_local=False))))
        for _radius_btn in (self._tpo_btn, self._os_topo_btn,
                            self._pond_place_btn):
            _radius_btn.toggled.connect(
                lambda _checked=False: (
                    None if getattr(self, "_syncing_ds_buttons", False)
                    else QTimer.singleShot(
                        0, self._sync_auto_search_areas)))

        # ── Survey Weather (collapsible) ───────────────────────────────────
        wx_content = QWidget()
        wx_vl = QVBoxLayout(wx_content)
        wx_vl.setContentsMargins(4, 4, 4, 4)
        wx_vl.setSpacing(4)
        wx_hint = QLabel("Look up the temperature, wind and rainfall at the time of "
                         "survey, ready for the weather-conditions table.")
        wx_hint.setWordWrap(True)
        wx_hint.setStyleSheet("color:#54635A; font-size:10px; background:transparent;")
        wx_vl.addWidget(wx_hint)
        self._weather_btn = _tool_btn(
            "Survey weather lookup", "cloud-sun", color="#2F5D8A",
            tooltip="Set a date, time and location and fetch the hourly weather "
                    "(temperature °C, wind mph, rainfall mm) for the survey hour — "
                    "past dates or up to a week ahead — fills the report "
                    "weather-conditions table")
        self._weather_btn.clicked.connect(self._on_weather_checker)
        wx_vl.addWidget(self._weather_btn)
        constraints_vl.addWidget(_collapsible_section("Survey Weather", "#2F5D8A", wx_content, collapsed=True,
                                          icon="cloud-sun"))

        # ── External Resources ─────────────────────────────────────────────
        vl.addWidget(_section_hdr("External Resources"))

        def _link_btn(icon_name, label, url, tooltip):
            btn = _tool_btn(label, icon_name, tooltip=tooltip, trailing="external-link")
            btn.clicked.connect(lambda: __import__(
                'qgis.PyQt.QtGui', fromlist=['QDesktopServices']
            ).QDesktopServices.openUrl(
                __import__('qgis.PyQt.QtCore', fromlist=['QUrl']).QUrl(url)
            ))
            return btn

        lpa_map_btn = _tool_btn("LPA website", "building",
                                tooltip="Open the local planning authority's website — resolves "
                                        "the council from the redline centroid and opens that "
                                        "council's homepage",
                                trailing="external-link")
        lpa_map_btn.clicked.connect(self._open_lpa_map)

        # MAGIC in England; Scotland's environment map for a Scottish site
        # (relabelled by _sync_region_ui when the redline moves).
        magic_ext_btn = _tool_btn(_region.magic_label(), "map-search",
                                  tooltip=_region.magic_tooltip(),
                                  trailing="external-link")
        magic_ext_btn.clicked.connect(self._open_magic_map)
        self._magic_ext_btn = magic_ext_btn

        nbn_ext_btn = _tool_btn("NBN Atlas", "leaf",
                                tooltip="Open the NBN Atlas — with a redline present, jumps "
                                        "straight to a species-records search using the saved "
                                        "Desk Study search radius",
                                trailing="external-link")
        nbn_ext_btn.clicked.connect(self._open_nbn_atlas)

        nls_ext_btn = _tool_btn("Historic maps", "map-2",
                                tooltip="Open the NLS georeferenced historic OS maps side-by-"
                                        "side with modern satellite, centred on the site — "
                                        "good for land-use / habitat-history change",
                                trailing="external-link")
        nls_ext_btn.clicked.connect(self._open_nls_maps)

        earth_ext_btn = _tool_btn("Google Earth", "world",
                                  tooltip="Open Google Earth centred on the site with historical "
                                          "imagery on — drag the time slider to check for recent "
                                          "change, e.g. a felled tree",
                                  trailing="external-link")
        earth_ext_btn.clicked.connect(self._open_google_earth)

        earth_kml_btn = _tool_btn("Earth Pro – Historical Imagery", "camera",
                                  tooltip="Open the redline in Google Earth Pro in a top-down "
                                          "view — select Historical Imagery (the clock icon) "
                                          "to compare available imagery dates",
                                  trailing="external-link")
        earth_kml_btn.clicked.connect(self._export_google_earth_kml)

        spotlight_3d_btn = _tool_btn("3D Spotlight (browser)", "world",
                                     tooltip="Export an interactive 3D spotlight viewer (HTML) — "
                                             "Google's 3D mesh clipped to the redline: the site in "
                                             "full colour, everything outside fogged in true 3D "
                                             "(trees and buildings included). Orbit, tilt and save "
                                             "a screenshot in your browser. Needs a Google Maps "
                                             "API key with the Map Tiles API enabled",
                                     trailing="download")
        spotlight_3d_btn.clicked.connect(self._export_3d_spotlight_html)

        sv_ext_btn = _tool_btn("Street View", "map-pin",
                               tooltip="Open Google Maps zoomed on the site — grab the pegman "
                                       "(bottom-right) to light up the blue Street View streets, "
                                       "then click any of them for a ground-level check of "
                                       "boundaries, buildings and access",
                               trailing="external-link")
        sv_ext_btn.clicked.connect(self._open_street_view)

        external_group, external_vl = _section_group()
        external_vl.addWidget(_grid_widget([
            lpa_map_btn,
            magic_ext_btn,
            nbn_ext_btn,
            nls_ext_btn,
            earth_ext_btn,
            earth_kml_btn,
            spotlight_3d_btn,
            sv_ext_btn,
        ]))
        vl.addWidget(external_group)


        vl.addStretch()
        inner.setLayout(vl)
        scroll.setWidget(inner)
        # Now that every desk-study map button exists, wire the right-click
        # "rename map (this session)" menu onto each. (This tab is built lazily,
        # so this can't run from __init__ — the buttons don't exist yet there.)
        self._wire_map_rename_menus()
        return outer

    # Desk-study map buttons that support the session-only right-click rename:
    # (panel attribute, map-type key into map_exporter.DESK_TITLES).
    _DESK_MAP_BTNS = (
        ("_ds_map_btn",   "Designated Sites Map"),
        ("_phi_map_btn",  "Priority Habitats Map"),
        ("_irz_map_btn",  "SSSI IRZ Map"),
        ("_lws_map_btn",  "LWS Map"),
        ("_pond_map_btn", "Pond Map"),
        ("_tpo_map_btn",  "TPO Map"),
        ("_ca_map_btn",   "Conservation Areas Map"),
        ("_epsl_map_btn", "EPSL Map"),
        ("_gcn_map_btn",  "GCN Risk Zones Map"),
    )

    def _wire_map_rename_menus(self):
        """Wire each desk-study map button's little pencil button to the rename
        popup, and re-apply any live override to the freshly (re)built buttons."""
        for attr, mtype in self._DESK_MAP_BTNS:
            btn = getattr(self, attr, None)
            if btn is None:
                continue
            btn._biome_default_label = btn._label_lbl.text()
            tb = getattr(btn, "_trail_btn", None)
            if tb is not None:
                self._wire_hint(tb, "Rename map")
                tb.clicked.connect(
                    lambda _c=False, b=btn, mt=mtype:
                        self._map_rename_menu(mt, b, b.rect().bottomLeft()))
            self._refresh_map_btn_label(mtype)


    # ------------------------------------------------------------------ #
    #  SCREENS                                                             #
    # ------------------------------------------------------------------ #

    def _render_ds_area(self):
        """Paint the area tile's caption + value for the currently selected unit.
        A ↻ glyph on the caption hints that right-click rotates the unit."""
        if not hasattr(self, '_ds_area_caption'):
            return
        idx = self._ds_area_unit_idx
        self._ds_area_caption.setText(self._DS_AREA_UNITS[idx] + "  ↻")
        self._ds_area_value.setText(self._ds_area_vals[idx])
        if hasattr(self, '_ds_blue_area_caption'):
            self._ds_blue_area_caption.setText(self._DS_AREA_UNITS[idx] + "  ↻")
            self._ds_blue_area_value.setText(self._ds_blue_area_vals[idx])

    def _on_field_survey_pushed(self, payload):
        """A survey arrived from the phone over the live tunnel.

        Runs on the GUI thread (queued from the HTTP worker via
        _FieldSyncBridge), so it can safely write layers and show a dialog.
        """
        from qgis.PyQt.QtWidgets import QMessageBox
        from . import field_mapper_import as fmi
        try:
            summary = fmi.import_survey(
                self, payload, photo_map=payload.get('photos'))
        except Exception as exc:
            import traceback
            traceback.print_exc()
            QMessageBox.warning(
                self.iface.mainWindow(), "Biome — Field Survey",
                "The phone sent a survey but it could not be imported:\n\n%s\n\n"
                "Ask the surveyor to use Download ZIP instead — nothing has "
                "been lost from the phone." % exc)
            return
        self._after_field_survey_import(summary)

    def _cycle_ds_area_unit(self):
        """Right-click on the area tile rotates the displayed unit ha → m² → km²."""
        self._ds_area_unit_idx = (self._ds_area_unit_idx + 1) % len(self._DS_AREA_UNITS)
        try:
            QSettings().setValue("biome/desk_study/area_unit_idx", self._ds_area_unit_idx)
        except Exception:
            pass
        self._render_ds_area()

    def _refresh_desk_study_stats(self):
        """Recalculate and display redline area and centroid in the Desk Study panel."""
        from qgis.core import (QgsCoordinateReferenceSystem, QgsCoordinateTransform,
                               QgsDistanceArea, QgsUnitTypes, QgsPointXY, QgsProject)
        from .gridref_utils import osgb_en_to_gridref

        self._ensure_tab_built(1)   # desk tab is semi-lazy
        if not hasattr(self, '_ds_area_value'):
            return

        self._refresh_blue_line_stats()

        redline_layer = self.layer_manager.get_layer("redline")
        has_redline = redline_layer is not None and redline_layer.featureCount() > 0
        if hasattr(self, '_ds_hint_card'):
            self._ds_hint_card.setVisible(
                bool(getattr(self, "_ds_show_first_run_hint", False))
                and not has_redline)
        if redline_layer is None or redline_layer.featureCount() == 0:
            self._ds_area_vals = ["—", "—", "—"]
            self._ds_centroid.setText("—")
            if hasattr(self, "_ds_lpa"):
                self._ds_lpa.setText("—")
            self._render_ds_area()
            return

        src_crs = redline_layer.crs()
        bng = QgsCoordinateReferenceSystem("EPSG:27700")
        need_transform = src_crs.authid() != "EPSG:27700"
        to_bng = QgsCoordinateTransform(src_crs, bng, QgsProject.instance()) if need_transform else None

        # Use ellipsoidal area in the layer's native CRS (most accurate)
        da = QgsDistanceArea()
        da.setSourceCrs(src_crs, QgsProject.instance().transformContext())
        da.setEllipsoid("GRS80")

        total_m2 = 0.0
        cx_sum = 0.0
        cy_sum = 0.0
        feat_count = 0

        for feat in redline_layer.getFeatures():
            geom = feat.geometry()
            if geom.isEmpty():
                continue

            area_m2 = abs(da.convertAreaMeasurement(
                da.measureArea(geom), QgsUnitTypes.AreaUnit.AreaSquareMeters
            ))
            total_m2 += area_m2

            # Transform the whole geometry to BNG first, then compute centroid.
            # Computing centroid in EPSG:3857 then projecting the point gives
            # a distorted result at UK latitudes (~60% N-S Mercator stretch).
            if need_transform:
                from qgis.core import QgsGeometry as _QG
                bng_geom = _QG(geom)
                bng_geom.transform(to_bng)
                bng_centroid = bng_geom.centroid().asPoint()
            else:
                bng_centroid = geom.centroid().asPoint()

            cx_sum += bng_centroid.x()
            cy_sum += bng_centroid.y()
            feat_count += 1

        if feat_count == 0 or total_m2 == 0:
            self._ds_area_vals = ["No valid geometry"] * 3
            self._ds_centroid.setText("No valid geometry")
            if hasattr(self, "_ds_lpa"):
                self._ds_lpa.setText("—")
            self._render_ds_area()
            return

        cx = cx_sum / feat_count  # BNG easting
        cy = cy_sum / feat_count  # BNG northing

        grid_ref = osgb_en_to_gridref(cx, cy, digits=10)

        self._ds_area_vals = [
            f"{total_m2 / 10000:,.4f}",      # ha
            f"{total_m2:,.1f}",              # m²
            f"{total_m2 / 1_000_000:,.6f}",  # km²
        ]
        self._ds_centroid.setText(grid_ref)
        self._render_ds_area()
        self._update_ds_lpa(cx, cy)

    def _refresh_blue_line_stats(self):
        """Show area and BNG centroid for the optional Blue Line Boundary."""
        from qgis.core import (QgsCoordinateReferenceSystem, QgsCoordinateTransform,
                               QgsDistanceArea, QgsUnitTypes, QgsProject, QgsGeometry)
        from .gridref_utils import osgb_en_to_gridref

        card = getattr(self, "_ds_blue_stats_card", None)
        if card is None:
            return
        layer = self.layer_manager.get_layer("site_boundary")
        has_blue = layer is not None and layer.featureCount() > 0
        card.setVisible(has_blue)
        if not has_blue:
            self._ds_blue_area_vals = ["—", "—", "—"]
            self._ds_blue_centroid.setText("—")
            self._render_ds_area()
            return

        src_crs = layer.crs()
        bng = QgsCoordinateReferenceSystem("EPSG:27700")
        need_transform = src_crs.authid() != "EPSG:27700"
        to_bng = (QgsCoordinateTransform(src_crs, bng, QgsProject.instance())
                  if need_transform else None)
        da = QgsDistanceArea()
        da.setSourceCrs(src_crs, QgsProject.instance().transformContext())
        da.setEllipsoid("GRS80")

        total_m2 = 0.0
        cx_sum = cy_sum = 0.0
        feat_count = 0
        for feat in layer.getFeatures():
            geom = feat.geometry()
            if not geom or geom.isNull() or geom.isEmpty():
                continue
            total_m2 += abs(da.convertAreaMeasurement(
                da.measureArea(geom), QgsUnitTypes.AreaUnit.AreaSquareMeters))
            if need_transform:
                bng_geom = QgsGeometry(geom)
                bng_geom.transform(to_bng)
                centroid = bng_geom.centroid().asPoint()
            else:
                centroid = geom.centroid().asPoint()
            cx_sum += centroid.x()
            cy_sum += centroid.y()
            feat_count += 1

        if feat_count == 0 or total_m2 == 0:
            self._ds_blue_area_vals = ["No valid geometry"] * 3
            self._ds_blue_centroid.setText("No valid geometry")
        else:
            self._ds_blue_area_vals = [
                f"{total_m2 / 10000:,.4f}",
                f"{total_m2:,.1f}",
                f"{total_m2 / 1_000_000:,.6f}",
            ]
            self._ds_blue_centroid.setText(osgb_en_to_gridref(
                cx_sum / feat_count, cy_sum / feat_count, digits=10))
        self._render_ds_area()

    def _update_ds_lpa(self, cx, cy):
        """Resolve and show the local planning authority for the redline centroid
        in the Redline boundary box. Cached per ~10 m bucket and resolved just
        after the refresh so the area/centroid render instantly; a stale resolve
        (redline moved again) is discarded via the pending-key guard."""
        if not hasattr(self, "_ds_lpa"):
            return
        key = (round(cx, -1), round(cy, -1))
        cache = getattr(self, "_ds_lpa_cache", None)
        if cache is None:
            cache = self._ds_lpa_cache = {}
        if key in cache:
            self._ds_lpa.setText(cache[key] or "Not found")
            self._store_ds_lpa(cache[key])
            self._ds_lpa_pending = None
            return
        self._ds_lpa.setText("Resolving…")
        self._ds_lpa_pending = key

        # Transform on the GUI thread because QgsProject and coordinate
        # transforms are QGIS objects. The slow HTTP-only lookup then runs in a
        # daemon thread, so a network timeout cannot freeze red-line placement
        # or delay the GCN DLL coverage notification.
        try:
            from qgis.core import (
                QgsCoordinateReferenceSystem, QgsCoordinateTransform,
                QgsPointXY, QgsProject,
            )
            bng = QgsCoordinateReferenceSystem("EPSG:27700")
            wgs = QgsCoordinateReferenceSystem("EPSG:4326")
            point = QgsCoordinateTransform(
                bng, wgs, QgsProject.instance()).transform(QgsPointXY(cx, cy))
            lon, lat = point.x(), point.y()
        except Exception:
            lon = lat = None

        bridge = getattr(self, "_ds_lpa_bridge", None)
        if bridge is None:
            bridge = _LpaLookupBridge(self)
            bridge.resolved.connect(self._apply_ds_lpa_result)
            self._ds_lpa_bridge = bridge

        inflight = getattr(self, "_ds_lpa_inflight", None)
        if inflight is None:
            inflight = self._ds_lpa_inflight = set()
        if key not in inflight:
            inflight.add(key)

            def _resolve_http():
                name = None
                if lon is not None and lat is not None:
                    try:
                        from .designated_sites import get_lpa_for_wgs84
                        name, _ref = get_lpa_for_wgs84(lon, lat)
                    except Exception:
                        name = None
                try:
                    bridge.resolved.emit((key, name))
                except RuntimeError:
                    pass   # panel closed while the request was in flight

            import threading
            threading.Thread(
                target=_resolve_http, name="BiomeLpaLookup", daemon=True,
            ).start()

        # Store BNG centroid for LPA finder
        self._centroid_bng = (cx, cy)
        # Tell the region switch where the site is (England vs Scotland data
        # services) and relabel the desk study to match.
        try:
            _region.set_site(cx, cy)
            self._sync_region_ui()
        except Exception:
            pass

    def _apply_ds_lpa_result(self, result):
        """Apply a background LPA result on the Qt GUI thread."""
        try:
            key, name = result
        except (TypeError, ValueError):
            return
        inflight = getattr(self, "_ds_lpa_inflight", None)
        if inflight is not None:
            inflight.discard(key)
        full = lpa_full_name(name) if name else ""
        cache = getattr(self, "_ds_lpa_cache", None)
        if cache is None:
            cache = self._ds_lpa_cache = {}
        cache[key] = full or ""
        if getattr(self, "_ds_lpa_pending", None) != key:
            return   # the redline moved; a newer lookup owns the visible tile
        self._store_ds_lpa(full)
        self._ds_lpa.setText(full or "Not found")
        self._ds_lpa_pending = None

    def _store_ds_lpa(self, full):
        """Persist the resolved desk-study LPA per-project so the filled-metric
        export pre-fills its Planning authority field from it (xlsm_exporter
        reads the same "desk/lpa" project entry back)."""
        if not full:
            return
        try:
            from .map_exporter import set_project_value
            set_project_value("desk/lpa", full)
        except Exception:
            pass

    def _open_lpa_finder(self):
        """Open the Planning Portal LPA Finder, pre-resolving the site's nearest
        postcode + local-authority district from the redline centroid and copying
        the postcode to the clipboard to paste into the finder.

        The Planning Portal finder has no documented URL parameter to pre-fill the
        postcode, and its form wants a postcode rather than an easting/northing —
        so we copy the postcode (falling back to E/N only if no postcode is found
        nearby, e.g. a rural site). The browser is opened first so the UI never
        appears to hang on the postcodes.io lookup."""
        from qgis.PyQt.QtCore import QUrl
        from qgis.PyQt.QtGui import QDesktopServices
        from qgis.PyQt.QtWidgets import QApplication

        QDesktopServices.openUrl(
            QUrl("https://www.planningportal.co.uk/find-your-local-planning-authority")
        )

        cen = self._redline_centroid_bng() or getattr(self, "_centroid_bng", None)
        if not cen:
            return
        e, n = cen

        postcode = district = None
        try:
            postcode, district = self._nearest_postcode_for_bng(e, n)
        except Exception:
            postcode = district = None

        head = f"LPA: {district} — " if district else ""
        if postcode:
            QApplication.clipboard().setText(postcode)
            self.iface.messageBar().pushMessage(
                "LPA Finder",
                f"{head}nearest postcode {postcode} copied to clipboard — "
                "paste it into the postcode box",
                level=Qgis.MessageLevel.Info, duration=8
            )
        else:
            QApplication.clipboard().setText(f"{int(e)}, {int(n)}")
            self.iface.messageBar().pushMessage(
                "LPA Finder",
                f"{head}no nearby postcode found; easting/northing "
                f"({int(e)}, {int(n)}) copied instead",
                level=Qgis.MessageLevel.Info, duration=8
            )

    def _open_lpa_map(self):
        """Open the local planning authority's WEBSITE homepage for the site.

        Resolves the council from the redline centroid (get_lpa_for_bng), then
        opens that council's homepage — derived from the registrable domain of the
        council's planning page we have on file (e.g. a Croydon page →
        www.croydon.gov.uk). Falls back to a web search for the council where we
        have no page on file. (The TPO 'no results' flow is unchanged — it still
        deep-links into the council's interactive map; this button is just the
        council website.)"""
        from qgis.PyQt.QtCore import QUrl
        from qgis.PyQt.QtGui import QDesktopServices
        from urllib.parse import urlparse, quote_plus

        cen = self._redline_centroid_bng() or getattr(self, "_centroid_bng", None)
        lpa_name = None
        if cen:
            try:
                lpa_name, _ref = get_lpa_for_bng(cen[0], cen[1])
            except Exception:
                lpa_name = None

        homepage = None
        if lpa_name:
            # 1. Audited/derived homepage table (covers ~335 LPAs incl. tricky
            #    slugs like Swindon, BCP, Dorset, abolished districts).
            try:
                homepage = lpa_homepage(lpa_name)
            except Exception:
                homepage = None
            # 2. Fallback: derive from the council's planning page domain on file.
            if not homepage:
                try:
                    _disp, portal_url, is_direct = get_tpo_council_portal(lpa_name)
                    if is_direct and portal_url:
                        host = urlparse(portal_url).netloc.lower()
                        parts = [p for p in host.split(".") if p]
                        if host.endswith(".gov.uk") and len(parts) >= 3:
                            reg = ".".join(parts[-3:])      # e.g. croydon.gov.uk
                        elif len(parts) >= 2:
                            reg = ".".join(parts[-2:])
                        else:
                            reg = host
                        if reg:
                            homepage = "https://www.%s/" % reg
                except Exception:
                    homepage = None

        if homepage:
            QDesktopServices.openUrl(QUrl(homepage))
            note = "Opening %s's website (%s)." % (lpa_name, homepage)
        elif lpa_name:
            QDesktopServices.openUrl(QUrl(
                "https://www.google.com/search?q=%s" % quote_plus(lpa_name + " council")))
            note = "No website on file for %s — opening a web search for it." % lpa_name
        else:
            QDesktopServices.openUrl(QUrl(
                "https://www.google.com/search?q=%s" % quote_plus("find your local council gov.uk")))
            note = "Couldn't determine the council from the redline — opening a council search."
        try:
            self.iface.messageBar().pushMessage("LPA", note, level=Qgis.MessageLevel.Info, duration=8)
        except Exception:
            pass

    def _open_magic_map(self):
        """Open the Defra MAGIC map framed on the redline boundary, or at its
        default view if there's no redline.

        MAGIC's documented 'box' parameter is left:bottom:right:top (min/max
        easting:northing, colon-separated) and opens in OSGB36 by default — the
        same CRS as the redline (EPSG:27700), so no srs override is needed. We use
        a square box centred on the site (min ~250 m across) so small sites aren't
        over-zoomed."""
        from qgis.PyQt.QtCore import QUrl
        from qgis.PyQt.QtGui import QDesktopServices

        box = None
        redline_layer = self.layer_manager.get_layer("redline")
        if redline_layer is not None and redline_layer.featureCount() > 0:
            try:
                from qgis.core import (QgsCoordinateReferenceSystem,
                                       QgsCoordinateTransform, QgsProject,
                                       QgsRectangle)
                ext = redline_layer.extent()
                bng = QgsCoordinateReferenceSystem("EPSG:27700")
                if redline_layer.crs() != bng:
                    ext = QgsCoordinateTransform(
                        redline_layer.crs(), bng, QgsProject.instance()
                    ).transformBoundingBox(ext)
                ext.grow(60)   # margin so the site isn't flush to the frame edge
                cx, cy = ext.center().x(), ext.center().y()
                half   = max(ext.width(), ext.height(), 250) / 2.0
                box = (cx - half, cy - half, cx + half, cy + half)
            except Exception:
                box = None

        # England → MAGIC (?&box=l:b:r:t); Scotland → Scotland's environment
        # map (?extent=xmin,ymin,xmax,ymax). Both EPSG:27700; see region.py.
        label = _region.magic_label()
        url = _region.magic_url(box)
        if box:
            self.iface.messageBar().pushMessage(
                label, f"Opening {label} framed on the redline boundary",
                level=Qgis.MessageLevel.Info, duration=5
            )

        QDesktopServices.openUrl(QUrl(url))

    def _redline_bbox_bng(self, min_span=250.0, margin=60.0):
        """(bbox, centroid) of the current redline in EPSG:27700 for map deep-links.

        bbox = (min_e, min_n, max_e, max_n): a square box centred on the site,
        grown by `margin` m and at least `min_span` m across so small sites
        aren't over-zoomed. centroid = (easting, northing). Same framing as
        _open_magic_map, returned as numbers for the council/LNRS deep-link
        builders. Returns (None, None) if there's no redline geometry."""
        from qgis.core import (QgsCoordinateReferenceSystem,
                               QgsCoordinateTransform, QgsProject, QgsRectangle)
        bng = QgsCoordinateReferenceSystem("EPSG:27700")
        ext = None
        redline_layer = self.layer_manager.get_layer("redline")
        if redline_layer is not None and redline_layer.featureCount() > 0:
            try:
                ext = redline_layer.extent()
                if redline_layer.crs() != bng:
                    ext = QgsCoordinateTransform(
                        redline_layer.crs(), bng, QgsProject.instance()
                    ).transformBoundingBox(ext)
            except Exception:
                ext = None
        if ext is not None and not ext.isEmpty():
            ext = QgsRectangle(ext)
            ext.grow(margin)
            cx, cy = ext.center().x(), ext.center().y()
            half   = max(ext.width(), ext.height(), min_span) / 2.0
        else:
            cen = self._redline_centroid_bng() or getattr(self, "_centroid_bng", None)
            if not cen:
                return None, None
            cx, cy = cen
            half   = min_span / 2.0
        bbox = (cx - half, cy - half, cx + half, cy + half)
        return bbox, (cx, cy)

    def _redline_centroid_bng(self):
        """(easting, northing) centroid of the current redline in EPSG:27700,
        computed live from the geometry, or None if there's no redline. Computing
        live keeps it correct even if the redline moved since the Desk Study stats
        last refreshed self._centroid_bng."""
        redline_layer = self.layer_manager.get_layer("redline")
        if redline_layer is None or redline_layer.featureCount() == 0:
            return None
        try:
            from qgis.core import (QgsCoordinateReferenceSystem,
                                   QgsCoordinateTransform, QgsProject, QgsGeometry)
            geoms = [f.geometry() for f in redline_layer.getFeatures()
                     if f.geometry() and not f.geometry().isEmpty()]
            if not geoms:
                return None
            c   = QgsGeometry.unaryUnion(geoms).centroid().asPoint()
            bng = QgsCoordinateReferenceSystem("EPSG:27700")
            if redline_layer.crs() != bng:
                c = QgsCoordinateTransform(
                    redline_layer.crs(), bng, QgsProject.instance()).transform(c)
            return (c.x(), c.y())
        except Exception:
            return None

    def _redline_centroid_wgs84(self):
        """(lat, lon) centroid of the current redline in WGS84, computed live,
        or None if there's no redline. Shared by the NBN / NLS / Google Earth
        deep-links, which all want decimal lat/lon."""
        en = self._redline_centroid_bng()
        if en is None:
            return None
        try:
            from qgis.core import (QgsCoordinateReferenceSystem,
                                   QgsCoordinateTransform, QgsPointXY, QgsProject)
            bng = QgsCoordinateReferenceSystem("EPSG:27700")
            wgs = QgsCoordinateReferenceSystem("EPSG:4326")
            pt  = QgsCoordinateTransform(bng, wgs, QgsProject.instance()).transform(
                QgsPointXY(en[0], en[1]))
            return (pt.y(), pt.x())   # lat, lon
        except Exception:
            return None

    def _nearest_postcode_for_bng(self, easting, northing):
        """(postcode, admin_district) nearest to a BNG (EPSG:27700) coordinate via
        a single postcodes.io reverse-geocode call, or (None, None). One call
        yields both the postcode to paste into the LPA finder and the local-
        authority district name. Mirrors the postcodes.io usage in
        designated_sites.get_lpa_for_bng's fallback."""
        import json
        from urllib.request import urlopen, Request
        from urllib.parse import urlencode
        from qgis.core import (QgsCoordinateReferenceSystem,
                               QgsCoordinateTransform, QgsPointXY, QgsProject)
        bng = QgsCoordinateReferenceSystem("EPSG:27700")
        wgs = QgsCoordinateReferenceSystem("EPSG:4326")
        pt  = QgsCoordinateTransform(bng, wgs, QgsProject.instance()).transform(
            QgsPointXY(easting, northing))
        params = urlencode({"lon": f"{pt.x():.6f}", "lat": f"{pt.y():.6f}",
                            "limit": 1, "radius": 2000})
        req = Request(f"https://api.postcodes.io/postcodes?{params}",
                      headers={"User-Agent": "BiomeQGISPlugin/1.0"})
        with urlopen(req, timeout=6) as resp:
            data = json.loads(resp.read())
        if data.get("status") == 200 and data.get("result"):
            r = data["result"][0]
            return ((r.get("postcode") or "").strip() or None,
                    (r.get("admin_district") or "").strip() or None)
        return None, None

    def _open_nbn_atlas(self):
        """Open NBN Atlas using the shared saved Desk Study search radius.
        species-records search centred on the redline centroid; otherwise open
        the 'Explore your area' page for manual address/coordinate entry.

        NBN's radius search is driven by decimal lat/lon (radius in km), so the
        centroid is computed live from the current redline geometry and
        transformed to WGS84 — this stays correct even if the redline has moved
        since the Desk Study stats were last refreshed."""
        from qgis.PyQt.QtCore import QUrl
        from qgis.PyQt.QtGui import QDesktopServices

        cen = self._redline_centroid_wgs84()
        if cen is not None:
            from .config import get_buffer_radius
            lat, lon = cen
            radius_m = get_buffer_radius()
            radius_km = radius_m / 1000.0
            radius_arg = (str(int(radius_km)) if radius_km.is_integer()
                          else f"{radius_km:g}")
            radius_label = self._buffer_preset_label(radius_m)
            url = (f"https://records.nbnatlas.org/occurrences/search"
                   f"?q=*:*&lat={lat:.6f}&lon={lon:.6f}&radius={radius_arg}")
            self.iface.messageBar().pushMessage(
                "NBN Atlas",
                f"Opening a {radius_label} species-records search centred on the redline "
                f"({lat:.5f}, {lon:.5f})",
                level=Qgis.MessageLevel.Info, duration=6
            )
        else:
            url = "https://records.nbnatlas.org/explore/your-area"
            self.iface.messageBar().pushMessage(
                "NBN Atlas",
                "No redline found — opening 'Explore your area' for manual "
                "address/coordinate entry",
                level=Qgis.MessageLevel.Info, duration=6
            )

        QDesktopServices.openUrl(QUrl(url))

    def _open_nls_maps(self):
        """Open the National Library of Scotland georeferenced maps side-by-side
        viewer, centred on the redline — historic OS on the left, modern satellite
        (Bing hybrid) on the right, ideal for spotting land-use / habitat-history
        change. Falls back to the viewer's default view with no redline.

        NLS encodes the view in the URL *hash* (#zoom&lat&lon), zoom 1–25."""
        from qgis.PyQt.QtCore import QUrl
        from qgis.PyQt.QtGui import QDesktopServices

        base = "https://maps.nls.uk/geo/explore/side-by-side/"
        cen  = self._redline_centroid_wgs84()
        if cen is not None:
            lat, lon = cen
            url = (f"{base}#zoom=16&lat={lat:.5f}&lon={lon:.5f}"
                   f"&layers=1&right=BingHyb")
            self.iface.messageBar().pushMessage(
                "Historic Maps",
                "Opening NLS historic OS vs modern satellite, centred on the site "
                "— use the layer menu to pick an era",
                level=Qgis.MessageLevel.Info, duration=6
            )
        else:
            url = base

        QDesktopServices.openUrl(QUrl(url))

    def _open_google_earth(self):
        """Open Google Earth (web) centred tightly on the redline so the
        historical-imagery slider can be used to check for recent change (e.g. a
        felled tree). Falls back to Earth's default view with no redline.

        Earth's URL camera form is @lat,lon,<alt>a,<dist>d,<fov>y,<head>h,<tilt>t,
        <roll>r — 'd' (eye distance in metres) sets the zoom. The historical-
        imagery timeline is opened via an undocumented '/data=' UI-state token
        (see _EARTH_TIMELINE_ON below)."""
        from qgis.PyQt.QtCore import QUrl
        from qgis.PyQt.QtGui import QDesktopServices

        # Google Earth's "/data=" segment carries view/UI state as a base64url
        # protobuf. This exact token — decoded by diffing a timeline-off vs
        # timeline-on URL — is just the "historical-imagery timeline shown" flag
        # (f5{1,"",1}); it encodes no location and no date, so it can be appended
        # to any @lat,lon view to open with the timeline already on. Undocumented
        # format: if a future Earth build ignores it, centring still works because
        # the location lives in the @ part, not here.
        _EARTH_TIMELINE_ON = "CgwqBggBEgAYAUICCAE6AwoBMEICCABKDQj___________8BEAA"

        cen = self._redline_centroid_wgs84()
        if cen is not None:
            lat, lon = cen
            url = (f"https://earth.google.com/web/@{lat:.7f},{lon:.7f},"
                   f"0a,500d,35y,0h,0t,0r/data={_EARTH_TIMELINE_ON}")
            self.iface.messageBar().pushMessage(
                "Google Earth",
                "Centred on the site with historical imagery on — drag the time "
                "slider to check for recent change, e.g. a felled tree",
                level=Qgis.MessageLevel.Info, duration=7
            )
        else:
            url = "https://earth.google.com/web/"

        QDesktopServices.openUrl(QUrl(url))

    def _export_google_earth_kml(self):
        """Open a temporary KML containing only the site redline in Earth Pro.

        The saved view is north-up and top-down for checking the site's available
        historical aerial imagery. Earth Pro does not expose a KML setting that
        switches Historical Imagery on, so the user selects its clock button.
        """
        import json
        import math
        import os
        import subprocess
        import tempfile
        from qgis.PyQt.QtCore import QUrl
        from qgis.PyQt.QtGui import QDesktopServices
        from qgis.PyQt.QtWidgets import QMessageBox
        from qgis.core import (QgsCoordinateReferenceSystem,
                               QgsCoordinateTransform, QgsProject)

        redline_layer = self.layer_manager.get_layer("redline")
        if redline_layer is None or redline_layer.featureCount() == 0:
            QMessageBox.warning(
                self, "No Red Line", "Please draw a Red Line Boundary first.")
            return

        wgs84 = QgsCoordinateReferenceSystem("EPSG:4326")
        to_wgs = QgsCoordinateTransform(
            redline_layer.crs(), wgs84, QgsProject.instance())
        rings = []

        for feat in redline_layer.getFeatures():
            geom = feat.geometry()
            if geom.isEmpty():
                continue
            geom.transform(to_wgs)
            gj = json.loads(geom.asJson())
            if gj.get("type") == "Polygon":
                source_rings = gj.get("coordinates", [])[:1]
            elif gj.get("type") == "MultiPolygon":
                source_rings = [poly[0] for poly in gj.get("coordinates", []) if poly]
            else:
                source_rings = []
            for ring in source_rings:
                points = [(float(p[0]), float(p[1])) for p in ring]
                if len(points) < 3:
                    continue
                if points[0] != points[-1]:
                    points.append(points[0])
                rings.append(points)

        if not rings:
            QMessageBox.warning(
                self, "No Geometry",
                "No valid polygon boundary found in the Redline layer.")
            return

        longitudes = [lon for ring in rings for lon, _lat in ring]
        latitudes = [lat for ring in rings for _lon, lat in ring]
        lon = (min(longitudes) + max(longitudes)) / 2.0
        lat = (min(latitudes) + max(latitudes)) / 2.0
        width_m = ((max(longitudes) - min(longitudes)) * 111320.0
                   * max(0.2, math.cos(math.radians(lat))))
        height_m = (max(latitudes) - min(latitudes)) * 110540.0
        view_range = max(250.0, width_m, height_m) * 1.8

        placemarks = "".join(
            "<Placemark><name>Site boundary</name><styleUrl>#redline</styleUrl>"
            "<LineString><tessellate>1</tessellate>"
            "<altitudeMode>clampToGround</altitudeMode><coordinates>"
            + " ".join(f"{x:.7f},{y:.7f},0" for x, y in ring)
            + "</coordinates></LineString></Placemark>"
            for ring in rings)

        kml = (
            '<?xml version="1.0" encoding="UTF-8"?>\n'
            '<kml xmlns="http://www.opengis.net/kml/2.2"><Document>'
            '<name>Site redline – historical imagery</name>'
            '<LookAt>'
            f'<longitude>{lon:.7f}</longitude><latitude>{lat:.7f}</latitude>'
            '<altitude>0</altitude><heading>0</heading><tilt>0</tilt>'
            f'<range>{view_range:.1f}</range>'
            '<altitudeMode>relativeToGround</altitudeMode></LookAt>'
            '<Style id="redline"><LineStyle><color>ff0000ff</color>'
            '<width>5</width></LineStyle></Style>'
            f'{placemarks}</Document></kml>\n')

        try:
            fd, path = tempfile.mkstemp(prefix="biome_site_redline_", suffix=".kml")
            with os.fdopen(fd, "w", encoding="utf-8") as handle:
                handle.write(kml)
        except OSError as exc:
            QMessageBox.critical(
                self, "Google Earth Pro", f"Could not create the KML:\n{exc}")
            return

        candidates = (
            os.path.join(os.environ.get("ProgramFiles", ""),
                         "Google", "Google Earth Pro", "client", "googleearth.exe"),
            os.path.join(os.environ.get("ProgramFiles(x86)", ""),
                         "Google", "Google Earth Pro", "client", "googleearth.exe"),
            os.path.join(os.environ.get("LOCALAPPDATA", ""),
                         "Google", "Google Earth Pro", "client", "googleearth.exe"),
        )
        earth_exe = next((candidate for candidate in candidates
                          if candidate and os.path.isfile(candidate)), None)
        try:
            if earth_exe:
                subprocess.Popen([earth_exe, path])
            else:
                QDesktopServices.openUrl(QUrl.fromLocalFile(path))
        except OSError as exc:
            QMessageBox.critical(
                self, "Google Earth Pro", f"Could not open Google Earth Pro:\n{exc}")
            return

        self.iface.messageBar().pushMessage(
            "Google Earth Pro",
            "Opened the site redline in a top-down view — select the clock icon "
            "to show the Historical Imagery slider.",
            level=Qgis.MessageLevel.Info, duration=8)

    def _export_google_earth_spotlight_kml_legacy(self):
        """Export the redline as a 'spotlight' KML for Google Earth Pro: a
        translucent foggy off-white ground-drape covering everything EXCEPT
        the site (cut out as a hole), a fog 'funnel' canopy above it that fogs
        3D buildings, plus a red site outline. Opened in Earth Pro, tilted and
        screenshotted, it gives the classic site-location figure — fogged
        context, clear site, red boundary.

        Notes learned the hard way (2026-07-09):
        - The clampToGround drape dims the ground cleanly but cannot cover 3D
          buildings (they poke through), and extruding it to a volume creates
          visible chimney walls around the hole.
        - FLAT elevated fog sheets with a redline hole don't work either:
          from an oblique camera the hole parallax-shifts off the site, so
          the site itself gets fogged (seen in Earth Pro on a small site).
        - Hence the funnel: a canopy at FUNNEL_H whose hole slopes down at
          tan(MIN_ELEV) to meet the ground at the redline. Any view ray into
          the site from a camera above MIN_ELEV elevation passes inside the
          funnel — clear site at any azimuth — while rays to anything outside
          the redline cross the surface once and get fogged. Still a surface
          (no vertical faces), so no wall. Buildings taller than FUNNEL_H, or
          cameras tilted below MIN_ELEV, break the illusion.
        - The cone base is sunk BASE_SINK below terrain: Earth's 3D
          photogrammetric mesh floats above its terrain model, and
          relativeToGround follows the terrain, so a base ring at altitude 0
          visibly hovered off the ground (Earth Pro, 2026-07-09).
        - relativeToGround altitude interpolates between VERTICES; the funnel
          rings are dense near the site so the surface hugs local ground
          there. Far away, hills may rise through the canopy; the ground
          drape still fogs them.
        - KML colours are aabbggrr; the site ring is reversed so it reads as a
          hole in Earth's even-odd fill."""
        import json
        import math
        import os
        from qgis.PyQt.QtCore import QUrl
        from qgis.PyQt.QtGui import QDesktopServices
        from qgis.PyQt.QtWidgets import QFileDialog, QMessageBox
        from qgis.core import (QgsCoordinateReferenceSystem,
                               QgsCoordinateTransform, QgsGeometry,
                               QgsPointXY, QgsProject)

        redline_layer = self.layer_manager.get_layer("redline")
        if redline_layer is None or redline_layer.featureCount() == 0:
            QMessageBox.warning(
                self, "No Red Line",
                "Please draw a Red Line Boundary first.")
            return

        wgs84 = QgsCoordinateReferenceSystem("EPSG:4326")
        to_wgs = QgsCoordinateTransform(
            redline_layer.crs(), wgs84, QgsProject.instance())

        # Collect every exterior ring (Polygon and MultiPolygon parts) as
        # [(lon, lat), ...] in WGS84, plus untransformed copies for buffering.
        rings, all_lon, all_lat, src_geoms = [], [], [], []

        def _rings_from_geojson(gj):
            t = gj.get("type")
            coords = gj.get("coordinates", [])
            if t == "Polygon" and coords:
                yield coords[0]
            elif t == "MultiPolygon":
                for poly in coords:
                    if poly:
                        yield poly[0]

        for feat in redline_layer.getFeatures():
            geom = feat.geometry()
            if geom.isEmpty():
                continue
            src_geoms.append(QgsGeometry(geom))
            geom.transform(to_wgs)
            for ring in _rings_from_geojson(json.loads(geom.asJson())):
                pts = [(float(c[0]), float(c[1])) for c in ring]
                if len(pts) < 3:
                    continue
                if pts[0] != pts[-1]:
                    pts.append(pts[0])
                rings.append(pts)
                all_lon.extend(p[0] for p in pts)
                all_lat.extend(p[1] for p in pts)

        if not rings:
            QMessageBox.warning(
                self, "No Geometry",
                "No valid polygon boundary found in the Redline layer.")
            return

        # Curtain rectangle: site bbox expanded by ~0.12 deg (~9-13 km) so it
        # fills any framing the user is likely to screenshot in Earth.
        M = 0.12
        w, e = min(all_lon) - M, max(all_lon) + M
        s, n = min(all_lat) - M, max(all_lat) + M
        outer = [(w, s), (e, s), (e, n), (w, n), (w, s)]

        FOG_RGB = "EEEDEA"      # foggy off-white, KML bbggrr order
        DIM = "80" + FOG_RGB    # ~50% ground drape (aabbggrr)
        FOG = "73" + FOG_RGB    # ~45% fog canopy / funnel surface
        GLOW = "ff0000ff"       # red outline

        def _coords(pts, alt=0):
            return " ".join(f"{x:.7f},{y:.7f},{alt}" for x, y in pts)

        # Fog funnel: a canopy sheet at FUNNEL_H whose hole slopes down and
        # touches the ground exactly at the redline, at slope tan(MIN_ELEV).
        # From any camera above MIN_ELEV degrees of elevation, no view ray
        # into the site crosses the fog — the site stays clear even at
        # oblique angles (a flat sheet's hole parallax-shifts off the site).
        # Approximated by STEPS concentric ring strips between buffered
        # outlines of the redline.
        FUNNEL_H = 60.0     # m: canopy height — above typical UK buildings
        MIN_ELEV = 20.0     # deg: lowest camera elevation with a clear site
        BASE_SINK = 5.0     # m: cone base sunk below terrain at the redline —
        #                     Earth's 3D mesh floats above its terrain model,
        #                     so a base at 0 can visibly hover off the ground
        STEPS = 4
        r_true = FUNNEL_H / math.tan(math.radians(MIN_ELEV))   # ~165 m
        lat0 = (min(all_lat) + max(all_lat)) / 2.0
        # Buffers run in the layer CRS (EPSG:3857), whose metres are inflated
        # by 1/cos(lat) at UK latitudes — scale so the slope is true-metre.
        m_scale = 1.0 / max(0.2, math.cos(math.radians(lat0)))

        site_src = QgsGeometry.unaryUnion(src_geoms)
        shells, alts = [rings], [-BASE_SINK]    # level 0: redline, sunk
        for k in range(1, STEPS + 1):
            buf = site_src.buffer(r_true * k / STEPS * m_scale, 12)
            buf.transform(to_wgs)
            lvl = []
            for ring in _rings_from_geojson(json.loads(buf.asJson())):
                pts = [(float(c[0]), float(c[1])) for c in ring]
                if pts and pts[0] != pts[-1]:
                    pts.append(pts[0])
                if len(pts) >= 4:
                    lvl.append(pts)
            shells.append(lvl if lvl else shells[-1])
            alts.append(FUNNEL_H * k / STEPS)

        def _fog_pm(name, outers, alt_o, holes, alt_h):
            # One Placemark per outer ring; holes assigned by containment so
            # multi-part redlines pair up correctly.
            out = []
            for o in outers:
                o_geom = QgsGeometry.fromPolygonXY(
                    [[QgsPointXY(x, y) for x, y in o]])
                inner = "".join(
                    "<innerBoundaryIs><LinearRing><coordinates>"
                    f"{_coords(h[::-1], alt_h)}"
                    "</coordinates></LinearRing></innerBoundaryIs>"
                    for h in holes
                    if o_geom.contains(
                        QgsGeometry.fromPointXY(QgsPointXY(*h[0]))))
                out.append(
                    f"<Placemark><name>{name}</name><styleUrl>#fog</styleUrl>"
                    "<Polygon><altitudeMode>relativeToGround</altitudeMode>"
                    "<outerBoundaryIs><LinearRing><coordinates>"
                    f"{_coords(o, alt_o)}"
                    "</coordinates></LinearRing></outerBoundaryIs>"
                    f"{inner}</Polygon></Placemark>")
            return "".join(out)

        funnel_xml = "".join(
            _fog_pm(f"Fog funnel {k}", shells[k], alts[k],
                    shells[k - 1], alts[k - 1])
            for k in range(1, STEPS + 1))
        canopy_xml = _fog_pm("Fog canopy", [outer], FUNNEL_H,
                             shells[STEPS], FUNNEL_H)

        ground_inner = "".join(
            "<innerBoundaryIs><LinearRing><coordinates>"
            f"{_coords(r[::-1])}"
            "</coordinates></LinearRing></innerBoundaryIs>"
            for r in rings)
        ground_xml = (
            '<Placemark><name>Ground fog</name><styleUrl>#curtain</styleUrl>'
            '<Polygon><tessellate>1</tessellate>'
            '<altitudeMode>clampToGround</altitudeMode>'
            '<outerBoundaryIs><LinearRing><coordinates>'
            f'{_coords(outer)}'
            '</coordinates></LinearRing></outerBoundaryIs>'
            f'{ground_inner}</Polygon></Placemark>')

        outline_xml = "".join(
            "<Placemark><name>Site boundary</name><styleUrl>#siteline</styleUrl>"
            "<Polygon><tessellate>1</tessellate>"
            "<altitudeMode>clampToGround</altitudeMode>"
            "<outerBoundaryIs><LinearRing><coordinates>"
            f"{_coords(r)}"
            "</coordinates></LinearRing></outerBoundaryIs></Polygon></Placemark>"
            for r in rings)

        kml = (
            '<?xml version="1.0" encoding="UTF-8"?>\n'
            '<kml xmlns="http://www.opengis.net/kml/2.2"><Document>'
            '<name>Site spotlight</name>'
            '<Style id="curtain">'
            '<LineStyle><width>0</width><color>00000000</color></LineStyle>'
            f'<PolyStyle><color>{DIM}</color><fill>1</fill><outline>0</outline></PolyStyle>'
            '</Style>'
            '<Style id="fog">'
            '<LineStyle><width>0</width><color>00000000</color></LineStyle>'
            f'<PolyStyle><color>{FOG}</color><fill>1</fill><outline>0</outline></PolyStyle>'
            '</Style>'
            '<Style id="siteline">'
            f'<LineStyle><color>{GLOW}</color><width>5</width></LineStyle>'
            '<PolyStyle><fill>0</fill><outline>1</outline></PolyStyle>'
            '</Style>'
            f'{ground_xml}'
            f'{funnel_xml}'
            f'{canopy_xml}'
            f'{outline_xml}'
            '</Document></kml>\n'
        )

        # Default next to the project if saved, else the user's home.
        proj_path = QgsProject.instance().fileName()
        base_dir = (os.path.dirname(proj_path) if proj_path
                    else os.path.expanduser("~"))
        default = os.path.join(base_dir, "site_spotlight.kml")
        path, _ = QFileDialog.getSaveFileName(
            self, "Export site spotlight KML", default, "KML files (*.kml)")
        if not path:
            return
        if not path.lower().endswith(".kml"):
            path += ".kml"

        try:
            with open(path, "w", encoding="utf-8") as fh:
                fh.write(kml)
        except OSError as exc:
            QMessageBox.critical(
                self, "Export failed", f"Could not write the KML:\n{exc}")
            return

        self.iface.messageBar().pushMessage(
            "Google Earth",
            "Spotlight KML saved — opening it now. In Google Earth Pro: tilt "
            "(hold Shift + drag), frame the site, then File → Save → "
            "Save Image for a high-res figure.",
            level=Qgis.MessageLevel.Info, duration=9)
        QDesktopServices.openUrl(QUrl.fromLocalFile(path))

    def _export_3d_spotlight_html(self):
        """Export the interactive 3D spotlight viewer (standalone HTML,
        CesiumJS + Google Photorealistic 3D Tiles — see earth3d_export.py).
        Unlike the Google Earth Pro KML, the fog here is applied to the 3D
        mesh itself, clipped exactly to the redline, so off-site trees and
        buildings fog properly from any angle with no wall artefacts. Needs
        a Google Maps Platform API key with the Map Tiles API enabled."""
        import json
        import os
        from qgis.PyQt.QtCore import QUrl
        from qgis.PyQt.QtGui import QDesktopServices
        from qgis.PyQt.QtWidgets import QFileDialog, QMessageBox
        from qgis.core import (QgsCoordinateReferenceSystem,
                               QgsCoordinateTransform, QgsProject)
        from .earth3d_export import build_spotlight_html

        redline_layer = self.layer_manager.get_layer("redline")
        if redline_layer is None or redline_layer.featureCount() == 0:
            QMessageBox.warning(
                self, "No Red Line",
                "Please draw a Red Line Boundary first.")
            return

        key = get_google_maps_api_key()
        if not key:
            self._open_google_key_dialog()
            key = get_google_maps_api_key()
            if not key:
                return

        wgs84 = QgsCoordinateReferenceSystem("EPSG:4326")
        to_wgs = QgsCoordinateTransform(
            redline_layer.crs(), wgs84, QgsProject.instance())

        def _rings_from_geojson(gj):
            t = gj.get("type")
            coords = gj.get("coordinates", [])
            if t == "Polygon" and coords:
                yield coords[0]
            elif t == "MultiPolygon":
                for poly in coords:
                    if poly:
                        yield poly[0]

        def _collect_rings(layer, transform):
            out = []
            for feat in layer.getFeatures():
                geom = feat.geometry()
                if geom.isEmpty():
                    continue
                geom.transform(transform)
                for ring in _rings_from_geojson(json.loads(geom.asJson())):
                    pts = [(float(c[0]), float(c[1])) for c in ring]
                    if len(pts) >= 3:
                        out.append(pts)
            return out

        rings = _collect_rings(redline_layer, to_wgs)
        if not rings:
            QMessageBox.warning(
                self, "No Geometry",
                "No valid polygon boundary found in the Redline layer.")
            return

        # Blue Line Boundary, when drawn, adds a red/blue switch in the viewer.
        blue_rings = []
        blue_layer = self.layer_manager.get_layer("site_boundary")
        if blue_layer is not None and blue_layer.featureCount() > 0:
            blue_to_wgs = QgsCoordinateTransform(
                blue_layer.crs(), wgs84, QgsProject.instance())
            blue_rings = _collect_rings(blue_layer, blue_to_wgs)

        title = QgsProject.instance().baseName() or "Site"
        html = build_spotlight_html(rings, key, title=title,
                                    blue_rings=blue_rings)

        proj_path = QgsProject.instance().fileName()
        base_dir = (os.path.dirname(proj_path) if proj_path
                    else os.path.expanduser("~"))
        # Timestamped name: every export is a distinct file, so an old tab or
        # stale copy can never masquerade as the new one.
        from datetime import datetime
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        default = os.path.join(base_dir, f"site_3d_spotlight_{stamp}.html")
        path, _ = QFileDialog.getSaveFileName(
            self, "Export 3D spotlight viewer", default,
            "HTML files (*.html)")
        if not path:
            return
        if not path.lower().endswith((".html", ".htm")):
            path += ".html"

        try:
            with open(path, "w", encoding="utf-8") as fh:
                fh.write(html)
        except OSError as exc:
            QMessageBox.critical(
                self, "Export failed", f"Could not write the file:\n{exc}")
            return

        self.iface.messageBar().pushMessage(
            "3D Spotlight",
            "Viewer saved — opening in your browser. Orbit and tilt freely, "
            "then use 'Save screenshot'. Keep the file private: your Google "
            "API key is embedded in it.",
            level=Qgis.MessageLevel.Info, duration=10)
        # NEVER open this as file:// — Chrome refuses to run web workers on
        # file:// pages, Cesium builds the globe surface in a worker, and the
        # result is a black-space globe with no imagery, fog or boundary
        # lines (debugged live 2026-07-09). Serve it over loopback instead.
        try:
            url = self._serve_spotlight_file(path)
        except Exception:
            url = None
        if url:
            QDesktopServices.openUrl(QUrl(url))
        else:
            QDesktopServices.openUrl(QUrl.fromLocalFile(path))

    def _serve_spotlight_file(self, path):
        """Serve the exported spotlight HTML over 127.0.0.1 and return its
        URL. One quiet singleton server per QGIS session, re-rooted when an
        export lands in a different folder; the daemon thread dies with
        QGIS."""
        import functools
        import http.server
        import os
        import threading
        from urllib.parse import quote

        directory = os.path.dirname(path)
        srv = getattr(self, "_spotlight_httpd", None)
        if srv is not None and getattr(self, "_spotlight_http_dir", "") != directory:
            try:
                srv.shutdown()
                srv.server_close()
            except Exception:
                pass
            srv = None
        if srv is None:
            class _QuietHandler(http.server.SimpleHTTPRequestHandler):
                def log_message(self, *args):
                    pass
            handler = functools.partial(_QuietHandler, directory=directory)
            srv = http.server.ThreadingHTTPServer(("127.0.0.1", 0), handler)
            threading.Thread(target=srv.serve_forever, daemon=True).start()
            self._spotlight_httpd = srv
            self._spotlight_http_dir = directory
        port = srv.server_address[1]
        return f"http://127.0.0.1:{port}/{quote(os.path.basename(path))}"

    def _open_street_view(self):
        """Open Google Maps zoomed tightly on the redline centroid so the user can
        grab the pegman and click into ground-level Street View on any covered
        street — without being forced into a single panorama that may not exist at
        the exact spot. Falls back to plain Google Maps with no redline.

        Google exposes NO stable URL parameter for the blue Street View coverage
        overlay — it's a pegman-drag UI state, not URL state (the official Maps
        URLs API only supports the traffic / transit / bicycling layers). The old
        '/data=!5m1!1e5' "Photos" layer code that this used to append was retired
        by Google; a dead enum now makes Maps reject the whole '/data=…' path and
        bounce to the homepage, which is why the link previously opened unzoomed.
        So we keep the URL to the one thing that's reliable — the '@lat,lon,zoom'
        map view — and let the user drop the pegman (bottom-right) to light up the
        blue streets, then click one."""
        from qgis.PyQt.QtCore import QUrl
        from qgis.PyQt.QtGui import QDesktopServices

        cen = self._redline_centroid_wgs84()
        if cen is not None:
            lat, lon = cen
            url = f"https://www.google.com/maps/@{lat:.7f},{lon:.7f},17z"
            self.iface.messageBar().pushMessage(
                "Street View",
                "Opening Google Maps zoomed on the site — grab the pegman "
                "(bottom-right) to show the blue Street View streets, then click "
                "one to drop into ground level",
                level=Qgis.MessageLevel.Info, duration=8
            )
        else:
            url = "https://www.google.com/maps"

        QDesktopServices.openUrl(QUrl(url))

    def _on_gcn_opacity_changed(self, value: int):
        """Update opacity on all loaded GCN Risk Zone layers and refresh the label."""
        self._gcn_opacity_val_lbl.setText(f"{value}%")
        opacity = value / 100.0
        from qgis.core import QgsProject
        root = QgsProject.instance().layerTreeRoot()
        group = root.findGroup(_GCN_GROUP_NAME)
        if group is None:
            return
        for node in group.children():
            lyr = QgsProject.instance().mapLayer(node.layerId()) if hasattr(node, 'layerId') else None
            if lyr is not None:
                try:
                    lyr.setOpacity(opacity)
                    lyr.triggerRepaint()
                except Exception:
                    pass

    def _on_georef_opacity_changed(self, value: int):
        """Apply opacity to all raster layers that are not part of the plugin's
        own layer groups (Base Maps, Baseline Habitats, Proposed Habitats,
        Ecological Features). These are the user's georeferenced/imported rasters."""
        self._georef_opacity_val_lbl.setText(f"{value}%")
        opacity = value / 100.0
        from qgis.core import QgsProject, QgsRasterLayer
        project = QgsProject.instance()
        root = project.layerTreeRoot()

        # Collect IDs of every layer inside the plugin's own groups
        _PLUGIN_GROUPS = (
            "Base Maps", "Baseline Habitats",
            "Proposed Habitats", "Ecological Features",
        )
        plugin_ids = set()
        for grp_name in _PLUGIN_GROUPS:
            grp = root.findGroup(grp_name)
            if grp is not None:
                for node in grp.findLayers():
                    if node.layer():
                        plugin_ids.add(node.layer().id())

        # Layer name fragments that should never be affected by this slider
        _EXCLUDE_NAMES = (
            "property boundaries",
            "hmlr",
            "inspire",
            "land registry",
            "mastermap",
            "cadastral",
            "title register",
        )

        changed = False
        for lyr in project.mapLayers().values():
            if not isinstance(lyr, QgsRasterLayer):
                continue
            if lyr.id() in plugin_ids:
                continue
            if any(frag in lyr.name().lower() for frag in _EXCLUDE_NAMES):
                continue
            try:
                lyr.setOpacity(opacity)
                lyr.triggerRepaint()
                changed = True
            except Exception:
                pass
        if changed:
            self.iface.mapCanvas().refresh()

    def _on_epsl_toggle(self, checked):
        """Show or hide (load if needed) EPSL point layer and distance lines within the search radius."""
        if not checked:
            # Just hide the layers — don't remove them
            from qgis.core import QgsProject
            for layer_name in ("EPSL Locations", "EPSL Distance Lines"):
                for lyr in QgsProject.instance().mapLayersByName(layer_name):
                    node = QgsProject.instance().layerTreeRoot().findLayer(lyr.id())
                    if node is not None:
                        node.setItemVisibilityChecked(False)
            self._epsl_btn.setText("🦇  Show EPSL\nLocations")
            self._epsl_lines_btn.setEnabled(False)
            self._epsl_lines_btn.setChecked(False)
            self._epsl_lines_btn.setText("〰  Distance lines: Off")
            self._epsl_label_btn.setEnabled(False)
            self._epsl_filter_btn.setEnabled(False)
            self._epsl_style_btn.setEnabled(False)
            self.iface.mapCanvas().refresh()
            return
        if not hasattr(self, "_centroid_bng"):
            QMessageBox.warning(self, "No Redline",
                "Please draw a redline boundary first.")
            self._epsl_btn.setChecked(False)
            return
        if _region.is_scotland():
            self._epsl_btn.setChecked(False)
            self._epsl_scotland_notice()
            return
        # If the layer already exists, just make it visible again
        from qgis.core import QgsProject
        existing = QgsProject.instance().mapLayersByName("EPSL Locations")
        if existing:
            node = QgsProject.instance().layerTreeRoot().findLayer(existing[0].id())
            if node is not None:
                node.setItemVisibilityChecked(True)
            self._epsl_btn.setText("🦇  Hide EPSL\nLocations")
            self._epsl_lines_btn.setEnabled(True)
            self._epsl_label_btn.setEnabled(True)
            self._epsl_filter_btn.setEnabled(True)
            self._epsl_style_btn.setEnabled(True)
            # Reapply the current default so an EPSL memory layer retained
            # across a plugin reload does not keep the former star renderer.
            self._epsl_apply_display()
            return
        from qgis.PyQt.QtWidgets import QApplication
        from qgis.PyQt.QtCore import Qt as _Qt
        e, n = self._centroid_bng
        QApplication.setOverrideCursor(_Qt.CursorShape.WaitCursor)
        redline_geom = self._get_redline_geom()
        try:
            ok, result = load_epsl_layer(e, n, redline_geom)
            if ok:
                # Apply current label mode immediately
                set_epsl_label_mode(self._epsl_label_mode)
                # Load distance lines but keep hidden (lines are off by default)
                load_epsl_distance_lines(redline_geom)
                from qgis.core import QgsProject
                for lyr in QgsProject.instance().mapLayersByName("EPSL Distance Lines"):
                    node = QgsProject.instance().layerTreeRoot().findLayer(lyr.id())
                    if node is not None:
                        node.setItemVisibilityChecked(False)
        except Exception as exc:
            import traceback
            from qgis.core import QgsMessageLog, Qgis
            QgsMessageLog.logMessage("EPSL load error:\n" + traceback.format_exc(),
                                     "Biome", Qgis.MessageLevel.Critical)
            ok, result = False, f"{type(exc).__name__}: {exc}"
        finally:
            QApplication.restoreOverrideCursor()
        if not ok:
            if result == "no_results":
                QMessageBox.information(self, "EPSL",
                    "No granted EPSL licences found within the search radius of the redline centroid.")
            else:
                QMessageBox.critical(self, "EPSL Load Failed", result)
            self._epsl_btn.setChecked(False)
            self._epsl_btn.setText("🦇  Show EPSL\nLocations")
        else:
            self._epsl_btn.setText("🦇  Hide EPSL\nLocations")
            self._epsl_lines_btn.setEnabled(True)
            self._epsl_lines_btn.setChecked(False)
            self._epsl_lines_btn.setText("〰  Distance lines: Off")
            self._epsl_label_btn.setEnabled(True)
            self._epsl_filter_btn.setEnabled(True)
            self._epsl_style_btn.setEnabled(True)
            self._reset_epsl_display_state()
            self.iface.mapCanvas().refresh()

    def _on_epsl_lines_toggle(self, checked):
        """Show or hide the EPSL distance lines layer."""
        from qgis.core import QgsProject
        lyrs = QgsProject.instance().mapLayersByName("EPSL Distance Lines")
        for lyr in lyrs:
            node = QgsProject.instance().layerTreeRoot().findLayer(lyr.id())
            if node is not None:
                node.setItemVisibilityChecked(checked)
        self._epsl_lines_btn.setText("〰  Distance lines: On" if checked else "〰  Distance lines: Off")
        self.iface.mapCanvas().refresh()

    def _on_epsl_label_selected(self, _checked=False):
        """Apply the explicit label choice selected from the Labels menu."""
        action = self.sender()
        mode = action.data() if action is not None else "species"
        _labels = {
            "species": "Labels: Species",
            "licence": "Labels: Licence reference",
            "both":    "Labels: Species and reference",
            "off":     "Labels: None",
        }
        try:
            self._epsl_label_mode = mode if mode in _labels else "species"
            self._epsl_label_btn.setText(_labels[self._epsl_label_mode])
            set_epsl_label_mode(self._epsl_label_mode)
            # set_epsl_label_mode rebuilds simple labeling, dropping any active
            # filter restriction — re-apply it so labels stay in step with the
            # filtered icons.
            mode = getattr(self, "_epsl_filter_mode", "all")
            if mode != "all":
                from qgis.core import QgsProject
                for lyr in QgsProject.instance().mapLayersByName("EPSL Locations"):
                    match_ids = [f.id() for f in lyr.getFeatures()
                                 if self._epsl_feature_matches(f, mode)]
                    id_expr = ("$id IN (%s)" % ",".join(str(i) for i in match_ids)
                               if match_ids else "FALSE")
                    self._epsl_apply_label_filter(lyr, id_expr)
                    lyr.triggerRepaint()
            self.iface.mapCanvas().refresh()
        except Exception as e:
            from qgis.core import QgsMessageLog, Qgis
            QgsMessageLog.logMessage(f"[EPSL Label] error: {e}", "Biome", Qgis.MessageLevel.Critical)

    def _epsl_feature_matches(self, feat, mode):
        """Return True if an EPSL feature matches the active filter `mode`,
        using the SAME recent/historic rule as the EPSL report:
          • recent  = current year and preceding nine calendar years,
            or no start date (treated as still active)
          • historic = older than that
          • bats     = species_gr contains 'bat' (case-insensitive)
        """
        from datetime import datetime, timezone
        if mode == "all":
            return True
        sg = (feat["species_gr"] or "").lower()
        is_bat = "bat" in sg
        # Ten inclusive calendar years: in 2026 this is 2017–2026; on
        # 1 January 2027 it automatically becomes 2018–2027.
        current_year = datetime.now(tz=timezone.utc).year
        cutoff = datetime(current_year - 9, 1, 1, tzinfo=timezone.utc)
        s = (feat["start_date"] or "").strip()
        if not s:
            is_recent = True
        else:
            try:
                dt = datetime.strptime(s, "%d/%m/%Y").replace(tzinfo=timezone.utc)
                is_recent = dt >= cutoff
            except (ValueError, TypeError):
                is_recent = True   # unparseable date → treat as active, like report
        if mode == "bats":
            return is_bat
        if mode == "recent":
            return is_recent
        if mode == "bats_recent":
            return is_bat and is_recent
        return True

    def _on_epsl_filter_toggle(self, checked):
        """Apply the direct All species/Bats only toggle.

        Bats are matched with the SAME rule as the EPSL report (species group
        contains 'bat'). The filter composes with the active Style — non-bat
        points are hidden while the bats keep whichever marker style is set —
        because both are rebuilt together in _epsl_apply_display().
        """
        self._epsl_filter_mode = "bats" if checked else "all"
        self._epsl_filter_btn.setText(
            "Species: Bats only" if checked else "Species: All species")
        self._epsl_apply_display()

    def _epsl_apply_label_filter(self, lyr, id_expr):
        """Wrap the layer's current simple labeling in a rule-based labeling
        whose single rule carries `id_expr` (or remove the filter when
        id_expr is None). Preserves the current label field/format set by
        set_epsl_label_mode."""
        from qgis.core import (
            QgsRuleBasedLabeling, QgsVectorLayerSimpleLabeling,
            QgsPalLayerSettings,
        )
        labeling = lyr.labeling()
        if labeling is None or not lyr.labelsEnabled():
            return
        # Extract the QgsPalLayerSettings currently in use. Copy-construct —
        # QgsPalLayerSettings has no .clone(), and a silent failure here means
        # filtered-out points keep their labels.
        try:
            if isinstance(labeling, QgsRuleBasedLabeling):
                # Already rule-based — pull the first rule's settings.
                kids = labeling.rootRule().children()
                pal = QgsPalLayerSettings(kids[0].settings()) if kids else None
            else:
                pal = QgsPalLayerSettings(labeling.settings())
        except Exception as e:
            from qgis.core import QgsMessageLog, Qgis
            QgsMessageLog.logMessage(f"[EPSL] label-filter settings copy failed: {e}",
                                     "Biome", Qgis.MessageLevel.Warning)
            pal = None
        if pal is None:
            return

        if id_expr is None:
            lyr.setLabeling(QgsVectorLayerSimpleLabeling(pal))
            lyr.setLabelsEnabled(True)
            return

        root = QgsRuleBasedLabeling.Rule(None)
        rule = QgsRuleBasedLabeling.Rule(pal)
        rule.setFilterExpression(id_expr)
        root.appendChild(rule)
        lyr.setLabeling(QgsRuleBasedLabeling(root))
        lyr.setLabelsEnabled(True)

    def _on_epsl_display_selected(self, _checked=False):
        """Apply the explicit marker choice selected from the Display menu."""
        action = self.sender()
        mode = action.data() if action is not None else "icons"
        _labels = {
            "star":   "Display: Standard",
            "icons":  "Display: Species icons",
            "recent": "Display: Recent/historic",
        }
        self._epsl_style_mode = mode if mode in _labels else "icons"
        self._epsl_style_btn.setText(_labels[self._epsl_style_mode])
        self._epsl_apply_display()

    def _reset_epsl_display_state(self):
        """Reset Style→Species icons and Filter→All, syncing both controls."""
        self._epsl_filter_mode = "all"
        self._epsl_style_mode  = "icons"
        if hasattr(self, "_epsl_filter_btn"):
            self._epsl_filter_btn.setChecked(False)
            self._epsl_filter_btn.setText("Species: All species")
        if hasattr(self, "_epsl_style_btn"):
            self._epsl_style_btn.setText("Display: Species icons")
        for mode, action in getattr(self, "_epsl_display_actions", {}).items():
            action.setChecked(mode == "icons")

    def _epsl_apply_display(self):
        """Single source of truth for the EPSL symbology — rebuilds the renderer
        AND the label filter together from the two orthogonal display states so
        they compose instead of overwriting each other:
            _epsl_style_mode  ∈ {star, icons, recent}  → base marker style
            _epsl_filter_mode ∈ {all, bats}            → which points are shown
        Choosing 'Bats' hides non-bat points (and their labels) while keeping
        whichever style is active for the ones that remain. setSubsetString is
        avoided — the EPSL layer is a memory layer whose subset support is
        unreliable across QGIS builds — so the bat restriction is baked into the
        renderer plus an `$id IN (...)` label filter instead."""
        from qgis.core import QgsProject
        bats_only = (getattr(self, "_epsl_filter_mode", "all") == "bats")
        style = getattr(self, "_epsl_style_mode", "icons")
        for lyr in QgsProject.instance().mapLayersByName("EPSL Locations"):
            # Clear any stale subset string from older builds.
            try:
                if lyr.subsetString():
                    lyr.setSubsetString("")
            except Exception:
                pass
            if style == "icons":
                apply_epsl_emoji_renderer(lyr, bats_only=bats_only)
            elif style == "recent":
                apply_epsl_recent_historic_renderer(lyr, bats_only=bats_only)
            else:
                apply_epsl_star_renderer(lyr, bats_only=bats_only)
            # Keep labels in step with the visible icons.
            if bats_only:
                match_ids = [f.id() for f in lyr.getFeatures()
                             if self._epsl_feature_matches(f, "bats")]
                id_expr = ("$id IN (%s)" % ",".join(str(i) for i in match_ids)
                           if match_ids else "FALSE")
            else:
                id_expr = None
            self._epsl_apply_label_filter(lyr, id_expr)
            lyr.triggerRepaint()
        self.iface.mapCanvas().refresh()

    # ── Tree Preservation Order handlers ──────────────────────────────────

    def _on_tpo_toggle(self, checked):
        """Show or hide TPO trees (points) AND zones (polygons) within the search radius.

        The 'tree' dataset only holds individual trees; woodland / area / group
        TPOs live in the 'tree-preservation-zone' polygon dataset, so both are
        loaded together — otherwise area/woodland TPOs never appear on the map.
        """
        from qgis.core import QgsProject
        root = QgsProject.instance().layerTreeRoot()
        _tpo_layers = ("Tree Preservation Orders", "TPO Zones",
                       "TPO Distance Lines")

        if not checked:
            if getattr(self, "_active_local_buffer_preset", None) == "tpo":
                self._active_local_buffer_preset = None
            # Just hide — don't remove
            for layer_name in _tpo_layers:
                for lyr in QgsProject.instance().mapLayersByName(layer_name):
                    node = root.findLayer(lyr.id())
                    if node is not None:
                        node.setItemVisibilityChecked(False)
            self._tpo_btn.setText("🌳  Show TPO\nLocations")
            self._tpo_lines_btn.setEnabled(False)
            self._tpo_lines_btn.setChecked(False)
            self._tpo_lines_btn.setText("〰  Distance lines: Off")
            self._tpo_label_btn.setEnabled(False)
            self.iface.mapCanvas().refresh()
            return
        if not hasattr(self, "_centroid_bng"):
            QMessageBox.warning(self, "No Redline",
                "Please draw a redline boundary first.")
            self._tpo_btn.setChecked(False)
            return
        if _region.is_scotland() and not self._ensure_spatialhub_key():
            self._tpo_btn.setChecked(False)
            return
        # TPO and standing-water searches use mutually exclusive local-buffer
        # presets. Starting a TPO search replaces any 250/500 m water preset.
        self._active_local_buffer_preset = "tpo"
        self._manual_local_buffers_enabled = False
        # Re-show existing layers if already loaded
        existing_pts_all = QgsProject.instance().mapLayersByName("Tree Preservation Orders")
        existing_pts = [lyr for lyr in existing_pts_all if lyr.featureCount() > 0]
        existing_zones = QgsProject.instance().mapLayersByName("TPO Zones")
        if existing_pts or existing_zones:
            for lyr in (existing_pts + existing_zones):
                node = root.findLayer(lyr.id())
                if node is not None:
                    node.setItemVisibilityChecked(True)
            self._tpo_btn.setText("🌳  Hide TPO\nLocations")
            self._tpo_lines_btn.setEnabled(bool(existing_pts))
            self._tpo_label_btn.setEnabled(bool(existing_pts))
            self.iface.mapCanvas().refresh()
            return
        # The empty point layer is retained so exported map legends can still
        # explain the TPO symbol. It is not a successful cached result, though:
        # reusing it must still show the council-link guidance.
        if existing_pts_all:
            self._show_tpo_no_results_dialog()
            self._tpo_btn.setChecked(False)
            self._tpo_btn.setText("🌳  Show TPO\nLocations")
            return
        e, n = self._centroid_bng
        from qgis.PyQt.QtWidgets import QApplication
        from qgis.PyQt.QtCore import Qt as _Qt
        from concurrent.futures import ThreadPoolExecutor
        from .designated_sites import (
            _planning_bbox_wkt, fetch_planning_entities, _TPO_RADIUS_M,
        )
        QApplication.setOverrideCursor(_Qt.CursorShape.WaitCursor)
        redline_geom = self._get_redline_geom()

        # Phase 1 — fetch the two planning.data.gov.uk datasets concurrently
        # (points 'tree' + polygon 'tree-preservation-zone'). HTTP only; the
        # shared WGS84 bbox is computed here on the main thread (needs QGIS).
        try:
            bbox_wkt = _planning_bbox_wkt(e, n, redline_geom, _TPO_RADIUS_M)
        except Exception:
            bbox_wkt = None

        if bbox_wkt is None:
            # Rare transform failure — let the loaders fetch synchronously.
            pts_ok,   pts_res   = load_tpo_layer(e, n, redline_geom)
            zones_ok, zones_res = load_tpo_zones_layer(e, n, redline_geom)
        else:
            def _safe_fetch(dataset):
                try:
                    return fetch_planning_entities(dataset, bbox_wkt), None
                except Exception as exc:
                    return None, str(exc)
            with ThreadPoolExecutor(max_workers=2) as pool:
                f_pts   = pool.submit(_safe_fetch, "tree")
                f_zones = pool.submit(_safe_fetch, "tree-preservation-zone")
                pts_entities,   pts_err   = f_pts.result()
                zones_entities, zones_err = f_zones.result()

            # Phase 2 — build both layers on the main thread (no second download)
            if pts_err is not None:
                pts_ok, pts_res = False, pts_err
            else:
                pts_ok, pts_res = load_tpo_layer(e, n, redline_geom, entities=pts_entities)
            if zones_err is not None:
                zones_ok, zones_res = False, zones_err
            else:
                zones_ok, zones_res = load_tpo_zones_layer(e, n, redline_geom, entities=zones_entities)

        if pts_ok:
            load_tpo_distance_lines(redline_geom)
            for lyr in QgsProject.instance().mapLayersByName("TPO Distance Lines"):
                node = root.findLayer(lyr.id())
                if node is not None:
                    node.setItemVisibilityChecked(False)
        QApplication.restoreOverrideCursor()

        # A non-'no_results' failure is a real (e.g. network) error
        hard_err = next((r for ok, r in ((pts_ok, pts_res), (zones_ok, zones_res))
                         if not ok and r != "no_results"), None)

        if not pts_ok and not zones_ok:
            if pts_res == "no_results" and zones_res == "no_results":
                # Genuinely nothing in the national dataset → council link
                self._show_tpo_no_results_dialog()
            else:
                QMessageBox.critical(self, "TPO Load Failed", hard_err or pts_res)
            self._tpo_btn.setChecked(False)
            self._tpo_btn.setText("🌳  Show TPO\nLocations")
            return

        # At least one layer loaded
        if hard_err:
            self.iface.messageBar().pushMessage(
                "TPO", f"Some TPO data could not be loaded: {hard_err}",
                level=Qgis.MessageLevel.Warning, duration=6)

        self._tpo_btn.setText("🌳  Hide TPO\nLocations")
        self._tpo_lines_btn.setEnabled(pts_ok)
        self._tpo_lines_btn.setChecked(False)
        self._tpo_lines_btn.setText("〰  Distance lines: Off")
        self._tpo_label_btn.setEnabled(pts_ok)

        n_pts   = pts_res.featureCount()   if pts_ok   else 0
        n_zones = zones_res.featureCount() if zones_ok else 0
        self.iface.messageBar().pushMessage(
            "TPO",
            f"Loaded {n_pts} protected tree(s) and {n_zones} TPO zone(s) "
            f"(woodland/area/group) within the search radius.",
            level=Qgis.MessageLevel.Info, duration=6)
        self.iface.mapCanvas().refresh()

    def _on_conservation_toggle(self, checked):
        """Show or hide (load if needed) conservation areas within the search radius."""
        from qgis.core import QgsProject
        root = QgsProject.instance().layerTreeRoot()
        name = "Conservation Areas"

        if not checked:
            for lyr in QgsProject.instance().mapLayersByName(name):
                node = root.findLayer(lyr.id())
                if node is not None:
                    node.setItemVisibilityChecked(False)
            self._ca_btn.setText("🏛  Show Conservation\nAreas")
            self.iface.mapCanvas().refresh()
            return

        if not hasattr(self, "_centroid_bng"):
            QMessageBox.warning(self, "No Redline",
                "Please draw a redline boundary first.")
            self._ca_btn.setChecked(False)
            return

        existing = QgsProject.instance().mapLayersByName(name)
        if existing:
            node = root.findLayer(existing[0].id())
            if node is not None:
                node.setItemVisibilityChecked(True)
            self._ca_btn.setText("🏛  Hide Conservation\nAreas")
            self.iface.mapCanvas().refresh()
            return

        e, n = self._centroid_bng
        from qgis.PyQt.QtWidgets import QApplication
        from qgis.PyQt.QtCore import Qt as _Qt
        QApplication.setOverrideCursor(_Qt.CursorShape.WaitCursor)
        ok, result = load_conservation_areas_layer(e, n, self._get_redline_geom())
        QApplication.restoreOverrideCursor()

        if not ok:
            self._ca_btn.setChecked(False)
            self._ca_btn.setText("🏛  Show Conservation\nAreas")
            if result == "no_results":
                self.iface.messageBar().pushMessage(
                    "Conservation Areas",
                    "No conservation areas found within the search radius of the site boundary.",
                    level=Qgis.MessageLevel.Info, duration=6)
            else:
                QMessageBox.critical(self, "Conservation Areas Load Failed", result)
            return

        self._ca_btn.setText("🏛  Hide Conservation\nAreas")
        self.iface.messageBar().pushMessage(
            "Conservation Areas",
            f"Loaded {result.featureCount()} conservation area(s) within the search radius.",
            level=Qgis.MessageLevel.Info, duration=6)
        self.iface.mapCanvas().refresh()

    def _run_desk_study_async(self, busy_text, fetch, on_success, error_title):
        """Run a desk-study network search on a background QgsTask so the QGIS UI
        stays responsive, showing a cancelable progress dialog.

        `fetch` is a no-argument callable executed ON A WORKER THREAD: it must do
        only network I/O plus pure-Python / QgsGeometry work and MUST NOT touch
        QgsProject, map layers or any UI. Capture the redline geometry / centroid
        on the main thread and close over it (see the callers). It returns the
        result rows.

        `on_success(rows)` runs back on the MAIN thread (via the task's finished
        signal) to build the results dialog. An exception raised inside `fetch`
        surfaces in a message box titled `error_title`; Cancel abandons the
        result. This replaces the old synchronous WaitCursor-blocked fetch that
        froze QGIS for the duration of the query.
        """
        from qgis.core import QgsTask, QgsApplication
        from qgis.PyQt.QtWidgets import QProgressDialog
        from qgis.PyQt.QtCore import Qt as _Qt

        dlg = QProgressDialog(busy_text, "Cancel", 0, 0, self)
        dlg.setWindowTitle("Desk Study")
        dlg.setWindowModality(_Qt.WindowModality.WindowModal)
        dlg.setMinimumDuration(0)
        dlg.setAutoClose(False)
        dlg.setAutoReset(False)

        state = {"canceled": False, "task": None, "finished": False}

        import threading as _thr
        _dbg("ASYNC[%s]: helper entered, main tid=%s" % (busy_text, _thr.get_ident()))

        def work(task):
            _dbg("ASYNC[%s]: work() START on tid=%s" % (busy_text, _thr.get_ident()))
            try:
                res = fetch()
                _dbg("ASYNC[%s]: work() END, n=%s" % (busy_text, len(res) if res is not None else "None"))
                return res
            except Exception as _exc:
                _dbg("ASYNC[%s]: work() EXC %r" % (busy_text, _exc))
                raise

        def done(exception, rows=None):
            # rows MUST default: QgsTask.fromFunction's wrapper calls
            # on_finished(exception) with NO result argument whenever the
            # returned value is falsy (empty result set) or an exception was
            # raised — a two-arg signature dies with a swallowed TypeError and
            # the progress dialog never closes.
            # Read the user's cancel intent BEFORE closing the dialog: closing a
            # QProgressDialog programmatically re-emits canceled(), which would
            # otherwise flip our flag and wrongly suppress the results table.
            was_canceled = state["canceled"]
            state["finished"] = True
            _dbg("ASYNC[%s]: done() on tid=%s exc=%r canceled=%s n=%s" % (
                busy_text, _thr.get_ident(), exception, was_canceled,
                len(rows) if rows is not None else "None"))
            try:
                dlg.close()
            except Exception:
                pass
            try:
                self._ds_tasks.discard(state["task"])
            except Exception:
                pass
            if was_canceled:
                return
            if exception is not None:
                QMessageBox.critical(self, error_title, str(exception))
                return
            try:
                on_success(rows if rows is not None else [])
            except Exception as _exc:
                _dbg("ASYNC[%s]: on_success EXC %r" % (busy_text, _exc))
                raise

        task = QgsTask.fromFunction(busy_text, work, on_finished=done)
        state["task"] = task

        def _on_cancel():
            # Ignore the canceled() that fires when we close the dialog ourselves
            # on completion — only a click while the task is still running counts.
            if state["finished"]:
                return
            state["canceled"] = True
            try:
                task.cancel()
            except Exception:
                pass
        dlg.canceled.connect(_on_cancel)

        if not hasattr(self, "_ds_tasks"):
            self._ds_tasks = set()
        self._ds_tasks.add(task)
        dlg.show()
        # Force an immediate paint so a fast-returning task can't beat the
        # dialog to the screen (it would otherwise just flash imperceptibly).
        try:
            from qgis.PyQt.QtWidgets import QApplication as _QApp
            _QApp.processEvents()
        except Exception:
            pass
        _dbg("ASYNC[%s]: dialog shown visible=%s; adding task" % (busy_text, dlg.isVisible()))
        QgsApplication.taskManager().addTask(task)

    def _run_ca_report(self):
        """Fetch conservation areas within the search radius and show the report dialog."""
        if not hasattr(self, "_centroid_bng"):
            QMessageBox.warning(self, "No Redline",
                "Please draw a redline boundary first.")
            return
        e, n = self._centroid_bng
        from qgis.PyQt.QtWidgets import QApplication
        from qgis.PyQt.QtCore import Qt as _Qt
        QApplication.setOverrideCursor(_Qt.CursorShape.WaitCursor)
        try:
            rows = query_conservation_areas_2km(e, n, self._get_redline_geom())
        except Exception as exc:
            QApplication.restoreOverrideCursor()
            QMessageBox.critical(self, "Conservation Areas Query Failed", str(exc))
            return
        QApplication.restoreOverrideCursor()
        self._show_ca_report_dialog(rows)

    def _show_ca_report_dialog(self, rows):
        """Display conservation areas in a styled, sortable table with copy/export."""
        from qgis.PyQt.QtWidgets import (
            QDialog, QVBoxLayout, QHBoxLayout, QTableWidget,
            QTableWidgetItem, QHeaderView, QPushButton, QLabel,
            QAbstractItemView,
        )
        from qgis.PyQt.QtCore import Qt

        try:
            from .gridref_utils import osgb_en_to_gridref
            e, n = self._centroid_bng
            site_gr = osgb_en_to_gridref(int(e), int(n), 6)
        except Exception:
            site_gr = ""

        # Pre-compute a distance string and sort nearest-first
        for r in rows:
            d = f"{r['distance_m'] / 1000:.2f} km"
            r["dist_dir"] = f"{d} {r['direction']}".strip() if r.get("direction") else d
        rows.sort(key=lambda r: r["distance_m"])

        dlg = QDialog(self.iface.mainWindow())
        _rad = self._radius_label("conservation_areas")
        dlg.setWindowTitle(f"Conservation Areas within {_rad}")
        dlg.setMinimumSize(860, 480)
        dlg.resize(960, 540)
        layout = QVBoxLayout(dlg)
        layout.setSpacing(8)

        hdr = QLabel(
            f"<b>Conservation Areas within {_rad} of site boundary</b>"
            + (f" ({site_gr})" if site_gr else "") +
            "<br><span style='font-size:10px;color:#717E75;'>"
            "Trees within a conservation area are protected: 6 weeks' written "
            "notice to the LPA is required before works.<br>"
            "Source: planning.data.gov.uk (MHCLG), Open Government Licence v3.0</span>"
        )
        hdr.setStyleSheet("font-size:12px; color:#1A2420; padding-bottom:4px;")
        layout.addLayout(self._report_header_row(hdr))

        cols = ["Conservation Area", "Distance to Redline", "Reference",
                "Area (ha)", "Entry Date", "Grid Ref"]
        DIST_COL = 1
        NAME_COL = 0

        if not rows:
            none_lbl = QLabel(f"No conservation areas found within {_rad} of the site boundary.")
            none_lbl.setStyleSheet(
                "font-size:12px; color:#7A5200; padding:12px;"
                "background:#FBF3E3; border-radius:6px; border:1px solid #E6CF95;"
            )
            layout.addWidget(none_lbl)
        else:
            summary = QLabel(f"Found <b>{len(rows)}</b> conservation area(s) within {_rad}")
            summary.setStyleSheet("font-size:11px; color:#54635A; padding-bottom:4px;")
            layout.addWidget(summary)

            tbl = QTableWidget(len(rows), len(cols))
            tbl.setHorizontalHeaderLabels(cols)
            tbl.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
            self._setup_copyable_table(tbl)

            for r_idx, r in enumerate(rows):
                area = r.get("area_ha")
                vals = [
                    r.get("name", ""),
                    r["dist_dir"],
                    r.get("reference", ""),
                    (f"{area:.2f}" if area else ""),
                    r.get("entry_date", ""),
                    r.get("grid_ref", ""),
                ]
                for c_idx, val in enumerate(vals):
                    if c_idx == DIST_COL:
                        item = _NumericSortItem(str(val) if val else "-")
                        item.setData(Qt.ItemDataRole.UserRole, r["distance_m"])
                    else:
                        item = QTableWidgetItem(str(val) if val else "-")
                    item.setTextAlignment(
                        Qt.AlignmentFlag.AlignVCenter | (Qt.AlignmentFlag.AlignCenter if c_idx == DIST_COL else Qt.AlignmentFlag.AlignLeft)
                    )
                    tbl.setItem(r_idx, c_idx, item)

            tbl.horizontalHeader().setSectionResizeMode(NAME_COL, QHeaderView.ResizeMode.Stretch)
            for ci in [i for i in range(len(cols)) if i != NAME_COL]:
                tbl.horizontalHeader().setSectionResizeMode(ci, QHeaderView.ResizeMode.ResizeToContents)
            tbl.verticalHeader().setVisible(False)
            tbl.setSortingEnabled(True)
            tbl.sortItems(DIST_COL)
            layout.addWidget(tbl)

        # Button row
        btn_row = QHBoxLayout()
        if rows:
            copy_btn = QPushButton("\U0001f4cb  Copy")
            copy_btn.setStyleSheet("""
                QPushButton {
                    background:#F4F6F2; color:#1A2420;
                    border:1px solid #B7BFB9; border-radius:4px;
                    font-size:11px; padding:4px 12px;
                }
                QPushButton:hover { background:#E3E7E1; }
            """)
            def _copy_csv():
                from .clipboard_utils import set_table_clipboard
                grid = [list(cols)]
                for r in rows:
                    area = r.get("area_ha")
                    grid.append([
                        r.get("name", ""), r["dist_dir"],
                        r.get("reference", ""),
                        (f"{area:.2f}" if area else ""),
                        r.get("entry_date", ""), r.get("grid_ref", ""),
                    ])
                set_table_clipboard(grid, header_rows=1)
            copy_btn.clicked.connect(_copy_csv)
            btn_row.addWidget(copy_btn)

            excel_btn = QPushButton("\U0001f4ca  Export to Excel")
            excel_btn.setStyleSheet("""
                QPushButton {
                    background:#F4F6F2; color:#7A5200;
                    border:1px solid #9A6700; border-radius:4px;
                    font-size:11px; padding:4px 12px;
                }
                QPushButton:hover { background:#FBF3E3; }
            """)
            _ca_excel_cols = [
                ("name",       "Conservation Area"),
                ("dist_dir",   "Distance to Redline"),
                ("reference",  "Reference"),
                ("area_ha",    "Area (ha)"),
                ("entry_date", "Entry Date"),
                ("grid_ref",   "Grid Ref"),
            ]
            def _excel_export():
                self._export_report_excel(
                    rows, _ca_excel_cols,
                    f"Conservation Areas within {_rad}", "Conservation_Areas_Report"
                )
            excel_btn.clicked.connect(_excel_export)
            btn_row.addWidget(excel_btn)

        btn_row.addStretch()
        close_btn = QPushButton("Close")
        close_btn.setStyleSheet("""
            QPushButton {
                background:#175E3B; color:white; border:none;
                border-radius:4px; font-size:11px; padding:4px 16px;
            }
            QPushButton:hover { background:#1F7A4D; }
        """)
        close_btn.clicked.connect(dlg.accept)
        btn_row.addWidget(close_btn)
        layout.addLayout(btn_row)
        dlg.setModal(False)
        dlg.show()

    def _show_tpo_no_results_dialog(self):
        """Explain a national-data miss and expose every useful local check.

        A blank planning.data.gov.uk response is not evidence that no TPO
        exists.  Keep the council map/register, independent postcode map,
        downloadable data and council website visible as separate routes rather
        than hiding all but whichever one happens to be the primary button.
        """
        from qgis.PyQt.QtWidgets import (QDialog, QVBoxLayout, QLabel, QPushButton,
                                     QHBoxLayout, QApplication)
        from qgis.PyQt.QtCore import Qt, QUrl
        from qgis.PyQt.QtGui import QDesktopServices

        if _region.is_scotland():
            # The council deep-link machinery is England-only; for Scotland
            # point at the national dataset and remind that not every council
            # contributes to it.
            from .designated_sites import SPATIALHUB_TPO_URL
            box = QMessageBox(self)
            box.setIcon(QMessageBox.Icon.Information)
            box.setWindowTitle("No National TPO Records Found")
            box.setTextFormat(Qt.TextFormat.RichText)
            box.setText(
                "<b>No Tree Preservation Orders were found</b> within the search "
                "radius in the Spatial Hub Scotland national dataset.<br><br>"
                "Not every Scottish council has digitised its TPOs into the "
                "national dataset, so this does <i>not</i> confirm the site is "
                "TPO-free — check the council's own register too.<br><br>"
                f"<a href='{SPATIALHUB_TPO_URL}'>Tree Preservation Orders - "
                "Scotland (Spatial Hub)</a>")
            box.setStandardButtons(QMessageBox.StandardButton.Ok)
            box.exec()
            return

        # Frame on the redline so we can deep-link straight into the council's
        # interactive map, and resolve which council this falls within
        # (authoritative planning.data.gov.uk point-in-polygon, postcodes.io
        # fallback).
        bbox, centroid = self._redline_bbox_bng()
        lpa_name = lpa_ref = None
        if centroid:
            e, n = centroid
            lpa_name, lpa_ref = get_lpa_for_bng(e, n)
            # Copy E/N — handy for searching the council's TPO map.
            QApplication.clipboard().setText(f"{int(e)}, {int(n)}")

        if bbox and centroid:
            portal_url, kind, council_display, _conf = build_tpo_deeplink(
                lpa_name, bbox, centroid)
        else:
            # No redline geometry — fall back to the landing-page resolution.
            council_display, portal_url, is_direct = get_tpo_council_portal(
                lpa_name, lpa_ref)
            kind = "portal" if is_direct else ("search" if lpa_name else "finder")

        # Some councils with no public map still publish a downloadable TPO
        # dataset (open data / FeatureServer / CSV) the user can import to QGIS.
        data_url = None
        try:
            data_url = tpo_data_download(lpa_name)
        except Exception:
            data_url = None

        # tposearch.co.uk aggregates ~179 councils' TPO data into one interactive
        # map searchable by postcode. Offer a site-zoomed link, but only where it
        # actually holds data for this council (so it can't show a misleadingly
        # empty result). Needs a nearby postcode from the redline centroid.
        tpo_search_url = None
        postcode = None
        try:
            if centroid:
                postcode, _ad = self._nearest_postcode_for_bng(e, n)
                tpo_search_url = tposearch_link(postcode, lpa_name)
        except Exception:
            tpo_search_url = None

        # The council homepage is a useful backstop where a TPO page has moved
        # or an authority publishes its register under an unexpected title.
        council_home_url = None
        try:
            council_home_url = lpa_homepage(lpa_name)
        except Exception:
            council_home_url = None

        intro = (
            "<b>No Tree Preservation Order records were returned by the national "
            "planning.data.gov.uk dataset</b> within the search radius.<br><br>"
        )
        if kind == "map":
            body = (
                intro +
                f"Not all councils publish TPO data nationally, so this does "
                f"<i>not</i> confirm the site is TPO-free. Opening "
                f"<b>{council_display}</b>'s interactive map zoomed to your site "
                f"— switch on the TPO / trees layer if it isn't already:"
            )
            link_text  = f"Open {council_display}'s map at your site"
            open_label = "Open Council Map at Site"
        elif kind == "portal":
            body = (
                intro +
                f"Not all councils publish TPO data nationally, so this does "
                f"<i>not</i> confirm the site is TPO-free. Check "
                f"<b>{council_display}</b>'s own TPO register directly:"
            )
            link_text  = "Open the council's TPO register"
            open_label = "Open Council TPO Page"
        elif kind == "search":
            body = (
                intro +
                f"No direct TPO link is on file for <b>{council_display}</b>. "
                f"This search will take you to their TPO register "
                f"<i>(site easting/northing copied to clipboard)</i>:"
            )
            link_text  = f"Search for {council_display}'s TPO register"
            open_label = "Search Council TPO Page"
        else:  # finder
            body = (
                intro +
                "The local planning authority for this location could not be "
                "determined automatically. Use the Planning Portal to find it "
                "<i>(site easting/northing copied to clipboard)</i>:"
            )
            link_text  = "Find your local planning authority"
            open_label = "Open LPA Finder"

        # Choose the best first action, but keep every distinct fallback visible
        # below it.  Previously the chosen primary route caused other useful
        # checks to disappear, which made gaps in the direct-map table look like
        # gaps in the entire linker.
        if kind == "map":
            primary_url, primary_label = portal_url, open_label
        elif tpo_search_url:
            primary_url, primary_label = tpo_search_url, "Open TPO Map at Site"
        else:
            primary_url, primary_label = portal_url, open_label

        dlg = QDialog(self)
        dlg.setWindowTitle("No National TPO Records Found")
        dlg.setMinimumWidth(500)
        layout = QVBoxLayout(dlg)
        layout.setSpacing(12)

        icon_label = QLabel()
        icon_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        icon_label.setPixmap(_svg_pixmap("tree", "#1F7A4D", 30))
        icon_label.setStyleSheet("background:transparent;")
        layout.addWidget(icon_label)

        msg = QLabel(body)
        msg.setWordWrap(True)
        msg.setTextFormat(Qt.TextFormat.RichText)
        layout.addWidget(msg)

        # Always retain the council route when it is not already the primary
        # action.  This is the authoritative check even when tposearch is the
        # quickest site-centred map.
        if portal_url and portal_url != primary_url:
            link_label = QLabel(f'<a href="{portal_url}">{link_text}</a>')
            link_label.setWordWrap(True)
            link_label.setTextFormat(Qt.TextFormat.RichText)
            link_label.setOpenExternalLinks(False)
            link_label.linkActivated.connect(
                lambda url: QDesktopServices.openUrl(QUrl(url))
            )
            link_label.setStyleSheet("color: #2F5D8A; font-size: 12px;")
            layout.addWidget(link_label)

        if data_url:
            dl_label = QLabel(
                '<a href="%s">Download this council\'s TPO dataset</a> '
                "&mdash; open data you can import into QGIS as a layer." % data_url
            )
            dl_label.setWordWrap(True)
            dl_label.setTextFormat(Qt.TextFormat.RichText)
            dl_label.setOpenExternalLinks(False)
            dl_label.linkActivated.connect(
                lambda url: QDesktopServices.openUrl(QUrl(url))
            )
            dl_label.setStyleSheet("color:#175E3B; font-size:12px;")
            layout.addWidget(dl_label)

        if tpo_search_url and tpo_search_url != primary_url:
            ts_label = QLabel(
                '<a href="%s">Check protected trees at this postcode on '
                "tposearch.co.uk</a> &mdash; an independent map that aggregates "
                "many councils' TPO data, opened at your site." % tpo_search_url
            )
            ts_label.setWordWrap(True)
            ts_label.setTextFormat(Qt.TextFormat.RichText)
            ts_label.setOpenExternalLinks(False)
            ts_label.linkActivated.connect(
                lambda url: QDesktopServices.openUrl(QUrl(url))
            )
            ts_label.setStyleSheet("color:#2F5D8A; font-size:12px;")
            layout.addWidget(ts_label)

        if (council_home_url and council_home_url != primary_url
                and council_home_url != portal_url):
            home_label = QLabel(
                '<a href="%s">Open %s\'s main website</a> &mdash; use its site '
                "search if the published TPO page or map has moved."
                % (council_home_url, council_display)
            )
            home_label.setWordWrap(True)
            home_label.setTextFormat(Qt.TextFormat.RichText)
            home_label.setOpenExternalLinks(False)
            home_label.linkActivated.connect(
                lambda url: QDesktopServices.openUrl(QUrl(url))
            )
            home_label.setStyleSheet("color:#2F5D8A; font-size:12px;")
            layout.addWidget(home_label)

        if centroid:
            location_text = f"Site location copied: {int(e)}, {int(n)}"
            if postcode:
                location_text += f"  |  {postcode}"
            location_label = QLabel(location_text)
            location_label.setWordWrap(True)
            location_label.setStyleSheet("color:#54635A; font-size:11px;")
            layout.addWidget(location_label)

        btn_row = QHBoxLayout()
        open_btn = QPushButton(primary_label)
        open_btn.setStyleSheet(
            "background:#2F5D8A; color:white; padding:6px 14px; "
            "border-radius:4px; font-weight:500;"
        )
        open_btn.clicked.connect(
            lambda: QDesktopServices.openUrl(QUrl(primary_url))
        )
        btn_row.addWidget(open_btn)
        if data_url:
            dl_btn = QPushButton("Download TPO Data")
            dl_btn.setStyleSheet(
                "background:#175E3B; color:white; padding:6px 14px; "
                "border-radius:4px; font-weight:500;"
            )
            dl_btn.clicked.connect(
                lambda: QDesktopServices.openUrl(QUrl(data_url))
            )
            btn_row.addWidget(dl_btn)
        close_btn = QPushButton("Close")
        close_btn.clicked.connect(dlg.accept)
        btn_row.addStretch()
        btn_row.addWidget(close_btn)
        layout.addLayout(btn_row)

        dlg.exec()

    def _on_tpo_lines_toggle(self, checked):
        """Show or hide the TPO distance lines layer."""
        from qgis.core import QgsProject
        lyrs = QgsProject.instance().mapLayersByName("TPO Distance Lines")
        for lyr in lyrs:
            node = QgsProject.instance().layerTreeRoot().findLayer(lyr.id())
            if node is not None:
                node.setItemVisibilityChecked(checked)
        self._tpo_lines_btn.setText("〰  Distance lines: On" if checked else "〰  Distance lines: Off")
        self.iface.mapCanvas().refresh()

    def _on_tpo_label_cycle(self):
        """Cycle TPO label: reference → name → address → off → reference …"""
        _cycle  = {"reference": "name", "name": "address", "address": "off", "off": "reference"}
        _labels = {
            "reference": "🏷  Label:\nReference",
            "name":      "🏷  Label:\nName",
            "address":   "🏷  Label:\nAddress",
            "off":       "🏷  Label:\nOff",
        }
        _fields = {
            "reference": "reference",
            "name":      "name",
            "address":   "address",
        }
        self._tpo_label_mode = _cycle.get(self._tpo_label_mode, "reference")
        self._tpo_label_btn.setText(_labels[self._tpo_label_mode])
        from qgis.core import QgsProject, QgsPalLayerSettings, QgsVectorLayerSimpleLabeling
        lyrs = QgsProject.instance().mapLayersByName("Tree Preservation Orders")
        if not lyrs:
            return
        lyr = lyrs[0]
        if self._tpo_label_mode == "off":
            lyr.setLabelsEnabled(False)
        else:
            pal = QgsPalLayerSettings()
            pal.enabled      = True
            pal.isExpression = False
            pal.fieldName    = _fields[self._tpo_label_mode]
            from qgis.core import QgsTextFormat, QgsTextBufferSettings, QgsVectorLayerSimpleLabeling
            from qgis.PyQt.QtGui import QColor
            fmt = QgsTextFormat()
            fmt.setSize(7)
            fmt.setColor(QColor("#175E3B"))
            buf = QgsTextBufferSettings()
            buf.setEnabled(True); buf.setSize(0.8); buf.setColor(QColor("white"))
            fmt.setBuffer(buf)
            pal.setFormat(fmt)
            lyr.setLabeling(QgsVectorLayerSimpleLabeling(pal))
            lyr.setLabelsEnabled(True)
        lyr.triggerRepaint()
        self.iface.mapCanvas().refresh()

    # -----------------------------------------------------------------------
    # Local Wildlife Sites handlers
    # -----------------------------------------------------------------------

    def _on_lws_add(self):
        """Open a spreadsheet-style table dialog to enter Local Wildlife Sites."""
        from qgis.PyQt.QtWidgets import (
            QDialog, QVBoxLayout, QHBoxLayout,
            QPushButton, QLabel, QTableWidget, QTableWidgetItem,
            QHeaderView, QAbstractItemView, QApplication,
        )
        from qgis.PyQt.QtGui import QColor, QKeySequence
        from qgis.PyQt.QtCore import Qt, QEvent

        INIT_ROWS = 10   # blank rows shown on open

        dlg = QDialog(self)
        dlg.setWindowTitle("Local Wildlife Sites — Enter Data")
        dlg.setMinimumSize(680, 420)
        dlg.resize(760, 480)
        layout = QVBoxLayout(dlg)
        layout.setSpacing(8)

        hdr = QLabel(
            "Enter or paste site data directly into the table below.  "
            "Rows with no Site Name or Grid Reference will be ignored.\n"
            "You can paste rows copied from Excel or a spreadsheet (Tab-separated).  "
            "Right-click any row to insert or delete rows."
        )
        hdr.setWordWrap(True)
        hdr.setStyleSheet(
            "font-size:11px; color:#004d40; background:#e0f2f1;"
            " border:1px solid #80cbc4; border-radius:6px; padding:8px 10px;"
        )
        layout.addWidget(hdr)

        # ── Table ─────────────────────────────────────────────────────────
        COLS = ["Site Name", "Grid Reference", "Area (ha)"]
        tbl  = QTableWidget(INIT_ROWS, len(COLS))
        tbl.setHorizontalHeaderLabels(COLS)
        tbl.setStyleSheet("""
            QTableWidget {
                font-size:11px; gridline-color:#b2dfdb;
                border:1px solid #80cbc4; border-radius:4px;
            }
            QTableWidget::item { padding:3px 6px; }
            QHeaderView::section {
                background:#00897b; color:white; font-weight:600;
                font-size:11px; padding:5px 6px;
                border:none; border-right:1px solid #80cbc4;
            }
            QTableWidget::item:selected { background:#b2dfdb; color:#004d40; }
        """)
        tbl.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeMode.Stretch)
        tbl.horizontalHeader().setSectionResizeMode(1, QHeaderView.ResizeMode.ResizeToContents)
        tbl.horizontalHeader().setSectionResizeMode(2, QHeaderView.ResizeMode.ResizeToContents)
        tbl.verticalHeader().setDefaultSectionSize(24)
        tbl.verticalHeader().setStyleSheet(
            "QHeaderView::section { background:#f0faf8; color:#004d40;"
            " font-size:10px; border:none; border-bottom:1px solid #b2dfdb; }"
        )
        tbl.setSelectionMode(QAbstractItemView.SelectionMode.ContiguousSelection)
        tbl.setContextMenuPolicy(Qt.ContextMenuPolicy.CustomContextMenu)

        # Placeholder text colour for empty cells
        def _set_placeholder(row):
            placeholders = ["e.g. Runnymede Meadows LWS", "e.g. TQ 443 806", "0.00"]
            for c, ph in enumerate(placeholders):
                if tbl.item(row, c) is None:
                    item = QTableWidgetItem("")
                    item.setData(Qt.ItemDataRole.UserRole, ph)  # store placeholder text
                    tbl.setItem(row, c, item)

        for r in range(INIT_ROWS):
            _set_placeholder(r)

        # Paint placeholder text grey when cell is empty and not being edited
        class _Delegate(type(tbl.itemDelegate())):
            def paint(self, painter, option, index):
                item = tbl.item(index.row(), index.column())
                # PyQt6 flags don't support "& int" — go via .value there.
                state_val = (option.state.value if hasattr(option.state, "value")
                             else int(option.state))
                if item and not item.text() and not (state_val & 0x20):
                    from qgis.PyQt.QtWidgets import QStyle
                    from qgis.PyQt import QtGui as _G
                    painter.save()
                    painter.setPen(_G.QColor("#A6AFA9"))
                    ph = item.data(Qt.ItemDataRole.UserRole) or ""
                    painter.drawText(option.rect.adjusted(6, 0, -6, 0),
                                     Qt.AlignmentFlag.AlignVCenter | Qt.AlignmentFlag.AlignLeft, ph)
                    painter.restore()
                else:
                    super().paint(painter, option, index)

        tbl.setItemDelegate(_Delegate(tbl))

        # ── Paste support (Ctrl+V) ────────────────────────────────────────
        def _handle_paste():
            clipboard = QApplication.clipboard().text()
            if not clipboard:
                return
            paste_rows = [r.split("\t") for r in clipboard.splitlines() if r.strip()]
            if not paste_rows:
                return

            sel = tbl.selectedRanges()
            start_row = sel[0].topRow() if sel else tbl.currentRow()
            start_col = sel[0].leftColumn() if sel else 0
            if start_row < 0:
                start_row = 0

            # Expand table if needed
            needed = start_row + len(paste_rows)
            if needed > tbl.rowCount():
                tbl.setRowCount(needed)
                for r in range(tbl.rowCount() - len(paste_rows), tbl.rowCount()):
                    _set_placeholder(r)

            for r_off, cols in enumerate(paste_rows):
                for c_off, val in enumerate(cols):
                    col = start_col + c_off
                    if col >= tbl.columnCount():
                        break
                    row = start_row + r_off
                    existing = tbl.item(row, col)
                    if existing:
                        existing.setText(val.strip())
                    else:
                        tbl.setItem(row, col, QTableWidgetItem(val.strip()))

        # Intercept Ctrl+V at widget level. NOTE: must be an event filter —
        # reassigning tbl.__class__ to a runtime subclass corrupts the sip
        # wrapper under PyQt6 (access violation on the next attribute lookup).
        class _PasteFilter(QObject):
            def eventFilter(self_, obj, event):
                if (event.type() == QEvent.Type.KeyPress
                        and event.matches(QKeySequence.StandardKey.Paste)):
                    _handle_paste()
                    return True
                return False

        tbl._paste_filter = _PasteFilter(tbl)
        tbl.installEventFilter(tbl._paste_filter)

        # ── Right-click context menu ──────────────────────────────────────
        def _context_menu(pos):
            from qgis.PyQt.QtWidgets import QMenu
            menu   = QMenu(tbl)
            row    = tbl.rowAt(pos.y())
            act_ins_above = menu.addAction("Insert row above")
            act_ins_below = menu.addAction("Insert row below")
            menu.addSeparator()
            act_del = menu.addAction("Delete selected row(s)")
            action = menu.exec(tbl.viewport().mapToGlobal(pos))
            if action == act_ins_above and row >= 0:
                tbl.insertRow(row)
                _set_placeholder(row)
            elif action == act_ins_below:
                tbl.insertRow(row + 1)
                _set_placeholder(row + 1)
            elif action == act_del:
                sel_rows = sorted(
                    {i.row() for i in tbl.selectedItems()}, reverse=True
                )
                for r in sel_rows:
                    tbl.removeRow(r)
                if tbl.rowCount() == 0:
                    tbl.setRowCount(1)
                    _set_placeholder(0)

        tbl.customContextMenuRequested.connect(_context_menu)
        layout.addWidget(tbl)

        # ── Add-row button ────────────────────────────────────────────────
        add_row_btn = QPushButton("Add Row")
        add_row_btn.setIcon(_svg_icon("plus", "#00695c", 13))
        add_row_btn.setIconSize(QSize(13, 13))
        add_row_btn.setStyleSheet(
            "QPushButton { background:#f0faf8; color:#00695c; border:1px solid #80cbc4;"
            " border-radius:4px; font-size:10px; padding:3px 10px; }"
            "QPushButton:hover { background:#b2dfdb; }"
        )
        add_row_btn.setMaximumWidth(120)

        def _add_row():
            r = tbl.rowCount()
            tbl.setRowCount(r + 1)
            _set_placeholder(r)
            tbl.scrollToBottom()

        add_row_btn.clicked.connect(_add_row)

        err_lbl = QLabel("")
        err_lbl.setStyleSheet("font-size:10px; color:#B3372F;")
        err_lbl.setWordWrap(True)

        row_btns = QHBoxLayout()
        row_btns.addWidget(add_row_btn)
        row_btns.addStretch()
        row_btns.addWidget(err_lbl)
        layout.addLayout(row_btns)

        # ── Dialog buttons ────────────────────────────────────────────────
        btn_row = QHBoxLayout()
        ok_btn = QPushButton("Add Sites to Map")
        ok_btn.setIcon(_svg_icon("check", "white", 14))
        ok_btn.setIconSize(QSize(14, 14))
        ok_btn.setStyleSheet(
            "QPushButton { background:#00897b; color:white; border:none;"
            " border-radius:6px; font-size:11px; font-weight:600; padding:7px 18px; }"
            "QPushButton:hover { background:#00695c; }"
        )
        cancel_btn = QPushButton("Cancel")
        cancel_btn.setStyleSheet(
            "QPushButton { background:#F4F6F2; color:#1A2420; border:1px solid #B7BFB9;"
            " border-radius:6px; font-size:11px; padding:7px 14px; }"
            "QPushButton:hover { background:#E3E7E1; }"
        )
        btn_row.addStretch()
        btn_row.addWidget(cancel_btn)
        btn_row.addWidget(ok_btn)
        layout.addLayout(btn_row)

        cancel_btn.clicked.connect(dlg.reject)

        def _on_ok():
            from .designated_sites import add_lws_site
            errors   = []
            added    = 0

            for r in range(tbl.rowCount()):
                name_item = tbl.item(r, 0)
                gr_item   = tbl.item(r, 1)
                area_item = tbl.item(r, 2)

                name = name_item.text().strip() if name_item else ""
                gr   = gr_item.text().strip()   if gr_item   else ""
                area_raw = area_item.text().strip() if area_item else ""

                # Skip completely blank rows
                if not name and not gr:
                    continue

                if not name:
                    errors.append(f"Row {r+1}: missing site name")
                    continue
                if not gr:
                    errors.append(f"Row {r+1}: missing grid reference")
                    continue

                try:
                    area = float(area_raw.replace(",", "")) if area_raw else 0.0
                except ValueError:
                    area = 0.0

                ok_add, result = add_lws_site(name, gr, area)
                if not ok_add:
                    errors.append(f"Row {r+1} ({name}): {result}")
                    continue

                display = f"{name}  |  {gr.upper()}  |  {area:.2f} ha"
                self._lws_list.addItem(display)
                self._lws_feature_ids.append(result)
                added += 1

            if errors:
                err_lbl.setText("Errors: " + "  ·  ".join(errors))
                if added > 0:
                    err_lbl.setText(
                        f"{added} site(s) added.  Errors: " + "  ·  ".join(errors)
                    )
                return   # keep dialog open so user can fix errors

            if added == 0:
                err_lbl.setText("No valid rows found — please enter at least one site.")
                return

            if hasattr(self, "_centroid_bng"):
                self._lws_hide_btn.setEnabled(True)
            self.iface.mapCanvas().refresh()
            dlg.accept()

        ok_btn.clicked.connect(_on_ok)
        dlg.exec()

    def _on_lws_list_double_click(self, item):
        """Remove a site when the user double-clicks it in the list."""
        from qgis.PyQt.QtWidgets import QMessageBox
        idx = self._lws_list.row(item)
        if idx < 0 or idx >= len(self._lws_feature_ids):
            return
        reply = QMessageBox.question(
            self, "Remove Site",
            f"Remove '{item.text().split('|')[0].strip()}' from the map?",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
        )
        if reply != QMessageBox.StandardButton.Yes:
            return False
        fid = self._lws_feature_ids.pop(idx)
        self._lws_list.takeItem(idx)
        from .designated_sites import remove_lws_site, remove_lws_distance_lines
        remove_lws_site(fid)
        remove_lws_distance_lines()
        if not self._lws_feature_ids:
            self._lws_hide_btn.setEnabled(False)
            self._lws_hide_btn.setChecked(False)
            self._lws_hide_btn.setText("👁  Hide Sites")
        self.iface.mapCanvas().refresh()

    def _on_lws_import(self):
        """Import LWS polygons, optionally limited to 2 km from the redline."""
        from qgis.PyQt.QtWidgets import QFileDialog, QMessageBox
        from qgis.core import QgsProject
        path, _ = QFileDialog.getOpenFileName(
            self, "Add Local Wildlife Sites from file", "",
            "Spatial data (*.gpkg *.geojson *.json *.shp);;GeoPackage (*.gpkg);;All files (*.*)",
        )
        if not path:
            return
        choice = QMessageBox.question(
            self, "Import extent",
            "Only import Local Wildlife Sites within 2 km of the site boundary?",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No |
            QMessageBox.StandardButton.Cancel,
            QMessageBox.StandardButton.Yes,
        )
        if choice == QMessageBox.StandardButton.Cancel:
            return
        within_2km = choice == QMessageBox.StandardButton.Yes
        redline_geom = self._get_redline_geom() if within_2km else None
        if within_2km and (redline_geom is None or redline_geom.isEmpty()):
            QMessageBox.warning(self, "No site boundary",
                "Draw or load the site boundary before importing sites within 2 km.")
            return
        from .designated_sites import import_lws_file
        ok, result = import_lws_file(path, redline_geom, within_2km)
        if not ok:
            QMessageBox.warning(self, "LWS import failed", str(result))
            return
        self._lws_hide_btn.setEnabled(True)
        self._lws_hide_btn.setChecked(False)
        self._lws_hide_btn.setText("Hide sites")
        self._lws_list.clear()
        self._lws_feature_ids.clear()
        layer = result["layer"]
        for feat in layer.getFeatures():
            name = str(feat["name"] or "Unnamed site")
            ref = str(feat["site_ref"] or "")
            area = float(feat["area_ha"] or 0)
            self._lws_list.addItem(f"{name}  |  {ref}  |  {area:.2f} ha")
        QgsProject.instance().setDirty(True)
        # Convert the BNG layer extent to the canvas CRS before zooming.
        # Passing layer.extent() straight to a Web Mercator canvas sends the
        # view to the Gulf of Guinea because the numeric coordinates are then
        # interpreted in the wrong CRS.
        canvas = self.iface.mapCanvas()
        canvas_extent = canvas.mapSettings().layerExtentToOutputExtent(
            layer, layer.extent())
        canvas.setExtent(canvas_extent)
        canvas.refresh()
        suffix = (f" {result['excluded']} outside the search area were excluded."
                  if result["excluded"] else "")
        QMessageBox.information(self, "Local Wildlife Sites imported",
            f"{result['included']} site boundaries were added to the project.{suffix}\n\n"
            "They will be saved with the QGIS project.")

    def _on_lws_visibility_toggle(self, checked):
        """Show or hide the Local Wildlife Sites layer (and lines if active)."""
        from qgis.core import QgsProject
        lyr_names = ["Local Wildlife Sites", "Local Wildlife Sites (Imported)",
                     "LWS Distance Lines"]
        for layer_name in lyr_names:
            for lyr in QgsProject.instance().mapLayersByName(layer_name):
                node = QgsProject.instance().layerTreeRoot().findLayer(lyr.id())
                if node is not None:
                    node.setItemVisibilityChecked(not checked)
        self._lws_hide_btn.setText("👁  Show Sites" if checked else "👁  Hide Sites")
        self.iface.mapCanvas().refresh()

    def _on_lws_lines_toggle(self, checked):
        """Show or hide LWS distance lines."""
        from qgis.core import QgsProject
        from .designated_sites import load_lws_distance_lines, remove_lws_distance_lines
        if checked:
            redline_geom = self._get_redline_geom()
            if redline_geom is None or redline_geom.isEmpty():
                self._lws_lines_btn.setChecked(False)
                return
            load_lws_distance_lines(redline_geom)
        else:
            remove_lws_distance_lines()
        self._lws_lines_btn.setText("〰  Distance lines: On" if checked else "〰  Distance lines: Off")
        self.iface.mapCanvas().refresh()

    def _on_lws_clear(self):
        """Remove all LWS sites."""
        from qgis.PyQt.QtWidgets import QMessageBox
        if not self._lws_feature_ids:
            QMessageBox.information(self, "No Sites", "No Local Wildlife Sites have been added yet.")
            return
        reply = QMessageBox.question(
            self, "Clear All Sites",
            "Remove all Local Wildlife Sites from the map?",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
        )
        if reply != QMessageBox.StandardButton.Yes:
            return
        from .designated_sites import clear_lws_sites, remove_lws_distance_lines
        clear_lws_sites()
        remove_lws_distance_lines()
        self._lws_list.clear()
        self._lws_feature_ids.clear()
        self._lws_hide_btn.setEnabled(False)
        self._lws_hide_btn.setChecked(False)
        self._lws_hide_btn.setText("👁  Hide Sites")
        self.iface.mapCanvas().refresh()

    # ── SSSI Impact Risk Zones handlers ────────────────────────────────────

    def _on_irz_toggle(self, checked):
        """Load or hide SSSI Impact Risk Zone layers."""
        from qgis.core import QgsProject
        if not checked:
            # Hide all IRZ layers
            for name in ("SSSI IRZ — Intersecting Zones", "SSSI IRZ — WMS Overlay"):
                for lyr in QgsProject.instance().mapLayersByName(name):
                    node = QgsProject.instance().layerTreeRoot().findLayer(lyr.id())
                    if node is not None:
                        node.setItemVisibilityChecked(False)
            self._irz_btn.setText("Show SSSI impact risk zones")
            self.iface.mapCanvas().refresh()
            return

        if not hasattr(self, "_centroid_bng"):
            from qgis.PyQt.QtWidgets import QMessageBox
            QMessageBox.warning(self, "No Redline",
                "Please draw a redline boundary first.")
            self._irz_btn.setChecked(False)
            self._irz_btn.setText("Show SSSI impact risk zones")
            return

        # Re-show existing layers if already loaded
        existing = QgsProject.instance().mapLayersByName("SSSI IRZ — Intersecting Zones")
        if existing:
            for name in ("SSSI IRZ — Intersecting Zones", "SSSI IRZ — WMS Overlay"):
                for lyr in QgsProject.instance().mapLayersByName(name):
                    node = QgsProject.instance().layerTreeRoot().findLayer(lyr.id())
                    if node is not None:
                        node.setItemVisibilityChecked(True)
            self._irz_btn.setText("Hide SSSI impact risk zones")
            self.iface.mapCanvas().refresh()
            return

        from qgis.PyQt.QtWidgets import QApplication
        from qgis.PyQt.QtCore import Qt as _Qt
        QApplication.setOverrideCursor(_Qt.CursorShape.WaitCursor)
        e, n = self._centroid_bng
        redline_geom = self._get_redline_geom()
        ok, result = load_irz_layer(e, n, redline_geom)
        QApplication.restoreOverrideCursor()

        if not ok:
            from qgis.PyQt.QtWidgets import QMessageBox
            if result == "no_results":
                QMessageBox.information(self, "SSSI IRZ",
                    "No SSSI Impact Risk Zones intersect the redline boundary.\n\n"
                    "This means no terrestrial SSSI IRZ consultation triggers have been "
                    "identified for this site location.")
            else:
                QMessageBox.critical(self, "IRZ Load Failed", result)
            self._irz_btn.setChecked(False)
            self._irz_btn.setText("Show SSSI impact risk zones")
            return

        # Store records for the report
        self._irz_records = result
        # Try to get IRZURL directly from ArcGIS attributes
        primary = next((r for r in result if r.get("irzurl")), None) or result[0]
        self._irz_best_url = primary.get("irzurl") or ""

        # Always enable the button — fall back to ArcGIS Hub page centred on site
        if not self._irz_best_url:
            try:
                from qgis.core import (
                    QgsCoordinateReferenceSystem, QgsCoordinateTransform,
                    QgsPointXY, QgsProject,
                )
                bng = QgsCoordinateReferenceSystem("EPSG:27700")
                wgs = QgsCoordinateReferenceSystem("EPSG:4326")
                xf  = QgsCoordinateTransform(bng, wgs, QgsProject.instance())
                pt  = xf.transform(QgsPointXY(e, n))
                self._irz_best_url = (
                    f"https://naturalengland-defra.opendata.arcgis.com/datasets/"
                    f"Defra::sssi-impact-risk-zones-england/explore"
                    f"?location={pt.y():.6f},{pt.x():.6f},15"
                )
            except Exception:
                self._irz_best_url = (
                    "https://naturalengland-defra.opendata.arcgis.com/datasets/"
                    "Defra::sssi-impact-risk-zones-england/explore"
                )

        self._irz_url_btn.setEnabled(True)
        self._irz_btn.setText("Hide SSSI impact risk zones")
        self.iface.mapCanvas().refresh()


    def _open_irz_url(self):
        """Open the IRZURL — direct link if available from ArcGIS, else ArcGIS Hub."""
        from qgis.PyQt.QtGui import QDesktopServices
        from qgis.PyQt.QtCore import QUrl
        url = getattr(self, "_irz_best_url", "")
        if url:
            QDesktopServices.openUrl(QUrl(url))
        else:
            QDesktopServices.openUrl(QUrl(
                "https://naturalengland-defra.opendata.arcgis.com/datasets/"
                "Defra::sssi-impact-risk-zones-england/explore"
            ))

    def _run_lws_report(self):
        """Generate and display the LWS report table."""
        from .designated_sites import get_lws_sites_with_distances
        if not self._lws_feature_ids:
            from qgis.PyQt.QtWidgets import QMessageBox
            QMessageBox.information(
                self, "No Sites",
                "No Local Wildlife Sites have been entered yet.\n"
                "Use the 'Add Wildlife Site' button to enter sites manually."
            )
            return
        redline_geom = self._get_redline_geom()
        rows = get_lws_sites_with_distances(redline_geom)
        self._show_lws_report_dialog(rows)

    def _show_lws_report_dialog(self, rows):
        """Display LWS results in a styled sortable table."""
        from qgis.PyQt.QtWidgets import (
            QDialog, QVBoxLayout, QHBoxLayout, QTableWidget,
            QTableWidgetItem, QHeaderView, QPushButton, QLabel,
            QAbstractItemView,
        )
        from qgis.PyQt.QtCore import Qt

        dlg = QDialog(self.iface.mainWindow())
        dlg.setWindowTitle("Local Wildlife Sites")
        dlg.setMinimumSize(760, 420)
        dlg.resize(860, 500)
        layout = QVBoxLayout(dlg)
        layout.setSpacing(8)

        hdr = QLabel(
            "<b>Local Wildlife Sites recorded within a 2 km radius of the survey site</b>"
            "<br><span style='font-size:10px;color:#717E75;'>*Data from DEFRA MAGIC / Local Records Centre</span>"
        )
        hdr.setStyleSheet("font-size:12px; color:#1A2420; padding-bottom:4px;")
        layout.addLayout(self._report_header_row(hdr))

        if not rows:
            layout.addWidget(_banner(
                "status-check", "No Local Wildlife Sites have been entered.",
                fg="#1F7A4D", bg="#E9F4EC", border="1px solid #BFE0CC",
                font_px=12, radius=6, margins=(12, 12, 12, 12)))
        else:
            summary = QLabel(f"Found <b>{len(rows)}</b> Local Wildlife Site(s)")
            summary.setStyleSheet("font-size:11px; color:#54635A; padding-bottom:4px;")
            layout.addWidget(summary)

            # Columns match the table in the screenshot:
            # Site Name | Grid Reference | Area (ha) | Approx. Closest Distance from Site (km)
            cols = ["Site Name", "Grid Reference", "Area (ha)", "Approx. Closest Distance\nfrom Site (km)"]

            tbl = QTableWidget(len(rows), len(cols))
            tbl.setHorizontalHeaderLabels(cols)
            tbl.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
            self._setup_copyable_table(tbl)

            for r_idx, r in enumerate(rows):
                dist_m = r["distance_m"]
                if dist_m == 0:
                    dist_str = "Within/adjacent to site"
                elif dist_m < 1000:
                    dist_str = f"{dist_m} m {r['direction']}"
                else:
                    dist_str = f"{dist_m / 1000:.2f} km {r['direction']}"

                vals = [
                    r["name"],
                    r["grid_ref"],
                    f"{r['area_ha']:.2f}" if r["area_ha"] else "—",
                    dist_str,
                ]
                for c_idx, val in enumerate(vals):
                    if c_idx == 3:
                        item = _NumericSortItem(str(val) if val else "—")
                        item.setData(Qt.ItemDataRole.UserRole, r["distance_m"])
                    else:
                        item = QTableWidgetItem(str(val) if val else "—")
                    item.setTextAlignment(
                        Qt.AlignmentFlag.AlignVCenter | (Qt.AlignmentFlag.AlignCenter if c_idx in (2, 3) else Qt.AlignmentFlag.AlignLeft)
                    )
                    tbl.setItem(r_idx, c_idx, item)

            tbl.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeMode.Stretch)
            for ci in (1, 2, 3):
                tbl.horizontalHeader().setSectionResizeMode(ci, QHeaderView.ResizeMode.ResizeToContents)
            tbl.verticalHeader().setVisible(False)
            tbl.setSortingEnabled(True)
            tbl.sortItems(3, Qt.SortOrder.AscendingOrder)   # closest first
            layout.addWidget(tbl)

        # Buttons
        btn_row = QHBoxLayout()

        if rows:
            copy_btn = QPushButton("Copy")
            copy_btn.setIcon(_svg_icon("copy", "#1A2420", 14))
            copy_btn.setIconSize(QSize(14, 14))
            copy_btn.setStyleSheet("""
                QPushButton {
                    background:#F4F6F2; color:#1A2420;
                    border:1px solid #B7BFB9; border-radius:4px;
                    font-size:11px; padding:4px 12px;
                }
                QPushButton:hover { background:#E3E7E1; }
            """)
            def _copy_csv():
                from .clipboard_utils import set_table_clipboard
                header = ["Site Name", "Grid Reference", "Area (ha)",
                          "Approx. Closest Distance from Site (km)"]
                grid = [header]
                for r in rows:
                    dist_m = r["distance_m"]
                    if dist_m == 0:
                        dist_str = "Within/adjacent to site"
                    elif dist_m < 1000:
                        dist_str = f"{dist_m} m {r['direction']}"
                    else:
                        dist_str = f"{dist_m / 1000:.2f} km {r['direction']}"
                    grid.append([
                        r["name"],
                        r["grid_ref"],
                        f"{r['area_ha']:.2f}" if r["area_ha"] else "",
                        dist_str,
                    ])
                set_table_clipboard(grid, header_rows=1)
            copy_btn.clicked.connect(_copy_csv)
            btn_row.addWidget(copy_btn)

        btn_row.addStretch()
        close_btn = QPushButton("Close")
        close_btn.setStyleSheet("""
            QPushButton {
                background:#F4F6F2; color:#1A2420;
                border:1px solid #B7BFB9; border-radius:4px;
                font-size:11px; padding:4px 12px;
            }
            QPushButton:hover { background:#E3E7E1; }
        """)
        close_btn.clicked.connect(dlg.accept)
        btn_row.addWidget(close_btn)
        layout.addLayout(btn_row)
        dlg.exec()

    def _run_tpo_report(self):
        """Fetch TPO data and show the report dialog."""
        if not hasattr(self, "_centroid_bng"):
            QMessageBox.warning(self, "No Redline",
                "Please draw a redline boundary first.")
            return
        if _region.is_scotland() and not self._ensure_spatialhub_key():
            return
        e, n = self._centroid_bng
        from qgis.PyQt.QtWidgets import QApplication
        from qgis.PyQt.QtCore import Qt as _Qt
        QApplication.setOverrideCursor(_Qt.CursorShape.WaitCursor)
        try:
            redline_geom = self._get_redline_geom()
            rows  = query_tpo_2km(e, n, redline_geom)
            zones = query_tpo_zones_2km(e, n, redline_geom)
        except Exception as exc:
            QApplication.restoreOverrideCursor()
            QMessageBox.critical(self, "TPO Query Failed", str(exc))
            return
        QApplication.restoreOverrideCursor()
        # Fold woodland/area/group zone polygons into the same table, mapped to
        # the tree-row shape. Without this the points-only 'tree' dataset hides
        # all area/woodland TPOs. The "[Zone]" prefix distinguishes them.
        for z in zones:
            area = z.get("area_ha")
            desc = (z.get("name") or z.get("notes") or z.get("description")
                    or z.get("zone_type") or "Woodland / area / group TPO")
            # A recognised zone type refines a name/notes description too
            # ("1/1972" alone says nothing about what kind of order it is).
            zt = z.get("zone_type", "")
            if zt and zt not in desc:
                desc = f"{desc} — {zt}"
            if area:
                desc = f"{desc} ({area:.2f} ha)"
            rows.append({
                "reference":  z.get("reference", ""),
                "tpo_ref":    z.get("tpo_ref", ""),
                "name":       f"[Zone] {desc}",
                "species":    z.get("species", ""),
                "address":    z.get("address", ""),
                "entry_date": z.get("entry_date", ""),
                "start_date": z.get("start_date", ""),
                "grid_ref":   z.get("grid_ref", ""),
                "distance_m": z["distance_m"],
                "direction":  z.get("direction", ""),
            })
        rows.sort(key=lambda r: r["distance_m"])
        self._show_tpo_report_dialog(rows)

    def _show_tpo_report_dialog(self, rows):
        """Display TPO results in a styled sortable table with copy/export — matches EPSL style."""
        from qgis.PyQt.QtWidgets import (
            QDialog, QVBoxLayout, QHBoxLayout, QTableWidget,
            QTableWidgetItem, QHeaderView, QPushButton, QLabel,
            QAbstractItemView,
        )
        from qgis.PyQt.QtCore import Qt

        try:
            from .gridref_utils import osgb_en_to_gridref
            e, n = self._centroid_bng
            site_gr = osgb_en_to_gridref(int(e), int(n), 6)
        except Exception:
            site_gr = ""

        dlg = QDialog(self.iface.mainWindow())
        _rad = self._radius_label("tpo")
        dlg.setWindowTitle(f"Tree Preservation Orders within {_rad}")
        dlg.setMinimumSize(1000, 520)
        dlg.resize(1100, 580)
        layout = QVBoxLayout(dlg)
        layout.setSpacing(8)

        # Header
        hdr = QLabel(
            f"<b>Tree Preservation Orders within {_rad} of site boundary</b>"
            + (f" ({site_gr})" if site_gr else "") +
            "<br><span style='font-size:10px;color:#717E75;'>"
            f"Source: {_region.tpo_source_name()}</span>"
        )
        hdr.setStyleSheet("font-size:12px; color:#1A2420; padding-bottom:4px;")
        layout.addLayout(self._report_header_row(hdr))

        if not rows:
            none_lbl = QLabel("\u2705  No Tree Preservation Orders found within the search radius of the site boundary.")
            none_lbl.setStyleSheet(
                "font-size:12px; color:#1F7A4D; padding:12px;"
                "background:#E9F4EC; border-radius:6px; border:1px solid #BFE0CC;"
            )
            layout.addWidget(none_lbl)
        else:
            summary = QLabel(f"Found <b>{len(rows)}</b> TPO record(s) — individual trees and "
                             f"woodland/area/group zones — within the search radius")
            summary.setStyleSheet("font-size:11px; color:#54635A; padding-bottom:4px;")
            layout.addWidget(summary)

            cols = ["Tree Ref", "Distance to Redline", "TPO Order", "Name / Description",
                    "Species", "Address", "Entry Date", "Start Date",
                    "Grid Ref"]

            tbl = QTableWidget(len(rows), len(cols))
            tbl.setHorizontalHeaderLabels(cols)
            tbl.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
            self._setup_copyable_table(tbl)

            for r_idx, r in enumerate(rows):
                dist_m = r["distance_m"]
                dist_str = f"{dist_m / 1000:.2f} km"
                dist_dir = f"{dist_str} {r['direction']}".strip() if r.get("direction") else dist_str

                vals = [
                    r.get("reference", ""),
                    dist_dir,
                    r.get("tpo_ref", ""),
                    r.get("name", ""),
                    r.get("species", ""),
                    r.get("address", ""),
                    r.get("entry_date", ""),
                    r.get("start_date", ""),
                    r.get("grid_ref", ""),
                ]
                for c_idx, val in enumerate(vals):
                    if c_idx == 1:
                        item = _NumericSortItem(str(val) if val else "\u2014")
                        item.setData(Qt.ItemDataRole.UserRole, r["distance_m"])
                    else:
                        item = QTableWidgetItem(str(val) if val else "\u2014")
                    item.setTextAlignment(
                        Qt.AlignmentFlag.AlignVCenter | (Qt.AlignmentFlag.AlignCenter if c_idx == 1 else Qt.AlignmentFlag.AlignLeft)
                    )
                    tbl.setItem(r_idx, c_idx, item)

            # Column widths — stretch name, resize others
            tbl.horizontalHeader().setSectionResizeMode(3, QHeaderView.ResizeMode.Stretch)
            for ci in [i for i in range(len(cols)) if i != 3]:
                tbl.horizontalHeader().setSectionResizeMode(ci, QHeaderView.ResizeMode.ResizeToContents)
            tbl.verticalHeader().setVisible(False)
            tbl.setSortingEnabled(True)
            tbl.sortItems(1)   # sort by distance ascending — 0.00 km (within redline) first
            layout.addWidget(tbl)

        # Button row
        btn_row = QHBoxLayout()

        if rows:
            copy_btn = QPushButton("\U0001f4cb  Copy")
            copy_btn.setStyleSheet("""
                QPushButton {
                    background:#F4F6F2; color:#1A2420;
                    border:1px solid #B7BFB9; border-radius:4px;
                    font-size:11px; padding:4px 12px;
                }
                QPushButton:hover { background:#E3E7E1; }
            """)
            def _copy_csv():
                from .clipboard_utils import set_table_clipboard
                header = [
                    "Tree Ref", "Distance to Redline", "TPO Order", "Name / Description",
                    "Species", "Address", "Entry Date", "Start Date",
                    "Grid Ref"
                ]
                grid = [header]
                for r in rows:
                    dist_m = r["distance_m"]
                    dist_str = f"{dist_m / 1000:.2f} km"
                    dist_dir = f"{dist_str} {r['direction']}".strip() if r.get("direction") else dist_str
                    grid.append([
                        r.get("reference", ""),
                        dist_dir,
                        r.get("tpo_ref", ""),
                        r.get("name", ""),
                        r.get("species", ""),
                        r.get("address", ""),
                        r.get("entry_date", ""),
                        r.get("start_date", ""),
                        r.get("grid_ref", ""),
                    ])
                set_table_clipboard(grid, header_rows=1)
            copy_btn.clicked.connect(_copy_csv)
            btn_row.addWidget(copy_btn)

            excel_btn = QPushButton("\U0001f4ca  Export to Excel")
            excel_btn.setStyleSheet("""
                QPushButton {
                    background:#F4F6F2; color:#175E3B;
                    border:1px solid #1F7A4D; border-radius:4px;
                    font-size:11px; padding:4px 12px;
                }
                QPushButton:hover { background:#E9F4EC; }
            """)
            _tpo_excel_cols = [
                ("reference",  "Tree Ref"),
                ("dist_dir",   "Distance to Redline"),
                ("tpo_ref",    "TPO Order"),
                ("name",       "Name / Description"),
                ("species",    "Species"),
                ("address",    "Address"),
                ("entry_date", "Entry Date"),
                ("start_date", "Start Date"),
                ("grid_ref",   "Grid Ref"),
            ]
            def _excel_export():
                export_rows = []
                for r in rows:
                    dist_m = r["distance_m"]
                    dist_str = f"{dist_m / 1000:.2f} km"
                    dist_dir = f"{dist_str} {r['direction']}".strip() if r.get("direction") else dist_str
                    export_rows.append({**r, "dist_dir": dist_dir})
                self._export_report_excel(
                    export_rows, _tpo_excel_cols,
                    f"Tree Preservation Orders within {_rad}", "TPO_Report"
                )
            excel_btn.clicked.connect(_excel_export)
            btn_row.addWidget(excel_btn)

        btn_row.addStretch()
        close_btn = QPushButton("Close")
        close_btn.setStyleSheet("""
            QPushButton {
                background:#175E3B; color:white; border:none;
                border-radius:4px; font-size:11px; padding:4px 16px;
            }
            QPushButton:hover { background:#1F7A4D; }
        """)
        close_btn.clicked.connect(dlg.accept)
        btn_row.addWidget(close_btn)
        layout.addLayout(btn_row)
        dlg.setModal(False)
        dlg.show()


    # ── Pond Mapper handlers ───────────────────────────────────────────────

    def _refresh_pond_controls(self):
        """Refresh standing-water ranges, counts, and action availability."""
        from qgis.core import QgsProject
        from .config import get_search_radius, get_pond_inner_radius
        from .designated_sites import _SW_LAYER_NAME

        inner = get_pond_inner_radius()
        outer = get_search_radius("standing_water")
        range_lbl = getattr(self, "_pond_search_range_lbl", None)
        if range_lbl is not None:
            range_lbl.setText(
                f"Search buffers · {inner} m inner · {outer} m outer")

        project = QgsProject.instance()
        known_count = sum(
            layer.featureCount()
            for layer in project.mapLayersByName(_SW_LAYER_NAME))
        known_status = getattr(self, "_pond_known_status_lbl", None)
        if known_status is not None:
            requested = bool(getattr(
                self, "_standing_water_buffer_requested", False))
            if known_count:
                known_status.setText(f"{known_count} highlighted")
            elif requested:
                known_status.setText("0 highlighted")
            else:
                known_status.setText("Not searched")
        reference_btn = getattr(self, "_os_topo_btn", None)
        if reference_btn is not None:
            if known_count:
                reference_btn.setText(
                    "Ponds loaded" if reference_btn.isChecked()
                    else "Show loaded ponds")
                reference_btn.setToolTip(
                    "Saved pond outlines are loaded in this project. Click to "
                    + ("hide them." if reference_btn.isChecked()
                       else "show them without querying the API again."))
            else:
                reference_btn.setText("Highlight known ponds & lakes")
                reference_btn.setToolTip(
                    "Highlight reference mapping within the search radius; "
                    "this does not create the required water-body markers.")

        premium_enabled = bool(get_os_ngd_api_key())
        reference_hint = getattr(self, "_pond_reference_hint", None)
        if reference_hint is not None:
            reference_hint.setVisible(not premium_enabled)
        edit_btn = getattr(self, "_pond_edit_known_btn", None)
        if edit_btn is not None:
            edit_btn.setVisible(premium_enabled)
            edit_btn.setEnabled(bool(known_count))
            if not known_count and edit_btn.isChecked():
                edit_btn.blockSignals(True)
                edit_btn.setChecked(False)
                edit_btn.blockSignals(False)
            edit_btn.setText("Finish correcting outlines" if edit_btn.isChecked()
                             else "Correct pond outlines")

        placed_count = sum(
            layer.featureCount()
            for layer in project.mapLayersByName(POND_LAYER_NAME))
        placed_status = getattr(self, "_pond_placed_status_lbl", None)
        if placed_status is not None:
            placed_status.setText(f"{placed_count} marker" +
                                  ("" if placed_count == 1 else "s"))
        auto_btn = getattr(self, "_pond_auto_place_btn", None)
        if auto_btn is not None:
            auto_btn.setEnabled(bool(known_count))

        placing = bool(getattr(self, "_pond_tool", None)
                       and self.iface.mapCanvas().mapTool() is self._pond_tool)
        if getattr(self, "_pond_manage_btn", None) is not None:
            self._pond_manage_btn.setEnabled(placed_count > 0)
        if getattr(self, "_pond_lines_btn", None) is not None:
            self._pond_lines_btn.setEnabled(placed_count > 0)
        labels_btn = getattr(self, "_pond_labels_btn", None)
        if labels_btn is not None:
            mode = pond_label_mode()
            enabled = mode != "off"
            labels_btn.setEnabled(placed_count > 0)
            labels_btn.blockSignals(True)
            labels_btn.setChecked(enabled)
            labels_btn.blockSignals(False)
            if hasattr(labels_btn, "_sync_checked_style"):
                labels_btn._sync_checked_style(enabled)
            labels_btn.setText({
                "both": "Labels: Name + distance",
                "name": "Labels: Name",
                "distance": "Labels: Distance",
                "off": "Labels: Off",
            }.get(mode, "Labels: Name + distance"))
        if getattr(self, "_pond_place_btn", None) is not None:
            self._pond_place_btn.setText(
                "Finish placing markers" if placing
                else "Place manually")

    def _toggle_known_water_editing(self, checked):
        """Use QGIS's native vertex tool to correct downloaded water outlines."""
        from qgis.core import QgsProject
        from .designated_sites import _SW_LAYER_NAME

        button = getattr(self, "_pond_edit_known_btn", None)
        layers = QgsProject.instance().mapLayersByName(_SW_LAYER_NAME)
        layer = layers[0] if layers else None
        if checked:
            if layer is None or layer.featureCount() == 0:
                if button is not None:
                    button.blockSignals(True)
                    button.setChecked(False)
                    button.blockSignals(False)
                return
            if not layer.isEditable() and not layer.startEditing():
                QMessageBox.warning(
                    self, "Standing Water Bodies",
                    "The highlighted water outlines could not be opened for editing.")
                if button is not None:
                    button.blockSignals(True)
                    button.setChecked(False)
                    button.blockSignals(False)
                return
            # Snapshot the downloaded shapes so finishing only smooths features
            # the user actually changed, never every mapped pond in the layer.
            self._pond_edit_original_geometries = {
                feature.id(): bytes(feature.geometry().asWkb())
                for feature in layer.getFeatures()
                if feature.geometry() is not None
                and not feature.geometry().isEmpty()
            }
            self.iface.setActiveLayer(layer)
            try:
                action = self.iface.actionVertexToolActiveLayer()
            except (AttributeError, TypeError):
                action = self.iface.actionVertexTool()
            action.trigger()
            if button is not None:
                button.setText("Finish correcting outlines")
                button.setToolTip(
                    "Move, add, or delete vertices, then click here to finish")
            self.iface.messageBar().pushInfo(
                "Correct water outline",
                "Click the water polygon to show its vertices. Move, add, or "
                "delete vertices, then click Finish correcting outlines.")
            return

        smoothed_count = 0
        if layer is not None and layer.isEditable():
            originals = getattr(self, "_pond_edit_original_geometries", {})
            for feature in layer.getFeatures():
                geometry = feature.geometry()
                original = originals.get(feature.id())
                if (original is None or geometry is None or geometry.isEmpty()
                        or bytes(geometry.asWkb()) == original):
                    continue
                # One restrained Chaikin pass removes hand-edited kinks while
                # keeping the corrected shoreline close to the placed vertices.
                smoothed = geometry.smooth(
                    iterations=1, offset=0.12, minimumDistance=-1,
                    maxAngle=180)
                if smoothed is None or smoothed.isEmpty():
                    continue
                old_area = geometry.area()
                new_area = smoothed.area()
                if (not smoothed.isGeosValid()
                        or old_area <= 0
                        or not 0.9 <= new_area / old_area <= 1.1):
                    continue
                if layer.changeGeometry(feature.id(), smoothed):
                    smoothed_count += 1
            if not layer.commitChanges():
                QMessageBox.warning(
                    self, "Standing Water Bodies",
                    "The outline changes could not be saved. The layer remains "
                    "in edit mode so the changes are not lost.")
                if button is not None:
                    button.blockSignals(True)
                    button.setChecked(True)
                    button.blockSignals(False)
                    button.setText("Finish correcting outlines")
                return
        self._pond_edit_original_geometries = {}
        try:
            self.iface.actionPan().trigger()
        except Exception:
            pass
        if button is not None:
            button.setText("Correct pond outlines")
            button.setToolTip(
                "Move or delete vertices where the mapped OS NGD water "
                "outline does not match the current aerial imagery")
        self.iface.mapCanvas().refresh()
        if smoothed_count:
            self.iface.messageBar().pushSuccess(
                "Standing Water Bodies",
                f"Smoothed {smoothed_count} corrected pond outline" +
                ("" if smoothed_count == 1 else "s") + ".")

    def _toggle_os_topo_layer(self, checked):
        """Load or remove standing water body polygons (ponds & lakes) within
        500m of the redline EDGE — OS Zoomstack surface water + OpenStreetMap,
        plus the Habitat Map of Scotland's standing waters for a Scottish site.
        """
        from .designated_sites import (
            load_standing_water_layer, remove_standing_water_layer,
            _SW_LAYER_NAME,
        )
        from qgis.core import QgsProject

        if not checked:
            edit_btn = getattr(self, "_pond_edit_known_btn", None)
            if edit_btn is not None and edit_btn.isChecked():
                edit_btn.setChecked(False)
            if getattr(self, "_active_local_buffer_preset", None) == "water":
                self._active_local_buffer_preset = None
            self._standing_water_buffer_requested = False
            remove_standing_water_layer()
            self.iface.mapCanvas().refresh()
            self._os_topo_btn.setText("Highlight known ponds & lakes")
            self._refresh_pond_controls()
            return

        if not hasattr(self, "_centroid_bng"):
            self._standing_water_buffer_requested = False
            from qgis.PyQt.QtWidgets import QMessageBox
            QMessageBox.warning(
                self.iface.mainWindow(), "No Redline",
                "Please draw a redline boundary first."
            )
            self._os_topo_btn.setChecked(False)
            self._refresh_pond_controls()
            return

        # Re-showing an already loaded layer is still a standing-water search:
        # replace any TPO/manual preset before the early cached-layer return.
        self._active_local_buffer_preset = "water"
        self._manual_local_buffers_enabled = False

        # Re-show existing layer instantly without reloading
        existing = QgsProject.instance().mapLayersByName(_SW_LAYER_NAME)
        if existing:
            self._standing_water_buffer_requested = True
            node = QgsProject.instance().layerTreeRoot().findLayer(existing[0].id())
            if node is not None:
                node.setItemVisibilityChecked(True)
            self._os_topo_btn.setText("Highlight known ponds & lakes")
            self.iface.mapCanvas().refresh()
            self._refresh_pond_controls()
            return

        # Show the range before fetching. It remains visible even when the
        # completed search returns zero water bodies.
        self._standing_water_buffer_requested = True
        self._sync_auto_search_areas()

        self._os_topo_btn.setText("⏳  Loading…")
        self._os_topo_btn.setEnabled(False)
        self.iface.mainWindow().repaint()

        easting, northing = self._centroid_bng
        redline_geom = self._get_redline_geom()

        ok, result = load_standing_water_layer(
            easting, northing, redline_geom,
            api_key=get_os_api_key(), ngd_api_key=get_os_ngd_api_key())

        self._os_topo_btn.setEnabled(True)

        if not ok:
            from qgis.PyQt.QtWidgets import QMessageBox
            self._os_topo_btn.setChecked(False)
            self._os_topo_btn.setText("Highlight known ponds & lakes")
            if result == "no_features":
                QMessageBox.information(
                    self.iface.mainWindow(), "Standing Water Bodies",
                    "No ponds or lakes found (OS surface water, OpenStreetMap"
                    + (", Habitat Map of Scotland" if _region.is_scotland() else "")
                    + ") within the search radius of the redline boundary."
                )
            else:
                QMessageBox.warning(
                    self.iface.mainWindow(), "Standing Water Bodies",
                    f"Could not load water bodies layer:\n{result}"
                )
            self._refresh_pond_controls()
            return

        self.iface.mapCanvas().refresh()
        self._os_topo_btn.setText("Highlight known ponds & lakes")
        self._refresh_pond_controls()

    def _on_pond_place_toggle(self, checked):
        """Activate or deactivate the pond placement map tool."""
        if not checked:
            self._finish_pond_placing()
            return
        if not hasattr(self, "_centroid_bng"):
            QMessageBox.warning(self, "No Redline",
                "Please draw a redline boundary first.")
            self._pond_place_btn.setChecked(False)
            return
        self._active_local_buffer_preset = "water"
        self._manual_local_buffers_enabled = False
        self._standing_water_buffer_requested = True
        self._sync_auto_search_areas()
        # Ensure the layer exists before placing
        get_or_create_pond_layer()
        self._pond_tool = _PondPlaceTool(self.iface.mapCanvas(), self.iface, self)
        self.iface.mapCanvas().setMapTool(self._pond_tool)
        # Any other panel button pressed while armed = user moved on to another
        # activity → auto-finish first. Removed again in _pond_ui_reset.
        if getattr(self, "_pond_abort_filter", None) is None:
            self._pond_abort_filter = _PondAbortFilter(self)
        from qgis.PyQt.QtWidgets import QApplication
        QApplication.instance().installEventFilter(self._pond_abort_filter)
        self._pond_place_btn.setText("Finish placing markers")
        # Keep the bar item so teardown can pop JUST this message instead of
        # wiping every notification on the bar.
        bar = self.iface.messageBar()
        self._pond_msg_item = bar.createMessage(
            "Standing Water Bodies",
            "Click on the map to place ponds. Right-click or press the button to finish.")
        bar.pushWidget(self._pond_msg_item, Qgis.MessageLevel.Info, 0)
        self._refresh_pond_controls()

    def _on_pond_auto_place(self):
        """Place one marker per mapped water body at its nearest redline point."""
        from qgis.core import QgsGeometry, QgsPointXY, QgsProject
        from .designated_sites import _SW_LAYER_NAME

        redline_geom = self._get_redline_geom()
        if redline_geom is None or redline_geom.isEmpty():
            QMessageBox.warning(
                self, "No Redline", "Please draw a redline boundary first.")
            return
        water_layers = QgsProject.instance().mapLayersByName(_SW_LAYER_NAME)
        if not water_layers or water_layers[0].featureCount() == 0:
            QMessageBox.information(
                self, "Standing Water Bodies",
                "Highlight known ponds and lakes before placing markers automatically.")
            return

        existing_count = sum(
            layer.featureCount()
            for layer in QgsProject.instance().mapLayersByName(POND_LAYER_NAME))
        if existing_count:
            reply = QMessageBox.question(
                self, "Replace Water-body Markers",
                "Replace the existing markers with one automatically placed "
                "marker for each highlighted water body?",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.Cancel,
                QMessageBox.StandardButton.Cancel)
            if reply != QMessageBox.StandardButton.Yes:
                return

        self._finish_pond_placing()
        remove_pond_distance_lines()
        remove_pond_layer()
        placed = 0
        for feature in water_layers[0].getFeatures():
            geometry = feature.geometry()
            if geometry is None or geometry.isEmpty():
                continue
            try:
                nearest = geometry.nearestPoint(redline_geom)
                if nearest is None or nearest.isEmpty():
                    shortest = redline_geom.shortestLine(geometry)
                    points = shortest.asPolyline() if shortest else []
                    nearest = (QgsGeometry.fromPointXY(points[-1])
                               if points else None)
                if nearest is None or nearest.isEmpty():
                    continue
                point = nearest.asPoint()
                add_pond(QgsPointXY(point), redline_geom)
                placed += 1
            except Exception:
                continue

        if placed:
            load_pond_distance_lines(redline_geom)
            set_pond_lines_visible(False)
            self._pond_lines_btn.setChecked(False)
            self.iface.messageBar().pushSuccess(
                "Standing Water Bodies",
                f"{placed} marker" + ("" if placed == 1 else "s") +
                " placed at the closest points to the redline.")
        else:
            QMessageBox.warning(
                self, "Standing Water Bodies",
                "No markers could be placed from the highlighted outlines.")
        self.iface.mapCanvas().refresh()
        self._refresh_pond_controls()

    def _pond_ui_reset(self):
        """Reset the pond-placement UI (button + instruction bar). Called from
        _finish_pond_placing AND from _PondPlaceTool.deactivate(), so picking
        ANY other map tool (QGIS toolbar, another Biome tool, tab navigation)
        can't leave a stuck 'Finish Placing' button and an immortal message
        bar behind. Idempotent."""
        filt = getattr(self, "_pond_abort_filter", None)
        if filt is not None:
            # removeEventFilter is a safe no-op when not installed; keep the
            # filter object for reuse on the next arming.
            try:
                from qgis.PyQt.QtWidgets import QApplication
                QApplication.instance().removeEventFilter(filt)
            except Exception:
                pass
        btn = getattr(self, "_pond_place_btn", None)
        if btn is not None:
            btn.blockSignals(True)
            btn.setChecked(False)
            btn.blockSignals(False)
            if hasattr(btn, "_sync_checked_style"):
                btn._sync_checked_style(False)
            btn.setText("Place manually")
        item = getattr(self, "_pond_msg_item", None)
        if item is not None:
            self._pond_msg_item = None
            try:
                self.iface.messageBar().popWidget(item)
            except Exception:
                pass
        self._refresh_pond_controls()

    def _finish_pond_placing(self):
        """Exit pond placement mode and restore the pan tool."""
        canvas = self.iface.mapCanvas()
        tool = getattr(self, "_pond_tool", None)
        if tool is not None and canvas.mapTool() is tool:
            # Tool deactivation owns the UI reset; do not run it twice.
            self.iface.actionPan().trigger()
        else:
            self._pond_ui_reset()
        self.iface.mapCanvas().refresh()

    def _open_pond_marker_manager(self):
        """Show a compact manager for moving or deleting placed SWB markers."""
        from qgis.PyQt.QtCore import Qt
        from qgis.PyQt.QtWidgets import (
            QAbstractItemView, QDialog, QDialogButtonBox, QFrame,
            QComboBox, QHBoxLayout, QLabel, QListWidget, QListWidgetItem,
            QPushButton, QVBoxLayout,
        )
        from qgis.core import QgsProject

        layers = QgsProject.instance().mapLayersByName(POND_LAYER_NAME)
        if not layers or layers[0].featureCount() == 0:
            QMessageBox.information(
                self, "Manage Water-body Markers",
                "There are no water-body markers to manage.")
            return
        layer = layers[0]
        ensure_pond_fields(layer)

        dlg = QDialog(self)
        dlg.setWindowTitle("Manage Water-body Markers")
        dlg.setMinimumWidth(390)
        layout = QVBoxLayout(dlg)
        layout.setContentsMargins(14, 14, 14, 14)
        layout.setSpacing(9)

        intro = QLabel(
            "Select a marker to move or delete it. References are renumbered "
            "automatically after a deletion.")
        intro.setWordWrap(True)
        intro.setStyleSheet("color:#46564D;font-size:11px;")
        layout.addWidget(intro)

        marker_list = QListWidget()
        marker_list.setSelectionMode(
            QAbstractItemView.SelectionMode.SingleSelection)
        marker_list.setAlternatingRowColors(True)
        marker_list.setMinimumHeight(150)
        marker_list.setStyleSheet(
            "QListWidget{background:#FFFFFF;border:1px solid #DDE5DF;"
            "border-radius:7px;padding:3px;}"
            "QListWidget::item{padding:7px 8px;border-radius:5px;}"
            "QListWidget::item:selected{background:#DDECE4;color:#173D2D;}")
        layout.addWidget(marker_list)

        classification_row = QHBoxLayout()
        classification_row.setContentsMargins(0, 0, 0, 0)
        classification_row.setSpacing(8)
        classification_label = QLabel("GCN classification")
        classification_label.setStyleSheet(
            "color:#33473D;font-size:11px;font-weight:600;")
        classification_combo = QComboBox()
        classification_combo.addItem("Not classified", "")
        classification_combo.addItem("P — GCN present", "P")
        classification_combo.addItem("A — GCN absent", "A")
        classification_combo.addItem("NS — No survey", "NS")
        classification_combo.setMinimumHeight(32)
        classification_combo.setStyleSheet(
            "QComboBox{background:#FFFFFF;color:#1A2420;"
            "border:1px solid #C9D3CC;border-radius:6px;padding:5px 9px;}"
            "QComboBox:hover{border-color:#1F4A37;}"
            "QComboBox QAbstractItemView{background:#FFFFFF;color:#1A2420;"
            "selection-background-color:#DDECE4;selection-color:#173D2D;}")
        classification_row.addWidget(classification_label)
        classification_row.addWidget(classification_combo, 1)
        layout.addLayout(classification_row)

        move_btn = QPushButton(" Move selected")
        move_btn.setIcon(theme.icon("arrows-move", theme.ACCENT_DARK, 16))
        move_btn.setStyleSheet(theme.btn_secondary())
        move_btn.setMinimumHeight(34)
        delete_btn = QPushButton(" Delete selected")
        delete_btn.setIcon(theme.icon("trash", "#922B21", 16))
        delete_btn.setStyleSheet(theme.btn_secondary())
        delete_btn.setMinimumHeight(34)

        actions = QHBoxLayout()
        actions.setContentsMargins(0, 0, 0, 0)
        actions.setSpacing(6)
        actions.addWidget(move_btn, 1)
        actions.addWidget(delete_btn, 1)
        layout.addLayout(actions)

        boundary_export_btn = QPushButton(
            " Export DLL enquiry boundary shapefile")
        boundary_export_btn.setIcon(
            theme.icon("file-export", theme.ACCENT_DARK, 16))
        boundary_export_btn.setStyleSheet(theme.btn_secondary())
        boundary_export_btn.setMinimumHeight(36)
        boundary_export_btn.setToolTip(
            "Export the redline polygon with the required Site_Name and "
            "Site_Ref attributes")
        layout.addWidget(boundary_export_btn)

        export_btn = QPushButton(
            " Export DLL enquiry waterbody shapefile")
        export_btn.setIcon(theme.icon("file-export", theme.ACCENT_DARK, 16))
        export_btn.setStyleSheet(theme.btn_secondary())
        export_btn.setMinimumHeight(36)
        export_btn.setToolTip(
            "Export selected pond polygons with the required Pond_ID and "
            "GCN attributes")
        layout.addWidget(export_btn)

        divider = QFrame()
        divider.setFrameShape(QFrame.Shape.HLine)
        divider.setStyleSheet("color:#DDE5DF;")
        layout.addWidget(divider)

        footer = QHBoxLayout()
        clear_all_btn = QPushButton("Clear all markers")
        clear_all_btn.setStyleSheet(
            "QPushButton{background:transparent;color:#922B21;border:none;"
            "padding:5px 7px;text-align:left;}"
            "QPushButton:hover{background:#FCEBE8;border-radius:5px;}")
        close_box = QDialogButtonBox(QDialogButtonBox.StandardButton.Close)
        close_box.rejected.connect(dlg.reject)
        footer.addWidget(clear_all_btn)
        footer.addStretch(1)
        footer.addWidget(close_box)
        layout.addLayout(footer)

        def _reload(select_row=0):
            marker_list.clear()
            rows = sorted(
                list(layer.getFeatures()),
                key=lambda feature: str(feature["pond_id"] or ""))
            for feature in rows:
                pond_id = str(feature["pond_id"] or "Water body")
                distance = str(feature["distance_label"] or "").strip()
                status = str(feature["gcn_status"] or "").strip().upper()
                status_text = f"  [{status}]" if status in ("P", "A", "NS") else ""
                text = pond_id + status_text + (f"    {distance}" if distance else "")
                item = QListWidgetItem(text)
                item.setData(Qt.ItemDataRole.UserRole, int(feature.id()))
                item.setData(Qt.ItemDataRole.UserRole + 1, pond_id)
                item.setData(Qt.ItemDataRole.UserRole + 2, status)
                marker_list.addItem(item)
            has_rows = marker_list.count() > 0
            move_btn.setEnabled(has_rows)
            delete_btn.setEnabled(has_rows)
            clear_all_btn.setEnabled(has_rows)
            if has_rows:
                marker_list.setCurrentRow(
                    max(0, min(select_row, marker_list.count() - 1)))

        def _sync_classification():
            item = marker_list.currentItem()
            status = (str(item.data(Qt.ItemDataRole.UserRole + 2) or "")
                      if item is not None else "")
            index = classification_combo.findData(status)
            classification_combo.blockSignals(True)
            classification_combo.setCurrentIndex(max(0, index))
            classification_combo.setEnabled(item is not None)
            classification_combo.blockSignals(False)

        def _selected():
            item = marker_list.currentItem()
            if item is None:
                return None
            return (int(item.data(Qt.ItemDataRole.UserRole)),
                    str(item.data(Qt.ItemDataRole.UserRole + 1)))

        def _move():
            selected = _selected()
            if selected is None:
                return
            dlg.accept()
            self._start_pond_marker_move(layer, selected[0], selected[1])

        def _delete():
            selected = _selected()
            if selected is None:
                return
            row = marker_list.currentRow()
            if self._delete_pond_marker(layer, selected[0], selected[1]):
                _reload(row)

        def _set_classification(_index):
            selected = _selected()
            if selected is None:
                return
            field_index = layer.fields().indexOf("gcn_status")
            if field_index < 0:
                return
            was_editing = layer.isEditable()
            if not was_editing and not layer.startEditing():
                return
            status = str(classification_combo.currentData() or "")
            ok = bool(layer.changeAttributeValue(
                selected[0], field_index, status))
            if not was_editing:
                if ok:
                    ok = bool(layer.commitChanges())
                if not ok:
                    layer.rollBack()
            if not ok:
                QMessageBox.warning(
                    self, "Standing Water Bodies",
                    "The GCN classification could not be saved.")
                return
            row = marker_list.currentRow()
            apply_pond_label_mode(layer)
            self.iface.mapCanvas().refresh()
            _reload(row)

        def _clear_all():
            if self._on_pond_clear():
                dlg.accept()

        move_btn.clicked.connect(_move)
        delete_btn.clicked.connect(_delete)
        marker_list.currentItemChanged.connect(
            lambda _current, _previous: _sync_classification())
        classification_combo.currentIndexChanged.connect(_set_classification)
        marker_list.itemDoubleClicked.connect(lambda _item: _move())
        export_btn.clicked.connect(
            lambda: self._export_ne_pond_shapefile(layer, dlg))
        boundary_export_btn.clicked.connect(
            lambda: self._export_dll_boundary_shapefile(dlg))
        clear_all_btn.clicked.connect(_clear_all)
        _reload()
        dlg.exec()

    def _export_dll_boundary_shapefile(self, parent=None):
        """Export the redline with Natural England DLL enquiry attributes."""
        import os
        from qgis.PyQt.QtCore import QVariant
        from qgis.PyQt.QtWidgets import QFileDialog, QInputDialog
        from qgis.core import (
            QgsCoordinateReferenceSystem, QgsCoordinateTransform, QgsFeature,
            QgsField, QgsGeometry, QgsProject, QgsVectorFileWriter,
            QgsVectorLayer,
        )
        from .layer_manager import REDLINE_NAME_ALIASES

        project = QgsProject.instance()
        redline_layer = None
        for name in REDLINE_NAME_ALIASES:
            candidates = project.mapLayersByName(name)
            if candidates and candidates[0].featureCount():
                redline_layer = candidates[0]
                break
        if redline_layer is None:
            QMessageBox.warning(
                parent or self, "DLL Enquiry Boundary",
                "No Red Line Boundary is available to export.")
            return False

        suggested_name = project.baseName() or ""
        site_name, accepted = QInputDialog.getText(
            parent or self, "DLL Enquiry Site Name", "Site name:",
            text=suggested_name)
        site_name = str(site_name).strip()
        if not accepted:
            return False
        if not site_name:
            QMessageBox.warning(
                parent or self, "DLL Enquiry Boundary",
                "Enter the site name recorded on the DLL enquiry form.")
            return False

        bng = QgsCoordinateReferenceSystem("EPSG:27700")
        transform = None
        if redline_layer.crs() != bng:
            transform = QgsCoordinateTransform(
                redline_layer.crs(), bng, project.transformContext())
        geometries = []
        for source_feature in redline_layer.getFeatures():
            geometry = QgsGeometry(source_feature.geometry())
            if geometry.isEmpty():
                continue
            if transform is not None:
                geometry.transform(transform)
            if not geometry.isGeosValid():
                geometry = geometry.makeValid()
            if not geometry.isEmpty():
                geometries.append(geometry)
        if not geometries:
            QMessageBox.warning(
                parent or self, "DLL Enquiry Boundary",
                "The Red Line Boundary does not contain valid polygon geometry.")
            return False

        # A single multipart feature avoids internal overlaps between separate
        # redline parts while retaining every site island.
        boundary = QgsGeometry.unaryUnion(geometries)
        if boundary.isEmpty():
            QMessageBox.warning(
                parent or self, "DLL Enquiry Boundary",
                "The Red Line Boundary could not be prepared for export.")
            return False
        if not boundary.isMultipart():
            boundary.convertToMultiType()

        safe_name = "".join(
            character if character.isalnum() or character in "-_ " else "_"
            for character in site_name).strip().replace(" ", "_") or "site"
        default_dir = project.absolutePath() or os.path.expanduser("~")
        chosen_dir = QFileDialog.getExistingDirectory(
            parent or self, "Choose Where to Save the DLL Boundary Export",
            default_dir)
        if not chosen_dir:
            return False
        # Keep every mandatory shapefile component together in a dedicated
        # folder so the user can move, zip, or upload the folder as one item.
        export_name = safe_name + "_DLL_enquiry_boundary"
        folder_name, accepted = QInputDialog.getText(
            parent or self, "DLL Boundary Export Folder",
            "Folder name:", text=export_name)
        if not accepted:
            return False
        folder_name = str(folder_name).strip()
        safe_folder_name = "".join(
            "_" if character in '<>:"/\\|?*' else character
            for character in folder_name).rstrip(". ")
        if not safe_folder_name:
            QMessageBox.warning(
                parent or self, "DLL Enquiry Boundary",
                "Enter a valid name for the export folder.")
            return False
        output_folder = os.path.join(chosen_dir, safe_folder_name)
        try:
            os.makedirs(output_folder, exist_ok=True)
        except OSError as exc:
            QMessageBox.critical(
                parent or self, "DLL Enquiry Boundary",
                "The export folder could not be created.\n\n" + str(exc))
            return False
        path = os.path.join(output_folder, export_name + ".shp")

        output = QgsVectorLayer(
            "MultiPolygon?crs=EPSG:27700", "DLL Enquiry Boundary", "memory")
        output.dataProvider().addAttributes([
            QgsField("Site_Name", QVariant.String, len=254),
            QgsField("Site_Ref", QVariant.String, len=5),
        ])
        output.updateFields()
        feature = QgsFeature(output.fields())
        feature.setGeometry(boundary)
        feature["Site_Name"] = site_name
        feature["Site_Ref"] = "00000"
        output.dataProvider().addFeature(feature)
        output.updateExtents()

        options = QgsVectorFileWriter.SaveVectorOptions()
        options.driverName = "ESRI Shapefile"
        options.fileEncoding = "UTF-8"
        options.actionOnExistingFile = (
            QgsVectorFileWriter.ActionOnExistingFile.CreateOrOverwriteFile)
        error, message, _new_file, _new_layer = (
            QgsVectorFileWriter.writeAsVectorFormatV3(
                output, path, project.transformContext(), options))
        if error != QgsVectorFileWriter.WriterError.NoError:
            QMessageBox.critical(
                parent or self, "DLL Enquiry Boundary",
                "The boundary shapefile could not be created.\n\n" +
                str(message))
            return False

        stem = os.path.splitext(path)[0]
        required = [stem + extension for extension in
                    (".shp", ".shx", ".dbf", ".prj")]
        missing = [os.path.basename(item) for item in required
                   if not os.path.isfile(item)]
        if missing:
            QMessageBox.warning(
                parent or self, "DLL Enquiry Boundary",
                "The export completed, but the following required component "
                "files are missing: " + ", ".join(missing))
            return False

        self.iface.messageBar().pushSuccess(
            "DLL Enquiry Boundary",
            "Exported the boundary shapefile set into " +
            os.path.basename(output_folder) + ".")
        return True

    def _export_ne_pond_shapefile(self, marker_layer, parent=None):
        """Export a validated Natural England DLL waterbody polygon layer."""
        import os
        from qgis.PyQt.QtCore import QVariant
        from qgis.PyQt.QtWidgets import (
            QAbstractItemView, QDialog, QDialogButtonBox, QFileDialog,
            QHBoxLayout, QInputDialog, QLabel, QListWidget, QListWidgetItem,
            QPushButton, QVBoxLayout,
        )
        from qgis.core import (
            QgsCoordinateReferenceSystem, QgsCoordinateTransform, QgsFeature,
            QgsField, QgsFields, QgsGeometry, QgsProject, QgsVectorFileWriter,
            QgsVectorLayer,
        )
        from .designated_sites import _SW_LAYER_NAME

        markers = sorted(
            list(marker_layer.getFeatures()),
            key=lambda feature: str(feature["pond_id"] or ""))
        if not markers:
            QMessageBox.information(
                parent or self, "Natural England Pond Layer",
                "There are no water-body markers to export.")
            return False

        # Let the user choose the ponds for this particular submission.  The
        # choice is deliberately made from markers because the marker carries
        # the Pond_ID and GCN classification which must be written onto the
        # polygon containing it.
        choose = QDialog(parent or self)
        choose.setWindowTitle("Choose Ponds to Export")
        choose.setMinimumWidth(390)
        choose_layout = QVBoxLayout(choose)
        choose_layout.setContentsMargins(14, 14, 14, 14)
        choose_layout.setSpacing(9)
        choose_intro = QLabel(
            "Tick the ponds to include. Each selected point will tag the "
            "pond polygon containing it with its Pond_ID and GCN value.")
        choose_intro.setWordWrap(True)
        choose_layout.addWidget(choose_intro)
        choose_list = QListWidget()
        choose_list.setSelectionMode(
            QAbstractItemView.SelectionMode.NoSelection)
        choose_list.setMinimumHeight(180)
        choose_list.setStyleSheet(
            "QListWidget{background:#FFFFFF;border:1px solid #DDE5DF;"
            "border-radius:7px;padding:3px;}"
            "QListWidget::item{padding:7px 8px;}")
        for marker in markers:
            pond_id = str(marker["pond_id"] or "Water body")
            status = str(marker["gcn_status"] or "").strip().upper()
            distance = str(marker["distance_label"] or "").strip()
            details = status if status in ("P", "A", "NS") else "Not classified"
            if distance:
                details += " · " + distance
            item = QListWidgetItem(f"{pond_id}    {details}")
            item.setData(Qt.ItemDataRole.UserRole, int(marker.id()))
            item.setFlags(item.flags() | Qt.ItemFlag.ItemIsUserCheckable)
            item.setCheckState(Qt.CheckState.Checked)
            choose_list.addItem(item)
        choose_layout.addWidget(choose_list)

        selection_actions = QHBoxLayout()
        select_all = QPushButton("Select all")
        select_none = QPushButton("Select none")
        select_all.setStyleSheet(theme.btn_secondary())
        select_none.setStyleSheet(theme.btn_secondary())
        selection_actions.addWidget(select_all)
        selection_actions.addWidget(select_none)
        selection_actions.addStretch(1)
        choose_layout.addLayout(selection_actions)
        choose_buttons = QDialogButtonBox(
            QDialogButtonBox.StandardButton.Cancel |
            QDialogButtonBox.StandardButton.Ok)
        choose_buttons.button(
            QDialogButtonBox.StandardButton.Ok).setText("Export selected")
        choose_buttons.accepted.connect(choose.accept)
        choose_buttons.rejected.connect(choose.reject)
        choose_layout.addWidget(choose_buttons)

        def _set_all(check_state):
            for row in range(choose_list.count()):
                choose_list.item(row).setCheckState(check_state)

        select_all.clicked.connect(
            lambda: _set_all(Qt.CheckState.Checked))
        select_none.clicked.connect(
            lambda: _set_all(Qt.CheckState.Unchecked))
        if choose.exec() != QDialog.DialogCode.Accepted:
            return False
        selected_ids = {
            int(choose_list.item(row).data(Qt.ItemDataRole.UserRole))
            for row in range(choose_list.count())
            if choose_list.item(row).checkState() == Qt.CheckState.Checked
        }
        markers = [marker for marker in markers
                   if int(marker.id()) in selected_ids]
        if not markers:
            QMessageBox.information(
                parent or self, "Natural England Pond Layer",
                "Select at least one pond to export.")
            return False

        missing_status = [
            str(feature["pond_id"] or "Water body") for feature in markers
            if str(feature["gcn_status"] or "").strip().upper()
            not in ("P", "A", "NS")
        ]
        if missing_status:
            QMessageBox.warning(
                parent or self, "Natural England Pond Layer",
                "Assign P, A, or NS to every pond before exporting.\n\n"
                "Not classified: " + ", ".join(missing_status))
            return False

        project = QgsProject.instance()
        outline_layers = project.mapLayersByName(_SW_LAYER_NAME)
        if not outline_layers or outline_layers[0].featureCount() == 0:
            QMessageBox.warning(
                parent or self, "Natural England Pond Layer",
                "No mapped pond polygons are available. Load or correct the "
                "pond outlines before exporting; Natural England requires "
                "polygon geometry rather than marker points.")
            return False
        outline_layer = outline_layers[0]
        bng = QgsCoordinateReferenceSystem("EPSG:27700")

        def _bng_geometry(feature, source_layer):
            geometry = QgsGeometry(feature.geometry())
            if source_layer.crs() != bng:
                geometry.transform(QgsCoordinateTransform(
                    source_layer.crs(), bng, project.transformContext()))
            return geometry

        outlines = []
        for feature in outline_layer.getFeatures():
            geometry = _bng_geometry(feature, outline_layer)
            if geometry is not None and not geometry.isEmpty():
                outlines.append((feature.id(), geometry))

        # The point determines which polygon receives its Pond_ID and GCN
        # values.  Distance zero covers points inside or precisely on the pond
        # boundary and prevents a nearby, unrelated pond being guessed.
        matches = []
        used_outline_ids = set()
        unmatched = []
        for marker in markers:
            marker_geometry = _bng_geometry(marker, marker_layer)
            candidates = [
                (marker_geometry.distance(geometry), fid, geometry)
                for fid, geometry in outlines if fid not in used_outline_ids
            ]
            if not candidates:
                unmatched.append(str(marker["pond_id"] or "Water body"))
                continue
            distance, outline_fid, outline_geometry = min(
                candidates, key=lambda item: item[0])
            if distance > 0.001:
                unmatched.append(str(marker["pond_id"] or "Water body"))
                continue
            used_outline_ids.add(outline_fid)
            matches.append((marker, QgsGeometry(outline_geometry)))

        if unmatched:
            QMessageBox.warning(
                parent or self, "Natural England Pond Layer",
                "Each marker must have a matching mapped polygon. Correct the "
                "outlines or move the following markers inside their ponds:\n\n"
                + ", ".join(unmatched))
            return False

        default_dir = project.absolutePath() or os.path.expanduser("~")
        project_name = project.baseName() or "site"
        safe_project_name = "".join(
            character if character.isalnum() or character in "-_ " else "_"
            for character in project_name).strip().replace(" ", "_") or "site"
        chosen_dir = QFileDialog.getExistingDirectory(
            parent or self,
            "Choose Where to Save the DLL Waterbody Export", default_dir)
        if not chosen_dir:
            return False
        export_name = safe_project_name + "_DLL_enquiry_waterbodies"
        folder_name, accepted = QInputDialog.getText(
            parent or self, "DLL Waterbody Export Folder",
            "Folder name:", text=export_name)
        if not accepted:
            return False
        folder_name = str(folder_name).strip()
        safe_folder_name = "".join(
            "_" if character in '<>:"/\\|?*' else character
            for character in folder_name).rstrip(". ")
        if not safe_folder_name:
            QMessageBox.warning(
                parent or self, "Natural England Pond Layer",
                "Enter a valid name for the export folder.")
            return False
        output_folder = os.path.join(chosen_dir, safe_folder_name)
        try:
            os.makedirs(output_folder, exist_ok=True)
        except OSError as exc:
            QMessageBox.critical(
                parent or self, "Natural England Pond Layer",
                "The export folder could not be created.\n\n" + str(exc))
            return False
        path = os.path.join(output_folder, export_name + ".shp")

        output = QgsVectorLayer(
            "MultiPolygon?crs=EPSG:27700", "Natural England Pond Layer",
            "memory")
        fields = QgsFields()
        fields.append(QgsField("Pond_ID", QVariant.String, len=20))
        fields.append(QgsField("GCN", QVariant.String, len=2))
        output.dataProvider().addAttributes(list(fields))
        output.updateFields()
        export_features = []
        for marker, geometry in matches:
            if not geometry.isMultipart():
                geometry.convertToMultiType()
            feature = QgsFeature(output.fields())
            feature.setGeometry(geometry)
            feature["Pond_ID"] = str(marker["pond_id"] or "")
            feature["GCN"] = str(marker["gcn_status"] or "").strip().upper()
            export_features.append(feature)
        output.dataProvider().addFeatures(export_features)
        output.updateExtents()

        options = QgsVectorFileWriter.SaveVectorOptions()
        options.driverName = "ESRI Shapefile"
        options.fileEncoding = "UTF-8"
        options.actionOnExistingFile = (
            QgsVectorFileWriter.ActionOnExistingFile.CreateOrOverwriteFile)
        error, message, _new_file, _new_layer = (
            QgsVectorFileWriter.writeAsVectorFormatV3(
                output, path, project.transformContext(), options))
        if error != QgsVectorFileWriter.WriterError.NoError:
            QMessageBox.critical(
                parent or self, "Natural England Pond Layer",
                "The shapefile could not be created.\n\n" + str(message))
            return False

        stem = os.path.splitext(path)[0]
        required = [stem + ext for ext in (".shp", ".shx", ".dbf", ".prj")]
        missing_files = [os.path.basename(item) for item in required
                         if not os.path.isfile(item)]
        if missing_files:
            QMessageBox.warning(
                parent or self, "Natural England Pond Layer",
                "The export completed, but the following required component "
                "files are missing: " + ", ".join(missing_files))
            return False

        self.iface.messageBar().pushSuccess(
            "Natural England Pond Layer",
            f"Exported {len(export_features)} pond polygon"
            + ("" if len(export_features) == 1 else "s")
            + " with Pond_ID and GCN attributes into "
            + os.path.basename(output_folder) + ".")
        return True

    def _on_pond_clear(self):
        """Remove all ponds after confirmation without cancelling placement."""
        from qgis.PyQt.QtWidgets import QMessageBox as _MB
        canvas = self.iface.mapCanvas()
        placing = (getattr(self, "_pond_tool", None) is not None
                   and canvas.mapTool() is self._pond_tool)
        reply = _MB.question(
            self, "Clear Placed Markers",
            "Remove all placed water-body markers and distance lines?",
            _MB.StandardButton.Yes | _MB.StandardButton.No,
            _MB.StandardButton.No,
        )
        if reply != _MB.StandardButton.Yes:
            return False
        remove_pond_distance_lines()
        remove_pond_layer()
        self._pond_lines_btn.setEnabled(False)
        self._pond_lines_btn.setChecked(False)
        self._pond_lines_btn.setText("〰  Distance lines: Off")
        if placing:
            # The placement tool resolves the layer afresh on every click, so
            # recreate an empty target and leave the active tool/button intact.
            get_or_create_pond_layer()
        self.iface.mapCanvas().refresh()
        self._refresh_pond_controls()
        return True

    def _open_pond_labels_menu(self, anchor):
        """Choose which information is shown beside SWB marker symbols."""
        menu = QMenu(anchor)
        menu.setMinimumWidth(205)
        menu.setStyleSheet(
            "QMenu { background:#FFFFFF; color:#1A2420;"
            " border:1px solid #B7BFB9; border-radius:8px; padding:5px;"
            " font-size:11px; }"
            "QMenu::item { color:#1A2420; background:transparent;"
            " padding:7px 24px 7px 10px; border-radius:5px; }"
            "QMenu::item:selected { background:#E6F0EA; color:#1F4A37; }"
            "QMenu::item:checked { background:#DDECE4; color:#173D2D;"
            " font-weight:600; }")
        group = QActionGroup(menu)
        group.setExclusive(True)
        current = pond_label_mode()
        choices = (
            ("both", "Name and distance"),
            ("name", "Name only"),
            ("distance", "Distance only"),
            ("off", "Labels off"),
        )
        for mode, label in choices:
            action = menu.addAction(label)
            action.setCheckable(True)
            action.setChecked(mode == current)
            action.setData(mode)
            group.addAction(action)
            action.triggered.connect(
                lambda _checked=False, m=mode: self._set_pond_label_mode(m))
        menu.exec(anchor.mapToGlobal(anchor.rect().bottomLeft()))

    def _set_pond_label_mode(self, mode):
        if mode != "off":
            self._pond_label_mode_before_off = mode
        apply_pond_label_mode(mode=mode, save=True)
        self.iface.mapCanvas().refresh()
        self._refresh_pond_controls()

    def _on_pond_labels_toggle(self, checked):
        if checked:
            mode = getattr(self, "_pond_label_mode_before_off", "both")
            if mode == "off":
                mode = "both"
        else:
            current = pond_label_mode()
            if current != "off":
                self._pond_label_mode_before_off = current
            mode = "off"
        self._set_pond_label_mode(mode)

    def _on_pond_lines_toggle(self, checked):
        """Show or hide the pond distance lines layer."""
        if checked:
            from qgis.core import QgsProject
            if not QgsProject.instance().mapLayersByName(POND_LINES_LAYER_NAME):
                # Lines are session-only derived data; after a project reopen
                # only the (persisted) ponds exist — rebuild the lines now.
                load_pond_distance_lines(self._get_redline_geom())
        set_pond_lines_visible(checked)
        self._pond_lines_btn.setText("〰  Distance lines: On" if checked else "〰  Distance lines: Off")
        self.iface.mapCanvas().refresh()

    def _run_epsl_report(self):
        """Fetch EPSL data and show the report dialog."""
        if not hasattr(self, "_centroid_bng"):
            QMessageBox.warning(self, "No Redline",
                "Please draw a redline boundary first.")
            return
        if _region.is_scotland():
            self._epsl_scotland_notice()
            return
        e, n = self._centroid_bng
        redline_geom = self._get_redline_geom()
        self._run_desk_study_async(
            "Searching EPSL records…",
            lambda: query_epsl_2km(e, n, redline_geom),
            lambda rows: self._show_epsl_report_dialog(rows, e, n),
            "EPSL Query Failed",
        )

    def _epsl_scotland_notice(self):
        """Explain why there is no EPSL search for a Scottish site and offer
        NatureScot's licensing summary page instead."""
        from qgis.PyQt.QtWidgets import QMessageBox as _MB
        from qgis.PyQt.QtCore import Qt as _Qt
        box = _MB(self)
        box.setIcon(_MB.Icon.Information)
        box.setWindowTitle(_region.EPSL_SCOTLAND_TITLE)
        box.setTextFormat(_Qt.TextFormat.RichText)
        note = _region.EPSL_SCOTLAND_NOTE.replace("\n\n", "<br><br>")
        box.setText(
            f"{note}<br><br><a href='{_region.EPSL_SCOTLAND_URL}'>"
            f"NatureScot · Summary of Licensing Data</a>")
        box.setStandardButtons(_MB.StandardButton.Ok)
        box.exec()

    def _sync_region_ui(self):
        """Relabel region-specific desk study controls after the site (or the
        Site region setting) changes: the AONB button reads NSA for a
        Scottish site, and the EPSL controls say why they are off."""
        try:
            aonb = self._ds_btns.get("aonb")
        except Exception:
            aonb = None
        if aonb is not None:
            lbl = ds_label("aonb")
            try:
                aonb._label_lbl.setText(lbl)
            except Exception:
                pass
            try:
                self._wire_hint(aonb, f"Toggle {lbl} overlay")
            except Exception:
                pass
        scot = _region.is_scotland()
        for attr in ("_epsl_btn",):
            btn = getattr(self, attr, None)
            if btn is not None:
                try:
                    self._wire_hint(
                        btn, _region.EPSL_SCOTLAND_TITLE + " — NatureScot publishes "
                        "no licence locations" if scot else
                        "Show granted European Protected Species Licence locations "
                        "within the search radius")
                except Exception:
                    pass

        # Buttons whose wording names the data source: MAGIC → Scotland's
        # environment map; Priority habitats → Habitat Map of Scotland.
        def _relabel(attr, text, tip=None):
            btn = getattr(self, attr, None)
            if btn is None:
                return
            try:
                btn._label_lbl.setText(text)
            except Exception:
                pass
            if tip:
                try:
                    self._wire_hint(btn, tip)
                except Exception:
                    pass

        _relabel("_magic_ext_btn", _region.magic_label(), _region.magic_tooltip())
        _relabel("_phi_report_btn", _region.phi_report_label(),
                 "Query the NatureScot Habitat Map of Scotland (EUNIS habitats mapped "
                 "to Scottish Biodiversity List priority habitats) within the search radius"
                 if scot else
                 "Query the Natural England Priority Habitats Inventory within the search radius")
        _relabel("_phi_map_btn", _region.phi_map_label(),
                 "Create a print layout showing all loaded Habitat Map of Scotland layers"
                 if scot else
                 "Create a print layout showing all loaded priority habitat layers")
        _relabel("_tpo_btn", "Show TPO locations",
                 "Load Tree Preservation Orders within the search radius from Spatial Hub "
                 "Scotland (needs the free Spatial Hub key — Biome settings)"
                 if scot else
                 "Load Tree Preservation Orders within the search radius — individual trees "
                 "(points) and woodland/area/group zones (polygons)")

    def _ensure_spatialhub_key(self):
        """Scottish TPO searches need the Spatial Hub Scotland authkey. If it
        is missing, offer to open the settings editor; return True when a key
        is available afterwards."""
        from .config import get_spatialhub_authkey
        if get_spatialhub_authkey():
            return True
        from qgis.PyQt.QtWidgets import QMessageBox as _MB
        res = _MB.question(
            self, "Spatial Hub Scotland key needed",
            "Tree Preservation Orders for a Scottish site come from the "
            "Improvement Service's Spatial Hub Scotland dataset, whose web "
            "service needs a (free) registered user's key.\n\n"
            "Enter your Spatial Hub key now?",
            _MB.StandardButton.Yes | _MB.StandardButton.No,
            _MB.StandardButton.Yes)
        if res == _MB.StandardButton.Yes:
            try:
                self._open_spatialhub_key_dialog()
            except Exception:
                pass
        return bool(get_spatialhub_authkey())

    def _report_header_row(self, hdr_label):
        """Wrap a report dialog's header label in a row. The copy/paste table
        style is now edited from the main-panel settings cog (see
        `_open_report_style_dialog`), not from a per-report cog."""
        from qgis.PyQt.QtWidgets import QHBoxLayout
        row = QHBoxLayout()
        row.setContentsMargins(0, 0, 0, 0)
        row.setSpacing(6)
        row.addWidget(hdr_label, 1)
        return row

    def _open_report_style_dialog(self):
        """Edit the report-table copy/paste style from the main settings cog.
        Any open report tables restyle live as the settings change; the
        dialog's own preview shows the effect otherwise."""
        from qgis.PyQt import sip
        from .clipboard_style_dialog import ClipboardStyleDialog
        live = [t for t in getattr(self, "_report_tables", [])
                if not sip.isdeleted(t)]
        self._report_tables = live
        ClipboardStyleDialog(self, tables=live).exec()

    def _open_branding_dialog(self):
        """Edit final-map branding (theme colours, logo, footnote) from the
        main settings cog. Was the 'Customise…' button in the map-export
        dialog; the export still bootstraps a default theme on first run."""
        from .branding import BrandingDialog
        BrandingDialog(self).exec()

    def _setup_copyable_table(self, tbl, include_headers=True):
        """Make a QTableWidget easy to copy-paste from.

        - Switches to item-level (cell) selection with extended multi-select
        - Ctrl+C copies selected cells as tab-separated text (pastes into Excel/Word)
        - Right-click context menu with 'Copy selected cells' and 'Copy all rows'
        - Restyles the table to the saved copy/paste report style, so what's
          on screen is exactly what a paste into the report will look like
          (deferred a tick, so it lands after the caller's own styling/fills)

        Pass include_headers=False for tables that carry their column headings
        inline as body rows (e.g. the EPSL report tabs), so 'Copy all rows'
        doesn't prepend a duplicate header line.
        """
        from qgis.PyQt.QtWidgets import QAbstractItemView, QMenu, QAction, QApplication, QLabel
        from qgis.PyQt.QtCore import Qt, QObject, QEvent, QTimer
        from qgis.PyQt.QtGui import QKeySequence
        from .clipboard_utils import set_table_clipboard

        tbl.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectItems)
        tbl.setSelectionMode(QAbstractItemView.SelectionMode.ExtendedSelection)
        # Drag a column heading to reorder columns; the copy follows the new
        # order. (Row order follows sorting — click a column heading to sort.)
        _hh = tbl.horizontalHeader()
        _hh.setSectionsMovable(True)
        _hh.setToolTip("Drag a column heading to reorder columns")

        def _show_copy_toast(msg="Copied to clipboard"):
            toast = QLabel(msg, tbl)
            toast.setStyleSheet(
                "QLabel { background:#2F5D8A; color:white; font-size:11px; "
                "padding:5px 12px; border-radius:4px; }"
            )
            toast.setAlignment(Qt.AlignmentFlag.AlignCenter)
            toast.adjustSize()
            # Position bottom-centre of the table
            tx = (tbl.width() - toast.width()) // 2
            ty = tbl.height() - toast.height() - 8
            toast.move(tx, ty)
            toast.show()
            toast.raise_()
            QTimer.singleShot(1400, toast.deleteLater)

        def _selected_rows(all_rows=False):
            """Copy payload as a list of row-lists, honouring any drag-reordered
            column order (rows already follow the table's sort order)."""
            hh = tbl.horizontalHeader()
            vis_cols = [hh.logicalIndex(v) for v in range(tbl.columnCount())]
            if all_rows:
                grid = []
                if include_headers:
                    grid.append([(tbl.horizontalHeaderItem(c).text()
                                  if tbl.horizontalHeaderItem(c) else "")
                                 for c in vis_cols])
                for r in range(tbl.rowCount()):
                    grid.append([(tbl.item(r, c).text() if tbl.item(r, c) else "")
                                 for c in vis_cols])
                return grid
            sel = tbl.selectedIndexes()
            if not sel:
                return []
            rows_set = sorted({i.row() for i in sel})
            cols_set = sorted({i.column() for i in sel},
                              key=lambda c: hh.visualIndex(c))
            return [[(tbl.item(r, c).text() if tbl.item(r, c) else "")
                     for c in cols_set]
                    for r in rows_set]

        def _copy_selected():
            grid = _selected_rows(all_rows=False)
            if grid:
                set_table_clipboard(grid, header_rows=0)
                n_rows = len(grid)
                msg = f"Copied {n_rows} row{'s' if n_rows != 1 else ''}"
                _show_copy_toast(msg)

        def _copy_all():
            set_table_clipboard(_selected_rows(all_rows=True),
                                header_rows=1 if include_headers else 0)
            _show_copy_toast(f"Copied all {tbl.rowCount()} rows")

        class _KeyFilter(QObject):
            def eventFilter(self, obj, event):
                if event.type() == QEvent.Type.KeyPress:
                    if event.matches(QKeySequence.StandardKey.Copy):
                        _copy_selected()
                        return True
                return False

        _kf = _KeyFilter(tbl)
        tbl.installEventFilter(_kf)
        tbl._key_filter = _kf  # keep reference alive

        tbl.setContextMenuPolicy(Qt.ContextMenuPolicy.CustomContextMenu)

        def _context_menu(pos):
            menu = QMenu(tbl)
            menu.setStyleSheet(
                "QMenu { font-size:11px; } "
                "QMenu::item:selected { background:#2E6B4F; color:white; }"
            )
            act_sel = QAction("Copy selected cells", menu)
            act_sel.triggered.connect(_copy_selected)
            act_all = QAction("Copy all rows (with headers)", menu)
            act_all.triggered.connect(_copy_all)
            menu.addAction(act_sel)
            menu.addAction(act_all)
            menu.exec(tbl.viewport().mapToGlobal(pos))

        tbl.customContextMenuRequested.connect(_context_menu)

        # ── WYSIWYG: mirror the copy/paste report style on the live table ──
        # Registered so the settings-cog style dialog can restyle open report
        # dialogs live; applied on a 0-timer so it runs after the caller has
        # populated items (it clears per-item fills the paste won't carry).
        from qgis.PyQt import sip
        from .clipboard_utils import apply_table_style
        self._report_tables = [
            t for t in getattr(self, "_report_tables", []) if not sip.isdeleted(t)
        ]
        self._report_tables.append(tbl)

        def _apply_house_style():
            if not sip.isdeleted(tbl):
                apply_table_style(tbl)

        QTimer.singleShot(0, _apply_house_style)

    def _show_epsl_report_dialog(self, rows, easting, northing):
        """Display EPSL results in a tabbed sortable table — All, Bats, Recent, Historic."""
        from qgis.PyQt.QtWidgets import (
            QDialog, QVBoxLayout, QHBoxLayout, QTableWidget,
            QTableWidgetItem, QHeaderView, QPushButton, QLabel,
            QAbstractItemView, QTabWidget, QWidget,
        )
        from qgis.PyQt.QtCore import Qt
        from datetime import datetime, timezone

        try:
            from .gridref_utils import osgb_en_to_gridref
            site_gr = osgb_en_to_gridref(int(easting), int(northing), 6)
        except Exception:
            site_gr = f"{int(easting)} {int(northing)}"

        # ── Split rows into bat / recent / historic ────────────────────
        bat_rows = [r for r in rows if "bat" in r.get("species_gr", "").lower()]

        current_year = datetime.now(tz=timezone.utc).year
        recent_start_year = current_year - 9
        cutoff = datetime(recent_start_year, 1, 1, tzinfo=timezone.utc)

        def _start_dt(row):
            """Parse start_date string 'DD/MM/YYYY' → datetime, or None."""
            s = row.get("start_date", "")
            if not s:
                return None
            try:
                return datetime.strptime(s, "%d/%m/%Y").replace(tzinfo=timezone.utc)
            except ValueError:
                return None

        def _sort_by_start(row_list, reverse=True):
            """Sort rows by start_date descending (most recent first).
            Rows with no start date fall to the end."""
            def _key(r):
                s = r.get("start_date", "")
                try:
                    return datetime.strptime(s, "%d/%m/%Y").replace(tzinfo=timezone.utc)
                except (ValueError, TypeError):
                    return datetime.min.replace(tzinfo=timezone.utc)
            return sorted(row_list, key=_key, reverse=reverse)

        recent_rows   = []
        historic_rows = []
        for r in rows:
            dt = _start_dt(r)
            # No start date → treat as still active → recent
            if dt is None or dt >= cutoff:
                recent_rows.append(r)
            else:
                historic_rows.append(r)

        bat_recent_rows   = [r for r in recent_rows   if "bat" in r.get("species_gr", "").lower()]
        bat_historic_rows = [r for r in historic_rows if "bat" in r.get("species_gr", "").lower()]

        dlg = QDialog(self)
        _rad = self._radius_label("epsl")
        dlg.setWindowTitle(f"EPSL Licences within {_rad}")
        dlg.setMinimumSize(900, 520)
        dlg.resize(1000, 580)
        layout = QVBoxLayout(dlg)
        layout.setSpacing(8)

        hdr = QLabel(
            f"<b>Granted European Protected Species Licences within {_rad} of site centroid</b> "
            f"({site_gr})<br>"
            f"<span style='font-size:10px;color:#717E75;'>"
            f"Source: Natural England via MAGIC (Landmark ArcGIS) — public data</span>"
        )
        hdr.setStyleSheet("font-size:12px; color:#1A2420; padding-bottom:4px;")
        layout.addLayout(self._report_header_row(hdr))

        tabs = QTabWidget()
        tabs.setMovable(True)  # drag to reorder tabs
        tabs.setStyleSheet(
            "QTabBar::tab { padding:5px 14px; font-size:11px; }"
            "QTabBar::tab:selected { background:#FBF4E4; color:#8A4B12; font-weight:500; }"
        )

        def _make_tab(tab_rows, species_label):
            w = QWidget()
            vl = QVBoxLayout(w)
            vl.setContentsMargins(0, 6, 0, 0)
            vl.setSpacing(4)
            if not tab_rows:
                vl.addWidget(_banner(
                    "status-check", f"No {species_label} EPSL licences found within {_rad}.",
                    fg="#1F7A4D", bg="#E9F4EC", border="1px solid #BFE0CC",
                    font_px=12, radius=6, margins=(12, 12, 12, 12)))
                return w

            from collections import Counter
            sp_counts = Counter(r["species_gr"] for r in tab_rows)
            summary_parts = [f"<b>{v}</b> {k}" for k, v in sorted(sp_counts.items())]
            summary = QLabel(
                f"Found <b>{len(tab_rows)}</b> licence(s): " + " · ".join(summary_parts)
            )
            summary.setStyleSheet("font-size:11px; color:#54635A; padding-bottom:4px;")
            vl.addWidget(summary)

            cols = ["Licence", "Distance to redline (km)",
                    "Species Group", "County", "End Date", "Grid Ref",
                    "Breeding Site Affected", "Breeding Site Damaged",
                    "Breeding Site Destroyed", "Resting Place Damaged",
                    "Resting Place Destroyed", "Hibernation Site Affected"]
            tbl = QTableWidget(len(tab_rows), len(cols))
            tbl.setHorizontalHeaderLabels(cols)
            tbl.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
            self._setup_copyable_table(tbl)
            for r_idx, row in enumerate(tab_rows):
                licence_cell = (
                    f"{row['case_ref']} — {row['start_date']} — {row['species']}"
                    if row['start_date'] else
                    f"{row['case_ref']} — {row['species']}"
                )
                dist_dir = f"{row['distance_m'] / 1000:.2f} km {row['direction']}"
                vals = [licence_cell, dist_dir, row["species_gr"],
                        row["county"], row["end_date"], row["grid_ref"],
                        row.get("affects_br", ""), row.get("damage_bre", ""),
                        row.get("destroy_br", ""), row.get("damage_res", ""),
                        row.get("destroy_re", ""), row.get("affects_hi", "")]
                for c_idx, val in enumerate(vals):
                    if c_idx == 1:
                        item = _NumericSortItem(val)
                        item.setData(Qt.ItemDataRole.UserRole, row["distance_m"])
                    else:
                        item = QTableWidgetItem(val)
                    item.setTextAlignment(
                        Qt.AlignmentFlag.AlignVCenter | (Qt.AlignmentFlag.AlignCenter if c_idx == 1 else Qt.AlignmentFlag.AlignLeft))
                    tbl.setItem(r_idx, c_idx, item)

            tbl.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeMode.Stretch)
            for ci in range(1, len(cols)):
                tbl.horizontalHeader().setSectionResizeMode(ci, QHeaderView.ResizeMode.ResizeToContents)
            tbl.verticalHeader().setVisible(False)
            tbl.setSortingEnabled(True)
            tbl.sortItems(1)
            vl.addWidget(tbl)
            return w

        # ── Report-format tabs: reproduce the previous-licences table ──────
        def _by_distance(lst):
            return sorted(lst, key=lambda r: r["distance_m"])

        def _report_cell(row):
            """(Reference·Date·Species, Distance, Impact) for the report table."""
            ref = (f"{row['case_ref']} — {row['start_date']} — {row['species']}"
                   if row['start_date'] else f"{row['case_ref']} — {row['species']}")
            dist = f"{row['distance_m'] / 1000:.2f} km {row['direction']}"
            return ref, dist, epsl_impact_summary(row)

        def _report_rows(recent, historic):
            grid = []
            for suffix, rws in ((f"{recent_start_year}–{current_year}",
                                 _by_distance(recent)),
                                (f"before {recent_start_year}",
                                 _by_distance(historic))):
                grid.append([f"Reference, Date and Species ({suffix})",
                             "Approx. Distance from Site", "Impact"])
                for row in rws:
                    grid.append(list(_report_cell(row)))
            return grid

        def _make_report_tab(recent, historic, empty_label):
            """Build a tab that mirrors the previous protected-species licences
            table: three columns (Reference·Date·Species / Approx. Distance /
            Impact) split into explicit recent and historic calendar-year
            bands, each ordered by distance from site."""
            w = QWidget()
            vl = QVBoxLayout(w)
            vl.setContentsMargins(0, 6, 0, 0)
            vl.setSpacing(4)

            sections = [(f"{recent_start_year}–{current_year}",
                         _by_distance(recent)),
                        (f"before {recent_start_year}",
                         _by_distance(historic))]
            if len(recent) + len(historic) == 0:
                vl.addWidget(_banner(
                    "status-check", f"No {empty_label} EPSL licences found within {_rad}.",
                    fg="#1F7A4D", bg="#E9F4EC", border="1px solid #BFE0CC",
                    font_px=12, radius=6, margins=(12, 12, 12, 12)))
                return w

            n_rows = sum(1 + len(rws) for _, rws in sections)  # header band + rows
            tbl = QTableWidget(n_rows, 3)
            tbl.setHorizontalHeaderLabels(
                ["Reference, Date and Species", "Approx. Distance from Site", "Impact"])
            tbl.horizontalHeader().setVisible(False)  # headers are inline, per the report
            tbl.verticalHeader().setVisible(False)
            tbl.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
            tbl.setWordWrap(True)
            self._setup_copyable_table(tbl, include_headers=False)

            r = 0
            for suffix, rws in sections:
                heads = [f"Reference, Date and Species ({suffix})",
                         "Approx. Distance from Site", "Impact"]
                for c, h in enumerate(heads):
                    it = QTableWidgetItem(h)
                    f = it.font(); f.setBold(True); it.setFont(f)
                    it.setTextAlignment(Qt.AlignmentFlag.AlignVCenter | Qt.AlignmentFlag.AlignLeft)
                    tbl.setItem(r, c, it)
                r += 1
                for row in rws:
                    for c, val in enumerate(_report_cell(row)):
                        it = QTableWidgetItem(val)
                        it.setTextAlignment(
                            Qt.AlignmentFlag.AlignVCenter | (Qt.AlignmentFlag.AlignCenter if c == 1 else Qt.AlignmentFlag.AlignLeft))
                        tbl.setItem(r, c, it)
                    r += 1

            hh = tbl.horizontalHeader()
            hh.setSectionResizeMode(0, QHeaderView.ResizeMode.Stretch)
            hh.setSectionResizeMode(1, QHeaderView.ResizeMode.ResizeToContents)
            hh.setSectionResizeMode(2, QHeaderView.ResizeMode.Stretch)
            tbl.resizeRowsToContents()
            vl.addWidget(tbl)
            return w

        # Give each tab a stable identity.  Labels contain changing result counts,
        # so they are unsuitable for restoring a user-defined order.
        _tab_defs = [
            ("all", _make_tab(rows, "granted"), None,
             f"All Licences ({len(rows)})"),
            ("bats", _make_tab(bat_rows, "bat"), "bat",
             f"Bats Only ({len(bat_rows)})"),
            ("recent", _make_tab(_sort_by_start(recent_rows), "recent"), "clock",
             f"Recent · {recent_start_year}–{current_year} ({len(recent_rows)})"),
            ("historic", _make_tab(_sort_by_start(historic_rows), "historic"), "book",
             f"Historic · before {recent_start_year} ({len(historic_rows)})"),
            ("bats_recent", _make_tab(_sort_by_start(bat_recent_rows), "bat recent"), "bat",
             f"Bats recent · {recent_start_year}–{current_year} ({len(bat_recent_rows)})"),
            ("bats_historic", _make_tab(_sort_by_start(bat_historic_rows), "bat historic"), "bat",
             f"Bats historic · before {recent_start_year} ({len(bat_historic_rows)})"),
            ("bats_report", _make_report_tab(bat_recent_rows, bat_historic_rows, "bat"), "report",
             f"Bats Report ({len(bat_rows)})"),
            ("all_report", _make_report_tab(recent_rows, historic_rows, "granted"), "report",
             f"All-Species Report ({len(rows)})"),
        ]
        _tabs_by_id = {}
        for _tab_id, _widget, _icon_name, _label in _tab_defs:
            _widget.setProperty("biomeTabId", _tab_id)
            _tabs_by_id[_tab_id] = _widget
            if _icon_name:
                tabs.addTab(_widget, _svg_icon(_icon_name, "#1F4A37", 14), _label)
            else:
                tabs.addTab(_widget, _label)

        _tab_order_key = "biome/desk_study/epsl_report_tab_order"
        _saved_order = QSettings().value(_tab_order_key, [])
        if isinstance(_saved_order, str):
            _saved_order = [_saved_order]
        elif not isinstance(_saved_order, (list, tuple)):
            _saved_order = []
        _known_order = [tab_id for tab_id in _saved_order if tab_id in _tabs_by_id]
        _known_order.extend(tab_id for tab_id, *_ in _tab_defs
                            if tab_id not in _known_order)
        for _target_index, _tab_id in enumerate(_known_order):
            _current_index = tabs.indexOf(_tabs_by_id[_tab_id])
            if _current_index != _target_index:
                tabs.tabBar().moveTab(_current_index, _target_index)

        def _save_tab_order(_from_index=None, _to_index=None):
            QSettings().setValue(
                _tab_order_key,
                [tabs.widget(i).property("biomeTabId") for i in range(tabs.count())],
            )

        tabs.tabBar().tabMoved.connect(_save_tab_order)
        layout.addWidget(tabs)

        # Map stable tab identity → row data and sheet name for copy/export.
        _TAB_DATA = {
            "all":          (rows,             f"EPSL Licences within {_rad}"),
            "bats":         (bat_rows,         "EPSL Bats Only"),
            "recent":       (recent_rows,      f"EPSL Recent ({recent_start_year}-{current_year})"),
            "historic":     (historic_rows,    f"EPSL Historic (before {recent_start_year})"),
            "bats_recent":  (bat_recent_rows,  f"EPSL Bats Recent ({recent_start_year}-{current_year})"),
            "bats_historic": (bat_historic_rows,f"EPSL Bats Historic (before {recent_start_year})"),
        }

        btn_row = QHBoxLayout()
        if rows:
            copy_btn = QPushButton("Copy")
            copy_btn.setIcon(_svg_icon("copy", "#1A2420", 14))
            copy_btn.setIconSize(QSize(14, 14))
            copy_btn.setStyleSheet("""
                QPushButton {
                    background:#F4F6F2; color:#1A2420;
                    border:1px solid #B7BFB9; border-radius:4px;
                    font-size:11px; padding:4px 12px;
                }
                QPushButton:hover { background:#E3E7E1; }
            """)
            def _copy():
                from .clipboard_utils import set_table_clipboard
                tab_id = tabs.currentWidget().property("biomeTabId")
                if tab_id in ("bats_report", "all_report"):
                    rec, hist = ((bat_recent_rows, bat_historic_rows)
                                 if tab_id == "bats_report"
                                 else (recent_rows, historic_rows))
                    set_table_clipboard(_report_rows(rec, hist))
                    return
                active, _ = _TAB_DATA.get(tab_id, (rows, ""))
                header = ["Licence", "Distance to redline (km)", "Species Group",
                          "County", "End Date", "Grid Ref", "Breeding Site Affected",
                          "Breeding Site Damaged", "Breeding Site Destroyed",
                          "Resting Place Damaged", "Resting Place Destroyed",
                          "Hibernation Site Affected"]
                grid = [header]
                for r in active:
                    grid.append([
                        f'{r["case_ref"]} — {r["start_date"]} — {r["species"].replace(chr(34), "")}',
                        f'{r["distance_m"]/1000:.2f} km {r["direction"]}',
                        r["species_gr"], r["county"], r["end_date"], r["grid_ref"],
                        r.get("affects_br", ""), r.get("damage_bre", ""), r.get("destroy_br", ""),
                        r.get("damage_res", ""), r.get("destroy_re", ""), r.get("affects_hi", ""),
                    ])
                set_table_clipboard(grid, header_rows=1)
            copy_btn.clicked.connect(_copy)
            btn_row.addWidget(copy_btn)

            excel_btn = QPushButton("Export to Excel")
            excel_btn.setIcon(_svg_icon("table", "#175E3B", 14))
            excel_btn.setIconSize(QSize(14, 14))
            excel_btn.setStyleSheet("""
                QPushButton {
                    background:#F4F6F2; color:#175E3B;
                    border:1px solid #1F7A4D; border-radius:4px;
                    font-size:11px; padding:4px 12px;
                }
                QPushButton:hover { background:#E9F4EC; }
            """)
            _excel_cols_epsl = [
                ("licence_cell", "Licence"),
                ("dist_dir",     "Distance to redline (km)"),
                ("species_gr",   "Species Group"),
                ("county",       "County"),
                ("end_date",     "End Date"),
                ("grid_ref",     "Grid Ref"),
                ("affects_br",   "Breeding Site Affected"),
                ("damage_bre",   "Breeding Site Damaged"),
                ("destroy_br",   "Breeding Site Destroyed"),
                ("damage_res",   "Resting Place Damaged"),
                ("destroy_re",   "Resting Place Destroyed"),
                ("affects_hi",   "Hibernation Site Affected"),
            ]
            def _excel_export():
                tab_id = tabs.currentWidget().property("biomeTabId")
                if tab_id in ("bats_report", "all_report"):
                    if tab_id == "bats_report":
                        rec, hist, sheet, fname = (bat_recent_rows, bat_historic_rows,
                                                   "EPSL Bats Report", "EPSL_Bats_Report")
                    else:
                        rec, hist, sheet, fname = (recent_rows, historic_rows,
                                                   "EPSL Report", "EPSL_Report")
                    self._export_epsl_report_excel(
                        [(f"{recent_start_year}-{current_year}", rec),
                         (f"before {recent_start_year}", hist)], sheet, fname)
                    return
                active, sheet = _TAB_DATA.get(tab_id, (rows, f"EPSL Licences within {_rad}"))
                export_rows = [{
                    **r,
                    "licence_cell": f"{r['case_ref']} — {r['start_date']} — {r['species']}" if r['start_date'] else f"{r['case_ref']} — {r['species']}",
                    "dist_dir": f"{r['distance_m']/1000:.2f} km {r['direction']}",
                } for r in active]
                self._export_report_excel(export_rows, _excel_cols_epsl, sheet, "EPSL_Report")
            excel_btn.clicked.connect(_excel_export)
            btn_row.addWidget(excel_btn)

        btn_row.addStretch()
        close_btn = QPushButton("Close")
        close_btn.setStyleSheet("""
            QPushButton {
                background:#9A6700; color:white; border:none;
                border-radius:4px; font-size:11px; padding:4px 16px;
            }
            QPushButton:hover { background:#8A4B12; }
        """)
        close_btn.clicked.connect(dlg.accept)
        btn_row.addWidget(close_btn)
        layout.addLayout(btn_row)
        dlg.setModal(False)
        dlg.show()

    def _on_buffer_2km_toggle(self, checked):
        """Add or remove the shared search-area ring manually."""
        self._manual_search_area_enabled = bool(checked)
        if not checked:
            remove_2km_buffer()
            self.iface.mapCanvas().refresh()
            return
        if not hasattr(self, "_centroid_bng"):
            QMessageBox.warning(self, "No Redline",
                "Please draw a redline boundary first.")
            self._buffer_2km_btn.setChecked(False)
            return
        e, n = self._centroid_bng
        ok, result = add_2km_buffer(e, n, self._get_redline_geom())
        if not ok:
            QMessageBox.warning(self, "Buffer Failed", str(result))
            self._buffer_2km_btn.setChecked(False)
        else:
            self.iface.mapCanvas().refresh()

    def _update_search_area_button_labels(self):
        """Show the live saved radii directly on both mapping-tool buttons."""
        from .config import (get_buffer_radius, get_search_radius,
                             get_pond_inner_radius)
        if getattr(self, "_buffer_2km_btn", None) is not None:
            self._buffer_2km_btn._label_lbl.setText(
                "Desk study area\n%s" %
                self._buffer_preset_label(get_buffer_radius()))
        if getattr(self, "_buffers_btn", None) is not None:
            tpo = get_search_radius("tpo")
            pond_inner = get_pond_inner_radius()
            pond_outer = get_search_radius("standing_water")
            mode = getattr(self, "_active_local_buffer_preset", None)
            if mode not in ("tpo", "water", "manual"):
                # Legacy/reopened projects may not yet have an explicit mode.
                # Choose one preset only; never advertise an impossible
                # combined TPO + standing-water state.
                if (getattr(self, "_tpo_btn", None)
                        and self._tpo_btn.isChecked()):
                    mode = "tpo"
                elif self._standing_water_search_active():
                    mode = "water"
                elif getattr(self, "_manual_local_buffers_enabled", False):
                    mode = "manual"
            if mode == "manual":
                from .panel_map_tools import _RedLineBuffersDialog
                vals = _RedLineBuffersDialog(
                    self.iface.mainWindow()).get_values()
                detail = (f"{vals['dist2']} m outer" if vals["outerOnly"]
                          else f"{vals['dist1']}–{vals['dist2']} m")
            elif mode == "tpo":
                detail = f"TPO · {tpo} m"
            elif mode == "water":
                detail = f"Water · {pond_inner}–{pond_outer} m"
            else:
                from .panel_map_tools import _RedLineBuffersDialog
                vals = _RedLineBuffersDialog(
                    self.iface.mainWindow()).get_values()
                detail = (f"{vals['dist2']} m outer" if vals["outerOnly"]
                          else f"{vals['dist1']}–{vals['dist2']} m")
            self._buffers_btn._label_lbl.setText(
                "Local buffers" + ("\n" + detail if detail else ""))
        if getattr(self, "_pond_search_range_lbl", None) is not None:
            self._refresh_pond_controls()

    def _standing_water_search_active(self):
        """Return whether the standing-water search extent should be shown.

        The request flag deliberately survives a zero-result known-pond search:
        the rings document the area searched, not the number of results.
        """
        if getattr(self, "_standing_water_buffer_requested", False):
            return True
        if (getattr(self, "_os_topo_btn", None)
                and self._os_topo_btn.isChecked()):
            return True
        if (getattr(self, "_pond_place_btn", None)
                and self._pond_place_btn.isChecked()):
            return True
        try:
            from qgis.core import QgsProject
            return any(
                layer.featureCount() > 0
                for layer in QgsProject.instance().mapLayersByName(
                    POND_LAYER_NAME))
        except Exception:
            return False

    def _sync_auto_search_areas(self, sync_local=True):
        """Keep desk-study and local search areas independent.

        Shared desk-study toggles update only the shared search area. Local
        controls additionally apply exactly one local preset (TPO, water, or
        manual), replacing the previous geometry rather than combining rings.
        """
        if getattr(self, "_suspend_auto_search_area_sync", False):
            return
        if not hasattr(self, "_centroid_bng"):
            return
        shared_buttons = (
            list(getattr(self, "_ds_btns", {}).values()) +
            list(getattr(self, "_phi_btns", {}).values()) +
            [btn for btn in (
                getattr(self, "_aw_btn", None),
                getattr(self, "_epsl_btn", None),
                getattr(self, "_ca_btn", None),
                getattr(self, "_edna_btn", None),
            ) if btn is not None]
        )
        shared_on = (any(btn.isChecked() for btn in shared_buttons) or
                     bool(getattr(self, "_manual_search_area_enabled", False)))
        e, n = self._centroid_bng
        if shared_on:
            add_2km_buffer(e, n, self._get_redline_geom())
        else:
            remove_2km_buffer()
        if getattr(self, "_buffer_2km_btn", None) is not None:
            self._buffer_2km_btn.blockSignals(True)
            self._buffer_2km_btn.setChecked(shared_on)
            self._buffer_2km_btn.blockSignals(False)
            if hasattr(self._buffer_2km_btn, "_sync_checked_style"):
                self._buffer_2km_btn._sync_checked_style(shared_on)

        # Designated sites, PHI, EPSL, and the other shared searches must not
        # alter the separately controlled local buffers.
        if not sync_local:
            self._update_search_area_button_labels()
            return

        manual_local = bool(getattr(
            self, "_manual_local_buffers_enabled", False))
        mode = getattr(self, "_active_local_buffer_preset", None)
        if mode not in ("tpo", "water", "manual"):
            if manual_local:
                mode = "manual"
            elif (getattr(self, "_tpo_btn", None)
                  and self._tpo_btn.isChecked()):
                mode = "tpo"
            elif self._standing_water_search_active():
                mode = "water"
        local_on = mode is not None
        if mode in ("tpo", "water"):
            from .config import get_search_radius, get_pond_inner_radius
            from .panel_map_tools import _RedLineBuffersDialog
            vals = _RedLineBuffersDialog(self.iface.mainWindow()).get_values()
            tpo_radius = get_search_radius("tpo")
            pond_inner = get_pond_inner_radius()
            pond_radius = get_search_radius("standing_water")
            # Automatic presets always replace previous generated buffer
            # layers, even if a manual run had disabled "delete old".
            vals.update(deleteOld=True, extraDistances=[], quiet=True)
            if mode == "water":
                vals.update(dist1=pond_inner, dist2=pond_radius,
                            outerOnly=False, quiet=True)
            else:
                vals.update(dist2=tpo_radius, outerOnly=True, quiet=True)
            local_on = bool(self._create_redline_buffers(vals=vals))
        elif mode == "manual":
            self._set_redline_buffers_visible(True)
        else:
            self._set_redline_buffers_visible(False)
        if getattr(self, "_buffers_btn", None) is not None:
            self._buffers_btn.blockSignals(True)
            self._buffers_btn.setChecked(local_on)
            self._buffers_btn.blockSignals(False)
            if hasattr(self._buffers_btn, "_sync_checked_style"):
                self._buffers_btn._sync_checked_style(local_on)
        self._update_search_area_button_labels()

    def _buffer_radius_menu(self, pos):
        """Choose the shared desk-study search distance from the live button."""
        from qgis.PyQt.QtWidgets import QMenu
        from .config import get_buffer_radius

        cur = get_buffer_radius()
        menu = QMenu(self)
        head = menu.addAction("Desk study search area")
        head.setEnabled(False)
        menu.addSeparator()
        for metres in (250, 500, 1000, 2000, 5000, 10000):
            label = f"{metres} m" if metres < 1000 else f"{metres // 1000} km"
            act = menu.addAction(label)
            act.setCheckable(True)
            act.setChecked(metres == cur)
            act.triggered.connect(lambda _checked=False, m=metres: self._set_buffer_radius(m))
        menu.addSeparator()
        custom = menu.addAction("Custom distance…")
        custom.triggered.connect(self._prompt_buffer_radius)
        menu.exec(self._buffer_2km_btn.mapToGlobal(pos))

    def _prompt_buffer_radius(self):
        """Ask for a custom shared desk-study search distance in metres."""
        from qgis.PyQt.QtWidgets import QInputDialog
        from .config import get_buffer_radius
        metres, ok = QInputDialog.getInt(
            self, "Desk study search area",
            "Distance from the redline boundary (metres):",
            get_buffer_radius(), 50, 25000, 50)
        if ok:
            self._set_buffer_radius(metres)

    def _radius_label(self, key):
        """Human label for a desk study search radius, e.g. '2 km', '750 m'.
        Used so report titles/headers state the ACTUAL distance searched rather
        than a hard-coded '2 km'."""
        from .config import get_search_radius
        m = get_search_radius(key)
        if m < 1000:
            return f"{m} m"
        return f"{m // 1000} km" if m % 1000 == 0 else f"{m / 1000:.1f} km"

    def _set_buffer_radius(self, metres):
        """Persist the shared radius, invalidate linked data, and redraw."""
        from .config import set_buffer_radius, SHARED_RADIUS_KEYS
        set_buffer_radius(metres)
        self._update_search_area_button_labels()
        try:
            self._suspend_auto_search_area_sync = True
            self._apply_radius_changes(list(SHARED_RADIUS_KEYS))
            QTimer.singleShot(
                0, lambda: setattr(
                    self, "_suspend_auto_search_area_sync", False))
        except Exception:
            self._suspend_auto_search_area_sync = False
            pass
        if (getattr(self, "_buffer_2km_btn", None) is not None
                and self._buffer_2km_btn.isChecked()
                and hasattr(self, "_centroid_bng")):
            e, n = self._centroid_bng
            add_2km_buffer(e, n, self._get_redline_geom())
            self.iface.mapCanvas().refresh()

    # ── Constraint buffers (designated sites & priority habitats) ───────────
    # Distinct from the Search Buffer (which rings the redline). These ring an
    # individual constraint feature (info-popup button) or a whole constraint
    # layer (right-click its toggle). See designated_sites.add_constraint_buffer.
    _BUFFER_PRESETS = (50, 100, 250, 500, 1000)

    @staticmethod
    def _buffer_preset_label(metres):
        return f"{metres} m" if metres < 1000 else f"{metres // 1000} km"

    def _constraint_source_layer(self, kind, key):
        """Return (loaded QgsVectorLayer or None, type_label) for a designated
        site ('ds') or priority-habitat group ('phi') key. type_label matches
        designated_sites._constraint_type_and_colour so buffers group correctly."""
        from qgis.core import QgsProject
        if kind == "ds":
            label = ds_label(key)
            lyrs  = QgsProject.instance().mapLayersByName(label)
            return (lyrs[0] if lyrs else None), label
        from .designated_sites import _phi_layer_name
        name = _phi_layer_name(key)
        lyrs = QgsProject.instance().mapLayersByName(name)
        prefix = "Priority Habitat - "
        type_label = name[len(prefix):] if name.startswith(prefix) else name
        return (lyrs[0] if lyrs else None), type_label

    def _constraint_buffer_menu(self, kind, key, btn, pos):
        """Right-click menu on a constraint toggle — buffer the WHOLE layer at a
        preset/custom distance, inspect what falls inside, or clear it."""
        from qgis.PyQt.QtWidgets import QMenu
        from .designated_sites import active_constraint_buffers

        layer, type_label = self._constraint_source_layer(kind, key)

        menu = QMenu(self)
        head = menu.addAction(f"Buffer {type_label}")
        head.setEnabled(False)
        menu.addSeparator()

        if layer is None:
            info = menu.addAction("Load this layer first to buffer it")
            info.setEnabled(False)
            menu.exec(btn.mapToGlobal(pos))
            return

        for metres in self._BUFFER_PRESETS:
            act = menu.addAction(f"Whole layer — {self._buffer_preset_label(metres)}")
            act.triggered.connect(
                lambda _c=False, m=metres, lyr=layer: self._do_layer_buffer(lyr, m))
        custom = menu.addAction("Whole layer — custom distance…")
        custom.triggered.connect(
            lambda _c=False, lyr=layer: self._prompt_layer_buffer(lyr))

        if type_label in active_constraint_buffers():
            menu.addSeparator()
            whats = menu.addAction("What's within this buffer?")
            whats.triggered.connect(
                lambda _c=False, t=type_label: self._buffer_analysis(t))
            rem = menu.addAction("Remove buffer")
            rem.triggered.connect(
                lambda _c=False, t=type_label: self._remove_constraint_buffer(t))

        menu.exec(btn.mapToGlobal(pos))

    def _do_layer_buffer(self, layer, metres):
        from .designated_sites import add_layer_buffer
        ok, res = add_layer_buffer(layer, metres)
        self._after_buffer(ok, res, f"Buffered {layer.name()} by {metres} m.")

    def _prompt_layer_buffer(self, layer):
        from qgis.PyQt.QtWidgets import QInputDialog
        metres, ok = QInputDialog.getInt(
            self, "Buffer distance",
            "Distance from the feature edge (metres):", 250, 5, 25000, 25)
        if ok:
            self._do_layer_buffer(layer, metres)

    def _feature_buffer_menu(self, source_layer_name, geom, anchor_widget,
                             default_name=None):
        """Distance menu for buffering a single identified constraint feature
        (opened from the feature-info popup's 'Buffer this feature…' button)."""
        from qgis.PyQt.QtWidgets import QMenu, QInputDialog
        menu = QMenu(self)
        head = menu.addAction("Buffer this feature")
        head.setEnabled(False)
        menu.addSeparator()
        for metres in self._BUFFER_PRESETS:
            act = menu.addAction(self._buffer_preset_label(metres))
            act.triggered.connect(
                lambda _c=False, m=metres: self._apply_feature_buffer(
                    source_layer_name, geom, m, default_name))
        menu.addSeparator()
        custom = menu.addAction("Custom distance…")

        def _custom(_c=False):
            metres, ok = QInputDialog.getInt(
                self, "Buffer distance",
                "Distance from the feature edge (metres):", 250, 5, 25000, 25)
            if ok:
                self._apply_feature_buffer(
                    source_layer_name, geom, metres, default_name)

        custom.triggered.connect(_custom)
        menu.exec(anchor_widget.mapToGlobal(anchor_widget.rect().bottomLeft()))

    def _apply_feature_buffer(self, source_layer_name, geom, metres,
                              default_name=None):
        """Buffer a single feature — prompts for a name that becomes the buffer's
        legend key (a single named site reads better than a generic distance
        band). Cancelling the name prompt cancels the buffer."""
        from qgis.PyQt.QtWidgets import QInputDialog
        from .designated_sites import add_constraint_buffer
        base = (default_name or source_layer_name or "").strip()
        suggested = (f"{base} ({metres} m buffer)" if base
                     else f"{metres} m buffer")
        name, ok = QInputDialog.getText(
            self, "Name this buffer",
            "Name for this buffer (appears in the legend):", text=suggested)
        if not ok:
            return
        name = name.strip() or None
        ok2, res = add_constraint_buffer(source_layer_name, geom, metres, name=name)
        self._after_buffer(
            ok2, res, f"Buffered “{name or source_layer_name}” by {metres} m.")

    def _after_buffer(self, ok, res, success_msg):
        if ok:
            self.iface.mapCanvas().refresh()
            self.iface.messageBar().pushMessage(
                "Constraint Buffer", success_msg,
                level=Qgis.MessageLevel.Success, duration=4)
        else:
            reasons = {
                "no_geometry":   "that feature has no geometry to buffer.",
                "no_features":   "that layer has no features to buffer.",
                "no_layer":      "the source layer is not loaded.",
                "bad_radius":    "the distance was invalid.",
                "buffer_failed": "the buffer could not be built.",
            }
            self.iface.messageBar().pushMessage(
                "Constraint Buffer",
                "Could not create buffer — " + reasons.get(str(res), str(res)),
                level=Qgis.MessageLevel.Warning, duration=5)

    def _remove_constraint_buffer(self, type_label):
        from .designated_sites import remove_constraint_buffer
        remove_constraint_buffer(type_label)
        self.iface.mapCanvas().refresh()

    @staticmethod
    def _feature_display_name(lyr, feat):
        """Best human label for a feature across the constraint / habitat layers."""
        try:
            fields = {f.name(): feat[f.name()] for f in lyr.fields()}
        except Exception:
            return lyr.name()
        for k in ("name", "sac_name", "spa_name", "MainHabs", "band",
                  "Baseline Habitat Type", "Proposed Habitat Type",
                  "ftype", "label", "reference", "case_ref", "Risk_Zone",
                  "habitat_type", "sub_habitat_type"):
            v = fields.get(k)
            if v is not None and str(v).strip() and str(v) != "NULL":
                return str(v).strip()
        return lyr.name()

    def _buffer_analysis(self, type_label):
        """List loaded, visible features that fall within a type's buffer, with
        edge distance/direction to the redline — the 'analysis' half of the
        constraint-buffer feature."""
        from qgis.core import (
            QgsProject, QgsVectorLayer, QgsGeometry,
            QgsCoordinateTransform, QgsCoordinateReferenceSystem,
        )
        from .designated_sites import (
            constraint_buffer_geometry, _redline_distance_and_direction,
            _buffer_layer_name, active_constraint_buffers,
        )

        buf = constraint_buffer_geometry(type_label)
        if buf is None or buf.isEmpty():
            self.iface.messageBar().pushMessage(
                "Constraint Buffer", f"No buffer set for {type_label}.",
                level=Qgis.MessageLevel.Info, duration=4)
            return

        bng          = QgsCoordinateReferenceSystem("EPSG:27700")
        project      = QgsProject.instance()
        root         = project.layerTreeRoot()
        redline_geom = self._get_redline_geom()
        skip         = {_buffer_layer_name(t) for t in active_constraint_buffers()}

        hits = []
        for lyr in project.mapLayers().values():
            if not isinstance(lyr, QgsVectorLayer) or lyr.name() in skip:
                continue
            node = root.findLayer(lyr.id())
            if node is None or not node.isVisible():
                continue
            try:
                to_bng = (QgsCoordinateTransform(lyr.crs(), bng, project)
                          if lyr.crs() != bng else None)
                for feat in lyr.getFeatures():
                    g = feat.geometry()
                    if g is None or g.isEmpty():
                        continue
                    gc = QgsGeometry(g)
                    if to_bng is not None:
                        gc.transform(to_bng)
                    if not buf.intersects(gc):
                        continue
                    c = gc.centroid().asPoint()
                    dist, direction = _redline_distance_and_direction(
                        redline_geom, gc, c.x(), c.y())
                    hits.append({
                        "layer":      lyr.name(),
                        "name":       self._feature_display_name(lyr, feat),
                        "distance_m": dist,
                        "direction":  direction,
                    })
            except Exception:
                continue

        hits.sort(key=lambda h: (h["layer"], h["distance_m"]))
        self._show_buffer_analysis_dialog(type_label, hits)

    def _show_buffer_analysis_dialog(self, type_label, hits):
        from qgis.PyQt.QtWidgets import (
            QDialog, QVBoxLayout, QHBoxLayout, QLabel, QPushButton,
            QTableWidget, QTableWidgetItem, QHeaderView,
        )
        dlg = QDialog(self.iface.mainWindow())
        dlg.setWindowTitle(f"Within {type_label} buffer")
        dlg.setMinimumSize(480, 340)
        v = QVBoxLayout(dlg)

        hdr = QLabel(f"{len(hits)} feature(s) within the {type_label} buffer")
        hdr.setStyleSheet(
            "font-size:13px; font-weight:500; color:white;"
            "background:#175E3B; padding:8px 10px; border-radius:4px;")
        v.addWidget(hdr)

        tbl = QTableWidget(len(hits), 4)
        tbl.setHorizontalHeaderLabels(
            ["Layer", "Feature", "Dist. to redline", "Dir."])
        tbl.verticalHeader().setVisible(False)
        tbl.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        for r, h in enumerate(hits):
            d    = h["distance_m"]
            dtxt = f"{d} m" if d < 1000 else f"{d / 1000:.2f} km"
            for c, val in enumerate([h["layer"], h["name"], dtxt, h["direction"]]):
                tbl.setItem(r, c, QTableWidgetItem(str(val)))
        tbl.horizontalHeader().setSectionResizeMode(
            QHeaderView.ResizeMode.Stretch)
        v.addWidget(tbl)

        hl = QHBoxLayout()
        hl.addStretch()
        close = QPushButton("Close")
        close.clicked.connect(dlg.accept)
        hl.addWidget(close)
        v.addLayout(hl)
        dlg.exec()

    def _constraint_feature_at(self, map_point):
        """Nearest visible designated-site / priority-habitat feature to a canvas
        point. Returns (source_layer_name, geometry_in_BNG, feature_name) or
        None."""
        from qgis.core import (
            QgsProject, QgsVectorLayer, QgsGeometry, QgsPointXY,
            QgsFeatureRequest, QgsRectangle, QgsCoordinateTransform,
            QgsCoordinateReferenceSystem,
        )
        canvas     = self.iface.mapCanvas()
        canvas_crs = canvas.mapSettings().destinationCrs()
        bng        = QgsCoordinateReferenceSystem("EPSG:27700")
        project    = QgsProject.instance()
        root       = project.layerTreeRoot()
        px         = canvas.mapUnitsPerPixel() * 16
        click_g    = QgsGeometry.fromPointXY(QgsPointXY(map_point))
        search     = QgsRectangle(map_point.x() - px, map_point.y() - px,
                                  map_point.x() + px, map_point.y() + px)
        ds_labels  = all_ds_labels()

        best = None
        for lyr in project.mapLayers().values():
            if not isinstance(lyr, QgsVectorLayer):
                continue
            nm = lyr.name()
            if not (nm in ds_labels or nm.startswith("Priority Habitat - ")):
                continue
            node = root.findLayer(lyr.id())
            if node is None or not node.isVisible():
                continue
            try:
                lyr_crs = lyr.crs()
                s = (QgsCoordinateTransform(canvas_crs, lyr_crs, project)
                     .transformBoundingBox(search)
                     if lyr_crs != canvas_crs else search)
                for feat in lyr.getFeatures(QgsFeatureRequest().setFilterRect(s)):
                    g = feat.geometry()
                    if g is None or g.isNull():
                        continue
                    gc = QgsGeometry(g)
                    if lyr_crs != canvas_crs:
                        gc.transform(QgsCoordinateTransform(
                            lyr_crs, canvas_crs, project))
                    d = click_g.distance(gc)
                    if best is None or d < best[0]:
                        best = (d, nm, QgsGeometry(feat.geometry()), lyr_crs,
                                self._feature_display_name(lyr, feat))
            except Exception:
                continue

        if best is None:
            return None
        _d, nm, geom, lyr_crs, feat_name = best
        if lyr_crs != bng:
            try:
                geom.transform(QgsCoordinateTransform(lyr_crs, bng, project))
            except Exception:
                pass
        return nm, geom, feat_name

    def _plan_info_tool(self):
        """Lazily create (and reuse) a _FeatureInfoTool instance for the
        plan-mode right-click menu, WITHOUT activating it as the map tool — we
        only call its search / popup helpers."""
        tool = getattr(self, "_info_tool_instance", None)
        if tool is None:
            try:
                tool = _FeatureInfoTool(self.iface.mapCanvas(), self.iface, self)
                self._info_tool_instance = tool
            except Exception:
                return None
        return tool

    def _show_click_gridref(self, map_point):
        """Show (and copy) the BNG grid reference of a canvas click point."""
        from qgis.core import (
            QgsProject, QgsCoordinateTransform, QgsCoordinateReferenceSystem,
        )
        from qgis.PyQt.QtWidgets import QApplication
        canvas     = self.iface.mapCanvas()
        canvas_crs = canvas.mapSettings().destinationCrs()
        bng        = QgsCoordinateReferenceSystem("EPSG:27700")
        x, y = map_point.x(), map_point.y()
        if canvas_crs != bng:
            try:
                p = QgsCoordinateTransform(
                    canvas_crs, bng, QgsProject.instance()).transform(map_point)
                x, y = p.x(), p.y()
            except Exception:
                pass
        try:
            from .gridref_utils import osgb_en_to_gridref
            gr = osgb_en_to_gridref(int(x), int(y), 6)
        except Exception:
            gr = f"{int(x)}, {int(y)}"
        try:
            QApplication.clipboard().setText(gr)
        except Exception:
            pass
        self.iface.messageBar().pushMessage(
            "Grid reference",
            f"{gr}   (E {int(x)}, N {int(y)}) — copied to clipboard",
            level=Qgis.MessageLevel.Info, duration=6)

    def _measure_from_redline(self, map_point):
        """Edge-to-point distance from the redline boundary to a canvas click,
        copied to the clipboard (replaces the Measure-from-redline tool)."""
        from qgis.core import (
            QgsProject, QgsGeometry, QgsPointXY,
            QgsCoordinateTransform, QgsCoordinateReferenceSystem,
        )
        from qgis.PyQt.QtWidgets import QApplication
        rl = self._get_redline_geom()          # BNG
        if rl is None or rl.isEmpty():
            self.iface.messageBar().pushMessage(
                "Distance from redline", "No redline boundary drawn yet.",
                level=Qgis.MessageLevel.Info, duration=4)
            return
        canvas_crs = self.iface.mapCanvas().mapSettings().destinationCrs()
        bng = QgsCoordinateReferenceSystem("EPSG:27700")
        pt  = QgsPointXY(map_point)
        if canvas_crs != bng:
            try:
                pt = QgsCoordinateTransform(
                    canvas_crs, bng, QgsProject.instance()).transform(pt)
            except Exception:
                pass
        d = rl.distance(QgsGeometry.fromPointXY(pt))
        txt = f"{d:.0f} m" if d < 1000 else f"{d / 1000:.2f} km"
        try:
            QApplication.clipboard().setText(txt)
        except Exception:
            pass
        self.iface.messageBar().pushMessage(
            "Distance from redline",
            f"{txt} — copied to clipboard",
            level=Qgis.MessageLevel.Info, duration=6)

    def _open_point_in(self, map_point, service):
        """Deep-link a canvas click point into an external site. 'google' and
        'nbn' use WGS84 lat/lon; 'magic' uses a ~250 m BNG box (same forms as the
        External Resources buttons, but centred on the click, not the redline)."""
        from qgis.core import (
            QgsProject, QgsPointXY,
            QgsCoordinateTransform, QgsCoordinateReferenceSystem,
        )
        from qgis.PyQt.QtCore import QUrl
        from qgis.PyQt.QtGui import QDesktopServices

        canvas_crs = self.iface.mapCanvas().mapSettings().destinationCrs()
        bng  = QgsCoordinateReferenceSystem("EPSG:27700")
        wgs  = QgsCoordinateReferenceSystem("EPSG:4326")
        proj = QgsProject.instance()
        p    = QgsPointXY(map_point)

        e, n = p.x(), p.y()
        if canvas_crs != bng:
            try:
                bp = QgsCoordinateTransform(canvas_crs, bng, proj).transform(p)
                e, n = bp.x(), bp.y()
            except Exception:
                pass
        lat = lon = None
        try:
            wp = QgsCoordinateTransform(canvas_crs, wgs, proj).transform(p)
            lat, lon = wp.y(), wp.x()
        except Exception:
            pass

        if service == "magic":
            half = 125.0
            url  = _region.magic_url((e - half, n - half, e + half, n + half))
        elif service == "google" and lat is not None:
            url = f"https://www.google.com/maps/search/?api=1&query={lat:.6f},{lon:.6f}"
        elif service == "nbn" and lat is not None:
            url = (f"https://records.nbnatlas.org/occurrences/search"
                   f"?q=*:*&lat={lat:.6f}&lon={lon:.6f}&radius=2")
        else:
            return
        QDesktopServices.openUrl(QUrl(url))

    def _pond_marker_at(self, map_point):
        """Return the nearest visible SWB marker within a 12-pixel tolerance."""
        from qgis.core import (QgsCoordinateTransform, QgsGeometry, QgsPointXY,
                               QgsProject)
        canvas = self.iface.mapCanvas()
        canvas_crs = canvas.mapSettings().destinationCrs()
        click = QgsGeometry.fromPointXY(QgsPointXY(map_point))
        tolerance = canvas.mapUnitsPerPixel() * 12
        best = None
        root = QgsProject.instance().layerTreeRoot()
        for layer in QgsProject.instance().mapLayersByName(POND_LAYER_NAME):
            node = root.findLayer(layer.id())
            if node is None or not node.isVisible():
                continue
            for feature in layer.getFeatures():
                geometry = QgsGeometry(feature.geometry())
                if layer.crs() != canvas_crs:
                    try:
                        geometry.transform(QgsCoordinateTransform(
                            layer.crs(), canvas_crs, QgsProject.instance()))
                    except Exception:
                        continue
                distance = click.distance(geometry)
                if distance <= tolerance and (best is None or distance < best[0]):
                    best = (distance, layer, feature.id(),
                            str(feature["pond_id"] or "SWB marker"))
        return None if best is None else best[1:]

    def _start_pond_marker_move(self, layer, fid, label):
        """Arm a one-click tool which moves the selected marker."""
        self._pond_move_tool = _PondMoveTool(
            self.iface.mapCanvas(), self.iface, self, layer, fid)
        self.iface.mapCanvas().setMapTool(self._pond_move_tool)
        bar = self.iface.messageBar()
        self._pond_move_msg_item = bar.createMessage(
            "Move water-body marker",
            f"Click the new position for {label}; right-click or press Escape to cancel.")
        bar.pushWidget(self._pond_move_msg_item, Qgis.MessageLevel.Info, 0)

    def _pond_move_ui_reset(self):
        item = getattr(self, "_pond_move_msg_item", None)
        if item is not None:
            self._pond_move_msg_item = None
            try:
                self.iface.messageBar().popWidget(item)
            except Exception:
                pass
        self._refresh_pond_controls()

    def _delete_pond_marker(self, layer, fid, label):
        """Delete one marker selected from the canvas context menu."""
        reply = QMessageBox.question(
            self, "Delete Water-body Marker",
            f"Delete {label}?",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            QMessageBox.StandardButton.No)
        if reply != QMessageBox.StandardButton.Yes:
            return False
        from qgis.core import QgsProject
        from .pond_mapper import renumber_ponds_by_distance
        root = QgsProject.instance().layerTreeRoot()
        lines_visible = any(
            root.findLayer(line.id()) is not None
            and root.findLayer(line.id()).isVisible()
            for line in QgsProject.instance().mapLayersByName(
                POND_LINES_LAYER_NAME))
        if not layer.isEditable() and not layer.startEditing():
            return False
        if not layer.deleteFeature(fid) or not layer.commitChanges():
            layer.rollBack()
            QMessageBox.warning(
                self, "Standing Water Bodies", "The marker could not be deleted.")
            return False
        redline = self._get_redline_geom()
        renumber_ponds_by_distance(layer, redline)
        load_pond_distance_lines(redline)
        set_pond_lines_visible(lines_visible)
        self.iface.mapCanvas().refresh()
        self._refresh_pond_controls()
        return True

    def _canvas_context_menu(self, menu, event=None):
        """Plan-mode right-click menu on the map — feature info, buffer a
        constraint feature, and the grid ref of the click point. Shown only when
        a pan tool is active (plan mode) and no habitat placement is in progress;
        never mid-placement."""
        from qgis.gui import QgsMapToolPan
        from qgis.PyQt.QtWidgets import QInputDialog

        canvas = self.iface.mapCanvas()
        if not isinstance(canvas.mapTool(), QgsMapToolPan):
            return
        # A habitat form is "active" only while it is alive and still on the
        # Mapping stack. The flag is set when a form opens but only cleared by
        # a completed placement, so leaving a form via Back / a tab switch /
        # a stack rebuild used to leave it pointing at a dead form and hide
        # this menu section for the rest of the session.
        form = getattr(self, "_active_habitat_form", None)
        if form is not None:
            alive = True
            try:
                from qgis.PyQt import sip
                alive = not sip.isdeleted(form)
            except Exception:
                pass
            try:
                in_stack = alive and self.stack.indexOf(form) != -1
            except Exception:
                in_stack = alive
            if alive and in_stack:
                return
            self._active_habitat_form = None

        try:
            if event is not None:
                pt = event.mapPoint()
            else:
                pt = canvas.getCoordinateTransform().toMapCoordinates(
                    canvas.mouseLastXY())
        except Exception:
            return

        menu.addSeparator()
        header = menu.addAction("Biome")
        header.setEnabled(False)

        # A direct context menu for the numbered standing-water markers.
        try:
            pond_hit = self._pond_marker_at(pt)
        except Exception:
            pond_hit = None
        if pond_hit is not None:
            pond_layer, pond_fid, pond_label = pond_hit
            marker_menu = menu.addMenu(pond_label)
            move_marker = marker_menu.addAction("Move marker")
            move_marker.triggered.connect(
                lambda _c=False, l=pond_layer, f=pond_fid, n=pond_label:
                    self._start_pond_marker_move(l, f, n))
            delete_marker = marker_menu.addAction("Delete marker")
            delete_marker.triggered.connect(
                lambda _c=False, l=pond_layer, f=pond_fid, n=pond_label:
                    self._delete_pond_marker(l, f, n))

        # Feature info — nearest feature across all visible layers (identical to
        # the feature-info tool, reusing its search + popup).
        info_tool = self._plan_info_tool()
        if info_tool is not None:
            try:
                feat_hit = info_tool.best_feature_at(pt)
            except Exception:
                feat_hit = None
            if feat_hit is not None:
                act = menu.addAction("Feature info…")
                act.triggered.connect(
                    lambda _c=False, p=pt:
                        self._plan_info_tool().identify_at_point(p, quiet=True))

        # Buffer this feature — designated sites / priority habitats only.
        try:
            buf_hit = self._constraint_feature_at(pt)
        except Exception:
            buf_hit = None
        if buf_hit is not None:
            nm, geom, feat_name = buf_hit
            label = feat_name or nm
            sub = menu.addMenu(f"Buffer “{label}”")
            for metres in self._BUFFER_PRESETS:
                a = sub.addAction(self._buffer_preset_label(metres))
                a.triggered.connect(
                    lambda _c=False, m=metres, n=nm, g=geom, dn=feat_name:
                        self._apply_feature_buffer(n, g, m, dn))
            sub.addSeparator()
            cust = sub.addAction("Custom distance…")

            def _cust(_c=False, n=nm, g=geom, dn=feat_name):
                metres, ok = QInputDialog.getInt(
                    self, "Buffer distance",
                    "Distance from the feature edge (metres):", 250, 5, 25000, 25)
                if ok:
                    self._apply_feature_buffer(n, g, metres, dn)

            cust.triggered.connect(_cust)

        # Point tools — always available.
        menu.addSeparator()
        gr = menu.addAction("Copy grid reference of this point")
        gr.triggered.connect(lambda _c=False, p=pt: self._show_click_gridref(p))
        dist = menu.addAction("Copy distance from redline to here")
        dist.triggered.connect(lambda _c=False, p=pt: self._measure_from_redline(p))
        ext = menu.addMenu("Open this point in")
        for _svc, _lbl in (("google", "Google Maps"),
                           ("magic",  _region.magic_label()),
                           ("nbn",    "NBN Atlas")):
            act = ext.addAction(_lbl)
            act.triggered.connect(
                lambda _c=False, p=pt, s=_svc: self._open_point_in(p, s))
        menu.addSeparator()
        zoom = menu.addAction("Zoom to redline boundary")
        zoom.triggered.connect(lambda _c=False: self._zoom_to_redline())

    _PROP_LAYER_NAME = "Property Boundaries (HMLR INSPIRE)"
    _PROP_WMS_URL    = "https://inspire.landregistry.gov.uk/inspire/ows"
    _PROP_LAYER_ID   = "inspire:CP.CadastralParcel"

    def _on_property_boundaries_toggle(self, checked):
        """Load or remove the HM Land Registry INSPIRE property boundary WMS layer."""
        from qgis.core import QgsProject, QgsRasterLayer

        project = QgsProject.instance()

        if not checked:
            # "Property Addresses (OSM)" sweeps up leftovers from builds that
            # still had the removed OSM house-number labels feature.
            for lyr in (project.mapLayersByName(self._PROP_LAYER_NAME)
                        + project.mapLayersByName("Property Addresses (OSM)")):
                project.removeMapLayer(lyr.id())
            self.iface.mapCanvas().refresh()
            return

        # Already loaded — just make visible
        existing = project.mapLayersByName(self._PROP_LAYER_NAME)
        if existing:
            node = project.layerTreeRoot().findLayer(existing[0].id())
            if node is not None:
                node.setItemVisibilityChecked(True)
            self._switch_basemap("Google Hybrid")
            canvas = self.iface.mapCanvas()
            if canvas.scale() > 5000:
                canvas.zoomScale(5000)
            else:
                canvas.refresh()
            return

        # ── Step 1: probe GetCapabilities to confirm service is reachable
        # and discover the exact layer names the server exposes ──────────────
        self._prop_btn.setEnabled(False)
        from qgis.PyQt.QtWidgets import QApplication as _App
        _App.setOverrideCursor(Qt.CursorShape.WaitCursor)
        try:
            self._probe_hmlr_and_load()
        finally:
            _App.restoreOverrideCursor()
            self._prop_btn.setEnabled(True)

    def _probe_hmlr_and_load(self):
        """Load the HM Land Registry INSPIRE cadastral parcel WMS layer directly
        via QGIS's own WMS provider.

        The INSPIRE service blocks plain Python urllib requests with HTTP 403,
        so we skip the GetCapabilities probe entirely and let QGIS handle the
        connection using its own network stack (which sends appropriate headers).
        """
        from qgis.core import QgsProject, QgsRasterLayer, Qgis, QgsMessageLog
        from qgis.PyQt.QtWidgets import QMessageBox

        wms_url   = self._PROP_WMS_URL
        prop_layer = self._PROP_LAYER_ID   # "inspire:CP.CadastralParcel"

        QgsMessageLog.logMessage(
            f"[HMLR] Loading layer='{prop_layer}' directly via QGIS WMS provider",
            "Biome", Qgis.MessageLevel.Info
        )

        uri = (
            "contextualWMSLegend=0"
            "&crs=EPSG:27700"
            "&dpiMode=7"
            "&featureCount=10"
            "&format=image/png"
            f"&layers={prop_layer}"
            "&styles="
            "&version=1.3.0"
            "&IgnoreGetMapUrl=1"
            f"&url={wms_url}"
        )

        lyr = QgsRasterLayer(uri, self._PROP_LAYER_NAME, "wms")

        if not lyr.isValid():
            QgsMessageLog.logMessage(
                f"[HMLR] Layer invalid — trying WMS 1.1.1 fallback",
                "Biome", Qgis.MessageLevel.Warning
            )
            uri_fallback = (
                "contextualWMSLegend=0"
                "&crs=EPSG:27700"
                "&dpiMode=7"
                "&featureCount=10"
                "&format=image/png"
                f"&layers={prop_layer}"
                "&styles="
                "&version=1.1.1"
                "&IgnoreGetMapUrl=1"
                f"&url={wms_url}"
            )
            lyr = QgsRasterLayer(uri_fallback, self._PROP_LAYER_NAME, "wms")

        if not lyr.isValid():
            QMessageBox.critical(
                self, "Property Boundaries",
                "Could not load the HM Land Registry INSPIRE WMS layer.\n\n"
                "The service may be temporarily unavailable.\n"
                f"URL: {wms_url}"
            )
            self._prop_btn.setChecked(False)
            return

        # Desk-study overlays are never persisted (see bng_plugin's session-only
        # strip) — WMS layers need this explicit tag since they aren't memory.
        lyr.setCustomProperty("biome_session_only", 1)

        # Suppress the Ordnance Survey attribution overlay in the map corner.
        # QGIS reads attribution from WMS capabilities and stores it on the layer;
        # we clear every known location it can be stored.
        lyr.setAttribution("")
        lyr.setAttributionUrl("")
        lyr.setCustomProperty("attribution", "")
        lyr.setCustomProperty("attributionUrl", "")
        # Also clear via the layer's metadata object (QGIS 3.x)
        md = lyr.metadata()
        md.setRights([])
        md.setLicenses([])
        lyr.setMetadata(md)

        project = QgsProject.instance()
        root    = project.layerTreeRoot()
        bm      = root.findGroup("Base Maps")
        idx     = root.children().index(bm) if bm is not None else len(root.children())
        lyr.setCustomProperty("skipMemoryLayersCheck", 1)
        project.addMapLayer(lyr, False)
        root.insertLayer(idx, lyr)

        # ── Ensure OSTN15 datum transform is active ───────────────────
        # The INSPIRE WMS data is in EPSG:27700 (BNG/OSGB36). Reprojecting
        # it to the canvas CRS (EPSG:3857, WGS84) with only a Helmert
        # 7-parameter transform leaves a ~1-2 m northward residual error.
        # OSTN15 (PROJ operation 7952) corrects this to <10 cm nationally.
        # We re-apply it here and warn the user if the grid file is absent.
        self._ensure_ostn15_for_property_boundaries()

        # Disable the CopyrightLabel decoration that QGIS auto-populates from
        # WMS capabilities attribution text and renders in the map corner.
        project.writeEntry("CopyrightLabel", "/Enabled", "false")
        project.writeEntry("CopyrightLabel", "/Label", "")
        project.writeEntry("WMSAttribution", "Title", "")
        project.writeEntry("WMSAttribution", "OnlineResource", "")
        for decoration in self.iface.activeDecorations():
            dname = type(decoration).__name__.lower()
            if "copyright" in dname or "attribution" in dname:
                decoration.setEnabled(False)
                if hasattr(decoration, "setLabel"):
                    decoration.setLabel("")
                if hasattr(decoration, "setText"):
                    decoration.setText("")

        self._switch_basemap("Google Hybrid")
        canvas = self.iface.mapCanvas()
        if canvas.scale() > 5000:
            canvas.zoomScale(5000)
        else:
            canvas.refresh()
        self.iface.messageBar().pushMessage(
            "Property Boundaries",
            f"Loaded '{prop_layer}'. Zoom to 1:5,000 or closer to see boundaries.",
            level=Qgis.MessageLevel.Success, duration=8,
        )

    def _ensure_ostn15_for_property_boundaries(self):
        """Verify the OSTN15 grid shift file is installed and active.

        OSTN15 (PROJ coordinate operation 7952) corrects the residual ~1-2 m
        northward shift seen when displaying HMLR INSPIRE boundaries (EPSG:27700)
        over Google satellite (EPSG:3857 / WGS84).  Without it, QGIS falls back
        to a 7-parameter Helmert transform which leaves that visible offset.

        This method:
        1. Re-applies the OSTN15 operation to the project transform context so it
           takes effect immediately for the newly loaded layer.  "OSTN15" only
           appears in the operation's proj pipeline string (its EPSG name is
           "OSGB36 to ETRS89 (2)"), and registration must use the PROJ-6 API
           addCoordinateOperation(proj_string) — TransformDetails has no .id
           and the deprecated addSourceDestinationDatumTransform(int, int)
           form is a no-op on current QGIS builds.
        2. Uses PROJ's own op.isAvailable flag to determine whether the grid
           file is actually usable (it is False when the grid is missing).
        3. If the grid is missing, shows a one-time info bar with a link to the
           QGIS Datum Grid Downloader so the user can install it in two clicks.
        """
        from qgis.core import (
            QgsCoordinateReferenceSystem, QgsCoordinateTransformContext,
            QgsProject, QgsDatumTransform, Qgis,
        )

        proj    = QgsProject.instance()
        ctx     = proj.transformContext()
        bng     = QgsCoordinateReferenceSystem("EPSG:27700")
        merc    = QgsCoordinateReferenceSystem("EPSG:3857")
        wgs84   = QgsCoordinateReferenceSystem("EPSG:4326")

        # ── 1+2. Force OSTN15 into the transform context; PROJ tells us
        # via isAvailable whether the grid file is actually installed ──────
        ostn15_present = False
        updated = False
        try:
            for src, dst in [(bng, merc), (bng, wgs84), (merc, bng), (wgs84, bng)]:
                ops = QgsDatumTransform.operations(src, dst)
                ostn15_op = next(
                    (op for op in ops if
                     "OSTN15" in op.proj.upper() or "OSTN15" in op.name.upper()),
                    None,
                )
                if ostn15_op is not None and ostn15_op.isAvailable:
                    ostn15_present = True
                    ctx.addCoordinateOperation(src, dst, ostn15_op.proj)
                    updated = True
                else:
                    # Best available (list is accuracy-ordered)
                    avail = next((op for op in ops if op.isAvailable), None)
                    if avail is not None:
                        ctx.addCoordinateOperation(src, dst, avail.proj)
                        updated = True
            if updated:
                proj.setTransformContext(ctx)
        except Exception:
            pass

        # ── 3. Warn if missing — show once per session ─────────────────
        if not ostn15_present and not getattr(self, "_ostn15_warning_shown", False):
            self._ostn15_warning_shown = True
            self.iface.messageBar().pushMessage(
                "Property Boundaries — Datum Accuracy",
                "The OSTN15 grid shift file is not installed. Boundaries may appear "
                "offset ~1-2 m from satellite imagery. To fix: Settings → Options → "
                "CRS and Transforms → Datum Transformations → Download datum grids, "
                "then search for 'OSTN15' and install.",
                level=Qgis.MessageLevel.Warning,
                duration=0,   # stays until dismissed
            )
        elif ostn15_present:
            # Silently confirm accuracy in the log
            from qgis.core import QgsMessageLog
            QgsMessageLog.logMessage(
                "[Biome] OSTN15 grid confirmed present — property boundary "
                "accuracy <10 cm.",
                "Biome", Qgis.MessageLevel.Info,
            )

    def _sync_desk_study_buttons_from_project(self, *args):
        """Reflect desk-study layers already present in the project in the
        toggle buttons — e.g. after reopening a saved project, reloading the
        plugin, or toggling a layer's eye in the Biome layer panel. Seeds the
        layer caches so toggling reuses the existing layers instead of
        re-downloading, and puts each button in the show/hide state matching
        what is actually visible on the canvas.

        Accepts *args so it can be wired straight to signals that pass a payload
        (readProject's document, visibilityChanged's node, or layersRemoved's
        layer-id list). Re-entrancy guarded:
        it calls setChecked(), and although the buttons are clicked-driven, the
        guard makes any signal feedback a definite no-op."""
        if getattr(self, "_syncing_ds_buttons", False):
            return
        if 1 in getattr(self, "_deferred_tabs", {}):
            # Desk tab (and its buttons) not built yet — _ensure_tab_built(1)
            # re-runs this sync the moment it is, so nothing is missed.
            return
        self._syncing_ds_buttons = True
        try:
            from qgis.core import QgsProject
            from .designated_sites import (
                adopt_project_layers, _phi_layer_name,
                _AW_LAYER_NAME, _SW_LAYER_NAME,
            )

            adopt_project_layers()
            project = QgsProject.instance()
            root = project.layerTreeRoot()

            def _flat(s):
                return " ".join((s or "").split())

            def _find(name):
                flat = _flat(name)
                return next((l for l in project.mapLayers().values()
                             if _flat(l.name()) == flat), None)

            def _node_visible(node):
                # Effective visibility: the node and all its ancestor groups.
                while node is not None and node.parent() is not None:
                    if not node.itemVisibilityChecked():
                        return False
                    node = node.parent()
                return node is not None

            def _visible(name):
                lyr = _find(name)
                return (lyr is not None
                        and _node_visible(root.findLayer(lyr.id())))

            # Designated sites — the "Hide/Show All" button reflects whether
            # ANY designation layer is visible, not whether all 8 types are.
            # Most sites lack some types within the search radius, so an "all visible" test
            # would leave the button unchecked while layers are loaded and
            # desync this checkable toggle (an extra click needed to hide).
            ds_any = False
            for entry in DESIGNATED_LAYERS:
                key = entry[0]
                vis = _visible(ds_label(key))
                ds_any = ds_any or vis
                if key in self._ds_btns:
                    self._ds_btns[key].setChecked(vis)
            self._ds_toggle_all_btn.setChecked(ds_any)
            self._ds_toggle_all_btn.setText(
                "⬛  Hide All Designated Sites" if ds_any
                else "⬛  Show All\nDesignated Sites")

            # Ancient Woodland + PHI habitat groups — same "any visible" rule
            # so the Hide/Show All toggle never desyncs when only some habitat
            # groups exist within the search radius.
            aw_vis = _visible(_AW_LAYER_NAME)
            self._aw_btn.setChecked(aw_vis)
            phi_any = aw_vis
            for gkey, btn in self._phi_btns.items():
                vis = _visible(_phi_layer_name(gkey))
                phi_any = phi_any or vis
                btn.setChecked(vis)
            self._phi_toggle_all_btn.setChecked(phi_any)
            self._phi_toggle_all_btn.setText(
                "🌿  Hide All Habitats" if phi_any else "🌿  Show All\nHabitats")

            # EPSL
            vis = _visible("EPSL Locations")
            self._epsl_btn.setChecked(vis)
            self._epsl_btn.setText("🦇  Hide EPSL\nLocations" if vis
                                   else "🦇  Show EPSL\nLocations")
            for w in (self._epsl_lines_btn, self._epsl_label_btn,
                      self._epsl_filter_btn, self._epsl_style_btn):
                w.setEnabled(vis)
            lines_vis = vis and _visible("EPSL Distance Lines")
            self._epsl_lines_btn.setChecked(lines_vis)
            self._epsl_lines_btn.setText(
                "〰  Distance lines: On" if lines_vis else "〰  Distance lines: Off")

            # TPO (button reflects trees OR zones; lines/labels are point-only)
            tpo_pts_layer = _find("Tree Preservation Orders")
            vis_pts = (_visible("Tree Preservation Orders")
                       and tpo_pts_layer is not None
                       and tpo_pts_layer.featureCount() > 0)
            vis = vis_pts or _visible("TPO Zones")
            self._tpo_btn.setChecked(vis)
            self._tpo_btn.setText("🌳  Hide TPO\nLocations" if vis
                                  else "🌳  Show TPO\nLocations")
            self._tpo_lines_btn.setEnabled(vis_pts)
            self._tpo_label_btn.setEnabled(vis_pts)
            lines_vis = vis_pts and _visible("TPO Distance Lines")
            self._tpo_lines_btn.setChecked(lines_vis)
            self._tpo_lines_btn.setText(
                "〰  Distance lines: On" if lines_vis else "〰  Distance lines: Off")

            # Conservation areas
            ca_vis = _visible("Conservation Areas")
            self._ca_btn.setChecked(ca_vis)
            self._ca_btn.setText("🏛  Hide Conservation\nAreas" if ca_vis
                                 else "🏛  Show Conservation\nAreas")

            # SSSI Impact Risk Zones
            vis = _visible("SSSI IRZ — Intersecting Zones")
            self._irz_btn.setChecked(vis)
            self._irz_btn.setText("Hide SSSI impact risk zones" if vis
                                  else "Show SSSI impact risk zones")

            # GCN risk zones (any visible layer in the GCN group) + coverage
            gcn_grp = root.findGroup(_GCN_GROUP_NAME)
            gcn_vis = bool(gcn_grp) and any(
                _node_visible(n) for n in gcn_grp.findLayers())
            self._gcn_btn.setChecked(gcn_vis)
            self._gcn_coverage_btn.setChecked(_visible("GCN DLL Coverage Areas"))
            self._edna_btn.setChecked(_visible("GCN eDNA Pond Surveys"))

            # Standing water bodies
            vis = _visible(_SW_LAYER_NAME)
            self._os_topo_btn.setChecked(vis)
            known_lyr = _find(_SW_LAYER_NAME)
            known_loaded = (known_lyr is not None
                            and known_lyr.featureCount() > 0)
            self._os_topo_btn.setText(
                ("Ponds loaded" if vis else "Show loaded ponds")
                if known_loaded else "Highlight known ponds & lakes")

            # Placed water bodies — persisted in the project GeoPackage, so a
            # reopened project can arrive with ponds the session never placed;
            # the distance-lines button must reflect them.
            pond_lyr = _find(POND_LAYER_NAME)
            has_ponds = pond_lyr is not None and pond_lyr.featureCount() > 0
            self._pond_lines_btn.setEnabled(has_ponds)
            lines_vis = has_ponds and _visible(POND_LINES_LAYER_NAME)
            self._pond_lines_btn.setChecked(lines_vis)
            self._pond_lines_btn.setText(
                "〰  Distance lines: On" if lines_vis else "〰  Distance lines: Off")

            # Buffers / property boundaries (their own toggle buttons, also
            # shown as trunks in the Biome layer panel).
            if hasattr(self, "_buffers_btn"):
                buf_grp = root.findGroup("Redline Boundary Buffers")
                self._buffers_btn.setChecked(
                    bool(buf_grp) and any(_node_visible(n)
                                          for n in buf_grp.findLayers()))
            if hasattr(self, "_buffer_2km_btn"):
                self._buffer_2km_btn.setChecked(_visible("Search Buffer"))
            if hasattr(self, "_prop_btn"):
                self._prop_btn.setChecked(_visible(self._PROP_LAYER_NAME))
        except Exception:
            import traceback
            from qgis.core import QgsMessageLog, Qgis
            QgsMessageLog.logMessage(
                "Desk-study button sync failed:\n" + traceback.format_exc(),
                "Biome", Qgis.MessageLevel.Warning)
        finally:
            self._syncing_ds_buttons = False

    def _on_ds_layer_toggle(self, key, checked):
        """Load or remove a single designated site layer within the search radius."""
        if not checked:
            remove_designated_layer(key)
            self.iface.mapCanvas().refresh()
            return
        if not hasattr(self, "_centroid_bng"):
            QMessageBox.warning(self, "No Redline",
                "Please draw a redline boundary first.")
            if key in self._ds_btns:
                self._ds_btns[key].setChecked(False)
            return
        e, n = self._centroid_bng
        redline_geom = self._get_redline_geom()
        from qgis.PyQt.QtWidgets import QApplication
        from qgis.PyQt.QtCore import Qt as _Qt
        QApplication.setOverrideCursor(_Qt.CursorShape.WaitCursor)
        ok, result = load_designated_layer(key, e, n, redline_geom)
        QApplication.restoreOverrideCursor()
        if not ok:
            if result == "no_features":
                # Look up the human-readable label for this key
                _label = ds_label(key)
                self.iface.messageBar().pushMessage(
                    "Designated Sites",
                    f"No {_label} found within the search radius of the redline.",
                    level=Qgis.MessageLevel.Info, duration=5,
                )
            elif result != "no_centroid":
                QMessageBox.warning(self, "Load Failed", result)
            if key in self._ds_btns:
                self._ds_btns[key].setChecked(False)
        else:
            self.iface.mapCanvas().refresh()

    def _on_ds_toggle_all(self, checked):
        """Toggle all designated site overlays on or off."""
        if checked and not hasattr(self, "_centroid_bng"):
            QMessageBox.warning(self, "No Redline",
                "Please draw a redline boundary first.")
            self._ds_toggle_all_btn.setChecked(False)
            return

        if not checked:
            # Hide: fast — just remove layers
            for key, btn in self._ds_btns.items():
                btn.setChecked(False)
                remove_designated_layer(key)
            self._ds_toggle_all_btn.setText("⬛  Show All Designated Sites")
            self.iface.mapCanvas().refresh()
            return

        # Show: fetch all GeoJSON on worker threads (HTTP only, no QGIS objects),
        # then create QGIS layers on the main thread, then show the summary dialog.
        from concurrent.futures import ThreadPoolExecutor
        from qgis.PyQt.QtWidgets import QApplication
        from qgis.PyQt.QtCore import Qt as _Qt, QTimer
        from .designated_sites import fetch_designated_geojson

        e, n = self._centroid_bng
        keys = list(self._ds_btns.keys())
        for btn in self._ds_btns.values():
            btn.setChecked(True)

        # Pre-compute bbox from redline buffer on main thread (needs QGIS).
        # redline_geom is reused by Phase 2, so define it up front — a failure
        # below must still leave it bound (None falls back to a centroid circle).
        redline_geom = None
        try:
            from .designated_sites import _build_search_buffer
            from .config import get_search_radius
            redline_geom = self._get_redline_geom()
            sb = _build_search_buffer(e, n, redline_geom,
                                      get_search_radius("designated_sites"))
            bb = sb.boundingBox()
            bbox_str = f"{bb.xMinimum()},{bb.yMinimum()},{bb.xMaximum()},{bb.yMaximum()},EPSG:27700"
        except Exception:
            bbox_str = None

        # Layers already cached and still in the project are simply re-shown —
        # no re-download on hide → show. Only the rest are fetched.
        from .designated_sites import (_DS_LAYER_CACHE, _ds_layer_in_project,
                                       adopt_project_layers)
        adopt_project_layers()
        held = set()
        for k in keys:
            cached = _DS_LAYER_CACHE.get(k)
            if cached is not None and _ds_layer_in_project(cached):
                held.add(k)
        fetch_keys = [k for k in keys if k not in held]

        # Phase 1 — parallel HTTP fetches on worker threads (thread-safe: no QGIS calls)
        fetch_results = {}
        if fetch_keys:
            QApplication.setOverrideCursor(_Qt.CursorShape.WaitCursor)
            try:
                with ThreadPoolExecutor(max_workers=9) as pool:
                    fetch_results = dict(
                        pool.map(lambda k: (k, fetch_designated_geojson(k, e, n, bbox_str)[1]), fetch_keys)
                    )
            finally:
                QApplication.restoreOverrideCursor()

        # Phase 2 — create QGIS layers on the main thread using pre-fetched data
        # (cached keys hit load_designated_layer's re-show path instantly; the
        # rest reuse the Phase 1 GeoJSON so there is NO second download here)
        absent_keys = []
        for key in keys:
            if key not in held and fetch_results.get(key) is None:
                absent_keys.append(key)
                self._ds_btns[key].setChecked(False)
            else:
                feats = None if key in held else fetch_results.get(key)
                ok, result = load_designated_layer(
                    key, e, n, redline_geom, features=feats)
                if not ok:
                    absent_keys.append(key)
                    self._ds_btns[key].setChecked(False)

        self._ds_toggle_all_btn.setText("⬛  Hide All Designated Sites")
        self._ds_full_search_completed = True
        self._ds_full_search_had_results = (len(absent_keys) < len(keys))
        self.iface.mapCanvas().refresh()

        # Phase 3 — show absent-layer summary via QTimer so the dialog opens
        # cleanly after the button-click event has fully unwound (avoids crashes)
        if absent_keys:
            absent_labels = [ds_label(k) for k in absent_keys]
            absent_list = "\n".join("  \u2022 " + lbl for lbl in absent_labels)
            msg = (
                "The following designation types were not found within the search radius of "
                "the redline boundary and have not been loaded:\n\n"
                + absent_list
                + "\n\nThis means no such sites exist within the search area, "
                "not that the data failed to load."
            )
            QTimer.singleShot(0, lambda: QMessageBox.information(
                self, "Designated Sites \u2014 Not Found Nearby", msg
            ))

    def _on_phi_toggle_all(self, checked):
        """Toggle all PHI habitat groups (and Ancient Woodland) on or off.

        When showing, every layer is fetched concurrently on worker threads
        (HTTP only, no QGIS objects) and built on the main thread afterwards,
        so the UI no longer blocks on a serial chain of paginated downloads.
        """
        if checked and not hasattr(self, "_centroid_bng"):
            QMessageBox.warning(self, "No Redline",
                "Please draw a redline boundary first.")
            self._phi_toggle_all_btn.setChecked(False)
            return

        from qgis.PyQt.QtWidgets import QApplication
        from qgis.PyQt.QtCore import Qt as _Qt

        # ── Hide — fast, no network ──────────────────────────────────────────
        if not checked:
            if hasattr(self, "_aw_btn"):
                self._aw_btn.setChecked(False)
                hide_ancient_woodland_layer()
            for gkey, btn in self._phi_btns.items():
                btn.setChecked(False)
                hide_phi_group_layer(gkey)
            self._phi_toggle_all_btn.setText("🌿  Show All Habitats")
            self.iface.mapCanvas().refresh()
            return

        e_all = getattr(self, "_centroid_bng", (None, None))
        if e_all[0] is None:
            return
        e, n = e_all
        redline_all = self._get_redline_geom()

        # Reflect the on-state up front
        if hasattr(self, "_aw_btn"):
            self._aw_btn.setChecked(True)
        for btn in self._phi_btns.values():
            btn.setChecked(True)

        from concurrent.futures import ThreadPoolExecutor, as_completed
        from qgis.core import QgsProject
        from .designated_sites import (
            _build_search_buffer, fetch_phi_group_geojson,
            fetch_ancient_woodland_geojson, _phi_layer_name,
            _AW_LAYER_CACHE, _ds_layer_in_project, adopt_project_layers,
        )

        # Pre-compute the shared bbox once on the main thread (needs QGIS).
        # PHI (ArcGIS) takes a bare envelope; AW (WFS) needs the CRS suffix.
        # Ancient woodland shares the Priority Habitats radius, so one value
        # sizes the bbox for both.
        try:
            from .config import get_search_radius
            sb = _build_search_buffer(e, n, redline_all,
                                      get_search_radius("priority_habitats"))
            bb = sb.boundingBox()
            phi_bbox = f"{bb.xMinimum()},{bb.yMinimum()},{bb.xMaximum()},{bb.yMaximum()}"
            aw_bbox  = phi_bbox + ",EPSG:27700"
        except Exception:
            phi_bbox = aw_bbox = None

        # Already-present layers are re-shown, not re-fetched.
        adopt_project_layers()
        gkeys = list(self._phi_btns.keys())
        phi_fetch = [k for k in gkeys
                     if not QgsProject.instance().mapLayersByName(_phi_layer_name(k))]
        aw_cached = _AW_LAYER_CACHE.get("ancient_wood")
        aw_held = (not hasattr(self, "_aw_btn")
                   or (aw_cached is not None and _ds_layer_in_project(aw_cached)))

        # Phase 1 — concurrent HTTP fetches on worker threads
        phi_feats, aw_feats = {}, None
        if phi_bbox is not None and (phi_fetch or not aw_held):
            QApplication.setOverrideCursor(_Qt.CursorShape.WaitCursor)
            try:
                with ThreadPoolExecutor(max_workers=8) as pool:
                    futs = {pool.submit(fetch_phi_group_geojson, k, phi_bbox): ("phi", k)
                            for k in phi_fetch}
                    if not aw_held:
                        futs[pool.submit(fetch_ancient_woodland_geojson, aw_bbox)] = ("aw", None)
                    for fut in as_completed(futs):
                        kind, k = futs[fut]
                        try:
                            res = fut.result()
                        except Exception:
                            res = None
                        if kind == "phi":
                            phi_feats[k] = res[1] if res else None
                        else:
                            aw_feats = res
            finally:
                QApplication.restoreOverrideCursor()

        # Phase 2 — build layers on the main thread from the pre-fetched data.
        # Cached/in-project layers hit the loaders' re-show path; freshly
        # fetched layers reuse the Phase 1 data (no second download); layers
        # with no features within the search radius are left unchecked.
        if hasattr(self, "_aw_btn"):
            if aw_held:
                ok = load_ancient_woodland_layer(e, n, redline_all)[0]
            elif aw_feats:
                ok = load_ancient_woodland_layer(e, n, redline_all, features=aw_feats)[0]
            else:
                ok = False
            if not ok:
                self._aw_btn.setChecked(False)

        for gkey in gkeys:
            if gkey not in phi_fetch:
                ok = load_phi_group_layer(gkey, e, n, redline_all)[0]
            elif phi_feats.get(gkey):
                ok = load_phi_group_layer(gkey, e, n, redline_all,
                                          features=phi_feats[gkey])[0]
            else:
                ok = False
            if not ok:
                self._phi_btns[gkey].setChecked(False)

        # Only the completed "Show All" search can support the categorical
        # statement that no priority habitats were found. Individual category
        # toggles do not establish absence across the whole inventory.
        self._phi_full_search_completed = True
        self._phi_full_search_had_results = (
            (hasattr(self, "_aw_btn") and self._aw_btn.isChecked())
            or any(btn.isChecked() for btn in self._phi_btns.values()))
        self._phi_toggle_all_btn.setText("🌿  Hide All Habitats")
        self.iface.mapCanvas().refresh()

    def _on_ancient_woodland_toggle(self, checked):
        """Load or hide the Ancient Woodland layer (WFS-backed, shown in Priority Habitats group)."""
        if not checked:
            hide_ancient_woodland_layer()
            self.iface.mapCanvas().refresh()
            return
        if not hasattr(self, "_centroid_bng"):
            QMessageBox.warning(self, "No Redline",
                "Please draw a redline boundary first.")
            self._aw_btn.setChecked(False)
            return
        e, n = self._centroid_bng
        redline_geom = self._get_redline_geom()
        from qgis.PyQt.QtWidgets import QApplication
        from qgis.PyQt.QtCore import Qt as _Qt
        QApplication.setOverrideCursor(_Qt.CursorShape.WaitCursor)
        ok, result = load_ancient_woodland_layer(e, n, redline_geom)
        QApplication.restoreOverrideCursor()
        if not ok:
            if result == "no_features":
                self.iface.messageBar().pushMessage(
                    "Biome",
                    "No Ancient Woodland found within the search radius.",
                    level=Qgis.MessageLevel.Info, duration=4,
                )
            else:
                QMessageBox.critical(self, "Ancient Woodland Load Failed", result)
            self._aw_btn.setChecked(False)
        else:
            self.iface.mapCanvas().refresh()

    def _on_phi_group_toggle(self, group_key, checked):
        """Load or hide the polygon layer for a PHI habitat group."""
        if not checked:
            hide_phi_group_layer(group_key)
            self.iface.mapCanvas().refresh()
            return
        if not hasattr(self, "_centroid_bng"):
            QMessageBox.warning(self, "No Redline",
                "Please draw a redline boundary first.")
            if group_key in self._phi_btns:
                self._phi_btns[group_key].setChecked(False)
            return
        e, n = self._centroid_bng
        redline_geom = self._get_redline_geom()
        from qgis.PyQt.QtWidgets import QApplication
        from qgis.PyQt.QtCore import Qt as _Qt
        QApplication.setOverrideCursor(_Qt.CursorShape.WaitCursor)
        ok, result = load_phi_group_layer(group_key, e, n, redline_geom)
        QApplication.restoreOverrideCursor()
        if not ok:
            if result == "no_features":
                self.iface.messageBar().pushMessage(
                    "Biome",
                    "No habitats of this type found within the search radius.",
                    level=Qgis.MessageLevel.Info, duration=4,
                )
            else:
                QMessageBox.critical(self, "PHI Load Failed", result)
            if group_key in self._phi_btns:
                self._phi_btns[group_key].setChecked(False)
        else:
            self.iface.mapCanvas().refresh()

    def _run_phi_report(self):
        """Query Priority Habitats Inventory within the search radius and show results."""
        if not hasattr(self, "_centroid_bng"):
            QMessageBox.warning(self, "No Redline",
                "Please draw a redline boundary first.")
            return

        group_nearby = get_phi_group_nearby()
        if group_nearby is None:
            # Ask once, then keep the choice in the QGIS profile.  It remains
            # editable under Biome settings → Priority habitat reports.
            from qgis.PyQt.QtWidgets import QDialog, QVBoxLayout, QHBoxLayout, QLabel, QPushButton, QCheckBox
            from qgis.PyQt.QtCore import Qt as _Qt2
            grp_dlg = QDialog(self)
            grp_dlg.setWindowTitle("Priority Habitats — Grouping Option")
            grp_dlg.setMinimumWidth(480)
            grp_dlg.setWindowFlags(grp_dlg.windowFlags() & ~_Qt2.WindowType.WindowContextHelpButtonHint)
            grp_vl = QVBoxLayout(grp_dlg)
            grp_vl.setSpacing(12)
            grp_vl.setContentsMargins(16, 16, 16, 16)

            title_lbl = QLabel("<b>Group nearby habitat polygons?</b>")
            title_lbl.setStyleSheet("font-size:13px; color:#1A2420;")
            grp_vl.addWidget(title_lbl)
            desc_lbl = QLabel(
                "The PHI dataset records habitat as individual polygons. A single woodland "
                "or meadow may be split across several polygons due to mapping boundaries, "
                "even though it forms one continuous habitat.\n\n"
                "Enabling this groups polygons of the same habitat type whose boundaries "
                "are <b>5m or less apart</b>, combining them into one row with a summed area. "
                "This gives a more meaningful count of distinct habitat patches.\n\n"
                "Leave off to review each polygon individually or to reflect the raw PHI data."
            )
            desc_lbl.setWordWrap(True)
            desc_lbl.setStyleSheet("font-size:11px; color:#1A2420; line-height:1.5;")
            grp_vl.addWidget(desc_lbl)
            chk = QCheckBox("Group same-habitat polygons within 5m of each other")
            chk.setChecked(False)
            chk.setStyleSheet("font-size:11px; color:#1A2420; padding-top:4px;")
            grp_vl.addWidget(chk)
            btn_row = QHBoxLayout()
            btn_row.addStretch()
            run_btn = QPushButton("Run Report")
            run_btn.setStyleSheet(
                "QPushButton { background:#2E6B4F; color:white; border:none; "
                "border-radius:4px; font-size:11px; padding:5px 18px; }"
                "QPushButton:hover { background:#2E6B4F; }"
            )
            run_btn.setDefault(True)
            cancel_btn = QPushButton("Cancel")
            cancel_btn.setStyleSheet(
                "QPushButton { background:#F4F6F2; color:#1A2420; border:1px solid #B7BFB9; "
                "border-radius:4px; font-size:11px; padding:5px 14px; }"
                "QPushButton:hover { background:#E3E7E1; }"
            )
            run_btn.clicked.connect(grp_dlg.accept)
            cancel_btn.clicked.connect(grp_dlg.reject)
            btn_row.addWidget(cancel_btn)
            btn_row.addWidget(run_btn)
            grp_vl.addLayout(btn_row)
            if grp_dlg.exec() != QDialog.DialogCode.Accepted:
                return
            group_nearby = chk.isChecked()
            set_phi_group_nearby(group_nearby)

        e, n = self._centroid_bng
        redline_geom = self._get_redline_geom()

        def _fetch():
            rows = query_priority_habitats_2km(e, n, redline_geom, group_nearby=group_nearby)
            # Append Ancient Woodland results (separate WFS dataset)
            try:
                from .designated_sites import query_ancient_woodland_2km
                aw_rows = query_ancient_woodland_2km(e, n, redline_geom)
                rows = list(rows) + list(aw_rows)
            except Exception:
                pass
            return rows

        self._run_desk_study_async(
            "Searching Priority Habitats…",
            _fetch,
            lambda rows: self._show_phi_report_dialog(rows, e, n, group_nearby=group_nearby),
            "PHI Query Failed",
        )

    def _show_phi_report_dialog(self, rows, easting, northing, group_nearby=False):
        """Display the PHI results in a sortable table dialog."""
        from qgis.PyQt.QtWidgets import (
            QDialog, QVBoxLayout, QHBoxLayout, QTableWidget,
            QTableWidgetItem, QHeaderView, QPushButton, QLabel,
            QAbstractItemView,
        )
        from qgis.PyQt.QtCore import Qt

        try:
            from .gridref_utils import osgb_en_to_gridref
            site_gr = osgb_en_to_gridref(int(easting), int(northing), 6)
        except Exception:
            site_gr = f"{int(easting)} {int(northing)}"

        dlg = QDialog(self)
        _rad = self._radius_label("priority_habitats")
        dlg.setWindowTitle(f"Priority Habitats within {_rad}")
        dlg.setMinimumSize(860, 540)
        dlg.resize(960, 600)
        layout = QVBoxLayout(dlg)
        layout.setSpacing(8)

        grouping_note = (
            " &nbsp;·&nbsp; <span style='color:#9A6700;'>Nearby polygons grouped (≤5m)</span>"
            if group_nearby else
            " &nbsp;·&nbsp; <span style='color:#717E75;'>Individual polygons (no grouping)</span>"
        )
        hdr = QLabel(
            f"<b>Priority habitats ({_region.phi_source_name()}) within {_rad} of site centroid</b> "
            f"({site_gr}){grouping_note}"
        )
        hdr.setStyleSheet("font-size:12px; color:#1A2420; padding-bottom:4px;")
        layout.addLayout(self._report_header_row(hdr))

        if not rows:
            layout.addWidget(_banner(
                "status-check", f"No priority habitats found within {_rad} of the site centroid.",
                fg="#1F7A4D", bg="#E9F4EC", border="1px solid #BFE0CC",
                font_px=12, radius=6, margins=(12, 12, 12, 12)))
        else:
            from qgis.PyQt.QtWidgets import QTabWidget
            habitat_types = {
                habitat_type
                for row in rows
                for habitat_type in _phi_type_keys(row)
            }
            total_ha = sum(r["area_ha"] for r in rows)
            summary  = QLabel(
                f"Found <b>{len(rows)}</b> feature(s) across "
                f"<b>{len(habitat_types)}</b> habitat type(s) — "
                f"total area within {_rad}: <b>{total_ha:,.2f} ha</b>"
            )
            summary.setStyleSheet("font-size:11px; color:#54635A; padding-bottom:4px;")
            layout.addWidget(summary)

            # ── shared helpers ────────────────────────────────────────────────
            cols = ["Habitat", "Distance to redline (km)", "Code", "Area (ha)",
                    "Description", "Centroid (GR)"]

            def _fill_table(tbl, data):
                for r_idx, row in enumerate(data):
                    desc     = row["description"] or (row["add_habs"] or "")
                    dist_dir = f"{row['distance_m'] / 1000:.2f} km {row['direction']}"
                    vals = [
                        row["habitat"],
                        dist_dir,
                        row["hab_code"],
                        f"{row['area_ha']:,.2f}" if row["area_ha"] else "—",
                        desc,
                        row["centroid_gr"],
                    ]
                    for c_idx, val in enumerate(vals):
                        if c_idx == 1:
                            item = _NumericSortItem(val)
                            item.setData(Qt.ItemDataRole.UserRole, row["distance_m"])
                        else:
                            item = QTableWidgetItem(val)
                        item.setTextAlignment(
                            Qt.AlignmentFlag.AlignVCenter | (
                                Qt.AlignmentFlag.AlignCenter if c_idx in (1, 3) else Qt.AlignmentFlag.AlignLeft
                            )
                        )
                        tbl.setItem(r_idx, c_idx, item)
                tbl.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeMode.ResizeToContents)
                tbl.horizontalHeader().setSectionResizeMode(1, QHeaderView.ResizeMode.ResizeToContents)
                tbl.horizontalHeader().setSectionResizeMode(2, QHeaderView.ResizeMode.ResizeToContents)
                tbl.horizontalHeader().setSectionResizeMode(3, QHeaderView.ResizeMode.ResizeToContents)
                tbl.horizontalHeader().setSectionResizeMode(4, QHeaderView.ResizeMode.Stretch)
                tbl.horizontalHeader().setSectionResizeMode(5, QHeaderView.ResizeMode.ResizeToContents)
                tbl.verticalHeader().setVisible(False)
                tbl.setSortingEnabled(True)
                tbl.sortItems(1)

            tabs = QTabWidget()
            tabs.setMovable(True)  # drag to reorder tabs
            tabs.setStyleSheet(
                "QTabBar::tab { font-size:11px; padding:5px 14px; }"
                "QTabBar::tab:selected { font-weight:500; color:#1A2420; }"
            )

            # ── Closest-per-type helper ──────────────────────────────────────
            # One row per underlying habitat type — the nearest row containing
            # that type. A combined label is not a separate classification.
            def _closest_per_type(data):
                seen = {}
                for row in data:
                    for habitat_type in _phi_type_keys(row):
                        if (habitat_type not in seen
                                or row["distance_m"] < seen[habitat_type]["distance_m"]):
                            seen[habitat_type] = row
                unique_rows = []
                chosen_ids = set()
                for row in seen.values():
                    if id(row) not in chosen_ids:
                        chosen_ids.add(id(row))
                        unique_rows.append(row)
                return sorted(unique_rows, key=lambda r: r["distance_m"])

            # Closest per type PLUS every feature within `threshold_m` of the redline.
            # The closest of each type is always included (even if beyond the
            # threshold); within the threshold, all matching features are added too.
            def _closest_plus_within(data, threshold_m):
                result = _closest_per_type(data)
                chosen = {id(r) for r in result}
                for row in data:
                    if row["distance_m"] <= threshold_m and id(row) not in chosen:
                        chosen.add(id(row))
                        result.append(row)
                return sorted(result, key=lambda r: r["distance_m"])

            closest_rows = _closest_per_type(rows)

            # Tab definitions: (tab title base, Excel export label, data, threshold_m)
            tab_defs = [
                ("All features",     "All Features",     rows,         None),
                ("Closest per type", "Closest per Type", closest_rows, None),
            ]
            for _thr in (100, 250, 500):
                tab_defs.append((
                    f"Closest + within {_thr} m",
                    f"Closest per Type + within {_thr} m",
                    _closest_plus_within(rows, _thr),
                    _thr,
                ))

            # Keyed by table widget, NOT tab index — the tabs are movable, so
            # an index into a parallel list would misroute after a drag.
            tab_info = {}
            for _base, _exp_label, _data, _thr in tab_defs:
                _tbl = QTableWidget(len(_data), len(cols))
                _tbl.setHorizontalHeaderLabels(cols)
                _tbl.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
                self._setup_copyable_table(_tbl)
                _fill_table(_tbl, _data)
                tabs.addTab(_tbl, f"{_base} ({len(_data)})")
                tab_info[_tbl] = (_data, _exp_label, _thr, _base)

            layout.addWidget(tabs)

        btn_row = QHBoxLayout()
        if rows:
            _grey_btn_qss = """
                QPushButton {
                    background:#F4F6F2; color:#1A2420;
                    border:1px solid #B7BFB9; border-radius:4px;
                    font-size:11px; padding:4px 12px;
                }
                QPushButton:hover { background:#E3E7E1; }
            """

            # ── Copy Summary: the report-prose sentence (section 4.2.1 style) ─
            _n_types, _n_parcels = len(habitat_types), len(rows)
            _summary_text = (
                f"There {'are' if _n_types != 1 else 'is'} {_n_types} priority "
                f"habitat type{'s' if _n_types != 1 else ''} across {_n_parcels} "
                f"parcel{'s' if _n_parcels != 1 else ''} within the {_rad} search area."
            )
            summary_btn = QPushButton("Copy Summary")
            summary_btn.setIcon(_svg_icon("copy", "#1A2420", 14))
            summary_btn.setIconSize(QSize(14, 14))
            summary_btn.setStyleSheet(_grey_btn_qss)
            summary_btn.setToolTip(
                f"Copy the report sentence:\n“{_summary_text}”")
            def _copy_summary():
                from qgis.PyQt.QtWidgets import QApplication
                QApplication.clipboard().setText(_summary_text)
                _flash_copied(summary_btn, "Copy Summary")
            summary_btn.clicked.connect(_copy_summary)
            btn_row.addWidget(summary_btn)

            # ── Copy Table Heading: the "Table 4.2:" line above the table —
            #    wording follows the active tab ──────────────────────────────
            def _heading_text():
                _tbl = tabs.currentWidget()
                _data, _exp_label, _thr, _base = tab_info[_tbl]
                if _thr is not None:
                    return (
                        f"Unique priority habitats mapped within a {_rad} radius, "
                        f"plus all priority habitats within {_thr} m."
                    )
                if _base == "Closest per type":
                    return (
                        f"Priority habitats mapped within a {_rad} radius "
                        f"(closest occurrence of each type)."
                    )
                return (
                    f"All priority habitats mapped within a {_rad} radius."
                )
            heading_btn = QPushButton("Copy Table Heading")
            heading_btn.setIcon(_svg_icon("copy", "#1A2420", 14))
            heading_btn.setIconSize(QSize(14, 14))
            heading_btn.setStyleSheet(_grey_btn_qss)
            def _update_heading_tip():
                heading_btn.setToolTip(
                    f"Copy the table heading line for the visible tab "
                    f"(goes after “Table 4.2:”):\n“{_heading_text()}”")
            _update_heading_tip()
            tabs.currentChanged.connect(lambda _i: _update_heading_tip())
            def _copy_heading():
                from qgis.PyQt.QtWidgets import QApplication
                QApplication.clipboard().setText(_heading_text())
                _flash_copied(heading_btn, "Copy Table Heading")
            heading_btn.clicked.connect(_copy_heading)
            btn_row.addWidget(heading_btn)

            copy_btn = QPushButton("Copy Full Table")
            copy_btn.setIcon(_svg_icon("copy", "#1A2420", 14))
            copy_btn.setIconSize(QSize(14, 14))
            copy_btn.setStyleSheet(_grey_btn_qss)
            copy_btn.setToolTip("Copy the visible tab's table to the clipboard")
            def _copy():
                from .clipboard_utils import set_table_clipboard
                # Export whichever tab is currently visible
                active_data = tab_info[tabs.currentWidget()][0]
                header = ["Habitat", "Distance to redline (km)", "Code",
                          "Area (ha)", "Description", "Centroid (GR)"]
                grid = [header]
                for r in active_data:
                    grid.append([
                        r["habitat"],
                        f'{r["distance_m"]/1000:.2f} km {r["direction"]}',
                        r["hab_code"],
                        r["area_ha"],
                        (r["description"] or r["add_habs"] or ""),
                        r["centroid_gr"],
                    ])
                set_table_clipboard(grid, header_rows=1)
            copy_btn.clicked.connect(_copy)
            btn_row.addWidget(copy_btn)

            excel_btn = QPushButton("Export to Excel")
            excel_btn.setIcon(_svg_icon("table", "#175E3B", 14))
            excel_btn.setIconSize(QSize(14, 14))
            excel_btn.setStyleSheet("""
                QPushButton {
                    background:#F4F6F2; color:#175E3B;
                    border:1px solid #1F7A4D; border-radius:4px;
                    font-size:11px; padding:4px 12px;
                }
                QPushButton:hover { background:#E9F4EC; }
            """)
            _excel_cols_phi = [
                ("habitat",     "Habitat"),
                ("dist_dir",    "Distance to redline (km)"),
                ("hab_code",    "Code"),
                ("area_ha",     "Area (ha)"),
                ("description", "Description"),
                ("centroid_gr", "Centroid (GR)"),
            ]
            def _export_excel():
                active_data, tab_label = tab_info[tabs.currentWidget()][:2]
                _excel_rows = [{**r, "dist_dir": f"{r['distance_m']/1000:.2f} km {r['direction']}"} for r in active_data]
                self._export_report_excel(
                    _excel_rows, _excel_cols_phi,
                    f"Priority Habitats within {_rad} — {tab_label}", "Priority_Habitats_Report"
                )
            excel_btn.clicked.connect(_export_excel)
            btn_row.addWidget(excel_btn)

        btn_row.addStretch()
        close_btn = QPushButton("Close")
        close_btn.setStyleSheet("""
            QPushButton {
                background:#1F7A4D; color:white; border:none;
                border-radius:4px; font-size:11px; padding:4px 16px;
            }
            QPushButton:hover { background:#1F7A4D; }
        """)
        close_btn.clicked.connect(dlg.accept)
        btn_row.addWidget(close_btn)
        layout.addLayout(btn_row)

        # ── Draw grouped patch outlines on the map (only when grouping is on) ─
        _OUTLINE_LAYER_NAME = "PHI Grouped Patches"
        outline_layer = None
        if group_nearby and rows:
            from qgis.core import (
                QgsVectorLayer, QgsFeature, QgsProject,
                QgsSingleSymbolRenderer, QgsFillSymbol,
                QgsPalLayerSettings, QgsVectorLayerSimpleLabeling,
                QgsTextFormat, QgsTextBufferSettings,
            )
            from qgis.PyQt.QtGui import QColor as _QColor

            # Remove any stale instance
            for old in QgsProject.instance().mapLayersByName(_OUTLINE_LAYER_NAME):
                QgsProject.instance().removeMapLayer(old.id())

            outline_layer = QgsVectorLayer(
                "Polygon?crs=EPSG:27700&field=habitat:string(120)&field=area_ha:double",
                _OUTLINE_LAYER_NAME, "memory"
            )
            dp = outline_layer.dataProvider()
            feats = []
            for r in rows:
                geom = r.get("geom")
                if geom and not geom.isEmpty():
                    f = QgsFeature(outline_layer.fields())
                    f.setGeometry(geom)
                    f.setAttributes([r["habitat"], r["area_ha"]])
                    feats.append(f)
            dp.addFeatures(feats)
            outline_layer.updateExtents()

            # Style: transparent fill, bold dashed orange outline
            sym = QgsFillSymbol.createSimple({
                "color":         "255,255,255,0",
                "outline_color": "#9A6700",
                "outline_width": "0.8",
                "outline_style": "dash",
            })
            outline_layer.setRenderer(QgsSingleSymbolRenderer(sym))
            outline_layer.setOpacity(1.0)

            # Labels showing habitat name
            pal = QgsPalLayerSettings()
            pal.enabled   = True
            pal.fieldName = "habitat"
            fmt = QgsTextFormat()
            fmt.setSize(8)
            fmt.setColor(_QColor("#8A4B12"))
            buf = QgsTextBufferSettings()
            buf.setEnabled(True); buf.setSize(0.8); buf.setColor(_QColor("white"))
            fmt.setBuffer(buf)
            pal.setFormat(fmt)
            outline_layer.setLabeling(QgsVectorLayerSimpleLabeling(pal))
            outline_layer.setLabelsEnabled(True)

            # Insert above base maps
            root = QgsProject.instance().layerTreeRoot()
            bm   = root.findGroup("Base Maps")
            idx  = root.children().index(bm) if bm is not None else 0
            outline_layer.setCustomProperty("skipMemoryLayersCheck", 1)   # session-only scratch
            QgsProject.instance().addMapLayer(outline_layer, False)
            root.insertLayer(idx, outline_layer)
            self.iface.mapCanvas().refresh()

        def _on_phi_dlg_closed():
            if outline_layer is not None:
                from qgis.core import QgsProject
                try:
                    QgsProject.instance().removeMapLayer(outline_layer.id())
                except Exception:
                    pass
                self.iface.mapCanvas().refresh()

        dlg.finished.connect(_on_phi_dlg_closed)
        dlg.setModal(False)
        dlg.show()

    def _export_report_excel(self, rows, columns, title, filename_hint):
        """Export a report table to an Excel (.xlsx) file.

        Args:
            rows:         list of dicts — each dict is one data row
            columns:      list of (dict_key, header_label) tuples in display order
            title:        string shown in cell A1 as a bold title
            filename_hint: suggested filename (without extension)
        """
        try:
            import openpyxl
        except ImportError:
            from qgis.PyQt.QtWidgets import QMessageBox
            QMessageBox.critical(self, "Export Failed",
                "openpyxl is not installed. Cannot export to Excel.")
            return

        from qgis.PyQt.QtWidgets import QFileDialog, QMessageBox, QApplication
        from qgis.PyQt.QtCore import Qt as _Qt
        import os

        path, _ = QFileDialog.getSaveFileName(
            self, "Save Excel Report",
            os.path.join(os.path.expanduser("~"), f"{filename_hint}.xlsx"),
            "Excel Files (*.xlsx)"
        )
        if not path:
            return

        QApplication.setOverrideCursor(_Qt.CursorShape.WaitCursor)
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            from openpyxl.utils import get_column_letter

            wb = Workbook()
            ws = wb.active
            ws.title = title[:31]  # Excel sheet name max 31 chars

            # Colour palette
            HDR_FILL  = PatternFill("solid", fgColor="1B5E20")   # dark green
            ALT_FILL  = PatternFill("solid", fgColor="E8F5E9")   # light green
            TITLE_FONT = Font(name="Calibri", bold=False, size=13, color="FFFFFF")
            HDR_FONT  = Font(name="Calibri", bold=False, size=10, color="FFFFFF")
            BODY_FONT = Font(name="Calibri", size=10)
            CENTRE    = Alignment(horizontal="center", vertical="center", wrap_text=True)
            LEFT      = Alignment(horizontal="left",   vertical="center", wrap_text=True)
            thin      = Side(border_style="thin", color="BDBDBD")
            BORDER    = Border(left=thin, right=thin, top=thin, bottom=thin)

            n_cols = len(columns)

            # ── Row 1: Title ──────────────────────────────────────────────
            ws.merge_cells(start_row=1, start_column=1,
                           end_row=1,   end_column=n_cols)
            title_cell = ws.cell(row=1, column=1, value=title)
            title_cell.font      = Font(name="Calibri", bold=False, size=14, color="FFFFFF")
            title_cell.fill      = PatternFill("solid", fgColor="1B5E20")
            title_cell.alignment = Alignment(horizontal="left", vertical="center")
            ws.row_dimensions[1].height = 22

            # ── Row 2: column headers ─────────────────────────────────────
            for c_idx, (_, header) in enumerate(columns, start=1):
                cell = ws.cell(row=2, column=c_idx, value=header)
                cell.font      = HDR_FONT
                cell.fill      = HDR_FILL
                cell.alignment = CENTRE
                cell.border    = BORDER
            ws.row_dimensions[2].height = 30

            # ── Data rows ─────────────────────────────────────────────────
            for r_idx, row in enumerate(rows, start=3):
                fill = ALT_FILL if r_idx % 2 == 1 else PatternFill()
                for c_idx, (key, _) in enumerate(columns, start=1):
                    val  = row.get(key, "")
                    cell = ws.cell(row=r_idx, column=c_idx, value=val)
                    cell.font      = BODY_FONT
                    cell.fill      = fill
                    cell.alignment = LEFT
                    cell.border    = BORDER

            # ── Auto-fit column widths ────────────────────────────────────
            for c_idx, (key, header) in enumerate(columns, start=1):
                max_len = len(str(header))
                for row in rows:
                    max_len = max(max_len, len(str(row.get(key, "") or "")))
                ws.column_dimensions[get_column_letter(c_idx)].width = min(max_len + 4, 50)

            # ── Freeze header rows ────────────────────────────────────────
            ws.freeze_panes = "A3"

            wb.save(path)
            QApplication.restoreOverrideCursor()
            QMessageBox.information(self, "Export Complete",
                f"Report saved to:\n{path}")
        except Exception as exc:
            QApplication.restoreOverrideCursor()
            QMessageBox.critical(self, "Export Failed", str(exc))

    def _export_epsl_report_excel(self, sections, title, filename_hint):
        """Export the previous-licences report table to .xlsx.

        Reproduces the three-column report layout (Reference·Date·Species /
        Approx. Distance / Impact) with an inline header band per period.
        `sections` is a list of (band_label, rows) pairs, rows being EPSL
        record dicts as returned by query_epsl_2km.
        """
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            from qgis.PyQt.QtWidgets import QMessageBox
            QMessageBox.critical(self, "Export Failed",
                "openpyxl is not installed. Cannot export to Excel.")
            return

        from qgis.PyQt.QtWidgets import QFileDialog, QMessageBox, QApplication
        from qgis.PyQt.QtCore import Qt as _Qt
        import os

        path, _ = QFileDialog.getSaveFileName(
            self, "Save Excel Report",
            os.path.join(os.path.expanduser("~"), f"{filename_hint}.xlsx"),
            "Excel Files (*.xlsx)"
        )
        if not path:
            return

        QApplication.setOverrideCursor(_Qt.CursorShape.WaitCursor)
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            from openpyxl.utils import get_column_letter

            wb = Workbook()
            ws = wb.active
            ws.title = title[:31]

            HDR_FILL  = PatternFill("solid", fgColor="FBF4E4")
            HDR_FONT  = Font(name="Calibri", bold=True, size=10)
            BODY_FONT = Font(name="Calibri", size=10)
            LEFT      = Alignment(horizontal="left",   vertical="center", wrap_text=True)
            CENTRE    = Alignment(horizontal="center", vertical="center", wrap_text=True)
            thin      = Side(border_style="thin", color="BDBDBD")
            BORDER    = Border(left=thin, right=thin, top=thin, bottom=thin)

            r = 1
            for label, sec_rows in sections:
                heads = [f"Reference, Date and Species ({label})",
                         "Approx. Distance from Site", "Impact"]
                for c, h in enumerate(heads, start=1):
                    cell = ws.cell(row=r, column=c, value=h)
                    cell.font      = HDR_FONT
                    cell.fill      = HDR_FILL
                    cell.alignment = CENTRE if c == 2 else LEFT
                    cell.border    = BORDER
                r += 1
                for row in sorted(sec_rows, key=lambda x: x.get("distance_m", 0)):
                    ref = (f"{row['case_ref']} — {row['start_date']} — {row['species']}"
                           if row.get('start_date') else
                           f"{row['case_ref']} — {row['species']}")
                    dist = f"{row['distance_m'] / 1000:.2f} km {row['direction']}"
                    for c, v in enumerate((ref, dist, epsl_impact_summary(row)), start=1):
                        cell = ws.cell(row=r, column=c, value=v)
                        cell.font      = BODY_FONT
                        cell.alignment = CENTRE if c == 2 else LEFT
                        cell.border    = BORDER
                    r += 1

            ws.column_dimensions[get_column_letter(1)].width = 48
            ws.column_dimensions[get_column_letter(2)].width = 22
            ws.column_dimensions[get_column_letter(3)].width = 38

            wb.save(path)
            QApplication.restoreOverrideCursor()
            QMessageBox.information(self, "Export Complete",
                f"Report saved to:\n{path}")
        except Exception as exc:
            QApplication.restoreOverrideCursor()
            QMessageBox.critical(self, "Export Failed", str(exc))

    def _create_desk_study_map(self, map_type):
        """Open a print layout for a desk study map. The title is preset per map
        type, unless the user set a session-only rename (right-click the button)."""
        from .map_exporter import build_desk_study_layout
        # No EPSL data exists for a Scottish site — say so rather than print
        # an empty map (see region.py).
        if map_type == "EPSL Map" and _region.is_scotland():
            self._epsl_scotland_notice()
            return
        # GCN map only: if the eDNA pond layer is switched on, let the user pick
        # whether/which results to include. No prompt when it isn't on.
        if map_type == "GCN Risk Zones Map" and not self._prompt_edna_map_options():
            return
        # A Pond Map with no manually placed ponds still needs its water-body
        # emblem and a clear empty-result statement in the legend.  Creating
        # the empty layer here gives the layout a styled legend source without
        # drawing any false feature on the map.
        if map_type == "Pond Map":
            get_or_create_pond_layer()
        meta = {}
        if (map_type == "Priority Habitats Map"
                and getattr(self, "_phi_full_search_completed", False)
                and not getattr(self, "_phi_full_search_had_results", False)):
            meta["empty_status"] = (
                "No priority habitats found within the search area")
        if (map_type == "Designated Sites Map"
                and getattr(self, "_ds_full_search_completed", False)
                and not getattr(self, "_ds_full_search_had_results", False)):
            meta["empty_status"] = (
                "No designated sites found within the search area")
        override = getattr(self, "_map_name_overrides", {}).get(map_type)
        if override:
            meta["title"] = override
        build_desk_study_layout(self.iface, map_type, meta)

    def _create_gcn_hsi_map(self):
        """Create the dedicated GCN pond-assessment map from saved HSI data."""
        from .map_exporter import build_desk_study_layout
        get_or_create_pond_layer()
        build_desk_study_layout(self.iface, "Pond Map", {
            "title": "GCN Pond Assessment",
            "include_pond_hsi_status": True,
        })

    def _prompt_pond_map_options(self):
        """Offer assessment-status labels when producing a Pond Map."""
        # Do not interrupt the normal Pond Map workflow until the user has
        # actually recorded an outcome for at least one pond in the HSI
        # calculator. Merely having mapped/unassessed ponds is not enough.
        try:
            import json
            from qgis.core import QgsProject
            raw, ok = QgsProject.instance().readEntry(
                "Biome", "gcn_hsi_ponds_json", "")
            records = json.loads(raw) if ok and raw else []
            has_assessment = any(
                isinstance(record, dict)
                and str(record.get("status") or "not_assessed") != "not_assessed"
                for record in records if isinstance(records, list))
        except Exception:
            has_assessment = False
        if not has_assessment:
            return {"include_pond_hsi_status": False}

        from qgis.PyQt.QtWidgets import (
            QDialog, QVBoxLayout, QLabel, QCheckBox, QDialogButtonBox)
        dlg = QDialog(self)
        dlg.setWindowTitle("Pond map options")
        layout = QVBoxLayout(dlg)
        intro = QLabel(
            "Choose whether the map labels should include the HSI result or survey status "
            "recorded in the HSI calculator.")
        intro.setWordWrap(True)
        layout.addWidget(intro)
        include = QCheckBox("Include HSI result and survey status for each pond")
        include.setChecked(False)
        layout.addWidget(include)
        buttons = QDialogButtonBox(
            QDialogButtonBox.StandardButton.Ok |
            QDialogButtonBox.StandardButton.Cancel)
        buttons.accepted.connect(dlg.accept)
        buttons.rejected.connect(dlg.reject)
        layout.addWidget(buttons)
        if dlg.exec() != QDialog.DialogCode.Accepted:
            return None
        return {"include_pond_hsi_status": include.isChecked()}

    def _map_rename_menu(self, map_type, btn, pos):
        """Rename popup for a desk-study map (opened from the button's pencil):
        the name is typed straight into the popup (Enter applies, Esc cancels).
        The button label mirrors the custom name so an active rename is always
        visible; the placeholder always shows the default it reverts to."""
        from qgis.PyQt.QtWidgets import (
            QMenu, QWidgetAction, QWidget, QVBoxLayout, QLabel, QLineEdit)
        from .map_exporter import desk_map_title
        default = desk_map_title(map_type)   # includes any settings-cog reword
        cur     = self._map_name_overrides.get(map_type)

        menu = QMenu(self)
        box  = QWidget(menu)
        vl   = QVBoxLayout(box)
        vl.setContentsMargins(10, 8, 10, 8)
        vl.setSpacing(4)
        cap = QLabel("Map name this session — Enter to apply")
        cap.setStyleSheet(
            f"color:{theme.MUTED}; font-size:10px; background:transparent;")
        edit = QLineEdit(cur or default)
        edit.setPlaceholderText(default)
        edit.setMinimumWidth(250)
        edit.selectAll()
        vl.addWidget(cap)
        vl.addWidget(edit)
        act = QWidgetAction(menu)
        act.setDefaultWidget(box)
        menu.addAction(act)
        if cur:
            menu.addSeparator()
            reset = menu.addAction(f'Reset to "{default}"')
            reset.triggered.connect(
                lambda _c=False: self._apply_map_rename(map_type, ""))

        def _commit():
            menu.close()
            self._apply_map_rename(map_type, edit.text())
        edit.returnPressed.connect(_commit)
        menu.aboutToShow.connect(edit.setFocus)
        menu.exec(btn.mapToGlobal(pos))

    def _apply_map_rename(self, map_type, name):
        """Set (or clear, when blank / typed back to default) the session-only
        name for a desk-study map, and mirror it onto the button label."""
        from .map_exporter import desk_map_title
        default = desk_map_title(map_type)
        name = (name or "").strip()
        if name and name != default:
            self._map_name_overrides[map_type] = name
        else:
            self._map_name_overrides.pop(map_type, None)
        self._refresh_map_btn_label(map_type)

    def _refresh_map_btn_label(self, map_type):
        """Compose the map button's card label as "Create <name>", where <name>
        is a session rename first, else a persistent reword from the settings
        cog, else the button's original descriptive label — so the label always
        says what the click DOES while still tracking any rename."""
        from .map_exporter import DESK_TITLES, desk_map_title
        for attr, mt in self._DESK_MAP_BTNS:
            if mt != map_type:
                continue
            btn = getattr(self, attr, None)
            if btn is None or not hasattr(btn, "_label_lbl"):
                return
            name = self._map_name_overrides.get(map_type)
            if not name:
                persist = desk_map_title(map_type)
                if persist != DESK_TITLES.get(map_type, map_type):
                    name = persist
            if not name:
                name = btn._biome_default_label
                # De-capitalise a plain leading word (not an acronym like SSSI/
                # TPO) so the composed label reads as one sentence.
                if len(name) > 1 and name[1].islower():
                    name = name[0].lower() + name[1:]
            btn._label_lbl.setText("Create " + name)
            return

    def _prompt_edna_map_options(self):
        """Before building the GCN risk zones map, and only when the eDNA pond
        survey layer is switched on, ask whether to include it and which result
        categories. Returns True to proceed, False if the user cancelled.

        When the eDNA layer isn't loaded or is toggled off, no dialog appears
        and the map is built exactly as before."""
        from qgis.core import QgsProject
        from .base_layers import (
            _EDNA_LAYER_NAME, EDNA_CATEGORIES, set_edna_categories,
        )

        proj = QgsProject.instance()
        root = proj.layerTreeRoot()
        lyrs = proj.mapLayersByName(_EDNA_LAYER_NAME)
        if not lyrs:
            return True
        edna_lyr = lyrs[0]
        node = root.findLayer(edna_lyr.id())

        def _eff_vis(n):
            while n is not None and n.parent() is not None:
                if not n.itemVisibilityChecked():
                    return False
                n = n.parent()
            return True

        if node is None or not _eff_vis(node):
            return True  # eDNA not on → no prompt, build as normal

        # Count points per category for the dialog
        counts = {k: 0 for k, _ in EDNA_CATEGORIES}
        for f in edna_lyr.getFeatures():
            v = str(f["GCN_occ"] or "")
            if v in counts:
                counts[v] += 1

        prev = getattr(self, "_edna_map_categories",
                       {k for k, _ in EDNA_CATEGORIES})

        from qgis.PyQt.QtWidgets import (
            QDialog, QVBoxLayout, QGridLayout, QCheckBox, QLabel,
            QDialogButtonBox,
        )
        from qgis.PyQt.QtCore import Qt

        dlg = QDialog(self)
        dlg.setWindowTitle("eDNA ponds on this map")
        v = QVBoxLayout(dlg)
        v.setContentsMargins(16, 14, 16, 12)
        v.setSpacing(10)

        intro = QLabel("The eDNA pond survey layer is switched on. Choose what "
                       "to show on the GCN risk zones map:")
        intro.setWordWrap(True)
        v.addWidget(intro)

        master = QCheckBox("Include eDNA pond surveys")
        master.setChecked(True)
        v.addWidget(master)

        # Category rows in a grid: checkboxes in column 0, counts right-aligned
        # in column 1 so the numbers line up in a neat column.
        grid = QGridLayout()
        grid.setContentsMargins(22, 0, 0, 0)   # indent under the master row
        grid.setHorizontalSpacing(14)
        grid.setVerticalSpacing(5)
        grid.setColumnStretch(0, 1)
        boxes, count_lbls = {}, []
        for r, (key, label) in enumerate(EDNA_CATEGORIES):
            cb = QCheckBox(label)
            cb.setChecked(key in prev)
            cnt = QLabel(str(counts.get(key, 0)))
            cnt.setAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
            cnt.setStyleSheet("color:#54635A;")
            grid.addWidget(cb,  r, 0)
            grid.addWidget(cnt, r, 1)
            boxes[key] = cb
            count_lbls.append(cnt)
        v.addLayout(grid)

        def _sync_enabled():
            on = master.isChecked()
            for cb in boxes.values():
                cb.setEnabled(on)
            for lb in count_lbls:
                lb.setEnabled(on)
        master.toggled.connect(_sync_enabled)
        _sync_enabled()

        bb = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel)
        bb.button(QDialogButtonBox.StandardButton.Ok).setText("Create map")
        bb.accepted.connect(dlg.accept)
        bb.rejected.connect(dlg.reject)
        v.addWidget(bb)

        if dlg.exec() != QDialog.DialogCode.Accepted:
            return False

        selected = {k for k, cb in boxes.items() if cb.isChecked()}
        if selected:
            self._edna_map_categories = selected

        if not master.isChecked() or not selected:
            # Exclude eDNA from the map entirely
            node.setItemVisibilityChecked(False)
            self._edna_btn.setChecked(False)
        else:
            node.setItemVisibilityChecked(True)
            set_edna_categories(selected)
        return True

    def _get_redline_geom(self):
        """Return the QgsGeometry of the current redline boundary in EPSG:27700 (BNG).

        The redline layer may be in EPSG:3857 (map canvas CRS default) so we
        always reproject to BNG before returning — this is required for
        QgsGeometry.distance() to return metres against BNG feature geometries.
        """
        try:
            from qgis.core import (
                QgsGeometry, QgsCoordinateReferenceSystem,
                QgsCoordinateTransform, QgsProject,
            )
            lyr = self.layer_manager.get_or_create_layer("redline")
            feats = list(lyr.getFeatures())
            if not feats:
                return None

            # Merge all redline features
            geom = QgsGeometry(feats[0].geometry())
            for f in feats[1:]:
                geom = geom.combine(f.geometry())

            # Reproject to BNG if necessary
            src_crs = lyr.crs()
            bng     = QgsCoordinateReferenceSystem("EPSG:27700")
            if src_crs != bng:
                xform = QgsCoordinateTransform(src_crs, bng, QgsProject.instance())
                geom.transform(xform)
            return geom
        except Exception:
            return None

    def _run_ds_report(self):
        """Query designated sites within the search radius and show a formatted results dialog."""
        if not hasattr(self, "_centroid_bng"):
            QMessageBox.warning(
                self, "No Redline",
                "Please draw a redline boundary first — "
                "the report uses its centroid as the search origin."
            )
            return

        easting, northing = self._centroid_bng
        redline_geom = self._get_redline_geom()
        self._run_desk_study_async(
            "Searching designated sites…",
            lambda: query_designated_sites_2km(easting, northing, redline_geom),
            lambda rows: self._show_ds_report_dialog(rows, easting, northing),
            "Query Failed",
        )

    def _show_ds_report_dialog(self, rows, easting, northing):
        """Display the designated sites report in a scrollable dialog."""
        from qgis.PyQt.QtWidgets import (
            QDialog, QVBoxLayout, QHBoxLayout, QTableWidget,
            QTableWidgetItem, QHeaderView, QPushButton, QLabel,
            QAbstractItemView,
        )
        from qgis.PyQt.QtCore import Qt

        # Defensive cleanup for malformed/empty source features: they should
        # never become a blank final row in the displayed or copied table.
        rows = [r for r in (rows or [])
                if str(r.get("name") or "").strip()
                or str(r.get("type") or "").strip()]

        def _distance_relationship(row):
            distance_m = float(row.get("distance_m") or 0.0)
            if distance_m <= 0.0:
                return "Within designation"
            direction = str(row.get("direction") or "").strip()
            return (f"{distance_m / 1000:.2f} km {direction}".strip())

        try:
            from .gridref_utils import osgb_en_to_gridref
            site_gr = osgb_en_to_gridref(int(easting), int(northing), 6)
        except Exception:
            site_gr = f"{int(easting)} {int(northing)}"

        dlg = QDialog(self)
        _rad = self._radius_label("designated_sites")
        dlg.setWindowTitle(f"Designated Sites within {_rad}")
        dlg.setMinimumSize(820, 520)
        dlg.resize(900, 580)
        layout = QVBoxLayout(dlg)
        layout.setSpacing(8)

        # Header
        hdr = QLabel(
            f"<b>Designated sites within {_rad} of site centroid</b> "
            f"({site_gr} | E{int(easting)} N{int(northing)})"
        )
        hdr.setStyleSheet("font-size:12px; color:#1A2420; padding-bottom:4px;")
        layout.addLayout(self._report_header_row(hdr))

        if not rows:
            layout.addWidget(_banner(
                "status-check", f"No designated sites found within {_rad} of the site centroid.",
                fg="#1F7A4D", bg="#E9F4EC", border="1px solid #BFE0CC",
                font_px=12, radius=6, margins=(12, 12, 12, 12)))
        else:
            # Summary counts by type
            from collections import Counter
            counts = Counter(r["type"] for r in rows)
            summary_parts = [f"<b>{v}</b> {k}" for k, v in sorted(counts.items())]
            summary = QLabel("Found: " + " · ".join(summary_parts))
            summary.setStyleSheet("font-size:11px; color:#54635A; padding-bottom:4px;")
            layout.addWidget(summary)

            # Table
            cols = ["Site Name", "Grid Reference", "Area (ha)",
                    "Relationship to Red Line", "Type"]
            tbl  = QTableWidget(len(rows), len(cols))
            tbl.setHorizontalHeaderLabels(cols)
            tbl.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
            self._setup_copyable_table(tbl)

            for r_idx, row in enumerate(rows):
                dist_dir = _distance_relationship(row)
                # Format site name: title case, then append type suffix if not already present
                _DESIG_ACRONYMS = ["SSSI", "SAC", "SPA", "NNR", "LNR", "AONB", "LWS", "SINC", "RAMSAR"]
                def _fix_acronyms(s):
                    import re
                    for acr in _DESIG_ACRONYMS:
                        s = re.sub(r'(?<![A-Za-z])' + acr.title() + r'(?![A-Za-z])', acr, s)
                    return s
                raw_name  = row["name"] or ""
                site_type = (row["type"] or "").strip()
                display_name = _fix_acronyms(raw_name.title())
                if site_type and not display_name.upper().rstrip().endswith(site_type.upper()):
                    display_name = f"{display_name} {site_type}"
                vals = [
                    display_name,
                    row["centroid_gr"],
                    f"{row['area_ha']:,.2f}" if row["area_ha"] else "—",
                    dist_dir,
                    row["type"],
                ]
                for c_idx, val in enumerate(vals):
                    if c_idx == 3:
                        item = _NumericSortItem(val)
                        item.setData(Qt.ItemDataRole.UserRole, row["distance_m"])
                    else:
                        item = QTableWidgetItem(val)
                    item.setTextAlignment(Qt.AlignmentFlag.AlignVCenter | (Qt.AlignmentFlag.AlignCenter if c_idx in (2, 3) else Qt.AlignmentFlag.AlignLeft))
                    tbl.setItem(r_idx, c_idx, item)

            tbl.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeMode.Stretch)
            for c in (1, 2, 3, 4):
                tbl.horizontalHeader().setSectionResizeMode(c, QHeaderView.ResizeMode.ResizeToContents)
            tbl.verticalHeader().setVisible(False)
            tbl.setSortingEnabled(True)
            tbl.sortItems(3)  # sort by distance by default
            layout.addWidget(tbl)

        # Buttons row
        btn_row = QHBoxLayout()

        if rows:
            _grey_btn_qss = """
                QPushButton {
                    background:#F4F6F2; color:#1A2420;
                    border:1px solid #B7BFB9; border-radius:4px;
                    font-size:11px; padding:4px 12px;
                }
                QPushButton:hover { background:#E3E7E1; }
            """

            # ── Copy Summary: the report-prose sentence (section 4.1.1 style) ─
            _n_sites = len(rows)
            _WORDS = {1: "one", 2: "two", 3: "three", 4: "four", 5: "five",
                      6: "six", 7: "seven", 8: "eight", 9: "nine"}
            _n_word = _WORDS.get(_n_sites, str(_n_sites))
            _summary_text = (
                f"There {'is' if _n_sites == 1 else 'are'} {_n_word} designated "
                f"site{'' if _n_sites == 1 else 's'} within the {_rad} search area."
            )
            summary_btn = QPushButton("Copy Summary")
            summary_btn.setIcon(_svg_icon("copy", "#1A2420", 14))
            summary_btn.setIconSize(QSize(14, 14))
            summary_btn.setStyleSheet(_grey_btn_qss)
            summary_btn.setToolTip(
                f"Copy the report sentence:\n“{_summary_text}”")
            def _copy_summary():
                from qgis.PyQt.QtWidgets import QApplication
                QApplication.clipboard().setText(_summary_text)
                _flash_copied(summary_btn, "Copy Summary")
            summary_btn.clicked.connect(_copy_summary)
            btn_row.addWidget(summary_btn)

            # ── Copy Table Caption: caption + acronym footnote for the types
            #    actually present in the results (Table 4.1 style) ────────────
            _TYPE_EXPANSIONS = [
                ("SSSI",   "Site of Special Scientific Interest"),
                ("SAC",    "Special Area of Conservation"),
                ("SPA",    "Special Protection Area"),
                ("Ramsar", "wetland of international importance designated "
                           "under the Ramsar Convention"),
                ("NNR",    "National Nature Reserve"),
                ("LNR",    "Local Nature Reserve"),
                ("AONB",   "Area of Outstanding Natural Beauty"),
                ("NSA",    "National Scenic Area"),
                # "National Parks" is not an acronym — nothing to expand
            ]
            _present = {r["type"] for r in rows}
            _expansions = [f"{acr} = {full}" for acr, full in _TYPE_EXPANSIONS
                           if acr in _present]
            _src = _region.ds_source_name()
            _footnote = (
                f"*Data from {_src} ({'; '.join(_expansions)})."
                if _expansions else f"*Data from {_src}."
            )
            caption_btn = QPushButton("Copy Footnote")
            caption_btn.setIcon(_svg_icon("copy", "#1A2420", 14))
            caption_btn.setIconSize(QSize(14, 14))
            caption_btn.setStyleSheet(_grey_btn_qss)
            caption_btn.setToolTip(
                f"Copy the data-source footnote for the table:\n“{_footnote}”")
            def _copy_caption():
                from qgis.PyQt.QtWidgets import QApplication
                QApplication.clipboard().setText(_footnote)
                _flash_copied(caption_btn, "Copy Footnote")
            caption_btn.clicked.connect(_copy_caption)
            btn_row.addWidget(caption_btn)

            copy_btn = QPushButton("Copy Full Table")
            copy_btn.setIcon(_svg_icon("copy", "#1A2420", 14))
            copy_btn.setIconSize(QSize(14, 14))
            copy_btn.setStyleSheet(_grey_btn_qss)
            copy_btn.setToolTip("Copy the whole table to the clipboard")
            def _copy_csv():
                from .clipboard_utils import set_table_clipboard
                header = ["Site Name", "Grid Reference", "Area (ha)",
                          "Relationship to Red Line", "Type"]
                def _fmt_name(r):
                    import re
                    _DESIG_ACRONYMS = ["SSSI", "SAC", "SPA", "NNR", "LNR", "AONB", "LWS", "SINC", "RAMSAR"]
                    def _fix_acr(s):
                        for acr in _DESIG_ACRONYMS:
                            s = re.sub(r'(?<![A-Za-z])' + acr.title() + r'(?![A-Za-z])', acr, s)
                        return s
                    raw  = r["name"] or ""
                    stype = (r["type"] or "").strip()
                    name = _fix_acr(raw.title())
                    if stype and not name.upper().rstrip().endswith(stype.upper()):
                        name = f"{name} {stype}"
                    return name
                grid = [header]
                for r in rows:
                    grid.append([
                        _fmt_name(r),
                        r["centroid_gr"],
                        r["area_ha"],
                        _distance_relationship(r),
                        r["type"],
                    ])
                set_table_clipboard(grid, header_rows=1)
            copy_btn.clicked.connect(_copy_csv)
            btn_row.addWidget(copy_btn)

            excel_btn = QPushButton("Export to Excel")
            excel_btn.setIcon(_svg_icon("table", "#175E3B", 14))
            excel_btn.setIconSize(QSize(14, 14))
            excel_btn.setStyleSheet("""
                QPushButton {
                    background:#F4F6F2; color:#175E3B;
                    border:1px solid #1F7A4D; border-radius:4px;
                    font-size:11px; padding:4px 12px;
                }
                QPushButton:hover { background:#E9F4EC; }
            """)
            def _fmt_name_excel(r):
                import re
                _DESIG_ACRONYMS = ["SSSI", "SAC", "SPA", "NNR", "LNR", "AONB", "LWS", "SINC", "RAMSAR"]
                def _fix_acr(s):
                    for acr in _DESIG_ACRONYMS:
                        s = re.sub(r'(?<![A-Za-z])' + acr.title() + r'(?![A-Za-z])', acr, s)
                    return s
                raw   = r["name"] or ""
                stype = (r["type"] or "").strip()
                name  = _fix_acr(raw.title())
                if stype and not name.upper().rstrip().endswith(stype.upper()):
                    name = f"{name} {stype}"
                return name
            _excel_rows_ds = [{**r, "name": _fmt_name_excel(r),
                               "dist_dir": _distance_relationship(r)}
                              for r in rows]
            _excel_cols_ds = [
                ("name",        "Site Name"),
                ("centroid_gr", "Grid Reference"),
                ("area_ha",     "Area (ha)"),
                ("dist_dir",    "Relationship to Red Line"),
                ("type",        "Type"),
            ]
            excel_btn.clicked.connect(lambda: self._export_report_excel(
                _excel_rows_ds, _excel_cols_ds,
                f"Designated Sites within {_rad}", "Designated_Sites_Report"
            ))
            btn_row.addWidget(excel_btn)

        btn_row.addStretch()
        close_btn = QPushButton("Close")
        close_btn.setStyleSheet("""
            QPushButton {
                background:#2E6B4F; color:white;
                border:none; border-radius:4px;
                font-size:11px; padding:4px 16px;
            }
            QPushButton:hover { background:#2E6B4F; }
        """)
        close_btn.clicked.connect(dlg.accept)
        btn_row.addWidget(close_btn)
        layout.addLayout(btn_row)

        dlg.setModal(False)
        dlg.show()

    def _check_gcn_coverage_warning(self):
        """Show a warning dialog if the redline centroid falls in a GCN DLL coverage area.

        Called automatically after a redline boundary is drawn. Only fires if the
        centroid is within one of the 21 county bounding boxes. Does nothing if the
        site is outside all coverage areas.
        """
        if not hasattr(self, "_centroid_bng"):
            return
        e, n = self._centroid_bng

        from .base_layers import _find_gcn_county
        match = _find_gcn_county(e, n)
        if match is None:
            return  # outside all DLL areas — no warning needed

        service_name, _ = match
        county_label = (
            service_name
            .replace("GCN_Risk_Zones_", "")
            .replace("_", " ")
        )

        from qgis.PyQt.QtWidgets import QMessageBox
        from qgis.PyQt.QtGui import QIcon
        msg = QMessageBox(self)
        msg.setWindowTitle("GCN District Licensing — Site in Coverage Area")
        msg.setIcon(QMessageBox.Icon.Warning)
        msg.setText(
            f"<b>This site is within the GCN DLL coverage area for {county_label}.</b>"
        )
        msg.setInformativeText(
            "Great Crested Newt District Licensing applies here.<br><br>"
            "Please go to the <b>Desk Study tab</b> and load the "
            "<b>GCN Risk Zones</b> layer to determine the risk zone "
            "for this site before proceeding."
        )
        msg.setStandardButtons(QMessageBox.StandardButton.Ok)
        msg.exec()

    def _on_gcn_hsi_calculator(self):
        """Open the pop-out GCN Habitat Suitability Index calculator (ARG UK
        Advice Note 5) as a modeless window so the map remains interactive."""
        try:
            from .gcn_hsi_dialog import GcnHsiDialog
        except Exception as e:
            from qgis.core import QgsMessageLog, Qgis
            QgsMessageLog.logMessage(
                "GCN HSI calculator failed to load: %s" % e, "biome", Qgis.MessageLevel.Critical)
            return
        existing = getattr(self, "_gcn_hsi_dialog", None)
        if existing is not None:
            try:
                existing.show()
                existing.raise_()
                existing.activateWindow()
                return
            except RuntimeError:
                self._gcn_hsi_dialog = None

        dlg = GcnHsiDialog(self)
        dlg.setModal(False)
        dlg.setAttribute(Qt.WidgetAttribute.WA_DeleteOnClose, True)
        self._gcn_hsi_dialog = dlg
        dlg.destroyed.connect(
            lambda _obj=None: setattr(self, "_gcn_hsi_dialog", None))
        dlg.show()
        dlg.raise_()
        dlg.activateWindow()

    def _on_weather_checker(self):
        """Open the pop-out Survey Weather checker (historical hourly conditions
        from Open-Meteo), pre-centred on the redline. It fills the report's
        weather-conditions table — temperature (°C), wind (mph) and precipitation
        (mm) for the survey hour. Standalone; nothing is written to any layer."""
        try:
            from .weather_dialog import WeatherDialog
        except Exception as e:
            from qgis.core import QgsMessageLog, Qgis
            QgsMessageLog.logMessage(
                "Survey Weather checker failed to load: %s" % e, "biome", Qgis.MessageLevel.Critical)
            return
        WeatherDialog(self, default_latlon=self._redline_centroid_wgs84()).exec()

    def _on_gcn_coverage_toggle(self, checked: bool):
        """Show/hide a memory layer of the TRUE outlines of every GCN DLL scheme
        (dissolved risk zones from gcn_dll_coverage.json); falls back to the
        lookup-table bounding boxes only if that file is unavailable.
        Useful for quickly checking whether a site falls within a covered area."""
        from qgis.core import (
            QgsVectorLayer, QgsFeature, QgsGeometry, QgsRectangle,
            QgsProject, QgsField, QgsSingleSymbolRenderer,
            QgsPalLayerSettings, QgsVectorLayerSimpleLabeling,
            QgsTextFormat, QgsCoordinateReferenceSystem,
        )
        from qgis.PyQt.QtCore import QVariant
        from qgis.PyQt.QtGui import QColor, QFont

        LAYER_NAME = "GCN DLL Coverage Areas"
        root = QgsProject.instance().layerTreeRoot()

        # If toggling off — hide
        if not checked:
            existing = QgsProject.instance().mapLayersByName(LAYER_NAME)
            if existing:
                node = root.findLayer(existing[0].id())
                if node is not None:
                    node.setItemVisibilityChecked(False)
            return

        # If already exists — show
        existing = QgsProject.instance().mapLayersByName(LAYER_NAME)
        if existing:
            node = root.findLayer(existing[0].id())
            if node is not None:
                node.setItemVisibilityChecked(True)
            return

        # Build from the offline true outlines (bounding boxes only as fallback)
        from .base_layers import _GCN_COUNTY_LOOKUP, gcn_coverage_outlines

        lyr = QgsVectorLayer(
            "Polygon?crs=EPSG:27700&field=county:string(80)",
            LAYER_NAME, "memory"
        )
        dp = lyr.dataProvider()

        features = []
        outlines = gcn_coverage_outlines()
        if outlines:
            # Largest first so the smaller overlapping schemes (Telford inside
            # Shropshire) draw on top and their labels stay visible.
            for county_label, _service, geom in sorted(
                    outlines, key=lambda o: o[2].area(), reverse=True):
                f = QgsFeature(lyr.fields())
                f.setGeometry(QgsGeometry(geom))
                f.setAttribute("county", county_label)
                features.append(f)
        else:
            for service_name, _lid, xmin, ymin, xmax, ymax in _GCN_COUNTY_LOOKUP:
                county_label = (
                    service_name
                    .replace("GCN_Risk_Zones_", "")
                    .replace("_", " ")
                )
                rect = QgsRectangle(xmin, ymin, xmax, ymax)
                f = QgsFeature(lyr.fields())
                f.setGeometry(QgsGeometry.fromRect(rect))
                f.setAttribute("county", county_label)
                features.append(f)

        dp.addFeatures(features)
        lyr.updateExtents()

        # Style — light green outline, very transparent fill
        from qgis.core import QgsFillSymbol, QgsSingleSymbolRenderer
        sym = QgsFillSymbol.createSimple({
            "color":         "100,200,100,30",
            "outline_color": "#4caf50",
            "outline_width": "0.7",
            "outline_style": "dash",
        })
        lyr.setRenderer(QgsSingleSymbolRenderer(sym))
        lyr.setOpacity(0.85)

        # Labels
        pal = QgsPalLayerSettings()
        pal.enabled = True
        pal.fieldName = "county"
        # Note: placement defaults to OverPoint for polygon layers — no need to set explicitly
        # (setting it directly is fragile across QGIS versions due to enum changes)
        fmt = QgsTextFormat()
        fmt.setSize(8)
        fmt.setColor(QColor("#1F7A4D"))
        from qgis.core import QgsTextBufferSettings
        buf = QgsTextBufferSettings()
        buf.setEnabled(True)
        buf.setSize(0.8)
        buf.setColor(QColor("white"))
        fmt.setBuffer(buf)
        pal.setFormat(fmt)
        lyr.setLabeling(QgsVectorLayerSimpleLabeling(pal))
        lyr.setLabelsEnabled(True)

        # Add to project above Base Maps group
        bm_group = root.findGroup("Base Maps")
        insert_idx = root.children().index(bm_group) if bm_group is not None else 0
        lyr.setCustomProperty("skipMemoryLayersCheck", 1)   # session-only scratch
        QgsProject.instance().addMapLayer(lyr, False)
        root.insertLayer(insert_idx, lyr)
        node = root.findLayer(lyr.id())
        if node is not None:
            node.setItemVisibilityChecked(True)

        self.iface.mapCanvas().refresh()

    def _on_gcn_toggle(self, checked: bool):
        """Toggle GCN Risk Zones for the county containing the redline centroid."""
        if not checked:
            remove_gcn_layers()
            return

        # Ensure we have a centroid — calculate if needed
        if not hasattr(self, "_centroid_bng"):
            self._refresh_desk_study_stats()

        if not hasattr(self, "_centroid_bng"):
            from qgis.PyQt.QtWidgets import QMessageBox
            QMessageBox.information(
                self.iface.mainWindow(),
                "GCN Risk Zones",
                "Please draw a Red Line Boundary first.\n\n"
                "The GCN overlay loads only the county your site falls in — "
                "a redline centroid is needed to identify the correct county."
            )
            self._gcn_btn.setChecked(False)
            return

        easting, northing = self._centroid_bng
        self.iface.messageBar().pushMessage(
            "Biome", "Loading GCN Risk Zones…", level=Qgis.MessageLevel.Info, duration=3
        )

        ok, result = load_gcn_for_point(easting, northing)

        if ok:
            self.iface.messageBar().pushMessage(
                "GCN Risk Zones", f"Loaded: {result}", level=Qgis.MessageLevel.Info, duration=5
            )
            # Apply current slider opacity to the freshly loaded layer
            self._on_gcn_opacity_changed(self._gcn_opacity_slider.value())
        elif result == "not_in_dll_area":
            from qgis.PyQt.QtWidgets import QMessageBox
            QMessageBox.information(
                self.iface.mainWindow(),
                "GCN Risk Zones",
                "This site is not within any active GCN District Level Licensing area.\n\n"
                "GCN DLL risk zones only cover counties with an active licensing scheme. "
                "Check MAGIC or Natural England's website for the current list of DLL areas."
            )
            self._gcn_btn.setChecked(False)
        else:
            self.iface.messageBar().pushMessage(
                "GCN Risk Zones",
                "Failed to load — check internet connection.",
                level=Qgis.MessageLevel.Critical, duration=6
            )
            self._gcn_btn.setChecked(False)

    def _on_edna_toggle(self, checked: bool):
        """Toggle GCN eDNA pond survey points within the search radius of the redline.

        Shows/hides like the other desk-study toggles. The eDNA layer is
        national (all of England), so it works even where the site is outside
        the 21 DLL risk-zone areas, and it rides along on the GCN Risk Zones
        print layout when left visible on the canvas."""
        from .base_layers import load_edna_ponds_for_geom, remove_edna_layers

        if not checked:
            remove_edna_layers()
            return

        geom = self._get_redline_geom()
        if geom is None or geom.isEmpty():
            from qgis.PyQt.QtWidgets import QMessageBox
            QMessageBox.information(
                self.iface.mainWindow(),
                "GCN eDNA Pond Surveys",
                "Please draw a Red Line Boundary first.\n\n"
                "eDNA pond surveys are loaded within the search radius of the redline."
            )
            self._edna_btn.setChecked(False)
            return

        self.iface.messageBar().pushMessage(
            "Biome", "Loading GCN eDNA pond surveys…", level=Qgis.MessageLevel.Info, duration=3
        )
        ok, result = load_edna_ponds_for_geom(geom)

        if ok:
            self.iface.messageBar().pushMessage(
                "GCN eDNA Pond Surveys", f"Loaded: {result}", level=Qgis.MessageLevel.Info, duration=5
            )
        elif result == "none_within_radius":
            from qgis.PyQt.QtWidgets import QMessageBox
            QMessageBox.information(
                self.iface.mainWindow(),
                "GCN eDNA Pond Surveys",
                "No Natural England eDNA pond survey records fall within the search radius of "
                "this redline.\n\nThese surveys exist only where DLL eDNA "
                "sampling has been carried out, so gaps are common."
            )
            self._edna_btn.setChecked(False)
        else:
            self.iface.messageBar().pushMessage(
                "GCN eDNA Pond Surveys",
                "Failed to load — check internet connection.",
                level=Qgis.MessageLevel.Critical, duration=6
            )
            self._edna_btn.setChecked(False)

    def _on_naturespace_link(self):
        """Open NatureSpace's public Impact Risk Zone map, with a brief note
        that their terms apply (the full terms are restated on their site)."""
        from qgis.PyQt.QtWidgets import QMessageBox
        from qgis.PyQt.QtGui import QDesktopServices
        from qgis.PyQt.QtCore import QUrl

        URL = "https://naturespaceuk.com/district-licensing/impact-map/"
        box = QMessageBox(self)
        box.setWindowTitle("NatureSpace Impact Risk Zone Map")
        box.setIcon(QMessageBox.Icon.Information)
        box.setText("<b>Opening NatureSpace's Impact Risk Zone map</b>")
        box.setInformativeText(
            "This covers the DLL areas run by NatureSpace, not Natural "
            "England.<br><br>"
            "<b>Their terms, in brief:</b> the maps are low-resolution and "
            "illustrative only — not suitable for small sites and no substitute "
            "for survey or a full desk study. © NatureSpace Partnership Ltd: "
            "don't share, reproduce, edit, or use commercially without their "
            "written permission, and no warranty of accuracy is given.<br><br>"
            "<i>The full terms are shown again on their website.</i>"
        )
        box.setStandardButtons(QMessageBox.StandardButton.Open | QMessageBox.StandardButton.Cancel)
        box.setDefaultButton(QMessageBox.StandardButton.Open)
        if box.exec() == QMessageBox.StandardButton.Open:
            QDesktopServices.openUrl(QUrl(URL))
